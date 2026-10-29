#!/usr/bin/env python3
"""
Excel: список участников KazBuild 2026, связанных с ЛКМ и строительной химией.
Источник: официальный PDF-список участников KB'26_List of Exhibitors_RUS.pdf
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = '/home/z/my-project/download/KazBuild_2026_ЛКМ_участники.xlsx'

HEADER_FONT = Font(name='Arial', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1E3A5F')
TITLE_FONT = Font(name='Arial', size=14, bold=True, color='1E3A5F')
NORMAL_FONT = Font(name='Arial', size=10, color='1A1F36')
BOLD_FONT = Font(name='Arial', size=10, bold=True, color='1A1F36')
RED_FONT = Font(name='Arial', size=10, bold=True, color='C62828')
ORANGE_FONT = Font(name='Arial', size=10, bold=True, color='F47920')
GRAY_FONT = Font(name='Arial', size=9, color='6B7280')

LIGHT_RED = PatternFill('solid', fgColor='FEE2E2')
LIGHT_ORANGE = PatternFill('solid', fgColor='FFEDD5')
LIGHT_YELLOW = PatternFill('solid', fgColor='FEF9C3')
LIGHT_GRAY = PatternFill('solid', fgColor='F3F4F6')

THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_cell(ws, row, col, value, font=NORMAL_FONT, fill=None, align=LEFT, num_format=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align
    cell.border = BORDER
    if num_format:
        cell.number_format = num_format
    return cell


wb = openpyxl.Workbook()

# ========================================================
# ЛИСТ 1: ВСЕ ЛКМ-УЧАСТНИКИ
# ========================================================
ws = wb.active
ws.title = 'ЛКМ участники'

ws.merge_cells('A1:G1')
ws['A1'] = 'KazBuild 2026 — участники, связанные с ЛКМ и строительной химией'
ws['A1'].font = TITLE_FONT
ws['A1'].alignment = CENTER
ws.row_dimensions[1].height = 25

ws.merge_cells('A2:G2')
ws['A2'] = 'Источник: официальный список участников KB\'26 (PDF от 20.07.2026). Алматы, 2-4 сентября 2026.'
ws['A2'].font = GRAY_FONT
ws['A2'].alignment = CENTER

# Заголовки
headers = ['№', 'Компания', 'Область деятельности', 'Страна', 'Категория', 'Статус', 'Примечание для Эмпилс']
for ci, h in enumerate(headers, 1):
    style_cell(ws, 4, ci, h, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER)
ws.row_dimensions[4].height = 35

# Данные — все ЛКМ и строительная химия из списка
companies = [
    # === ПРЯМЫЕ ЛКМ (краски, лаки) ===
    (1, 'EMPILS', 'Лакокрасочные материалы (TITAN, ОРЕОЛ, Расцвет, Простокрашено!, Profiwood)', 'Россия', '🎨 ЛКМ — производитель', 'Участник',
     'Стенд 10-679. Наш стенд. TITAN 3в1, быстросохнущая, молотковая + PROFIWOOD', LIGHT_RED),
    (2, 'FAB PAINT AND CHEMICAL INDUSTRY', 'Лакокрасочные материалы', 'Азербайджан', '🎨 ЛКМ — производитель', 'Участник',
     'Стенд 10-696. Бренд FAB Boya. Воднодисперсионные, синтетические, целлюлозные краски. ПРЯМОЙ КОНКУРЕНТ.', LIGHT_RED),
    (3, 'TAUDI', 'Лаки и краски', 'Казахстан', '🎨 ЛКМ — производитель', 'Впервые',
     'Локальный казахстанский производитель ЛКМ. Впервые на выставке. Изучить ассортимент, цены, объёмы.', LIGHT_RED),
    (4, 'SAHAR SHIMI ALVAN', 'Лакокрасочная промышленность и производство смол', 'Иран', '🎨 ЛКМ — производитель (смолы)', 'Впервые',
     'Иранский производитель. ЛКМ + смолы (сырьё). Возможный поставщик сырья или конкурент в нижнем сегменте.', LIGHT_RED),

    # === СТРОИТЕЛЬНАЯ ХИМИЯ (герметики, пена, клеи) — смежная категория ===
    (5, 'KAYSDAA THAILAND', 'Силиконовые герметики, герметики для стекла, пенопластовые клеи, герметики на водной основе, клеи и герметизирующие составы', 'Таиланд', '🧪 Строительная химия', 'Участник',
     'Тайский производитель герметиков. Не ЛКМ, но смежная категория — может конкурировать за полку.', LIGHT_ORANGE),
    (6, 'WD SOLUTIONS', 'Продажа строительной химии', 'Казахстан', '🧪 Строительная химия', 'Участник',
     'Казахстанский дистрибьютор стройхимии. Возможный партнёр для дистрибуции TITAN в РК.', LIGHT_ORANGE),
    (7, 'PINGDINGSHAN AOFENG NEW MATERIAL TECHNOLOGY', 'Силиконовый герметик, пенополиуретан, клей и герметизирующее средство', 'Китай', '🧪 Строительная химия', 'Участник',
     'Китайский производитель. Герметики, пена, клеи. Демпинг возможен.', LIGHT_ORANGE),
    (8, 'SHANDONG JUHUAN NEW MATERIAL TECHNOLOGY', 'Пенополиуретан', 'Китай', '🧪 Строительная химия', 'Участник',
     'Китайский производитель монтажной пены. Не ЛКМ.', LIGHT_ORANGE),
    (9, 'NAMI GOSTAR DIBA', 'Клеевые изделия', 'Иран', '🧪 Строительная химия', 'Впервые',
     'Иранский производитель клеёв. Не ЛКМ, но смежная категория.', LIGHT_ORANGE),
    (10, 'APOLLO', 'Пенопласт, герметики', 'Китай', '🧪 Строительная химия', 'Участник',
     'Китайский поставщик. Герметики + пенопласт. Не ЛКМ.', LIGHT_ORANGE),

    # === ГИДРО/ТЕПЛОИЗОЛЯЦИЯ — частично пересекается с ЛКМ ===
    (11, 'ТЕХНОНИКОЛЬ', 'Гидроизоляция, теплоизоляция, строительная химия, защитно-дренажные материалы, звукоизоляция, полиэфирные волокна', 'Россия', '🛡 Изоляция + химия', 'Участник',
     'Крупнейший российский производитель. Строительная химия — часть ассортимента. Не прямой конкурент по ЛКМ, но бренд известен.', LIGHT_YELLOW),
    (12, 'KAZ REZ', 'Гидроизоляционные материалы, гидроизоляция и теплоизоляция', 'Казахстан', '🛡 Изоляция', 'Впервые',
     'Казахстанский производитель гидроизоляции. Не ЛКМ.', LIGHT_YELLOW),
    (13, 'KESHUN WATERPROOF', 'Гидроизоляционные материалы, уплотнительные материалы, несущие и изоляционные материалы', 'Китай', '🛡 Изоляция', 'Участник',
     'Китайский производитель гидроизоляции. Не ЛКМ.', LIGHT_YELLOW),
    (14, 'UNITED ACADEMIC DESIGN', 'Герметики, теплоизоляционные материалы, проектирование', 'Казахстан', '🛡 Изоляция + химия', 'Впервые',
     'Казахстанская компания. Герметики + теплоизоляция. Не ЛКМ.', LIGHT_YELLOW),
    (15, 'MTC-SERVICE GROUP', 'Изоляционные и сопутствующие материалы', 'Казахстан', '🛡 Изоляция', 'Участник',
     'Казахстанская компания. Не ЛКМ.', LIGHT_YELLOW),
    (16, 'КОНСТЕР ИНЖИНИРИНГ', 'Гидроизоляция и теплоизоляция', 'Казахстан', '🛡 Изоляция', 'Участник',
     'Не ЛКМ.', LIGHT_YELLOW),

    # === ОТДЕЛОЧНЫЕ МАТЕРИАЛЫ (штукатурки, наливные полы) — смежная категория ===
    (17, 'ОКС-СТРОЙ', 'Сухие строительные смеси, гипсокартон, комплектующие, отделочные материалы, звукоизоляция', 'Казахстан', '🏗 Отделочные материалы', 'Участник',
     'Казахстанский производитель сухих смесей. Не ЛКМ, но может быть партнёром по комплексным поставкам.', LIGHT_GRAY),
    (18, 'EXCLUSIVE D', 'Наливные полы', 'Узбекистан', '🏗 Отделочные материалы', 'Впервые',
     'Узбекский производитель наливных полов. Не ЛКМ, но смежная категория.', LIGHT_GRAY),
    (19, 'ECOPLIT', 'Стеновые панели и гипсоакрил', 'Узбекистан', '🏗 Отделочные материалы', 'Участник',
     'Узбекский производитель. Не ЛКМ.', LIGHT_GRAY),

    # === АССОЦИАЦИИ ===
    (20, 'АССОЦИАЦИЯ ХИМИЧЕСКОЙ ПРОМЫШЛЕННОСТИ KAZХИМИЯ', 'Ассоциация', 'Казахстан', '🏛 Ассоциация', 'Участник',
     'Казахстанская ассоциация химической промышленности. Полезный контакт для лоббирования и нетворкинга.', LIGHT_GRAY),
    (21, 'АССОЦИАЦИЯ ТЕПЛОИЗОЛЯЦИОННЫХ МАТЕРИАЛОВ (ОЮЛ)', 'Ассоциация', 'Казахстан', '🏛 Ассоциация', 'Участник',
     'Ассоциация производителей теплоизоляции. Не ЛКМ, но смежная сфера.', LIGHT_GRAY),
]

# Заполнение данных
for ri, (n, name, activity, country, cat, status, note, fill) in enumerate(companies, 5):
    is_empils = name == 'EMPILS'
    name_font = Font(name='Arial', size=10, bold=True, color='F47920' if is_empils else '1A1F36')
    
    style_cell(ws, ri, 1, n, font=BOLD_FONT, fill=fill, align=CENTER)
    style_cell(ws, ri, 2, name, font=name_font, fill=fill, align=LEFT)
    style_cell(ws, ri, 3, activity, font=NORMAL_FONT, fill=fill, align=LEFT)
    style_cell(ws, ri, 4, country, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws, ri, 5, cat, font=BOLD_FONT, fill=fill, align=LEFT)
    style_cell(ws, ri, 6, status, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws, ri, 7, note, font=NORMAL_FONT, fill=fill, align=LEFT)

# ИТОГО
r_total = 5 + len(companies)
ws.merge_cells(start_row=r_total, start_column=1, end_row=r_total, end_column=7)
style_cell(ws, r_total, 1, f'ИТОГО: {len(companies)} компаний (4 — прямые ЛКМ, 6 — строительная химия, 6 — изоляция, 3 — отделка, 2 — ассоциации)', 
           font=Font(name='Arial', size=10, bold=True, color='FFFFFF'), fill=HEADER_FILL, align=LEFT)

# Ширина
col_widths = [4, 30, 45, 12, 25, 10, 50]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w


# ========================================================
# ЛИСТ 2: ТОЛЬКО ПРЯМЫЕ ЛКМ
# ========================================================
ws2 = wb.create_sheet('Только ЛКМ')

ws2.merge_cells('A1:F1')
ws2['A1'] = 'KazBuild 2026 — только производители ЛКМ (краски, лаки)'
ws2['A1'].font = TITLE_FONT
ws2['A1'].alignment = CENTER
ws2.row_dimensions[1].height = 25

ws2.merge_cells('A2:F2')
ws2['A2'] = '4 компании из официального списка, включая Эмпилс. Это все производители красок на выставке.'
ws2['A2'].font = GRAY_FONT
ws2['A2'].alignment = CENTER

headers2 = ['№', 'Компания', 'Что производит', 'Страна', 'Статус', 'Действие для Эмпилс']
for ci, h in enumerate(headers2, 1):
    style_cell(ws2, 4, ci, h, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER)
ws2.row_dimensions[4].height = 30

paint_only = [
    (1, 'EMPILS', 'ЛКМ: TITAN (3в1, быстросохнущая, молотковая), ОРЕОЛ, Расцвет, Простокрашено!, Profiwood', 'Россия', 'Участник',
     'Стенд 10-679. Наша позиция — единственный крупный российский производитель ЛКМ на выставке.', LIGHT_RED),
    (2, 'FAB PAINT AND CHEMICAL INDUSTRY', 'ЛКМ: воднодисперсионные, синтетические, целлюлозные краски, лаки, растворители, эпоксидные смолы (FAB Boya)', 'Азербайджан', 'Участник',
     'Стенд 10-696. ПРЯМОЙ КОНКУРЕНТ. Посетить стенд, изучить цены, ассортимент, объёмы. Баку — основан в 2001 г.', LIGHT_RED),
    (3, 'TAUDI', 'Лаки и краски', 'Казахстан', 'Впервые',
     'Локальный казахстанский производитель. Впервые на выставке. Изучить: объёмы производства, цены, каналы сбыта. Возможный партнёр или конкурент.', LIGHT_RED),
    (4, 'SAHAR SHIMI ALVAN', 'Лакокрасочная промышленность + производство смол', 'Иран', 'Впервые',
     'Иранский производитель. ЛКМ + смолы (сырьё). Может быть поставщиком сырья. Изучить экспортные возможности в РФ/РК.', LIGHT_RED),
]

for ri, (n, name, what, country, status, action, fill) in enumerate(paint_only, 5):
    is_empils = name == 'EMPILS'
    name_font = Font(name='Arial', size=10, bold=True, color='F47920' if is_empils else '1A1F36')
    
    style_cell(ws2, ri, 1, n, font=BOLD_FONT, fill=fill, align=CENTER)
    style_cell(ws2, ri, 2, name, font=name_font, fill=fill, align=LEFT)
    style_cell(ws2, ri, 3, what, font=NORMAL_FONT, fill=fill, align=LEFT)
    style_cell(ws2, ri, 4, country, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws2, ri, 5, status, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws2, ri, 6, action, font=NORMAL_FONT, fill=fill, align=LEFT)

col_widths2 = [4, 30, 55, 12, 10, 55]
for i, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w


# ========================================================
# ЛИСТ 3: СТРОИТЕЛЬНАЯ ХИМИЯ (смежная категория)
# ========================================================
ws3 = wb.create_sheet('Строительная химия')

ws3.merge_cells('A1:F1')
ws3['A1'] = 'KazBuild 2026 — строительная химия (герметики, пена, клеи)'
ws3['A1'].font = TITLE_FONT
ws3['A1'].alignment = CENTER
ws3.row_dimensions[1].height = 25

ws3.merge_cells('A2:F2')
ws3['A2'] = '6 компаний. Не ЛКМ, но смежная категория — могут конкурировать за полку или быть партнёрами.'
ws3['A2'].font = GRAY_FONT
ws3['A2'].alignment = CENTER

for ci, h in enumerate(['№', 'Компания', 'Продукция', 'Страна', 'Статус', 'Примечание'], 1):
    style_cell(ws3, 4, ci, h, font=HEADER_FONT, fill=HEADER_FILL, align=CENTER)
ws3.row_dimensions[4].height = 30

chem = [
    (1, 'KAYSDAA THAILAND', 'Силиконовые герметики, герметики для стекла, пенопластовые клеи, герметики на водной основе', 'Таиланд', 'Участник',
     'Тайский производитель. Не ЛКМ, но смежная категория.'),
    (2, 'WD SOLUTIONS', 'Продажа строительной химии', 'Казахстан', 'Участник',
     'Казахстанский дистрибьютор. Возможный партнёр для дистрибуции TITAN.'),
    (3, 'PINGDINGSHAN AOFENG', 'Силиконовый герметик, пенополиуретан, клей и герметизирующее средство', 'Китай', 'Участник',
     'Китайский производитель. Возможен демпинг.'),
    (4, 'SHANDONG JUHUAN', 'Пенополиуретан', 'Китай', 'Участник', 'Монтажная пена. Не ЛКМ.'),
    (5, 'NAMI GOSTAR DIBA', 'Клеевые изделия', 'Иран', 'Впервые', 'Иранский производитель клеёв.'),
    (6, 'APOLLO', 'Пенопласт, герметики', 'Китай', 'Участник', 'Герметики + пенопласт.'),
]

for ri, (n, name, prod, country, status, note) in enumerate(chem, 5):
    fill = LIGHT_ORANGE if ri % 2 == 1 else PatternFill('solid', fgColor='FFFFFF')
    style_cell(ws3, ri, 1, n, font=BOLD_FONT, fill=fill, align=CENTER)
    style_cell(ws3, ri, 2, name, font=BOLD_FONT, fill=fill, align=LEFT)
    style_cell(ws3, ri, 3, prod, font=NORMAL_FONT, fill=fill, align=LEFT)
    style_cell(ws3, ri, 4, country, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws3, ri, 5, status, font=NORMAL_FONT, fill=fill, align=CENTER)
    style_cell(ws3, ri, 6, note, font=NORMAL_FONT, fill=fill, align=LEFT)

col_widths3 = [4, 28, 50, 12, 10, 45]
for i, w in enumerate(col_widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w


# Сохранение
wb.save(OUT)
print(f'OK: {OUT}')
import os
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
print(f'\nЛисты:')
for sn in wb.sheetnames:
    print(f'  "{sn}"')
print(f'\nВсего компаний: {len(companies)}')
print(f'  Прямые ЛКМ: 4 (включая Эмпилс)')
print(f'  Строительная химия: 6')
print(f'  Изоляция: 6')
print(f'  Отделочные материалы: 3')
print(f'  Ассоциации: 2')
