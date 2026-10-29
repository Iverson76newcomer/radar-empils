#!/usr/bin/env python3
"""
Приглашение на KazBuild 2026 в стиле Эмпилс.
Двуязычное: русский + казахский.
Формат: PowerPoint, вертикальный А4.
"""

from pptx import Presentation
from pptx.util import Pt, Cm, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

OUT = '/home/z/my-project/download/Приглашение_KazBuild_2026_Эмпилс.pptx'

# Цвета Эмпилс
ORANGE = RGBColor(0xF4, 0x76, 0x29)
DARK = RGBColor(0x31, 0x31, 0x31)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT_GRAY = RGBColor(0xF3, 0xF4, 0xF6)
LIGHT_ORANGE = RGBColor(0xFF, 0xED, 0xD5)
DARK_BLUE = RGBColor(0x1E, 0x3A, 0x5F)

# А4 портрет
prs = Presentation()
prs.slide_width = Cm(21.0)
prs.slide_height = Cm(29.7)

slide = prs.slides.add_slide(prs.slide_layouts[6])
sw = prs.slide_width
sh = prs.slide_height

# Белый фон
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
bg.line.fill.background(); bg.shadow.inherit = False

# === ВЕРХНЯЯ ОРАНЖЕВАЯ ПОЛОСА ===
top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, Cm(1.0))
top_bar.fill.solid(); top_bar.fill.fore_color.rgb = ORANGE
top_bar.line.fill.background(); top_bar.shadow.inherit = False

# === ЛОГОТИП EMPILS ===
logo_tb = slide.shapes.add_textbox(Cm(1.5), Cm(1.3), Cm(18), Cm(1.2))
tf = logo_tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "EMPILS"
r.font.name = 'Arial'; r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = ORANGE
r2 = p.add_run()
r2.text = "  ·  ЗАО «Эмпилс»  ·  Ростов-на-Дону"
r2.font.name = 'Arial'; r2.font.size = Pt(12); r2.font.color.rgb = GRAY

# === ЗАГОЛОВОК ===
title_tb = slide.shapes.add_textbox(Cm(1.5), Cm(3.0), Cm(18), Cm(2.5))
tf = title_tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "ПРИГЛАШЕНИЕ"
r.font.name = 'Arial'; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = DARK

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(4)
r = p2.add_run()
r.text = "Шақыру"
r.font.name = 'Arial'; r.font.size = Pt(22); r.font.color.rgb = GRAY

# Оранжевая разделительная линия
ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(7.5), Cm(5.8), Cm(6.0), Cm(0.08))
ul.fill.solid(); ul.fill.fore_color.rgb = ORANGE
ul.line.fill.background(); ul.shadow.inherit = False

# === ВЫСТАВКА ===
exh_tb = slide.shapes.add_textbox(Cm(1.5), Cm(6.3), Cm(18), Cm(1.5))
tf = exh_tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "KazBuild 2026"
r.font.name = 'Arial'; r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = DARK_BLUE

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(2)
r = p2.add_run()
r.text = "Казахстанская международная строительная и интерьерная выставка"
r.font.name = 'Arial'; r.font.size = Pt(11); r.font.color.rgb = GRAY

p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
r = p3.add_run()
r.text = "Қазақстан халықаралық құрылыс және интерьер көрмесі"
r.font.name = 'Arial'; r.font.size = Pt(11); r.font.color.rgb = GRAY

# === ДАТЫ И МЕСТО ===
# Тёмная карточка
date_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(2.5), Cm(9.0), Cm(16.0), Cm(3.5))
date_card.fill.solid(); date_card.fill.fore_color.rgb = DARK
date_card.line.fill.background()
date_card.shadow.inherit = False

# Даты
date_tb = slide.shapes.add_textbox(Cm(3.0), Cm(9.3), Cm(15.0), Cm(2.8))
tf = date_tb.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "2–4 сентября 2026"
r.font.name = 'Arial'; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = ORANGE

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "2–4 қыркүйек 2026"
r.font.name = 'Arial'; r.font.size = Pt(16); r.font.color.rgb = WHITE

p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(6)
r = p3.add_run()
r.text = "КЦД «Атакент»  ·  Алматы, Казахстан"
r.font.name = 'Arial'; r.font.size = Pt(12); r.font.color.rgb = WHITE

# === ВРЕМЯ РАБОТЫ ===
time_tb = slide.shapes.add_textbox(Cm(3.0), Cm(13.0), Cm(15.0), Cm(1.5))
tf = time_tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Время работы: 2–3 сентября 10:00–18:00  ·  4 сентября 10:00–16:00"
r.font.name = 'Arial'; r.font.size = Pt(10); r.font.color.rgb = GRAY

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "Жұмыс уақыты: 2–3 қыркүйек 10:00–18:00  ·  4 қыркүйек 10:00–16:00"
r.font.name = 'Arial'; r.font.size = Pt(10); r.font.color.rgb = GRAY

# === СТЕНД EMPILS ===
# Оранжевая карточка
stand_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(3.5), Cm(15.2), Cm(14.0), Cm(4.5))
stand_card.fill.solid(); stand_card.fill.fore_color.rgb = LIGHT_ORANGE
stand_card.line.color.rgb = ORANGE; stand_card.line.width = Pt(2)
stand_card.shadow.inherit = False

# Заголовок
stand_tb = slide.shapes.add_textbox(Cm(4.0), Cm(15.5), Cm(13.0), Cm(3.5))
tf = stand_tb.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "ПОСЕТИТЕ НАШ СТЕНД"
r.font.name = 'Arial'; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = DARK

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "Біздің стендімізге шақырамыз!"
r.font.name = 'Arial'; r.font.size = Pt(12); r.font.color.rgb = GRAY

# Номер стенда — крупно
p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(8)
r = p3.add_run()
r.text = "Павильон 10  ·  Стенд 10-679"
r.font.name = 'Arial'; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = ORANGE

# === ПРОМОКОД ===
# Карточка промокода
promo_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(4.5), Cm(20.5), Cm(12.0), Cm(3.5))
promo_card.fill.solid(); promo_card.fill.fore_color.rgb = WHITE
promo_card.line.color.rgb = DARK; promo_card.line.width = Pt(1.5)
promo_card.shadow.inherit = False

# Текст промокода
promo_tb = slide.shapes.add_textbox(Cm(5.0), Cm(20.7), Cm(11.0), Cm(3.0))
tf = promo_tb.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "БЕСПЛАТНОЕ ПОСЕЩЕНИЕ"
r.font.name = 'Arial'; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = DARK

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "Тегін кіру үшін промокодты пайдаланыңыз"
r.font.name = 'Arial'; r.font.size = Pt(10); r.font.color.rgb = GRAY

p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(6)
r = p3.add_run()
r.text = "ПРОМОКОД:"
r.font.name = 'Arial'; r.font.size = Pt(11); r.font.color.rgb = GRAY

p4 = tf.add_paragraph(); p4.alignment = PP_ALIGN.CENTER
r = p4.add_run()
r.text = "KBA5GR5U4G"
r.font.name = 'Arial'; r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = ORANGE

p5 = tf.add_paragraph(); p5.alignment = PP_ALIGN.CENTER
p5.space_before = Pt(4)
r = p5.add_run()
r.text = "Получите e-ticket на kazbuild.kz"
r.font.name = 'Arial'; r.font.size = Pt(9); r.font.color.rgb = GRAY

# === QR-КОД ЗАГЛУШКА ===
qr_placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(7.5), Cm(24.5), Cm(6.0), Cm(3.0))
qr_placeholder.fill.solid(); qr_placeholder.fill.fore_color.rgb = LIGHT_GRAY
qr_placeholder.line.color.rgb = GRAY; qr_placeholder.line.width = Pt(0.5)
qr_placeholder.shadow.inherit = False

qr_tb = slide.shapes.add_textbox(Cm(7.5), Cm(24.5), Cm(6.0), Cm(3.0))
tf = qr_tb.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "[ QR-код ]\nkazbuild.kz"
r.font.name = 'Arial'; r.font.size = Pt(12); r.font.color.rgb = GRAY

# === НИЖНЯЯ ПОЛОСА ===
footer_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Cm(28.2), sw, Cm(1.5))
footer_bar.fill.solid(); footer_bar.fill.fore_color.rgb = DARK
footer_bar.line.fill.background(); footer_bar.shadow.inherit = False

footer_tb = slide.shapes.add_textbox(Cm(1.5), Cm(28.3), Cm(18.0), Cm(1.3))
tf = footer_tb.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "EMPILS  ·  "
r.font.name = 'Arial'; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = ORANGE
r2 = p.add_run()
r2.text = "empils.ru  ·  titanpaints.ru  ·  +7 (863) 203-71-30"
r2.font.name = 'Arial'; r2.font.size = Pt(10); r2.font.color.rgb = WHITE

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "Красишь 1 раз — служит 10 лет"
r.font.name = 'Arial'; r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = ORANGE

prs.save(OUT)
print(f'OK: {OUT}')
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
