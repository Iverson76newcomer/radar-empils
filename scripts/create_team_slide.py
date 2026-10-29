#!/usr/bin/env python3
"""
Отдельный слайд: Команда Эмпилс на вашей территории.
"""

from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

SRC = '/home/z/my-project/upload/Презентация_партнёрам_разделители.pptx'
OUT = '/home/z/my-project/download/Слайд_команда_Эмпилс.pptx'

COL_ORANGE = RGBColor(0xF4, 0x76, 0x29)
COL_DARK = RGBColor(0x31, 0x31, 0x31)
COL_BLACK = RGBColor(0x00, 0x00, 0x00)
COL_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COL_GRAY = RGBColor(0x6B, 0x72, 0x80)
COL_LIGHT_BG = RGBColor(0xF8, 0xF8, 0xF8)
COL_LIGHT_ORANGE = RGBColor(0xFF, 0xED, 0xD5)
COL_LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)
COL_LIGHT_GREEN = RGBColor(0xDC, 0xFC, 0xE7)

prs = Presentation(SRC)
slide = prs.slides.add_slide(prs.slide_layouts[0])

sw = prs.slide_width
sh = prs.slide_height

# Белый фон
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
bg.fill.solid(); bg.fill.fore_color.rgb = COL_WHITE
bg.line.fill.background(); bg.shadow.inherit = False

# === ЗАГОЛОВОК ===
title_box = slide.shapes.add_textbox(Cm(1.7), Cm(1.0), Cm(35), Cm(1.2))
tf = title_box.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "КОМАНДА ЭМПИЛС НА ВАШЕЙ ТЕРРИТОРИИ"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(32); r.font.color.rgb = COL_DARK

ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.7), Cm(2.2), Cm(6.5), Cm(0.1))
ul.fill.solid(); ul.fill.fore_color.rgb = COL_ORANGE
ul.line.fill.background(); ul.shadow.inherit = False

sub = slide.shapes.add_textbox(Cm(1.7), Cm(2.5), Cm(40), Cm(0.8))
tf = sub.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "Начните работать с Эмпилс — и наша команда поможет вам масштабировать продажи"
r.font.name = 'Inter Regular'; r.font.size = Pt(13); r.font.color.rgb = COL_GRAY

# === ЛЕВАЯ ЧАСТЬ: Структура команды (3 уровня) ===

sec1 = slide.shapes.add_textbox(Cm(1.5), Cm(4.0), Cm(20), Cm(0.7))
tf = sec1.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "СТРУКТУРА КОМАНДЫ"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = COL_ORANGE

team = [
    {
        'role': 'Региональный директор',
        'desc': 'Ответственный за территорию. Стратегическое планирование, согласование условий, решение escalated-вопросов.',
        'detail': 'Курирует всех дистрибьюторов на территории',
        'color': COL_LIGHT_BLUE,
        'icon': '🎯',
    },
    {
        'role': 'Региональный менеджер',
        'desc': 'Оперативное управление. Контроль отгрузок, MML, маркетинговые программы, аналитика продаж.',
        'detail': 'Точка контакта для дистрибьютора на ежедневной основе',
        'color': COL_LIGHT_GREEN,
        'icon': '📊',
    },
    {
        'role': 'ЭТП — Эксклюзивный торговый представитель',
        'desc': 'Полевая работа с АКБ. Визиты в РТТ, обучение ТП дистрибьютора, выкладка, POS-материалы, контроль полки.',
        'detail': '120 РТТ на 1 ЭТП',
        'color': COL_LIGHT_ORANGE,
        'icon': '🚗',
    },
]

for i, member in enumerate(team):
    y = Cm(5.0 + i * 5.5)
    
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.5), y, Cm(20.0), Cm(4.8))
    card.fill.solid(); card.fill.fore_color.rgb = member['color']
    card.line.color.rgb = COL_ORANGE; card.line.width = Pt(1)
    card.shadow.inherit = False
    
    role_box = slide.shapes.add_textbox(Cm(2.0), y + Cm(0.3), Cm(18.5), Cm(1.0))
    tf = role_box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"{member['icon']}  {member['role']}"
    r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = COL_DARK
    
    desc_box = slide.shapes.add_textbox(Cm(2.0), y + Cm(1.5), Cm(18.5), Cm(2.0))
    tf = desc_box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = member['desc']
    r.font.name = 'Inter Regular'; r.font.size = Pt(10); r.font.color.rgb = COL_BLACK
    
    detail_box = slide.shapes.add_textbox(Cm(2.0), y + Cm(3.5), Cm(18.5), Cm(0.8))
    tf = detail_box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = member['detail']
    r.font.name = 'Inter Regular'; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = COL_ORANGE
    
    if i < len(team) - 1:
        arrow = slide.shapes.add_textbox(Cm(11.0), y + Cm(4.6), Cm(1.0), Cm(1.0))
        tf = arrow.text_frame
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = p.add_run()
        r.text = "↓"
        r.font.name = 'Arial'; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = COL_ORANGE

# === ПРАВАЯ ЧАСТЬ: Принцип масштабирования ===

sec2 = slide.shapes.add_textbox(Cm(23.0), Cm(4.0), Cm(20), Cm(0.7))
tf = sec2.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "ПРИНЦИП МАСШТАБИРОВАНИЯ"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = COL_ORANGE

# Тёмная карточка с формулой
formula_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(23.0), Cm(5.0), Cm(20.0), Cm(5.0))
formula_card.fill.solid(); formula_card.fill.fore_color.rgb = COL_DARK
formula_card.line.fill.background()
formula_card.shadow.inherit = False

formula_box = slide.shapes.add_textbox(Cm(24.0), Cm(5.5), Cm(18.0), Cm(4.0))
tf = formula_box.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Больше АКБ с Эмпилс"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = COL_WHITE

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(8)
r = p2.add_run()
r.text = "↓"
r.font.name = 'Arial'; r.font.size = Pt(24); r.font.color.rgb = COL_ORANGE

p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(4)
r = p3.add_run()
r.text = "Больше наша команда на вашей территории"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = COL_ORANGE

# 120 РТТ — оранжевая карточка
norm_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(23.0), Cm(10.5), Cm(20.0), Cm(4.0))
norm_card.fill.solid(); norm_card.fill.fore_color.rgb = COL_LIGHT_ORANGE
norm_card.line.color.rgb = COL_ORANGE; norm_card.line.width = Pt(1)
norm_card.shadow.inherit = False

norm_box = slide.shapes.add_textbox(Cm(24.0), Cm(10.8), Cm(18.0), Cm(3.5))
tf = norm_box.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "120 РТТ"
r.font.name = 'Inter Extra Bold'; r.font.size = Pt(36); r.font.color.rgb = COL_ORANGE

p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run()
r.text = "на 1 ЭТП"
r.font.name = 'Inter Regular'; r.font.size = Pt(14); r.font.color.rgb = COL_DARK

p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(6)
r = p3.add_run()
r.text = "Каждые 120 активных торговых точек\n= +1 торговый представитель Эмпилс"
r.font.name = 'Inter Regular'; r.font.size = Pt(10); r.font.color.rgb = COL_BLACK

# 3 карточки поддержки
support_y = Cm(15.5)
support_title = slide.shapes.add_textbox(Cm(23.0), support_y, Cm(20), Cm(0.7))
tf = support_title.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "ЧТО ДЕЛАЕТ КОМАНДА:"
r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = COL_ORANGE

supports = [
    ("📋", "Контроль MML", "Слежение за присутствием\nвсех обязательных SKU"),
    ("🎓", "Обучение ТП", "Тренинги по продукту,\nпродажам, TCO"),
    ("📍", "Выкладка и POS", "Контроль полки,\nустановка POS-материалов"),
]

for i, (icon, title, desc) in enumerate(supports):
    x = Cm(23.0 + i * 7.0)
    y = support_y + Cm(1.0)
    
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Cm(6.3), Cm(4.5))
    card.fill.solid(); card.fill.fore_color.rgb = COL_LIGHT_BG
    card.line.color.rgb = COL_GRAY; card.line.width = Pt(0.5)
    card.shadow.inherit = False
    
    tb = slide.shapes.add_textbox(x + Cm(0.3), y + Cm(0.3), Cm(5.7), Cm(4.0))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = icon
    r.font.name = 'Arial'; r.font.size = Pt(20)
    
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)
    r = p2.add_run()
    r.text = title
    r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = COL_DARK
    
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(4)
    r = p3.add_run()
    r.text = desc
    r.font.name = 'Inter Regular'; r.font.size = Pt(9); r.font.color.rgb = COL_BLACK

# Нижняя полоса
footer_y = Cm(26.5)
footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, footer_y, sw, Cm(1.5))
footer.fill.solid(); footer.fill.fore_color.rgb = COL_DARK
footer.line.fill.background(); footer.shadow.inherit = False

ft = slide.shapes.add_textbox(Cm(1.5), footer_y + Cm(0.2), Cm(40), Cm(1.1))
tf = ft.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Начните работать с Эмпилс — и наша команда поможет вам масштабировать продажи. "
r.font.name = 'Inter Regular'; r.font.size = Pt(12); r.font.color.rgb = COL_WHITE
r2 = p.add_run()
r2.text = "Чем больше АКБ, тем больше команда на вашей стороне."
r2.font.name = 'Inter Regular'; r2.font.size = Pt(12); r2.font.bold = True; r2.font.color.rgb = COL_ORANGE

# Оставляем только новый слайд
sldIdLst = prs.slides._sldIdLst
all_elems = list(sldIdLst)
for elem in all_elems[:-1]:
    sldIdLst.remove(elem)

prs.save(OUT)
print(f'OK: {OUT}')
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
print(f'Слайдов: {len(prs.slides)}')
