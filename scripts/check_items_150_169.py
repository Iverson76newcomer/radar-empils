import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
ws = wb["300 SKU — детально"]

# Header
print("Header (row 3):", [c.value for c in ws[3]])

for r in range(4, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx in [37, 38, 56, 61, 62, 63, 150, 169]:
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"\nRow {r} (idx={idx}): {vals}")
