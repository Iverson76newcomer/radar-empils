"""
Add new slide: "Сложности и риски при работе с топовыми дистрибьюторами конкурентов"

Layout: 4 challenge cards in 2×2 grid, each describing a key difficulty:
1. График новинок и сопровождение — distributors want annual roadmap
2. Баланс наценки и уходимости — markup vs shelf velocity paradox
3. Узнаваемость бренда vs компании — brand recognition gap
4. Риск перехода (switching cost) — fear of changing supplier

Plus a bottom strip with key insight/takeaway.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import shutil

SRC = "/home/z/my-project/upload/Презентация_для_ГД_final.pptx"
OUT = "/home/z/my-project/download/Презентация_для_ГД_final.pptx"
shutil.copy(SRC, OUT)

prs = Presentation(OUT)

# Remove empty trailing slide (slide 12)
from pptx.oxml.ns import qn
sld_lst = prs.slides._sldIdLst
slide_ids = list(sld_lst)
# Check if last slide is mostly empty
last_slide = list(prs.slides)[-1]
n_shapes = sum(1 for _ in last_slide.shapes)
if n_shapes <= 2:
    print(f"Removing empty slide 12 ({n_shapes} shapes)")
    sld_lst.remove(slide_ids[-1])

# ── Palette ──────────────────────────────────────────────────────────
BG          = RGBColor(0xF5, 0xF2, 0xEC)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xD9, 0xD2, 0xC2)
DIVIDER     = RGBColor(0xEF, 0xE9, 0xDD)
INK         = RGBColor(0x2F, 0x3A, 0x48)
MUTE        = RGBColor(0x8A, 0x95, 0xA3)
RED         = RGBColor(0xC0, 0x00, 0x00)
SAND        = RGBColor(0xC9, 0xA7, 0x7C)
BLUE        = RGBColor(0x6B, 0x8F, 0xA8)
GREEN       = RGBColor(0x8B, 0xAA, 0x8A)
AMBER       = RGBColor(0xD4, 0xA5, 0x74)
PURPLE      = RGBColor(0x7B, 0x6F, 0xA8)

def cm(v):
    return Emu(int(v * 360000))

def add_text(slide, x, y, w, h, text, *,
             size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
             va=MSO_ANCHOR.TOP, font="Inter", spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = va
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if spacing:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(spacing))
    return tb

def add_rect(slide, x, y, w, h, fill, *,
             line=None, line_w=6350, corner_radius=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner_radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Emu(line_w)
    else:
        s.line.fill.background()
    if corner_radius is not None:
        try: s.adjustments[0] = corner_radius
        except: pass
    s.shadow.inherit = False
    return s


# Create new slide
layout = prs.slide_layouts[0]
slide_new = prs.slides.add_slide(layout)
for ph in list(slide_new.placeholders):
    sp = ph._element
    sp.getparent().remove(sp)

bg = slide_new.background
bg.fill.solid()
bg.fill.fore_color.rgb = BG

# ════════════════════════════════════════════════════════════════════════
# HEADER
# ════════════════════════════════════════════════════════════════════════
add_text(slide_new, cm(0.5), cm(0.5),
         cm(20.0), cm(0.95),
         "Сложности и риски работы с топовыми дистрибьюторами",
         size=20, bold=True, color=INK, font="Inter")
add_text(slide_new, cm(0.5), cm(1.55),
         cm(20.0), cm(0.5),
         "Ключевые наблюдения при переходе партнёров от конкурентов к Эмпилс",
         size=11, color=MUTE, font="Inter")
# Status pill
add_rect(slide_new, cm(20.0), cm(0.45),
         cm(5.0), cm(0.78),
         WHITE, line=RED, line_w=9525, corner_radius=0.5)
add_text(slide_new, cm(20.0), cm(0.58),
         cm(5.0), cm(0.65),
         "РИСКИ · НАБЛЮДЕНИЯ", size=10, bold=True,
         color=RED, font="Inter", align=PP_ALIGN.CENTER, spacing=150)
# Divider
add_rect(slide_new, cm(0.5), cm(2.22),
         cm(24.5), cm(0.08), CARD_BORDER)


# ════════════════════════════════════════════════════════════════════════
# 2×2 GRID of 4 challenge cards
# Each card: 11.9 cm wide × 4.8 cm tall, gap 0.3 cm
# Total: 24.1 × 9.9 cm
# ════════════════════════════════════════════════════════════════════════

CARD_W = cm(11.9)
CARD_H = cm(4.8)
CARD_GAP_X = cm(0.3)
CARD_GAP_Y = cm(0.3)
GRID_X0 = cm(0.5)
GRID_Y0 = cm(2.67)

CHALLENGES = [
    {
        "n": "01",
        "title": "ГРАФИК НОВИНОК И СОПРОВОЖДЕНИЕ",
        "subtitle": "Дистрибьютор хочет видеть план на год",
        "bullets": [
            "Годовой roadmap выхода новинок с датами",
            "Презентация новинки → определение точек продаж",
            "Маркетинговая поддержка на каждом этапе",
            "Дистрибуция с повышенной отсрочкой платежа",
        ],
        "insight": "Без годового плана дистрибьютор не рискует",
        "accent": RED,
    },
    {
        "n": "02",
        "title": "БАЛАНС НАЦЕНКИ И УХОДИМОСТИ",
        "subtitle": "Высокая наценка ≠ высокие продажи",
        "bullets": [
            "Хорошая наценка радует дистрибьютора",
            "Но цена на полке растёт → товар стоит",
            "Найти грань: маржа vs скорость оборота",
            "Точное определение этой грани — сложно",
        ],
        "insight": "Дистрибьютор застрял между маржой и оборотом",
        "accent": SAND,
    },
    {
        "n": "03",
        "title": "УЗНАВАЕМОСТЬ БРЕНДА, А НЕ КОМПАНИИ",
        "subtitle": "Потребитель знает бренд, не производителя",
        "bullets": [
            "Уходимость с полки = узнаваемость бренда ЛКМ",
            "Бренд Эмпилс ≠ бренд TITAN / ОРЕОЛ / NEWTONE",
            "Конкуренты (Hammerite, DALI) — уConsumersknowthebrand",
            "Нужно вкладывать в узнаваемость конкретного продукта",
        ],
        "insight": "Безbrand recognition товар не уходит с полки",
        "accent": BLUE,
    },
    {
        "n": "04",
        "title": "РИСК СМЕНЫ ПОСТАВЩИКА",
        "subtitle": "Страх перехода от проверенного конкурента",
        "bullets": [
            "Привыкли к продукту конкурента (Hammerite и др.)",
            "Боязнь перебоев в поставках нового поставщика",
            "Риск возвратов от недовольных потребителей",
            "Затраты на переобучение ТП под новый ассортимент",
        ],
        "insight": "Смена поставщика = операционный риск",
        "accent": PURPLE,
    },
]

for i, ch in enumerate(CHALLENGES):
    col = i % 2
    row = i // 2
    cx = GRID_X0 + col * (CARD_W + CARD_GAP_X)
    cy = GRID_Y0 + row * (CARD_H + CARD_GAP_Y)

    # Card background
    add_rect(slide_new, cx, cy, CARD_W, CARD_H,
             WHITE, line=CARD_BORDER, line_w=6350, corner_radius=0.03)
    # Accent top bar
    add_rect(slide_new, cx, cy, CARD_W, cm(0.06), ch['accent'])

    # Number + title row
    add_text(slide_new, cx + cm(0.35), cy + cm(0.25),
             cm(1.2), cm(0.6),
             ch['n'], size=20, bold=True, color=ch['accent'],
             font="Roboto Mono", va=MSO_ANCHOR.MIDDLE)
    add_text(slide_new, cx + cm(1.65), cy + cm(0.3),
             cm(9.8), cm(0.4),
             ch['title'], size=10.5, bold=True, color=INK,
             font="Inter", spacing=50)
    add_text(slide_new, cx + cm(1.65), cy + cm(0.75),
             cm(9.8), cm(0.35),
             ch['subtitle'], size=9, color=MUTE,
             font="Inter")

    # Divider
    add_rect(slide_new, cx + cm(0.35), cy + cm(1.25),
             CARD_W - cm(0.7), cm(0.04), DIVIDER)

    # Bullets
    bullet_y = cy + cm(1.45)
    for j, b in enumerate(ch['bullets']):
        add_text(slide_new, cx + cm(0.35), bullet_y,
                 cm(0.2), cm(0.3),
                 "·", size=12, bold=True, color=ch['accent'],
                 font="Inter")
        add_text(slide_new, cx + cm(0.6), bullet_y + cm(0.02),
                 CARD_W - cm(1.0), cm(0.3),
                 b, size=9, color=INK, font="Inter")
        bullet_y += cm(0.42)

    # Insight footer
    KPI_Y = cy + CARD_H - cm(0.7)
    add_rect(slide_new, cx + cm(0.35), KPI_Y - cm(0.08),
             CARD_W - cm(0.7), cm(0.03), DIVIDER)
    add_text(slide_new, cx + cm(0.35), KPI_Y,
             cm(1.5), cm(0.35),
             "ВЫВОД:", size=8, bold=True, color=MUTE,
             font="Inter", va=MSO_ANCHOR.MIDDLE, spacing=100)
    add_text(slide_new, cx + cm(2.0), KPI_Y,
             CARD_W - cm(2.5), cm(0.35),
             ch['insight'], size=9, bold=True, color=ch['accent'],
             font="Inter", va=MSO_ANCHOR.MIDDLE)
    # Icon at right
    add_text(slide_new, cx + CARD_W - cm(0.85), KPI_Y,
             cm(0.5), cm(0.35),
             "⚠", size=12, color=ch['accent'],
             font="Inter", va=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)

# ════════════════════════════════════════════════════════════════════════
# BOTTOM STRIP — key takeaway
# ════════════════════════════════════════════════════════════════════════
STRIP_Y = cm(12.65)
add_rect(slide_new, cm(0.5), STRIP_Y,
         cm(24.5), cm(1.2),
         INK, corner_radius=0.03)
add_text(slide_new, cm(0.8), STRIP_Y + cm(0.15),
         cm(23.9), cm(0.4),
         "СТРАТЕГИЧЕСКИЙ ВЫВОД",
         size=8, bold=True, color=SAND, font="Inter", spacing=200)
add_text(slide_new, cm(0.8), STRIP_Y + cm(0.55),
         cm(23.9), cm(0.6),
         "Успешный переход дистрибьютора требует не только выгодной наценки — нужен годовой план новинок, инвестиции в узнаваемость бренда продукта и снижение операционных рисков перехода.",
         size=10, color=WHITE, font="Inter", va=MSO_ANCHOR.MIDDLE)


# Save
prs.save(OUT)
print(f"\nSaved: {OUT}")
print(f"Total slides: {len(prs.slides)}")
