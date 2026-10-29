#!/usr/bin/env python3
"""
Добавить в ТЗ описание шелфтокера со стопером.
Вставляется после секции "1.б. Шелфтокер угловой" и перед "2. Воблер".
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT = '/home/z/my-project/download/ТЗ_POSM_TITAN_шелфтокер_воблер_ценник.docx'
OUTPUT = '/home/z/my-project/download/ТЗ_POSM_TITAN_шелфтокер_воблер_ценник.docx'

RED = RGBColor(0xC6, 0x28, 0x28)
DARK = RGBColor(0x1A, 0x1F, 0x36)
GRAY = RGBColor(0x60, 0x60, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document(INPUT)

# Найдём индекс параграфа "2. Воблер"
insert_before_idx = None
for i, p in enumerate(doc.paragraphs):
    if '2. Воблер' in p.text:
        insert_before_idx = i
        break

print(f'Вставка перед параграфом {insert_before_idx}: "{doc.paragraphs[insert_before_idx].text[:50]}"')


def insert_paragraph_before(doc, before_idx, text, size=11, bold=False, color=DARK,
                            align=WD_ALIGN_PARAGRAPH.LEFT, italic=False, space_after=6, space_before=0):
    ref_para = doc.paragraphs[before_idx]
    new_para = ref_para.insert_paragraph_before('')
    pf = new_para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.3
    new_para.alignment = align
    if text:
        run = new_para.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = 'Arial'
    return new_para


def insert_table_before(doc, before_idx, headers, rows, col_widths=None):
    ref_para = doc.paragraphs[before_idx]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = WHITE
        run.font.name = 'Arial'
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:fill'), '1E3A5F')
        cell._element.get_or_add_tcPr().append(shading)
    
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.size = Pt(10)
            run.font.name = 'Arial'
            run.font.color.rgb = DARK
            if ci == 0:
                run.font.bold = True
    
    ref_para._element.addprevious(table._tbl)
    return table


# === Заголовок ===
insert_paragraph_before(doc, insert_before_idx, '', space_after=4)

p = doc.paragraphs[insert_before_idx]
new_p = p.insert_paragraph_before('1.в. Шелфтокер со стопером — техническое задание')
new_p.paragraph_format.space_before = Pt(12)
new_p.paragraph_format.space_after = Pt(6)
for run in new_p.runs:
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RED
    run.font.name = 'Arial'

# === Назначение ===
insert_paragraph_before(doc, insert_before_idx, 'Назначение', size=12, bold=True, color=RED, space_before=8, space_after=4)

insert_paragraph_before(doc, insert_before_idx,
    'Шелфтокер со стопером устанавливается на торец полки с выкладкой TITAN. '
    'Основная полоса крепится к полке двусторонним скотчем, а по краям полосы '
    'расположены стоперы — выступающие элементы, которые привлекают внимание '
    'покупателя при взгляде вдоль полки и несут дополнительную информацию. '
    'Стоперы видны с торца полки — когда покупатель проходит мимо, '
    'он видит именно стопер, а не основную полосу.',
    size=11, space_after=6)

# === Текстовая информация ===
insert_paragraph_before(doc, insert_before_idx, 'Текстовая информация на макете', size=12, bold=True, color=RED, space_before=8, space_after=4)

insert_table_before(doc, insert_before_idx,
    ['Зона', 'Текст', 'Шрифт / Стиль'],
    [
        ['Основная полоса — логотип', 'TITAN', 'Крупно, белым на красном фоне. По центру полосы.'],
        ['Основная полоса — слоган', 'Красишь 1 раз — служит 10 лет', 'Жирно, белым. Под логотипом.'],
        ['Стопер левый — заголовок', '10 ЛЕТ', 'Очень крупно, жёлтым на красном. Видно с торца полки.'],
        ['Стопер левый — подзаголовок', 'без перекраски', 'Мелким шрифтом, белым. Под «10 ЛЕТ».'],
        ['Стопер правый — заголовок', '3в1', 'Очень крупно, жёлтым на красном. Видно с торца полки.'],
        ['Стопер правый — подзаголовок', 'грунт + эмаль + ржавчина', 'Мелким шрифтом, белым. Под «3в1».'],
        ['QR-код', 'QR-код → www.titanpaints.ru', 'На основной полосе, справа. Не менее 20×20 мм.'],
    ],
    col_widths=[4, 6, 6.5]
)

# === Визуальные элементы ===
insert_paragraph_before(doc, insert_before_idx, 'Визуальные элементы', size=12, bold=True, color=RED, space_before=8, space_after=4)

elements = [
    '• Основная полоса — красный фон (#C62828), крепится на торец полки',
    '• Стоперы — выступают влево и вправо от краёв полосы (по 30-40 мм)',
    '• Стопер — часть вырубки, отгибается на 90° от основной полосы',
    '• Левый стопер: «10 ЛЕТ» — акцент на долговечность',
    '• Правый стопер: «3в1» — акцент на универсальность',
    '• Логотип TITAN — по центру основной полосы',
    '• QR-код — на основной полосе, на белой подложке',
    '• Жёлтая полоса (#FBBF24) — разделитель на основной полосе',
    '• Форма стоперов — прямоугольная со скруглёнными углами или фигурная',
]
for el in elements:
    insert_paragraph_before(doc, insert_before_idx, el, size=11, space_after=2)

# === Технические параметры ===
insert_paragraph_before(doc, insert_before_idx, 'Технические параметры для типографии', size=12, bold=True, color=RED, space_before=8, space_after=4)

insert_table_before(doc, insert_before_idx,
    ['Параметр', 'Значение'],
    [
        ['Тип', 'Шелфтокер со стоперами (двусторонний выступ)'],
        ['Основная полоса (Ш × В)', '300 × 70 мм (крепится на торец полки)'],
        ['Стопер левый (Ш × В)', '40 × 100 мм (выступает влево, отгибается на 90°)'],
        ['Стопер правый (Ш × В)', '40 × 100 мм (выступает вправо, отгибается на 90°)'],
        ['Общий размер в развёртке', '380 × 100 мм (полоса + 2 стопера)'],
        ['Материал', 'Картон мелованный 350 г/м²'],
        ['Покрытие', 'Ламинация глянцевая (лицевая сторона)'],
        ['Печать', '4+0 (полноцвет, лицевая сторона)'],
        ['Биговка', '2 линии биговки — на стыке полосы со стоперами (для отгиба 90°)'],
        ['Вырубка', 'По контуру: полоса + 2 стопера (единой деталью)'],
        ['Крепление', 'Двусторонний скотч на оборот основной полосы'],
        ['Стоперы', 'Отгибаются вперёд на 90° после установки — видны с торца полки'],
        ['Тираж (ориентир.)', '500 шт (пилот) → 2000+ шт (масштаб)'],
        ['Срок изготовления', '7-10 рабочих дней после утверждения макета'],
    ],
    col_widths=[5, 12]
)

insert_paragraph_before(doc, insert_before_idx, '', space_after=4)

doc.save(OUTPUT)
print(f'OK: {OUTPUT}')
