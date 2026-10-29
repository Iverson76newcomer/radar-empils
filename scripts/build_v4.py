"""
Финальная сборка v4:
1. Заменить 5 жёлтых позиций AvantGarde (idx 1-5) на 5 топовых новинок
2. Обновить обе вкладки: '300 SKU — детально' и '300 SKU с ценами'
3. Пересчитать матрицы на 'Сводка'
4. Создать лист с журналом изменений
"""
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict, Counter
import shutil
import os

SRC_PATH = "/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал.xlsx"
OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v4.xlsx"

# Копируем оригинал
shutil.copy(SRC_PATH, OUT_PATH)
wb = load_workbook(OUT_PATH)

# === Стили ===
HEAD_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEAD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
BODY_FONT = Font(name="Calibri", size=10)
BODY_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
NUM_ALIGN = Alignment(horizontal="right", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
PRICE_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E79")
NEW_ROW_FILL = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")  # зелёный — новые
REMOVED_FILL = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")  # оранжевый — удалённые
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # исходный жёлтый (для удаления)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# === 5 НОВИНОК ДЛЯ ЗАМЕНЫ ===
NEW_ITEMS = [
    {
        'idx': 1,
        'sphere': 'Помещения (интерьер)',
        'purpose': 'Прочее',
        'base': 'Водно-дисперсионная',
        'brand': 'Ореол',
        'type': 'Краски',
        'name': 'Краска для плит OSB Ореол (база А) 1 кг',
        'grp_code': 88965,
        'abc': 'А',
        'отгрузочная': 237.90, 'базовая': 297.38, 'фасовка': 1,
        'лист': 'Ореол',
        'ценовое_имя': 'Краска для плит OSB Ореол (база А)',
        'status': 'НОВИНКА: топ-1 по продажам (вторичка 164 тыс.руб), новая фасовка 1 кг',
        'code_novinki': 88965,
        'vtor_rub': 164257,
    },
    {
        'idx': 2,
        'sphere': 'Помещения (интерьер)',
        'purpose': 'Прочее',
        'base': 'Водно-дисперсионная',
        'brand': 'Ореол',
        'type': 'Краски',
        'name': 'Краска для плит OSB Ореол (база А) 12 кг',
        'grp_code': 88967,
        'abc': 'А',
        'отгрузочная': 2330.20, 'базовая': 2912.75, 'фасовка': 12,
        'лист': 'Ореол',
        'ценовое_имя': 'Краска для плит OSB Ореол (база А)',
        'status': 'НОВИНКА: топ-2 по продажам (вторичка 49 тыс.руб), новая фасовка 12 кг',
        'code_novinki': 88967,
        'vtor_rub': 48656,
    },
    {
        'idx': 3,
        'sphere': 'Дерево',
        'purpose': 'Защита и декор дерева',
        'base': 'Водно-дисперсионная',
        'brand': 'PROFIWOOD',
        'type': 'Декоративные составы',
        'name': 'Антисептик кроющий PROFIWOOD горький шоколад 0.8 л 0.9 кг',
        'grp_code': 88852,
        'abc': 'А',
        'отгрузочная': 463.60, 'базовая': 579.50, 'фасовка': 0.9,
        'лист': 'PROFIWOOD',
        'ценовое_имя': 'Антисептик кроющий PROFIWOOD горький шоколад',
        'status': 'НОВИНКА: топ-3 по продажам (вторичка 19 тыс.руб), новая фасовка 0.9 кг',
        'code_novinki': 88852,
        'vtor_rub': 19080,
    },
    {
        'idx': 4,
        'sphere': 'Фасад',
        'purpose': 'Фасад',
        'base': 'Водно-дисперсионная',
        'brand': 'Ореол',
        'type': 'Краски',
        'name': 'Краска ВД резиновая Ореол (База А) 1.3 кг',
        'grp_code': 89855,
        'abc': 'А',
        'отгрузочная': 457.50, 'базовая': 571.88, 'фасовка': 1.3,
        'лист': 'Ореол',
        'ценовое_имя': 'Краска водно-дисперсионная резиновая Ореол (База А)',
        'status': 'НОВИНКА: топ-4 по продажам (вторичка 12 тыс.руб), новая фасовка 1.3 кг',
        'code_novinki': 89855,
        'vtor_rub': 11729,
    },
    {
        'idx': 5,
        'sphere': 'Дерево',
        'purpose': 'Защита и декор дерева',
        'base': 'Водно-дисперсионная',
        'brand': 'PROFIWOOD',
        'type': 'Декоративные составы',
        'name': 'Антисептик кроющий PROFIWOOD белый База А 2.1 л 2.6 кг',
        'grp_code': 88857,
        'abc': 'А',
        'отгрузочная': 1085.80, 'базовая': 1357.25, 'фасовка': 2.6,
        'лист': 'PROFIWOOD',
        'ценовое_имя': 'Антисептик кроющий PROFIWOOD белый База А',
        'status': 'НОВИНКА: топ-5 по продажам (вторичка 9 тыс.руб), новая фасовка 2.6 кг',
        'code_novinki': 88857,
        'vtor_rub': 8665,
    },
]

# === 1. Обновить лист "300 SKU — детально" ===
print("=== Обновление листа '300 SKU — детально' ===")
ws_det = wb["300 SKU — детально"]

# Найти строки по idx
row_by_idx = {}
for r in range(4, ws_det.max_row + 1):
    idx = ws_det.cell(row=r, column=1).value
    if idx is not None:
        row_by_idx[idx] = r

# Заменить 5 позиций
for item in NEW_ITEMS:
    r = row_by_idx[item['idx']]
    ws_det.cell(row=r, column=1, value=item['idx'])
    ws_det.cell(row=r, column=2, value=item['sphere'])
    ws_det.cell(row=r, column=3, value=item['purpose'])
    ws_det.cell(row=r, column=4, value=item['base'])
    ws_det.cell(row=r, column=5, value=item['brand'])
    ws_det.cell(row=r, column=6, value=item['type'])
    ws_det.cell(row=r, column=7, value=item['name'])
    ws_det.cell(row=r, column=8, value=item['grp_code'])
    ws_det.cell(row=r, column=9, value=item['abc'])
    print(f"  ✓ Заменён #{item['idx']} (строка {r}): {item['name'][:60]}")

# === 2. Обновить лист "300 SKU с ценами" ===
print("\n=== Обновление листа '300 SKU с ценами' ===")
ws_prices = wb["300 SKU с ценами"]

# Найти строки по idx (на этом листе)
row_by_idx_p = {}
for r in range(5, ws_prices.max_row + 1):
    idx = ws_prices.cell(row=r, column=1).value
    if idx is not None:
        row_by_idx_p[idx] = r

for item in NEW_ITEMS:
    r = row_by_idx_p[item['idx']]
    # Обновить данные
    ws_prices.cell(row=r, column=1, value=item['idx'])
    ws_prices.cell(row=r, column=2, value=item['sphere'])
    ws_prices.cell(row=r, column=3, value=item['purpose'])
    ws_prices.cell(row=r, column=4, value=item['base'])
    ws_prices.cell(row=r, column=5, value=item['brand'])
    ws_prices.cell(row=r, column=6, value=item['type'])
    ws_prices.cell(row=r, column=7, value=item['name'])
    ws_prices.cell(row=r, column=8, value=item['grp_code'])
    ws_prices.cell(row=r, column=9, value=item['abc'])
    ws_prices.cell(row=r, column=10, value=item['отгрузочная'])
    ws_prices.cell(row=r, column=11, value=item['базовая'])
    ws_prices.cell(row=r, column=12, value=item['фасовка'])
    ws_prices.cell(row=r, column=13, value=item['лист'])
    ws_prices.cell(row=r, column=14, value=item['status'])
    
    # Применить зелёную подсветку для новых позиций (вместо жёлтой)
    for c in range(1, 15):
        cell = ws_prices.cell(row=r, column=c)
        cell.fill = NEW_ROW_FILL
    
    # Установить формат цен
    ws_prices.cell(row=r, column=10).number_format = '# ##0.00'
    ws_prices.cell(row=r, column=10).font = PRICE_FONT
    ws_prices.cell(row=r, column=11).number_format = '# ##0.00'
    ws_prices.cell(row=r, column=12).number_format = '0.##'
    
    print(f"  ✓ Заменён #{item['idx']} (строка {r}) на листе '300 SKU с ценами': {item['name'][:50]}")

# Обновить заголовок (строка 2) листа "300 SKU с ценами"
ws_prices["A2"] = ("Источник цен: прайс-лист «16.10.2026 РОССИЯ с НДС, руб.» | "
                    "ВНЕСЕНО ИЗМЕНЕНИЕ: 5 позиций AvantGarde (грунт-эмаль, сняты с производства) "
                    "заменены на 5 топовых новинок 2026 года | ИТОГО: 300 SKU")

# === 3. Пересчёт матриц на листе "Сводка" ===
print("\n=== Пересчёт матриц 'Сводка' ===")
ws_sv = wb["Сводка"]

# Прочитать обновлённый детальный лист
items_all = []
for r in range(4, ws_det.max_row + 1):
    if ws_det.cell(row=r, column=1).value is None:
        continue
    items_all.append({
        'idx': ws_det.cell(row=r, column=1).value,
        'sphere': ws_det.cell(row=r, column=2).value,
        'purpose': ws_det.cell(row=r, column=3).value,
        'base': ws_det.cell(row=r, column=4).value,
        'brand': ws_det.cell(row=r, column=5).value,
        'type': ws_det.cell(row=r, column=6).value,
    })

print(f"  Всего SKU: {len(items_all)}")

# === Матрица 1: Основа × Сфера ===
spheres_order = ['Металл', 'Помещения (интерьер)', 'Дерево', 'Пол', 'Фасад']
bases_order = ['Органическая', 'Водно-дисперсионная']

matrix_base_sphere = {(b, s): 0 for b in bases_order for s in spheres_order}
for it in items_all:
    key = (it['base'], it['sphere'])
    if key in matrix_base_sphere:
        matrix_base_sphere[key] += 1

for row_offset, base in enumerate(bases_order):
    r = 5 + row_offset
    ws_sv.cell(row=r, column=1, value=base)
    for col_offset, sphere in enumerate(spheres_order):
        ws_sv.cell(row=r, column=2 + col_offset, value=matrix_base_sphere[(base, sphere)])
    total_row = sum(matrix_base_sphere[(base, s)] for s in spheres_order)
    ws_sv.cell(row=r, column=7, value=total_row)
    pct = total_row / len(items_all)
    ws_sv.cell(row=r, column=8, value=f"{pct*100:.1f}%")

r = 7
ws_sv.cell(row=r, column=1, value='ИТОГО')
for col_offset, sphere in enumerate(spheres_order):
    col_total = sum(matrix_base_sphere[(b, sphere)] for b in bases_order)
    ws_sv.cell(row=r, column=2 + col_offset, value=col_total)
ws_sv.cell(row=r, column=7, value=len(items_all))
ws_sv.cell(row=r, column=8, value='100%')

# === Матрица 2: СФЕРА × БРЕНД ===
brand_counts = Counter(it['brand'] for it in items_all)
original_brands = ['ОРЕОЛ TITAN', 'Ореол', 'Ореол Дисконт', 'NEWTONE', 'AvantGarde',
                    'Расцвет', 'Простокрашено! БАУ', 'PROFIWOOD']
new_brands = [b for b in brand_counts if b not in original_brands]
all_brands = original_brands + new_brands

# Очистить и пересоздать
for r in range(10, 17):
    for c in range(1, 15):
        ws_sv.cell(row=r, column=c, value=None)

ws_sv.cell(row=10, column=1, value='СФЕРА / БРЕНД')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=10, column=2 + col_offset, value=brand)
ws_sv.cell(row=10, column=2 + len(all_brands), value='ИТОГО')

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=10, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER

matrix_sphere_brand = {(s, b): 0 for s in spheres_order for b in all_brands}
for it in items_all:
    if it['brand'] in all_brands:
        matrix_sphere_brand[(it['sphere'], it['brand'])] += 1

for row_offset, sphere in enumerate(spheres_order):
    r = 11 + row_offset
    ws_sv.cell(row=r, column=1, value=sphere)
    for col_offset, brand in enumerate(all_brands):
        val = matrix_sphere_brand[(sphere, brand)]
        if val > 0:
            ws_sv.cell(row=r, column=2 + col_offset, value=val)
    total_row = sum(matrix_sphere_brand[(sphere, b)] for b in all_brands)
    ws_sv.cell(row=r, column=2 + len(all_brands), value=total_row)
    
    for c in range(1, 3 + len(all_brands)):
        cell = ws_sv.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c > 1:
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

r = 16
ws_sv.cell(row=r, column=1, value='ИТОГО')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=r, column=2 + col_offset, value=brand_counts[brand])
ws_sv.cell(row=r, column=2 + len(all_brands), value=len(items_all))

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=r, column=c)
    cell.font = Font(name="Calibri", size=10, bold=True)
    cell.fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    cell.border = BORDER
    if c > 1:
        cell.alignment = CENTER_ALIGN
    else:
        cell.alignment = Alignment(horizontal="left", vertical="center")

# Ширина колонок
for col_offset in range(len(all_brands) + 1):
    ws_sv.column_dimensions[get_column_letter(2 + col_offset)].width = 13
ws_sv.column_dimensions[get_column_letter(2 + len(all_brands))].width = 8
ws_sv.column_dimensions['A'].width = 24

# === Матрица 3: ДЕТАЛЬНАЯ КЛАССИФИКАЦИЯ ПО НАЗНАЧЕНИЮ ===
purposes_order = [
    'Грунт-эмаль по металлу', 'Эмаль по металлу', 'Пол',
    'Стены и потолки (интерьер)', 'Защита и декор дерева',
    'Эмали ВД прочие', 'Лаки прочие', 'Грунтовки', 'Фасад',
    'Прочее', 'Яхтный лак', 'Шпатлёвки', 'Паркетный лак',
    'Молотковая эмаль', 'Потолки', 'Резиновая краска',
]

matrix_purpose_sphere = {(p, s): 0 for p in purposes_order for s in spheres_order}
purpose_counts = Counter()
for it in items_all:
    key = (it['purpose'], it['sphere'])
    if key in matrix_purpose_sphere:
        matrix_purpose_sphere[key] += 1
        purpose_counts[it['purpose']] += 1

purposes_sorted = sorted(purposes_order, key=lambda p: -purpose_counts.get(p, 0))

for r in range(20, 36):
    for c in range(1, 8):
        ws_sv.cell(row=r, column=c, value=None)

ws_sv.cell(row=19, column=1, value='Сфера / Назначение')
sphere_cols = ['Металл', 'Помещения (интерьер)', 'Дерево', 'Пол', 'Фасад']
for col_offset, sp in enumerate(sphere_cols):
    ws_sv.cell(row=19, column=2 + col_offset, value=sp)
ws_sv.cell(row=19, column=7, value='ИТОГО')

for c in range(1, 8):
    cell = ws_sv.cell(row=19, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER

for row_offset, purpose in enumerate(purposes_sorted):
    r = 20 + row_offset
    ws_sv.cell(row=r, column=1, value=purpose)
    for col_offset, sp in enumerate(sphere_cols):
        val = matrix_purpose_sphere[(purpose, sp)]
        if val > 0:
            ws_sv.cell(row=r, column=2 + col_offset, value=val)
    total_row = purpose_counts.get(purpose, 0)
    if total_row > 0:
        ws_sv.cell(row=r, column=7, value=total_row)
    
    for c in range(1, 8):
        cell = ws_sv.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c > 1:
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

print(f"  ✓ Матрицы обновлены")

# === 4. Создать лист "Изменения v4 (замена AvantGarde)" ===
print("\n=== Создание листа 'Изменения v4 (замена AvantGarde)' ===")
sheet_name = "Изменения v4 (AvantGarde→новинки)"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws_log = wb.create_sheet(sheet_name)

# Заголовок
ws_log.merge_cells("A1:H1")
ws_log["A1"] = "Журнал изменений v4: замена 5 позиций AvantGarde на топ-5 новинок 2026 года"
ws_log["A1"].font = TITLE_FONT
ws_log.row_dimensions[1].height = 28

ws_log.merge_cells("A2:H2")
ws_log["A2"] = ("Пользователь выделил жёлтым 5 позиций AvantGarde (грунт-эмаль 0.8-0.9 кг, idx 1-5) — они отсутствуют "
                "в текущем прайс-листе 16.10.2026 (сняты с производства). Заменяем на 5 топовых новинок 2026 года "
                "с новыми фасовками, которых ранее не было в ассортименте, но которые уже показывают хорошие продажи "
                "(по данным файла «Новинки за 2026 год первичка и вторичка.xlsx», период 01.01-06.10.2026).")
ws_log["A2"].font = SUBTITLE_FONT
ws_log["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws_log.row_dimensions[2].height = 60

# Шапка
log_headers = [
    ("#", 5), ("Действие", 14), ("Старая позиция (AvantGarde)", 45),
    ("Новая позиция (новинка)", 50),
    ("Отгруз. цена, руб", 14), ("Фасовка", 9),
    ("Вторичка 2026, руб", 14), ("Комментарий", 55),
]
for col_idx, (title, width) in enumerate(log_headers, 1):
    cell = ws_log.cell(row=4, column=col_idx, value=title)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER
    ws_log.column_dimensions[get_column_letter(col_idx)].width = width
ws_log.row_dimensions[4].height = 38

# Сначала заголовок для удалённых
row_i = 5
ws_log.cell(row=row_i, column=1, value="Удалено (5 шт.):")
ws_log.cell(row=row_i, column=1).font = Font(name="Calibri", size=11, bold=True, color="C00000")
for c in range(1, 9):
    ws_log.cell(row=row_i, column=c).fill = REMOVED_FILL
ws_log.merge_cells(f"A{row_i}:H{row_i}")
row_i += 1

# Старые AvantGarde-позиции
old_items = [
    (1, 'AvantGarde', 'Грунт-эмаль AvantGarde (база А) глянцевая белая', 0.9, 987.94),
    (2, 'AvantGarde', 'Грунт-эмаль AvantGarde глянцевая зеленая RAL 6016', 0.9, 987.94),
    (3, 'AvantGarde', 'Грунт-эмаль AvantGarde глянцевая красно-коричневая RAL 8015', 0.9, 987.94),
    (4, 'AvantGarde', 'Грунт-эмаль AvantGarde глянцевая серая RAL 7042', 0.9, 987.94),
    (5, 'AvantGarde', 'Грунт-эмаль AvantGarde глянцевая черная RAL 9005', 0.8, 987.94),
]
for idx, brand, name, fas, price in old_items:
    vals = [idx, "Удалено", name, "—", price, fas, "—", "Сняты с производства, отсутствуют в прайс-листе 16.10.2026"]
    for col_idx, val in enumerate(vals, 1):
        cell = ws_log.cell(row=row_i, column=col_idx, value=val)
        cell.border = BORDER
        cell.font = BODY_FONT
        cell.fill = REMOVED_FILL
        if col_idx == 1:
            cell.alignment = CENTER_ALIGN
            cell.font = Font(name="Calibri", size=10, bold=True, color="C00000")
        elif col_idx in (5, 6, 7):
            cell.alignment = NUM_ALIGN
            if val is not None and isinstance(val, (int, float)):
                if col_idx == 5:
                    cell.number_format = '# ##0.00'
                else:
                    cell.number_format = '0.##'
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_log.row_dimensions[row_i].height = 36
    row_i += 1

# Заголовок для добавленных
row_i += 1
ws_log.cell(row=row_i, column=1, value="Добавлено (5 новинок 2026 года):")
ws_log.cell(row=row_i, column=1).font = Font(name="Calibri", size=11, bold=True, color="375623")
for c in range(1, 9):
    ws_log.cell(row=row_i, column=c).fill = NEW_ROW_FILL
ws_log.merge_cells(f"A{row_i}:H{row_i}")
row_i += 1

for item in NEW_ITEMS:
    vals = [
        item['idx'], "Новая", "—", item['name'],
        item['отгрузочная'], item['фасовка'],
        item['vtor_rub'], item['status'],
    ]
    for col_idx, val in enumerate(vals, 1):
        cell = ws_log.cell(row=row_i, column=col_idx, value=val)
        cell.border = BORDER
        cell.font = BODY_FONT
        cell.fill = NEW_ROW_FILL
        if col_idx == 1:
            cell.alignment = CENTER_ALIGN
            cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
        elif col_idx in (5, 6, 7):
            cell.alignment = NUM_ALIGN
            if val is not None and isinstance(val, (int, float)):
                if col_idx in (5, 7):
                    cell.number_format = '# ##0.00'
                else:
                    cell.number_format = '0.##'
                if col_idx == 5:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_log.row_dimensions[row_i].height = 50
    row_i += 1

ws_log.freeze_panes = "B5"

# === 5. Сохранить ===
wb.properties.creator = "Z.ai"
wb.properties.title = "Ассортимент Олимп 300 SKU — финал v4 (замена AvantGarde на новинки)"
wb.properties.subject = "5 позиций AvantGarde заменены на топ-5 новинок 2026 года"

print(f"\nЛисты в финальном файле: {wb.sheetnames}")
wb.save(OUT_PATH)
print(f"\n✓ Файл сохранён: {OUT_PATH}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
