#!/usr/bin/env python3
"""
Быстрый поиск данных о партнёрах Эмпилс.
Волна 1: Массовый веб-поиск по ИНН (сохранение в JSON)
"""

import openpyxl
import subprocess
import json
import time
import os
import sys

SEARCH_DIR = '/home/z/my-project/scripts/search_results'
os.makedirs(SEARCH_DIR, exist_ok=True)

src_file = '/home/z/my-project/upload/Задание на дополнение таблицы белые пятна.xlsx'

# Загружаем файл
wb = openpyxl.load_workbook(src_file)
ws = wb['Лист1']

# Собираем компании
companies = []
for row_num in range(3, ws.max_row + 1):
    client = ws.cell(row=row_num, column=4).value
    inn = ws.cell(row=row_num, column=5).value
    if client and inn:
        companies.append({
            'row': row_num,
            'client': str(client).strip(),
            'inn': str(inn).strip(),
        })

print(f"Найдено {len(companies)} компаний")

# Делаем поиск для каждой компании
for i, comp in enumerate(companies):
    out_file = os.path.join(SEARCH_DIR, f"{comp['inn']}.json")
    
    # Если уже искали — пропускаем
    if os.path.exists(out_file) and os.path.getsize(out_file) > 10:
        print(f"[{i+1}/{len(companies)}] SKIP {comp['client']} (уже есть)")
        continue
    
    query = f"ИНН {comp['inn']} {comp['client'][:30]} сайт адрес"
    print(f"[{i+1}/{len(companies)}] {comp['client'][:40]}...")
    
    try:
        result = subprocess.run(
            ['z-ai', 'function', '-n', 'web_search',
             '-a', json.dumps({"query": query, "num": 8}),
             '-o', out_file],
            capture_output=True, text=True, timeout=30
        )
    except subprocess.TimeoutExpired:
        print(f"  TIMEOUT")
        continue
    except Exception as e:
        print(f"  ERROR: {e}")
        continue
    
    # Небольшая пауза
    time.sleep(0.5)

print(f"\nГотово! Результаты в {SEARCH_DIR}")
