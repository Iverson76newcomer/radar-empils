"""Найти 5 позиций AvantGarde, выделенных жёлтым цветом на вкладке '300 SKU с ценами'."""
import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал.xlsx")
ws = wb["300 SKU с ценами"]

print(f"=== Вкладка '300 SKU с ценами': {ws.max_row} строк × {ws.max_column} колонок ===")
print("Шапка (строка 4):")
for c in range(1, ws.max_column + 1):
    print(f"  Col {c}: {ws.cell(row=4, column=c).value}")

# Найти все жёлтые ячейки в колонке с брендом (#5) или в любых колонках
print("\n=== Поиск жёлтых ячеек (fill.start_color.rgb) ===")
yellow_rgb_codes = ['FFFF00', 'FFFF99', 'FFFFE0', 'FFFFFF00', '00FFFF00', 'FFFFFF99']

# Iterate all rows
yellow_rows = set()
for r in range(5, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.start_color and fill.start_color.rgb:
            rgb = str(fill.start_color.rgb)
            # Yellow variations
            if any(y in rgb.upper() for y in ['FFFF00', 'FFFF99', 'FFFFE0', 'FFFFB0', 'FFFFCC']):
                yellow_rows.add(r)
                break

print(f"Найдено {len(yellow_rows)} жёлтых строк: {sorted(yellow_rows)}")

# Print details of yellow rows
print("\n=== Содержимое жёлтых строк ===")
for r in sorted(yellow_rows):
    idx = ws.cell(row=r, column=1).value
    sphere = ws.cell(row=r, column=2).value
    brand = ws.cell(row=r, column=5).value
    name = ws.cell(row=r, column=7).value
    grp_code = ws.cell(row=r, column=8).value
    print(f"\n--- Строка {r}, idx={idx}, бренд={brand}, гр.код={grp_code} ---")
    print(f"  Сфера: {sphere}")
    print(f"  Наименование: {name}")
    # Show fill color of each cell
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.start_color and fill.start_color.rgb:
            rgb = str(fill.start_color.rgb)
            if any(y in rgb.upper() for y in ['FFFF00', 'FFFF99', 'FFFFE0', 'FFFFB0', 'FFFFCC']):
                header = ws.cell(row=4, column=c).value
                print(f"  [ЖЁЛТЫЙ] Col {c} ({header}): {cell.value}")
                print(f"    RGB: {rgb}, patternType: {fill.patternType}")

# Also check tab "300 SKU — детально"
print("\n\n=== Проверим вкладку '300 SKU — детально' на жёлтые ===")
ws2 = wb["300 SKU — детально"]
yellow_rows2 = set()
for r in range(4, ws2.max_row + 1):
    for c in range(1, ws2.max_column + 1):
        cell = ws2.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.start_color and fill.start_color.rgb:
            rgb = str(fill.start_color.rgb)
            if any(y in rgb.upper() for y in ['FFFF00', 'FFFF99', 'FFFFE0', 'FFFFB0', 'FFFFCC']):
                yellow_rows2.add(r)
                break

print(f"Жёлтых строк: {len(yellow_rows2)}")
for r in sorted(yellow_rows2):
    idx = ws2.cell(row=r, column=1).value
    brand = ws2.cell(row=r, column=5).value
    name = ws2.cell(row=r, column=7).value
    print(f"  Строка {r}: idx={idx}, бренд={brand}, имя={name}")
