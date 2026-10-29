#!/usr/bin/env python3
"""Generate DOCX with all calendar mockups."""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = '/home/z/my-project/download/Макеты_календаря_2027_Цикл_защиты.docx'
IMG_DIR = '/home/z/my-project/download/calendar_mockups'

P = {
    'primary': RGBColor(0xC6, 0x28, 0x28),
    'dark': RGBColor(0x1A, 0x1F, 0x36),
    'secondary': RGBColor(0x5A, 0x60, 0x80),
    'green': RGBColor(0x16, 0xA3, 0x4A),
}

def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)

def add_h(doc, text, level=1, color=None):
    if color is None:
        color = P['primary'] if level == 1 else P['dark']
    h = doc.add_heading('', level=level)
    run = h.add_run(text)
    run.font.color.rgb = color
    run.font.name = 'Calibri'

def add_p(doc, text, size=11, bold=False, italic=False, color=None, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = color
    return p

def add_image(doc, path, width_cm=16):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Cm(width_cm))


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.25

# Portrait orientation for this document
section = doc.sections[0]
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

# Cover
cover = doc.add_table(rows=1, cols=1)
cell = cover.rows[0].cells[0]
set_cell_bg(cell, 'C62828')
cell.text = ''
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(10)
run = p.add_run('EMPILS 2027')
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p2 = cell.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(25)
run = p2.add_run('Макеты настенного календаря\n«Цикл защиты»')
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p3 = cell.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_after = Pt(30)
run = p3.add_run('Защити один раз — не перекрашивай 10 лет')
run.font.size = Pt(14)
run.font.italic = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_page_break()

# Concept
add_h(doc, 'Концепция', level=1)
add_p(doc,
    'Календарь настенный, вертикальный. 4 страницы (квартала) + обложка. '
    'На каждой странице — 3 месяца. Формат повторяет текущий календарь Эмпилс, '
    'но с интеграцией ТСО-философии через 3 элемента:', indent=True)
add_p(doc, '1. Символ года (вверху) — меняется по кварталам', bold=True, color=P['primary'])
add_p(doc, '2. Цифра квартала (между календарём и панелью) — одна крупная ТСО-цифра', bold=True, color=P['primary'])
add_p(doc, '3. Слоган (в нижней панели) — вместо цитаты, каждый квартал свой', bold=True, color=P['primary'])

doc.add_page_break()

# Cover page
add_h(doc, 'Обложка календаря', level=1)
add_image(doc, os.path.join(IMG_DIR, '0_cover.png'), width_cm=12)
add_p(doc, 'Тёмный фон + красный акцент. Щит с «10 лет» — символ года. '
           'Главный слоган: «Защити один раз — не перекрашивай 10 лет». '
           '6 брендов списком, каждый своим цветом.', italic=True, color=P['secondary'])

doc.add_page_break()

# Q1
add_h(doc, 'Q1. Металл (Январь — Март)', level=1, color=P['primary'])
add_image(doc, os.path.join(IMG_DIR, '1_q1_metal.png'), width_cm=14)
add_p(doc, 'Символ: щит с «10 лет» (защита металла)', bold=True)
add_p(doc, 'Цифра: 1 = 3 (1 банка TITAN вместо 3 продуктов)', bold=True)
add_p(doc, 'Слоган: «Защити металл один раз — забудь на 10 лет»', bold=True)
add_p(doc, 'Цвет квартала: красный (TITAN, Ореол)', italic=True, color=P['secondary'])

doc.add_page_break()

# Q2
add_h(doc, 'Q2. Дерево (Апрель — Июнь)', level=1, color=P['green'])
add_image(doc, os.path.join(IMG_DIR, '2_q2_wood.png'), width_cm=14)
add_p(doc, 'Символ: годовые кольца дерева с «7» (7 лет защиты)', bold=True, color=P['green'])
add_p(doc, 'Цифра: 7 лет (Profiwood защищает дерево)', bold=True, color=P['green'])
add_p(doc, 'Слоган: «Защити дерево — не меняй доски 7 лет»', bold=True, color=P['green'])
add_p(doc, 'Цвет квартала: зелёный (Profiwood, дерево)', italic=True, color=P['secondary'])

doc.add_page_break()

# Q3
add_h(doc, 'Q3. Интерьер (Июль — Сентябрь)', level=1, color=RGBColor(0x25, 0x63, 0xEB))
add_image(doc, os.path.join(IMG_DIR, '3_q3_interior.png'), width_cm=14)
add_p(doc, 'Символ: «1 слой» с зачёркнутыми 3 слоями', bold=True, color=RGBColor(0x25, 0x63, 0xEB))
add_p(doc, 'Цифра: 1 слой (Newtone — вместо 3 слоёв)', bold=True, color=RGBColor(0x25, 0x63, 0xEB))
add_p(doc, 'Слоган: «Покрась стены за выходные — без запаха и разводов»', bold=True, color=RGBColor(0x25, 0x63, 0xEB))
add_p(doc, 'Цвет квартала: синий (Newtone, водные краски)', italic=True, color=P['secondary'])

doc.add_page_break()

# Q4
add_h(doc, 'Q4. Премиум и план (Октябрь — Декабрь)', level=1, color=RGBColor(0x7C, 0x3A, 0xED))
add_image(doc, os.path.join(IMG_DIR, '4_q4_premium.png'), width_cm=14)
add_p(doc, 'Символ: «−61%» (экономия TITAN за 10 лет)', bold=True, color=RGBColor(0x7C, 0x3A, 0xED))
add_p(doc, 'Цифра: −61% (экономия vs ПФ-115+грунт)', bold=True, color=RGBColor(0x7C, 0x3A, 0xED))
add_p(doc, 'Слоган: «Заплатил больше сейчас — сэкономил на годы»', bold=True, color=RGBColor(0x7C, 0x3A, 0xED))
add_p(doc, 'Цвет квартала: фиолетовый (AvantGarde, премиум)', italic=True, color=P['secondary'])

doc.add_page_break()

# Summary table
add_h(doc, 'Сводка элементов по кварталам', level=1)

from docx.enum.table import WD_TABLE_ALIGNMENT

def make_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if col_widths:
        for i, w in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Cm(w)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, '1A1F36')
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(9)
            if r_idx % 2 == 0:
                set_cell_bg(cell, 'FAFAFA')

summary = [
    ['Обложка', 'Щит + «10 лет»', '—', 'Защити один раз — не перекрашивай 10 лет', 'Тёмный + красный'],
    ['Q1 Металл', 'Щит (защита)', '1 = 3', 'Защити металл один раз — забудь на 10 лет', 'Красный'],
    ['Q2 Дерево', 'Кольца дерева (7 лет)', '7 лет', 'Защити дерево — не меняй доски 7 лет', 'Зелёный'],
    ['Q3 Интерьер', '1 слой (вместо 3)', '1 слой', 'Покрась стены за выходные — без запаха и разводов', 'Синий'],
    ['Q4 Премиум', '−61% (экономия)', '−61%', 'Заплатил больше сейчас — сэкономил на годы', 'Фиолетовый'],
]
make_table(doc, ['Страница', 'Символ', 'Цифра', 'Слоган', 'Цвет'],
           summary, col_widths=[2.5, 3.5, 2, 6, 2.5])

add_p(doc, '')
add_p(doc, 'Принцип: минимум текста, максимум смысла. Календарь висит на стене 3 месяца — '
           'слоган и цифра запоминаются подсознательно, как рекламный билборд.',
      italic=True, color=P['secondary'])

# Final note
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(15)
run = p.add_run('Empils • TITAN • Календарь 2027 «Цикл защиты» • Макеты для дизайнера')
run.font.size = Pt(10)
run.font.italic = True
run.font.color.rgb = P['secondary']

doc.save(OUT_PATH)
print(f'Document saved: {OUT_PATH}')
print(f'Size: {os.path.getsize(OUT_PATH)} bytes')
