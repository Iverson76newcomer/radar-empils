#!/usr/bin/env python3
"""
Фильтрация и добавление туркменских компаний в Excel.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy

SRC_EXCEL = '/home/z/my-project/download/Экспорт_потенциальные_дистрибьюторы_с_Беларусью.xlsx'
OUT = '/home/z/my-project/download/Экспорт_потенциальные_дистрибьюторы_с_Туркменистаном.xlsx'


CLEAN_COMPANIES = [
    {
        'name': 'Ynamdar Store',
        'category': 'Розничные сети и магазины',
        'about': 'Сеть магазинов ЛКМ в Ашхабаде. Эмульсии, декоративные краски, краски для ткани.',
        'features': 'розница',
        'assortment': 'эмульсии, декоративные краски, краски для ткани',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'store.ynamdar.com',
    },
    {
        'name': 'Aýterek (маркетплейс)',
        'category': 'Розничные сети и магазины',
        'about': 'Маркетплейс ЛКМ в Туркменистане. Kent, Bio plus, Alvan, Polisan Perla.',
        'features': 'розница',
        'assortment': 'Kent, Bio plus, Alvan, Polisan Perla',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'ayterek.com',
    },
    {
        'name': 'BETEK TÜRKMENISTAN (EDA MARKET)',
        'category': 'Официальный дистрибьютор (BETEK)',
        'about': 'Дистрибьютор BETEK в Туркменистане. Магазин в Ашхабаде, ул. А. Ниязова 157.',
        'features': 'официальный дилер, розница',
        'assortment': 'BETEK, ЛКМ',
        'address': 'г. Ашхабад, ул. А. Ниязова 157',
        'phones': '',
        'emails': '',
        'site': 'tm.gorodwiki.ru',
    },
    {
        'name': 'PASYL BOÝAG',
        'category': 'Розничные сети и магазины',
        'about': 'Магазин красок в Ашхабаде.',
        'features': 'розница',
        'assortment': 'ЛКМ',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'tm.gorodwiki.ru',
    },
    {
        'name': 'AWTORENKLER TM / AUTOLUX',
        'category': 'Дистрибьютор',
        'about': 'Дистрибьютор ЛКМ в Туркменистане (автомобильные краски).',
        'features': 'дистрибьютор, опт',
        'assortment': 'автомобильные краски',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'tm.gorodwiki.ru',
    },
    {
        'name': 'Trade.com.tm (PAINT-AND-LACQUER COATINGS)',
        'category': 'Дистрибьютор',
        'about': 'Оптовая и розничная продажа мебельных красок. Ашхабат, Ахал.',
        'features': 'опт, розница',
        'assortment': 'мебельные краски, растворители',
        'address': 'г. Ашхабад, Ахал',
        'phones': '+993 12 760-322, +993 24 322-050',
        'emails': '',
        'site': 'trade.com.tm',
    },
    {
        'name': 'ALSAT (Arzan wholesale)',
        'category': 'Дистрибьютор',
        'about': 'Оптовая торговля для розничных продавцов и предпринимателей. Карtridges and paint wholesale, painting tools. Доставка из Ашхабада.',
        'features': 'опт',
        'assortment': 'ЛКМ, painting tools, cartridges',
        'address': 'г. Ашхабад, ул. Гарашсызлык 71',
        'phones': '',
        'emails': 'info@alsat',
        'site': 'alsat.com.tm',
    },
    {
        'name': 'EDA Economic Society',
        'category': 'Импортер / Дистрибьютор (BETEK)',
        'about': 'Импорт-экспорт лакокрасочной продукции BETEK и строительных материалов.',
        'features': 'импортёр, опт',
        'assortment': 'BETEK, стройматериалы',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'eda.company',
    },
    {
        'name': 'Bedew TM (интернет-магазин)',
        'category': 'Розничные сети и магазины',
        'about': 'Интернет-магазин ЛКМ в Туркменистане. Алкидные масляные глянцевые краски Alvan, Topas.',
        'features': 'розница',
        'assortment': 'Alvan, Topas, алкидные масляные краски',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'bedew.com.tm',
    },
    {
        'name': 'Новур (TMTrade)',
        'category': 'Производитель',
        'about': 'Производитель строительных материалов и красок. Мары/Туркменистан.',
        'features': 'производитель, опт',
        'assortment': 'строительные материалы, краски',
        'address': 'г. Мары',
        'phones': '+993 61 169-199, +993 56 469-199, +993 243 630-50',
        'emails': '',
        'site': 'trade.com.tm',
    },
    {
        'name': 'TMCARS (Мары)',
        'category': 'Розничные сети и магазины',
        'about': 'Стройматериалы и грунт-краска в Марыйском регионе. Гипсовые украшения и декоративные колонны.',
        'features': 'розница',
        'assortment': 'стройматериалы, грунт-краска',
        'address': 'г. Мары',
        'phones': '',
        'emails': '',
        'site': 'tmcars.info',
    },
    {
        'name': 'ИП Кушнаренко С.А. (экспортёр)',
        'category': 'Импортер / Дистрибьютор',
        'about': 'Экспорт грунтовок и грунт-эмалей в Туркменистан. ТМ: Poller, Престиж, Казачка, Krasava, Formula Q8.',
        'features': 'импортёр, опт',
        'assortment': 'грунтовки, грунт-эмали (Poller, Престиж, Казачка, Krasava, Formula Q8)',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'donexport.donland.ru',
    },
    {
        'name': 'ALVAN (через Flagma-tm)',
        'category': 'Производитель / Поставщик',
        'about': 'Производство и оптовая реализация красок ALVAN. Легко наносятся, для потолков и стен.',
        'features': 'производитель, опт',
        'assortment': 'ALVAN, водоэмульсионные краски',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'flagma-tm.com',
    },
    {
        'name': 'Made In Turkmenistan (Emulsion Paint)',
        'category': 'Производитель / Экспортёр',
        'about': 'Производство эмульсионных красок в Туркменистане. Экспортное оптовое направление. Соответствие международным стандартам.',
        'features': 'производитель, экспорт, опт',
        'assortment': 'эмульсионные краски',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'madeinturkmenistan.com',
    },
    {
        'name': 'Turkmenexporters (пластиковая тара)',
        'category': 'Производитель',
        'about': 'Производитель пластиковой тары для эмульсионных красок (3 000 шт. минимальная партия).',
        'features': 'производитель, опт',
        'assortment': 'пластиковая тара для ЛКМ',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'turkmenexporters.com.tm',
    },
    {
        'name': 'Стройматериалы Henkel (tm.all.biz)',
        'category': 'Импортер / Дистрибьютор',
        'about': 'Поставки строительных материалов Henkel в Ашхабад. Ceresit CТ 42 акриловая краска.',
        'features': 'импортёр, опт',
        'assortment': 'Henkel, Ceresit, акриловые краски',
        'address': 'г. Ашхабад',
        'phones': '',
        'emails': '',
        'site': 'tm.all.biz',
    },
]

print(f"Подготовлено компаний для добавления: {len(CLEAN_COMPANIES)}")

wb = openpyxl.load_workbook(SRC_EXCEL)
ws = wb['Лист1']

last_row = 4
for r in range(4, ws.max_row + 1):
    if ws.cell(r, 2).value:
        last_row = r

print(f"Последняя заполненная строка: {last_row}")

sample_row = None
for r in range(4, ws.max_row + 1):
    if ws.cell(r, 2).value == 'Туркменистан':
        sample_row = r
        break

print(f"Строка-образец для форматирования: {sample_row}")

sample_fill = None
if sample_row:
    sample_cell = ws.cell(sample_row, 2)
    if sample_cell.fill and sample_cell.fill.fgColor:
        sample_fill = copy(sample_cell.fill)
if not sample_fill:
    sample_fill = PatternFill('solid', fgColor='FEE2E2')

thin = Side(style='thin', color='D1D5DB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

current_row = last_row + 1
added = 0

for c in CLEAN_COMPANIES:
    cells_data = [
        '',
        'Туркменистан',
        c['category'],
        c['name'],
        c['about'],
        c['features'],
        c['assortment'],
        c['address'],
        c['phones'],
        c['emails'],
        c['site'],
    ]
    
    for ci, val in enumerate(cells_data, 1):
        cell = ws.cell(current_row, ci, val)
        if sample_row:
            sample_cell = ws.cell(sample_row, ci)
            cell.font = copy(sample_cell.font)
        else:
            cell.font = Font(name='Arial', size=10, bold=(ci in [2, 4]))
        cell.fill = copy(sample_fill)
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        cell.border = border
    
    current_row += 1
    added += 1

print(f"Добавлено строк: {added}")

col_widths = {2: 14, 3: 28, 4: 36, 5: 50, 6: 35, 7: 32, 8: 32, 9: 28, 10: 28, 11: 22}
for ci, w in col_widths.items():
    cur = ws.column_dimensions[get_column_letter(ci)].width
    if not cur or cur < w:
        ws.column_dimensions[get_column_letter(ci)].width = w

wb.save(OUT)
import os
print(f"\nOK: {OUT}")
print(f"Размер: {os.path.getsize(OUT) / 1024:.1f} KB")

wb2 = openpyxl.load_workbook(OUT, data_only=True)
ws2 = wb2['Лист1']
from collections import Counter
countries = []
for ri in range(4, ws2.max_row + 1):
    c = ws2.cell(ri, 2).value
    if c:
        countries.append(str(c).strip())
print(f"\n=== Итоговая статистика ===")
for c, cnt in Counter(countries).most_common():
    marker = ' ← +' if c == 'Туркменистан' else ''
    print(f"  {c}: {cnt} компаний{marker}")
print(f"Всего: {sum(Counter(countries).values())} компаний")
