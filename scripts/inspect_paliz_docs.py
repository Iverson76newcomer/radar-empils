"""Extract text from both docx files: exhibition report + Palizh partner file."""
from docx import Document

print("=" * 100)
print("FILE 1: 1_Сводный_отчёт_по_итогам_выставки.docx")
print("=" * 100)
doc1 = Document('/home/z/my-project/upload/1_Сводный_отчёт_по_итогам_выставки.docx')
for i, p in enumerate(doc1.paragraphs):
    if p.text.strip():
        print(p.text)
# Tables
for ti, t in enumerate(doc1.tables):
    print(f"\n[TABLE {ti+1}]")
    for r in t.rows:
        cells = [c.text.strip().replace('\n', ' | ') for c in r.cells]
        print(" | ".join(cells))

print("\n\n")
print("=" * 100)
print("FILE 2: Palizh.docx")
print("=" * 100)
doc2 = Document('/home/z/my-project/upload/Palizh.docx')
for i, p in enumerate(doc2.paragraphs):
    if p.text.strip():
        print(p.text)
for ti, t in enumerate(doc2.tables):
    print(f"\n[TABLE {ti+1}]")
    for r in t.rows:
        cells = [c.text.strip().replace('\n', ' | ') for c in r.cells]
        print(" | ".join(cells))
