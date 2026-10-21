import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v3.xlsx", data_only=True)

print("=" * 80)
print("ФИНАЛЬНАЯ ПРОВЕРКА v3 (с сохранением idx)")
print("=" * 80)

# === Лист Сводка ===
ws = wb["Сводка"]
print("\n--- Лист 'Сводка' ---")
print("\nМАТРИЦА Основа × Сфера:")
for r in range(4, 8):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 9)]
    print(f"  Row {r}: {vals}")

print("\nМАТРИЦА СФЕРА × БРЕНД:")
for r in range(10, 17):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 14)]
    print(f"  Row {r}: {vals}")

print("\nДЕТАЛЬНАЯ КЛАССИФИКАЦИЯ ПО НАЗНАЧЕНИЮ:")
for r in range(19, 36):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

# === Лист 300 SKU — детально ===
ws = wb["300 SKU — детально"]
print(f"\n--- Лист '300 SKU — детально' ---")
total = 0
for r in range(4, ws.max_row + 1):
    if ws.cell(row=r, column=1).value is not None:
        total += 1
print(f"Всего SKU: {total}")

# Проверим 3 заменённых SKU по их исходным idx
print("\n3 ЗАМЕНЁННЫХ SKU (по idx 37, 150, 169):")
for r in range(4, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx in [37, 150, 169]:
        vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
        print(f"  #{idx} (строка {r}):")
        print(f"    Сфера: {vals[1]} | Назначение: {vals[2]} | Основа: {vals[3]}")
        print(f"    Бренд: {vals[4]} | Тип: {vals[5]}")
        print(f"    Имя: {vals[6]}")
        print(f"    Гр.код: {vals[7]} | ABC: {vals[8]}")

# 5 новых SKU
print("\n5 НОВЫХ SKU (idx 296-300):")
for r in range(4, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx in [296, 297, 298, 299, 300]:
        vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
        print(f"  #{idx} (строка {r}): {vals[4]} | {vals[6][:70]}")

# Проверка: удалённые SKU не должны присутствовать
print("\nПРОВЕРКА: удалённые SKU не должны присутствовать:")
deleted_idx = [38, 56, 61, 62, 63]
found_deleted = []
for r in range(4, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx in deleted_idx:
        found_deleted.append((r, idx, ws.cell(row=r, column=7).value))
if found_deleted:
    for r, idx, name in found_deleted:
        print(f"  ⚠ #{idx} (строка {r}): {name}")
else:
    print("  ✓ Все 5 удалённых SKU отсутствуют (по idx)")

# === Лист 300 SKU с ценами ===
ws = wb["300 SKU с ценами"]
print(f"\n--- Лист '300 SKU с ценами' ---")

total = 0
has_price = 0
for r in range(5, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx is None:
        continue
    total += 1
    if ws.cell(row=r, column=10).value is not None:
        has_price += 1
print(f"Всего SKU: {total}")
print(f"С отгрузочной ценой: {has_price}")
print(f"Без цены: {total - has_price}")

# === Подсчёт по матрице: сравнить с матрицей ===
print("\n=== СВЕРКА МАТРИЦ ===")
# Считаем по обновлённому детальному листу
from collections import Counter, defaultdict
ws = wb["300 SKU — детально"]
items_all = []
for r in range(4, ws.max_row + 1):
    if ws.cell(row=r, column=1).value is None:
        continue
    items_all.append({
        'sphere': ws.cell(row=r, column=2).value,
        'purpose': ws.cell(row=r, column=3).value,
        'base': ws.cell(row=r, column=4).value,
        'brand': ws.cell(row=r, column=5).value,
    })

sphere_counts = Counter(it['sphere'] for it in items_all)
print(f"\nРаспределение по сферам (фактически): {dict(sphere_counts)}")
print(f"  Металл: {sphere_counts.get('Металл', 0)} (ожидалось 128)")
print(f"  Помещения: {sphere_counts.get('Помещения (интерьер)', 0)} (ожидалось 73)")
print(f"  Дерево: {sphere_counts.get('Дерево', 0)} (ожидалось 53)")
print(f"  Пол: {sphere_counts.get('Пол', 0)} (ожидалось 34)")
print(f"  Фасад: {sphere_counts.get('Фасад', 0)} (ожидалось 12)")
print(f"  ИТОГО: {len(items_all)} (ожидалось 300)")

brand_counts = Counter(it['brand'] for it in items_all)
print(f"\nРаспределение по брендам (фактически):")
for b, c in sorted(brand_counts.items(), key=lambda x: -x[1]):
    print(f"  {b}: {c}")

# Проверка статуса в журнале
print("\n--- Журнал изменений ---")
ws = wb["Изменения v3 (10 SKU)"]
for r in range(4, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
    if any(v is not None for v in vals):
        action = vals[0] or ''
        idx = vals[1] if vals[1] else ''
        brand = vals[2] if vals[2] else ''
        name = (vals[4] or '')[:50] if vals[4] else ''
        price = vals[5] if vals[5] else ''
        print(f"  Row {r}: {action:<10} #{idx} | {brand} | {name} | {price}")
