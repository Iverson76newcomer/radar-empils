"""
Add "Система бонусирования" table after the "Товарооборот и заработок партнёра" table
in the user's edited version of the KP document.

Bonus structure:
- Год 1: ТО 60 млн → 4% ассортимент + 4% ежемесячный ТО + 1% годовой план = 9% → 5.4 млн ₽
- Год 2: ТО 120 млн → 4% ассортимент + 4% ежемесячный ТО + 2% годовой план = 10% → 12 млн ₽
- Год 3: ТО 312 млн → 4% ассортимент + 4% ежемесячный ТО + 3% годовой план = 11% → 34.3 млн ₽
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import copy
import os

SRC = '/home/z/my-project/upload/КП_Олимп_коммерческое_предложение.docx'
OUT = '/home/z/my-project/download/КП_Олимп_коммерческое_предложение.docx'

INK = RGBColor(0x2F, 0x3A, 0x48)
MUTE = RGBColor(0x8A, 0x95, 0xA3)
RED = RGBColor(0xC0, 0x00, 0x00)
SAND = RGBColor(0xC9, 0xA7, 0x7C)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GREEN = RGBColor(0x1A, 0x7D, 0x44)

def shade_cell(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {qn('w:fill'): color_hex, qn('w:val'): 'clear'})
    shading.append(shd)

def set_cell(cell, text, size=8, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''; p = cell.paragraphs[0]; p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color; run.font.name = 'Inter'

# Load document
doc = Document(SRC)

# Find the position to insert: after Table 2 (Товарооборот и заработок)
# We need to insert a heading + table AFTER the paragraph that says "Цель: 312 млн..."
# and BEFORE the divider "———"
# 
# Strategy: find the paragraph with "Цель:" text, then insert after the next paragraph
# Actually, we need to insert AFTER Table 2 and before the divider line.
# Tables and paragraphs are separate elements in the body.
# We'll insert the new content right after Table 2's element.

# Find Table 2 (the forecast table with "Год 1", "Год 2", "Год 3")
target_table = None
target_table_idx = None
for i, table in enumerate(doc.tables):
    if len(table.rows) >= 2 and table.rows[1].cells[0].text.strip() == 'Год 1':
        target_table = table
        target_table_idx = i
        break

print(f"Found forecast table at index {target_table_idx}")

if target_table is None:
    print("ERROR: Could not find forecast table")
    exit(1)

# Get the table's XML element
table_element = target_table._element

# Create the bonus heading paragraph
bonus_heading = doc.add_paragraph()
bonus_heading.space_before = Pt(10)
bonus_heading.space_after = Pt(6)
run = bonus_heading.add_run('СИСТЕМА БОНУСИРОВАНИЯ')
run.font.size = Pt(11)
run.font.bold = True
run.font.color.rgb = DARK_GREEN
run.font.name = 'Inter'

# Create bonus table
bonus_data = [
    ['Год', 'Товарооборот', 'Бонус за ассортимент (4%)', 'Бонус за ежемес. ТО (4%)', 'Бонус за год. план', 'Итого бонус', 'Итого доход (наценка + бонус)'],
    ['Год 1', '60 млн ₽', '2,4 млн ₽', '2,4 млн ₽', '0,6 млн ₽ (1%)', '5,4 млн ₽ (9%)', '26,4 млн ₽'],
    ['Год 2', '120 млн ₽', '4,8 млн ₽', '4,8 млн ₽', '2,4 млн ₽ (2%)', '12,0 млн ₽ (10%)', '54,0 млн ₽'],
    ['Год 3', '312 млн ₽', '12,5 млн ₽', '12,5 млн ₽', '9,4 млн ₽ (3%)', '34,3 млн ₽ (11%)', '143,3 млн ₽'],
    ['ИТОГО', '492 млн ₽', '19,7 млн ₽', '19,7 млн ₽', '12,4 млн ₽', '51,7 млн ₽', '223,7 млн ₽'],
]

bonus_table = doc.add_table(rows=len(bonus_data), cols=7)
bonus_table.style = 'Table Grid'
bonus_table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, row_data in enumerate(bonus_data):
    row = bonus_table.rows[i]
    is_header = i == 0
    is_total = i == len(bonus_data) - 1
    for j, val in enumerate(row_data):
        set_cell(row.cells[j], val, size=7.5, bold=(is_header or is_total),
                 color=(WHITE if is_header else (DARK_GREEN if is_total else INK)),
                 align=(WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT))
        if is_header:
            shade_cell(row.cells[j], '1A7D44')
        elif is_total:
            shade_cell(row.cells[j], 'E8F5E9')
        elif i % 2 == 1:
            shade_cell(row.cells[j], 'F5F2EC')

# Create explanatory paragraph after bonus table
bonus_note = doc.add_paragraph()
bonus_note.space_before = Pt(4)
bonus_note.space_after = Pt(4)
run = bonus_note.add_run('Условия: выполнение плана по ассортименту (наличие обязательных SKU) + выполнение ежемесячного плана товарооборота + выполнение годового плана. Бонус выплачивается ретро-бонусом (кредит-нотой) по итогам периода.')
run.font.size = Pt(8)
run.font.italic = True
run.font.color.rgb = MUTE
run.font.name = 'Inter'

# Now move the bonus heading, table, and note to right after the forecast table
# Get the body element
body = doc.element.body

# Find the forecast table element in the body
# The bonus heading, table, and note were added at the END of the document
# We need to move them to right after the forecast table

# Get references to the elements we need to move
bonus_heading_elem = bonus_heading._element
bonus_table_elem = bonus_table._element
bonus_note_elem = bonus_note._element

# Remove them from end
body.remove(bonus_heading_elem)
body.remove(bonus_table_elem)
body.remove(bonus_note_elem)

# Insert after the forecast table
# table_element is the forecast table's XML element
# We insert: bonus_heading, then bonus_table, then bonus_note, right after table_element
table_element.addnext(bonus_note_elem)
table_element.addnext(bonus_table_elem)
table_element.addnext(bonus_heading_elem)

doc.save(OUT)
print(f"Saved: {OUT} ({os.path.getsize(OUT)} bytes)")

# Verify
doc2 = Document(OUT)
print(f"\nTables: {len(doc2.tables)}")
# Check the new bonus table
for i, table in enumerate(doc2.tables):
    if len(table.rows) >= 2 and 'бонус' in table.rows[0].cells[0].text.lower() or (len(table.rows) >= 2 and 'Бонус' in table.rows[0].cells[2].text):
        print(f"\nBonus table found at index {i}:")
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip()[:30] for cell in row.cells]
            print(f"  Row {ri}: {cells}")
        break

import os
print(f"\nFile size: {os.path.getsize(OUT)} bytes")
