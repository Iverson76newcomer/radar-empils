import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v3.xlsx", data_only=True)

print("=" * 80)
print("ФИНАЛЬНАЯ ПРОВЕРКА ФАЙЛА v3")
print("=" * 80)

# === Лист Сводка ===
ws = wb["Сводка"]
print("\n--- Лист 'Сводка' ---")
print("\nМАТРИЦА Основа × Сфера:")
for r in range(4, 8):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 9)]
    print(f"  Row {r}: {vals}")

print("\nМАТРИЦА СФЕРА × БРЕНД (с новыми брендами):")
for r in range(10, 17):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 14)]
    print(f"  Row {r}: {vals}")

print("\nДЕТАЛЬНАЯ КЛАССИФИКАЦИЯ ПО НАЗНАЧЕНИЮ:")
for r in range(19, 36):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

# === Лист 300 SKU — детально ===
ws = wb["300 SKU — детально"]
print(f"\n--- Лист '300 SKU — детально' ---")
print(f"Всего строк данных: {ws.max_row - 3}")
print(f"Первая SKU: {[c.value for c in ws[4]]}")
print(f"Последняя SKU: {[c.value for c in ws[ws.max_row]]}")

# === Лист 300 SKU с ценами ===
ws = wb["300 SKU с ценами"]
print(f"\n--- Лист '300 SKU с ценами' ---")
print(f"Всего SKU с ценами: {ws.max_row - 4}")

# Проверим 3 заменённых SKU
print("\n3 ЗАМЕНЁННЫХ SKU:")
for r in range(5, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    # После перенумерации #37 → #37 (если порядок сохранился), но лучше искать по содержимому
    name = ws.cell(row=r, column=7).value or ''
    if 'красная RAL 3020' in name and '3 в 1' in name:
        print(f"  Строка {r} (#{idx}): {name[:70]}...")
        print(f"    Отгруз: {ws.cell(row=r, column=10).value} | Базовая: {ws.cell(row=r, column=11).value} | Фасовка: {ws.cell(row=r, column=12).value}")
    if 'влагоизолирующий' in name.lower() and 'Грунт-концентрат' in name:
        print(f"  Строка {r} (#{idx}): {name[:70]}...")
        print(f"    Отгруз: {ws.cell(row=r, column=10).value} | Базовая: {ws.cell(row=r, column=11).value} | Фасовка: {ws.cell(row=r, column=12).value}")
    if 'чёрная RAL 9005' in name.lower() or 'черная RAL 9005' in name.lower():
        if 'NEWTONE' in name and '13.9' in name:
            print(f"  Строка {r} (#{idx}): {name[:70]}...")
            print(f"    Отгруз: {ws.cell(row=r, column=10).value} | Базовая: {ws.cell(row=r, column=11).value} | Фасовка: {ws.cell(row=r, column=12).value}")

# 5 новых SKU
print("\n5 НОВЫХ SKU (296-300):")
for r in range(5, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx in [296, 297, 298, 299, 300]:
        print(f"  #{idx} | {ws.cell(row=r, column=5).value} | {ws.cell(row=r, column=7).value}")
        print(f"    Отгруз: {ws.cell(row=r, column=10).value} | Базовая: {ws.cell(row=r, column=11).value} | Фасовка: {ws.cell(row=r, column=12).value}")

# Проверка: есть ли в файле удалённые SKU?
print("\nПРОВЕРКА: удалённые SKU не должны присутствовать:")
deleted_keywords = ['голубая 0.9 кг', 'тёмно-зелёная 0.9 кг', 'шоколадная 0.9 кг', 'шоколадная 1.9 кг', 'шоколадная 2.7 кг']
for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=7).value or ''
    if any(kw in name.lower() for kw in deleted_keywords):
        if 'РАСЦВЕТ' in name and '3 в 1' in name:
            print(f"  ⚠ НАЙДЕНО (старое имя): #{ws.cell(row=r, column=1).value} | {name}")
if not any('РАСЦВЕТ' in (ws.cell(row=r, column=7).value or '') and '3 в 1' in (ws.cell(row=r, column=7).value or '') and any(kw in (ws.cell(row=r, column=7).value or '').lower() for kw in deleted_keywords) for r in range(5, ws.max_row + 1)):
    print("  ✓ Все 5 удалённых SKU отсутствуют")

# === Журнал изменений ===
ws = wb["Изменения v3 (10 SKU)"]
print(f"\n--- Лист 'Изменения v3 (10 SKU)' ---")
print(f"Всего строк: {ws.max_row}")
print("\nСодержимое:")
for r in range(4, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
    if any(v is not None for v in vals):
        action = vals[0] or ''
        idx_val = vals[1] if vals[1] else ''
        brand = vals[2] if vals[2] else ''
        name = (vals[4] or '')[:50] if vals[4] else ''
        price = vals[5] if vals[5] else ''
        print(f"  Row {r}: {action:<10} #{idx_val} | {brand} | {name} | {price}")

# === Подсчёт итогов ===
print("\n" + "=" * 80)
print("ИТОГОВАЯ СТАТИСТИКА")
print("=" * 80)
# Подсчёт по листу с ценами
ws = wb["300 SKU с ценами"]
total = 0
has_price = 0
for r in range(5, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx is None:
        continue
    total += 1
    if ws.cell(row=r, column=10).value is not None:
        has_price += 1
print(f"Всего SKU: {total}")
print(f"С отгрузочной ценой: {has_price}")
print(f"Без цены: {total - has_price}")
