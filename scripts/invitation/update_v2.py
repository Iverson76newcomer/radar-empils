"""
Обновить приглашение:
1. Сохранить все правки пользователя
2. Добавить пункт про вечерний ужин в блок расходов
"""
import openpyxl
from openpyxl import load_workbook
import shutil
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from copy import deepcopy

SRC_PATH = "/home/z/my-project/upload/Приглашение_Олимп_Алексей_Анатольевич.docx"
OUT_PATH = "/home/z/my-project/download/Приглашение_Олимп_Алексей_Анатольевич_v2.docx"

# Копируем оригинал (с правками пользователя) — это будет основа
shutil.copy(SRC_PATH, OUT_PATH)

# Открыть для редактирования
doc = Document(OUT_PATH)

# Найти таблицу с расходами
target_table = None
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if 'Эмпилс берёт на себя' in cell.text:
                target_table = cell
                break

if not target_table:
    print("✗ Не найдена ячейка с расходами")
    exit(1)

print("=== Найдена ячейка с расходами ===")
print(f"Абзацев в ячейке: {len(target_table.paragraphs)}")
for i, p in enumerate(target_table.paragraphs):
    print(f"  [{i}] '{p.text}'")

# Найти последний пункт "трансфер из аэропорта"
last_para_with_check = None
last_para_idx = -1
for i, p in enumerate(target_table.paragraphs):
    if '✓' in p.text or 'трансфер' in p.text.lower():
        last_para_with_check = p
        last_para_idx = i

print(f"\nПоследний пункт с ✓: [{last_para_idx}] '{last_para_with_check.text}'")

# Создать новый параграф на основе последнего (скопировать формат)
new_para = deepcopy(last_para_with_check._element)

# Очистить текст и заменить на новый про ужин
# Найти все runs и обновить текст
from docx.text.run import Run
from docx.oxml.ns import qn

# Создать новый параграф (как копию последнего)
new_p = OxmlElement('w:p')
# Скопировать pPr (свойства параграфа)
src_pPr = last_para_with_check._element.find(qn('w:pPr'))
if src_pPr is not None:
    new_pPr = deepcopy(src_pPr)
    new_p.append(new_pPr)

# Создать 2 runs: "✓  " (жирный синий) и текст про ужин
# Беру стиль из существующих runs последнего параграфа
src_runs = last_para_with_check._element.findall(qn('w:r'))
print(f"  Найдено runs в источнике: {len(src_runs)}")

# First run — это "✓  "
if src_runs:
    # Копируем первый run (с галочкой)
    new_r1 = deepcopy(src_runs[0])
    new_p.append(new_r1)
    
    # Копируем второй run (с текстом) и меняем текст
    if len(src_runs) > 1:
        new_r2 = deepcopy(src_runs[1])
        # Найти t элемент и заменить текст
        t_elements = new_r2.findall(qn('w:t'))
        if t_elements:
            # Обновить текст на "ужин в ресторане в вечер визита"
            for t_elem in t_elements:
                t_elem.text = 'ужин в ресторане в вечер визита'
            new_p.append(new_r2)

# Вставить новый параграф после последнего пункта
last_para_with_check._element.addnext(new_p)

# Сохранить
doc.save(OUT_PATH)

# Проверить результат
print("\n=== Проверка результата ===")
doc2 = Document(OUT_PATH)
for table in doc2.tables:
    for row in table.rows:
        for cell in row.cells:
            if 'Эмпилс берёт на себя' in cell.text:
                print("Содержимое ячейки после обновления:")
                for i, p in enumerate(cell.paragraphs):
                    print(f"  [{i}] '{p.text}'")

print(f"\n✓ Файл сохранён: {OUT_PATH}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
