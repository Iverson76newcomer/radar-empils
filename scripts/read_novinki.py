import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Новинки за 2026 год первичка и вторичка.xlsx", data_only=True)
ws = wb["Первичные и вторичные продажи"]

print(f"=== Лист 'Первичные и вторичные продажи': {ws.max_row} строк × {ws.max_column} колонок ===\n")

# Print all rows
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r}: {vals}")
