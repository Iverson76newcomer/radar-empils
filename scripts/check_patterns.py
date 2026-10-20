import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

wb1 = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
ws1 = wb1['300 SKU — детально']

# Print sample items per brand
print("=== Примеры наименований по брендам ===")
brands_seen = {}
for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, values_only=True):
    if row[0] is None:
        continue
    brand = row[4]
    if brand not in brands_seen:
        brands_seen[brand] = []
    if len(brands_seen[brand]) < 5:
        brands_seen[brand].append(row[6])

for b, names in brands_seen.items():
    print(f"\n--- {b} ---")
    for n in names:
        print(f"  {n}")
