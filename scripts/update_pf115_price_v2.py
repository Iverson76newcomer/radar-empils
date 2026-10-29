"""
Финальное обновление Titan economic calculation.docx:
Применяем ОБЕ правки (TITAN 860 + ПФ-115 320) к исходному файлу за один проход.
Порядок замен важен: длинные/уникальные строки → короткие.
"""
import shutil
from docx import Document
from docx.oxml.ns import qn

SRC = '/home/z/my-project/upload/Titan economic calculation.docx'
OUT = '/home/z/my-project/download/Titan_economic_calculation.docx'
shutil.copy(SRC, OUT)

doc = Document(OUT)

# Порядок замен: сначала длинные уникальные строки, потом короткие
replacements = [
    # Вступительный текст
    ('TITAN 0,9 кг стоит ~700 ₽', 'TITAN 0,9 кг стоит ~860 ₽'),
    ('ПФ-115 0,9 кг — ~250 ₽',   'ПФ-115 0,9 кг — ~320 ₽'),

    # Плашка-вывод целиком (одной строкой, чтобы избежать конфликтов)
    ('5350 ₽ (8250 ₽ – 2900 ₽)', '5 954 ₽ (9 174 ₽ – 3 220 ₽)'),

    # Таблица — значения ПФ-115 (длинные с пробелами)
    ('8 250', '9 174'),     # ССВ 10 лет ПФ-115
    ('2 500', '2 780'),     # ИТОГО 1 перекраска ПФ-115
    ('1 000', '1 280'),     # Материалы ПФ-115
    # Таблица — значения TITAN (длинные с пробелами)
    ('1 400', '1 720'),     # Материалы TITAN
    ('2 900', '3 220'),     # ИТОГО 1 перекраска TITAN + ССВ 10 лет TITAN

    # Таблица — значения ССВ/год (короткие)
    ('825',   '917'),       # ССВ/год ПФ-115
    ('290',   '322'),       # ССВ/год TITAN

    # Цены банок (короткие, простые числа — применяем в самом конце)
    ('700',   '860'),       # Цена банки TITAN
    ('250',   '320'),       # Цена банки ПФ-115
]

# Утилиты
def replace_text_in_element(elem, old, new):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    t_nodes = elem.findall('.//w:t', ns)
    if not t_nodes:
        return 0
    full_text = ''.join(t.text or '' for t in t_nodes)
    if old not in full_text:
        return 0
    count = full_text.count(old)
    new_text = full_text.replace(old, new)
    if t_nodes:
        t_nodes[0].text = new_text
        for t in t_nodes[1:]:
            t.text = ''
    return count

counts = {}
for old, new in replacements:
    counts[old] = 0
    for p in doc.paragraphs:
        counts[old] += replace_text_in_element(p._p, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    counts[old] += replace_text_in_element(p._p, old, new)
                for nested_tbl in cell.tables:
                    for nested_row in nested_tbl.rows:
                        for nested_cell in nested_row.cells:
                            for p in nested_cell.paragraphs:
                                counts[old] += replace_text_in_element(p._p, old, new)

print('Замены в файле 2 (финальные):')
for old, new in replacements:
    c = counts.get(old, 0)
    print(f'  {old!r} → {new!r}: {c}')

doc.save(OUT)
print(f'\nOK saved: {OUT}')
import os
print(f'Size: {os.path.getsize(OUT)} bytes')
