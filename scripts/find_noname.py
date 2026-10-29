"""Найти все SKU с брендом NoName в файле v4."""
import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал_v4.xlsx", data_only=True)

# Проверить оба листа: '300 SKU — детально' и '300 SKU с ценами'
for sheet_name in ['300 SKU — детально', '300 SKU с ценами']:
    ws = wb[sheet_name]
    print(f"\n=== Лист '{sheet_name}' ===")
    
    # Шапка
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Шапка: {headers}")
    
    # Найти все строки с брендом NoName
    noname_rows = []
    for r in range(5, ws.max_row + 1):
        brand = ws.cell(row=r, column=5).value
        if brand and 'noname' in str(brand).lower():
            vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            noname_rows.append((r, vals))
    
    print(f"\nНайдено SKU с брендом 'NoName': {len(noname_rows)}")
    for r, vals in noname_rows:
        print(f"\n  Строка {r}:")
        for c, (header, val) in enumerate(zip(headers, vals), 1):
            if val is not None:
                print(f"    {header}: {val}")
