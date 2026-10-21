"""
Поиск альтернатив в новом прайс-листе (16.10.2026) для 8 ранее несопоставленных SKU.
"""
import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

PRICE_PATH = "/home/z/my-project/upload/16.10.2026 РОССИЯ с НДС, руб..xlsx"
wb2 = load_workbook(PRICE_PATH, data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) \
            and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name, 'fасовка': h_val, 'базовая': m_val, 'отгрузочная': n_val,
            })

# 8 несопоставленных SKU
unmatched_items = [
    {
        'idx': 37, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 вишневая 1.9 кг', 'fas': 1.9,
        'color_keyword': 'вишневая', 'ral': None,
    },
    {
        'idx': 38, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 голубая 0.9 кг', 'fas': 0.9,
        'color_keyword': 'голубая', 'ral': None,
    },
    {
        'idx': 56, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 тёмно-зелёная 0.9 кг', 'fas': 0.9,
        'color_keyword': 'тёмно-зелёная', 'ral': None,
    },
    {
        'idx': 61, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 0.9 кг', 'fas': 0.9,
        'color_keyword': 'шоколадная', 'ral': None,
    },
    {
        'idx': 62, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 1.9 кг', 'fas': 1.9,
        'color_keyword': 'шоколадная', 'ral': None,
    },
    {
        'idx': 63, 'brand': 'Расцвет', 'sphere': 'Металл', 'purpose': 'Грунт-эмаль по металлу',
        'name': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 2.7 кг', 'fas': 2.7,
        'color_keyword': 'шоколадная', 'ral': None,
    },
    {
        'idx': 150, 'brand': 'Ореол', 'sphere': 'Помещения (интерьер)', 'purpose': 'Грунтовка',
        'name': 'Грунт-концентрат АКВАСТОП ОРЕОЛ 1 кг', 'fas': 1.0,
        'color_keyword': None, 'ral': None,
    },
    {
        'idx': 169, 'brand': 'NEWTONE', 'sphere': 'Помещения (интерьер)', 'purpose': 'Краска интерьерная',
        'name': 'Краска ВД «NEWTONE ABSOLUTE BLACK» для интерьеров моющаяся глубокоматовая черная RAL 9005 10 л 13.9 кг',
        'fas': 13.9, 'color_keyword': 'черная', 'ral': 'RAL 9005',
    },
]

# Найти все позиции 3 в 1 в листе Расцвет
print("=== Все '3 в 1' в листе 'Расцвет' (новый прайс 16.10.2026) ===")
for p in price_dict['Расцвет']:
    if p['name'] and '3 в 1' in p['name'].lower():
        print(f"  '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Найти все позиции с цветом в Расцвет '3 в 1' по цветам
print("\n=== Все уникальные цвета '3 в 1' в Расцвет ===")
unique = set()
for p in price_dict['Расцвет']:
    if p['name'] and '3 в 1' in p['name'].lower():
        unique.add(p['name'])
for n in unique:
    print(f"  {n}")

# Все NEWTONE черная
print("\n=== NEWTONE: 'черная' / 'чёрная' ===")
for p in price_dict['NEWTONE']:
    if p['name'] and ('чёрная' in p['name'].lower() or 'черная' in p['name'].lower()):
        print(f"  '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Все Ореол 'грунт-концентрат'
print("\n=== Ореол: 'грунт-концентрат' ===")
for p in price_dict['Ореол']:
    if p['name'] and 'грунт-концентрат' in p['name'].lower():
        print(f"  '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Поиск 'АКВАСТОП' во всех листах
print("\n=== Все листы: 'аквастоп' ===")
found = False
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'аквастоп' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
            found = True
if not found:
    print("  (ничего не найдено)")

# Поиск 'absolute' во всех листах
print("\n=== Все листы: 'absolute' ===")
found = False
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'absolute' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
            found = True
if not found:
    print("  (ничего не найдено)")

# Поиск 'черная' / 'чёрная' во всех листах с фасовкой 13.9
print("\n=== Все листы: 'черная' или 'чёрная' с фасовкой близко к 13.9 ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and ('чёрная' in p['name'].lower() or 'черная' in p['name'].lower()):
            if p['fасовка'] and isinstance(p['fасовка'], (int, float)) and abs(p['fасовка'] - 13.9) < 1:
                print(f"  [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
