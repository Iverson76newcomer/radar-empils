"""
Выбор 5 популярных SKU из ABC-анализа (группа А), которых нет в текущем ассортименте 300 SKU.
"""
import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

# === Загрузить ассортимент 300 SKU ===
wb1 = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
ws1 = wb1["300 SKU — детально"]

assortment_names = []
for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, values_only=True):
    if row[0] is None:
        continue
    assortment_names.append({
        'idx': row[0],
        'sphere': row[1], 'purpose': row[2], 'base': row[3],
        'brand': row[4], 'type': row[5], 'name': row[6],
        'grp_code': row[7], 'abc': row[8]
    })
print(f"Ассортимент: {len(assortment_names)} SKU")

# Множество групповых кодов
assortment_grp_codes = set(item['grp_code'] for item in assortment_names if item['grp_code'])
print(f"Уникальных групповых кодов: {len(assortment_grp_codes)}")

# === Загрузить ABC-анализ ===
wb2 = load_workbook("/home/z/my-project/upload/Ассортимент для Олипма.xlsx", data_only=True)
ws2 = wb2["АBС анализ"]

abc_items = []
for row in ws2.iter_rows(min_row=4, max_row=ws2.max_row, values_only=True):
    if row[0] is None:
        continue
    abc_items.append({
        'sbu': row[0], 'brand': row[1], 'type': row[2], 'grp_code': row[3],
        'name': row[4], 'cumulative': row[5], 'group': row[6], 'mml': row[7]
    })
print(f"\nABC-анализ: {len(abc_items)} позиций")

# Группы
group_counts = defaultdict(int)
for it in abc_items:
    group_counts[it['group']] += 1
print(f"Распределение по группам:")
for g, c in sorted(group_counts.items()):
    print(f"  {g}: {c}")

# === Фильтрация: позиции группы А, которых НЕТ в ассортименте (по групповому коду) ===
candidates_A = []
for it in abc_items:
    if it['group'] != 'А':
        continue
    if it['grp_code'] in assortment_grp_codes:
        continue
    candidates_A.append(it)

print(f"\nКандидаты группы А (не в ассортименте): {len(candidates_A)}")

# Сортируем по нарастающему итогу (чем меньше cumulative — тем выше в топе)
candidates_A.sort(key=lambda x: x['cumulative'] if x['cumulative'] is not None else 1)

# Топ-30 кандидатов
print("\n=== ТОП-30 кандидатов (по убыванию популярности) ===")
print(f"{'#':<4} {'SBU':<6} {'Бренд':<22} {'Тип':<20} {'Код':<7} {'Нар.итог':<10} {'Наименование'}")
for i, it in enumerate(candidates_A[:30], 1):
    print(f"{i:<4} {it['sbu'][:5]:<6} {(it['brand'] or '')[:20]:<22} {(it['type'] or '')[:18]:<20} {str(it['grp_code']):<7} {(str(it['cumulative'])[:8] if it['cumulative'] else 'N/A'):<10} {(it['name'] or '')[:60]}")

# Распределение топ-30 по брендам
print("\n=== Распределение топ-30 по брендам ===")
brand_count = defaultdict(int)
for it in candidates_A[:30]:
    brand_count[it['brand']] += 1
for b, c in sorted(brand_count.items(), key=lambda x: -x[1]):
    print(f"  {b}: {c}")

# Распределение топ-30 по типам
print("\n=== Распределение топ-30 по типам продукции ===")
type_count = defaultdict(int)
for it in candidates_A[:30]:
    type_count[it['type']] += 1
for t, c in sorted(type_count.items(), key=lambda x: -x[1]):
    print(f"  {t}: {c}")
