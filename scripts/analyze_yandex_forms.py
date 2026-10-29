#!/usr/bin/env python3
"""
Универсальный скрипт анализа ответов из Яндекс.Форм через agent-browser.
Использует браузер для авторизованных API-запросов.

Использование:
    python3 analyze_yandex_forms.py <survey_id> <report_title> <output_docx>

Пример:
    python3 analyze_yandex_forms.py 6a4377dc49af47002b451dbe "Лемана ПРО" /home/z/my-project/download/Отчёт_Лемана_ПРО.docx
"""

import json
import sys
import os
import time
import subprocess
from collections import Counter
import matplotlib
matplotlib.use('Agg')
import matplotlib.font_manager as fm
# Liberation Sans supports Cyrillic and is available
fm.fontManager.addfont('/usr/share/fonts/truetype/chinese/LiberationSans-Regular.ttf')
fm.fontManager.addfont('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
import matplotlib.pyplot as plt
plt.rcParams['font.sans-serif'] = ['Liberation Sans', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CHARTS_DIR = '/tmp/yandex_charts'
COLORS = [
    '#2563eb', '#7c3aed', '#db2777', '#ea580c', '#16a34a',
    '#0891b2', '#4f46e5', '#c026d3', '#dc2626', '#ca8a04',
    '#059669', '#0284c7', '#9333ea', '#f59e0b', '#10b981',
]


def browser_eval(js_code):
    """Execute JS in browser via agent-browser, return parsed result."""
    cmd = ['agent-browser', 'eval', js_code]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=60,
                           encoding='utf-8', errors='replace')
    output = result.stdout.strip()
    # agent-browser returns JSON-encoded string, need to parse it
    if output.startswith('"') and output.endswith('"'):
        try:
            # Use json to properly decode escape sequences
            output = json.loads(output)
        except json.JSONDecodeError:
            # Fallback: strip quotes manually
            output = output[1:-1]
    return output


def get_form_info_via_browser(survey_id):
    """Get form info using browser session."""
    js = f"""
fetch('/admin/gateway/root/form/getSurveyInfo', {{
  method: 'POST', credentials: 'include',
  headers: {{'Content-Type': 'application/json', 'X-CSRF-Token': document.querySelector('meta[name=csrf-token]').content}},
  body: JSON.stringify({{surveyId: '{survey_id}'}})
}}).then(r => r.json()).then(j => JSON.stringify(j))
"""
    output = browser_eval(js)
    try:
        return json.loads(output)
    except json.JSONDecodeError:
        return None


def get_answers_via_browser(survey_id):
    """Get all answers using browser session."""
    js = f"""
fetch('/admin/gateway/root/form/getAnswers', {{
  method: 'POST', credentials: 'include',
  headers: {{'Content-Type': 'application/json', 'X-CSRF-Token': document.querySelector('meta[name=csrf-token]').content}},
  body: JSON.stringify({{
    surveyId: '{survey_id}',
    pageSize: 1000, view: 'cards', dateFrom: '', dateTo: '', questions: ''
  }})
}}).then(r => r.json()).then(j => JSON.stringify(j))
"""
    output = browser_eval(js)
    try:
        return json.loads(output)
    except json.JSONDecodeError:
        return None


def count_frequency(answers, col_idx):
    """Count frequency of each answer value for a column."""
    freq = Counter()
    total = 0
    for ans in answers:
        if col_idx >= len(ans['data']):
            continue
        item = ans['data'][col_idx]
        if not item:
            continue
        val = item.get('value')
        if val is None or val == '-' or val == '':
            continue
        if isinstance(val, list):
            for v in val:
                if v and v != '-' and v != 'Затрудняюсь':
                    freq[v] += 1
            total += 1
        else:
            if val != '-':
                freq[val] += 1
                total += 1
    return freq, total


def calc_scale_stats(answers, col_idx, is_matrix=False):
    """Calculate average for scale/matrix questions."""
    values = []
    if is_matrix:
        for ans in answers:
            if col_idx >= len(ans['data']):
                continue
            item = ans['data'][col_idx]
            if not item:
                continue
            val = item.get('value')
            if isinstance(val, list):
                for v in val:
                    try:
                        n = float(v)
                        if 0 <= n <= 10:
                            values.append(n)
                    except (ValueError, TypeError):
                        pass
    else:
        for ans in answers:
            if col_idx >= len(ans['data']):
                continue
            item = ans['data'][col_idx]
            if not item:
                continue
            val = item.get('value')
            if isinstance(val, list) and val:
                try:
                    values.append(float(val[0]))
                except (ValueError, TypeError):
                    pass
            elif val:
                try:
                    values.append(float(val))
                except (ValueError, TypeError):
                    pass
    if values:
        return sum(values) / len(values), values
    return 0, []


def make_bar_chart(freq_data, title, output_path):
    """Create a horizontal bar chart from frequency data."""
    if not freq_data:
        return None
    items = sorted(freq_data.items(), key=lambda x: -x[1])
    labels = [k[:50] + '...' if len(k) > 50 else k for k, _ in items]
    values = [v for _, v in items]
    total = sum(values)

    fig, ax = plt.subplots(figsize=(10, max(3, len(items) * 0.45 + 1)), constrained_layout=True)
    bars = ax.barh(range(len(labels)), values, color=COLORS[:len(labels)])
    ax.set_yticks(range(len(labels)))
    ax.set_yticklabels(labels, fontsize=9)
    ax.invert_yaxis()
    ax.set_xlabel('Количество ответов', fontsize=9)
    for bar, v in zip(bars, values):
        pct = v / total * 100 if total else 0
        ax.text(bar.get_width() + 0.1, bar.get_y() + bar.get_height()/2,
                f'{v} ({pct:.0f}%)', va='center', fontsize=8)
    ax.set_xlim(0, max(values) * 1.35 if values else 1)
    ax.set_title(title, fontsize=11, fontweight='bold', pad=10)
    ax.grid(axis='x', alpha=0.3)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return output_path


def make_scale_chart(values, avg, title, output_path, max_val=10):
    """Create a distribution chart for scale questions."""
    if not values:
        return None
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 4), constrained_layout=True,
                                    gridspec_kw={'width_ratios': [2, 1]})
    bins = list(range(0, int(max_val) + 2))
    counts, _, patches = ax1.hist(values, bins=bins, edgecolor='white', linewidth=1.2)
    for i, (c, p) in enumerate(zip(counts, patches)):
        if c > 0:
            color = '#ef4444' if i <= 3 else '#f59e0b' if i <= 7 else '#22c55e'
            p.set_facecolor(color)
    ax1.set_xlabel('Оценка', fontsize=9)
    ax1.set_ylabel('Количество', fontsize=9)
    ax1.set_title(f'{title}\nСреднее: {avg:.1f}', fontsize=10, fontweight='bold')
    ax1.set_xticks(bins[:-1])
    ax1.grid(axis='y', alpha=0.3)
    ax1.spines['top'].set_visible(False)
    ax1.spines['right'].set_visible(False)

    low = sum(1 for v in values if v <= 3)
    mid = sum(1 for v in values if 4 <= v <= 7)
    high = sum(1 for v in values if v >= 8)
    segs = [('Низкая (0-3)', low, '#ef4444'),
            ('Средняя (4-7)', mid, '#f59e0b'),
            ('Высокая (8-10)', high, '#22c55e')]
    segs = [(n, v, c) for n, v, c in segs if v > 0]
    if segs:
        ax2.pie([s[1] for s in segs], labels=[s[0] for s in segs],
                colors=[s[2] for s in segs], autopct='%1.0f%%', textprops={'fontsize': 9})
        ax2.set_title('Распределение', fontsize=10)
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return output_path


def add_chart_to_doc(doc, chart_path, width_cm=16):
    """Add chart image to document centered."""
    if chart_path and os.path.exists(chart_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(chart_path, width=Cm(width_cm))


def add_heading_styled(doc, text, level=1, color=(26, 31, 54)):
    """Add styled heading."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(*color)
    return h


def generate_docx_report(form_info, answers_data, output_path, report_title):
    """Generate full DOCX report with analysis."""
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Title page
    title = doc.add_heading('', level=0)
    run = title.add_run(f'Анализ ответов анкеты\n{report_title}')
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(198, 40, 40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'Форма: {form_info.get("name", "")}')
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(90, 96, 128)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f'Всего ответов: {len(answers_data.get("answers", []))}\n').bold = True
    p.add_run(f'Дата анализа: {time.strftime("%d.%m.%Y %H:%M")}\n')
    p.add_run(f'Источник: Яндекс.Формы')

    doc.add_page_break()

    columns = answers_data.get('columns', [])
    answers = answers_data.get('answers', [])
    total = len(answers)

    if total == 0:
        doc.add_paragraph('Ответов пока нет.')
        doc.save(output_path)
        return

    os.makedirs(CHARTS_DIR, exist_ok=True)

    # Summary
    add_heading_styled(doc, '1. Сводная информация', level=1)
    p = doc.add_paragraph()
    p.add_run(f'Всего получено ответов: ').bold = True
    p.add_run(f'{total}\n')
    p.add_run(f'Название формы: ').bold = True
    p.add_run(f'{form_info.get("name", "")}\n')
    p.add_run(f'Количество вопросов: ').bold = True
    p.add_run(f'{len(columns)}')

    # Per-question
    add_heading_styled(doc, '2. Анализ по вопросам', level=1)

    for idx, col in enumerate(columns):
        q_text = col.get('text', f'Вопрос {idx + 1}')
        q_type = col.get('type', 'unknown')
        widget = col.get('widget', '')

        add_heading_styled(doc, f'Вопрос {idx + 1}. {q_text[:200]}', level=2,
                          color=(198, 40, 40))

        type_label = {
            'enum': 'Выбор из вариантов' + (f' ({widget})' if widget else ''),
            'string': 'Текстовый ответ',
            'matrix': 'Шкала',
            'integer': 'Число',
        }.get(q_type, q_type)
        p = doc.add_paragraph()
        p.add_run(f'Тип вопроса: ').italic = True
        p.add_run(f'{type_label}').italic = True

        if q_type == 'matrix':
            avg, values = calc_scale_stats(answers, idx, is_matrix=True)
            max_val = max(values) if values else 10
            chart_path = os.path.join(CHARTS_DIR, f'q{idx + 1}.png')
            make_scale_chart(values, avg, f'Q{idx + 1}', chart_path, max_val=max_val)
            add_chart_to_doc(doc, chart_path)
            p = doc.add_paragraph()
            p.add_run(f'Среднее значение: ').bold = True
            p.add_run(f'{avg:.1f} (из {max_val:.0f})\n')
            p.add_run(f'Количество ответов: ').bold = True
            p.add_run(f'{len(values)}')

        elif q_type == 'string':
            text_answers = []
            for ans in answers:
                if idx < len(ans['data']):
                    item = ans['data'][idx]
                    if not item:
                        continue
                    val = item.get('value')
                    if val and val != '-' and isinstance(val, str):
                        text_answers.append(val)
                    elif isinstance(val, list):
                        text_answers.extend([v for v in val if v and v != '-'])
            if text_answers:
                p = doc.add_paragraph()
                p.add_run(f'Получено текстовых ответов: ').bold = True
                p.add_run(f'{len(text_answers)}\n')
                for ta in text_answers:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(ta)
            else:
                p = doc.add_paragraph()
                p.add_run('Текстовых ответов не получено.').italic = True
        else:
            is_multi = widget == 'checkbox'
            freq, freq_total = count_frequency(answers, idx)
            if freq:
                chart_path = os.path.join(CHARTS_DIR, f'q{idx + 1}.png')
                make_bar_chart(dict(freq), f'Q{idx + 1}', chart_path)
                add_chart_to_doc(doc, chart_path)
                p = doc.add_paragraph()
                if is_multi:
                    p.add_run(f'Всего респондентов: ').bold = True
                    p.add_run(f'{freq_total} (множественный выбор)\n')
                else:
                    p.add_run(f'Всего ответов: ').bold = True
                    p.add_run(f'{freq_total}\n')
                p = doc.add_paragraph()
                p.add_run('Топ ответов:').bold = True
                for i, (val, cnt) in enumerate(freq.most_common(3), 1):
                    pct = cnt / freq_total * 100 if freq_total else 0
                    p = doc.add_paragraph(style='List Number')
                    p.add_run(f'{val} — {cnt} ({pct:.0f}%)')
            else:
                p = doc.add_paragraph()
                p.add_run('Ответов не получено.').italic = True
        doc.add_paragraph()

    doc.save(output_path)
    print(f'Report saved: {output_path}')


def main():
    if len(sys.argv) < 4:
        print('Usage: python3 analyze_yandex_forms.py <survey_id> <report_title> <output_docx>')
        sys.exit(1)
    survey_id = sys.argv[1]
    report_title = sys.argv[2]
    output_path = sys.argv[3]

    # Make sure we're on the Yandex Forms page
    print('Refreshing browser session...')
    subprocess.run(['agent-browser', 'open', f'https://forms.yandex.ru/admin/{survey_id}/edit'],
                   capture_output=True, text=True, timeout=30)
    time.sleep(3)

    print(f'Fetching form info for {survey_id}...')
    form_info = get_form_info_via_browser(survey_id)
    if not form_info or 'name' not in form_info:
        print('ERROR: Failed to get form info.')
        sys.exit(1)
    print(f'Form: {form_info.get("name", "unknown")}')
    print(f'Answers: {form_info.get("answers", 0)}')

    print(f'\nFetching all answers...')
    answers_data = get_answers_via_browser(survey_id)
    if not answers_data:
        print('ERROR: Failed to get answers.')
        sys.exit(1)

    answers_count = len(answers_data.get('answers', []))
    columns_count = len(answers_data.get('columns', []))
    print(f'Got {answers_count} answers, {columns_count} columns')

    print(f'\nGenerating DOCX report: {output_path}')
    generate_docx_report(form_info, answers_data, output_path, report_title)
    print(f'\n✅ Done! Report saved to: {output_path}')


if __name__ == '__main__':
    main()
