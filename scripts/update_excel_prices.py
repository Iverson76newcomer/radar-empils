#!/usr/bin/env python3
"""
Обновление Excel-файла с найденными ценами.
Источник:
- prices_results.json — первый проход (6 SKU)
- prices_retry_results.json — второй проход (12 SKU, из них 11 найдены)
"""

import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_PATH = '/home/z/my-project/download/TITAN_конкуренты_цены_ВсеИнструменты.xlsx'

# Загрузка результатов первого прохода
with open('/home/z/my-project/scripts/prices_results.json', 'r', encoding='utf-8') as f:
    pass1 = json.load(f)

# Загрузка результатов второго прохода
with open('/home/z/my-project/scripts/prices_retry_results.json', 'r', encoding='utf-8') as f:
    pass2 = json.load(f)

# Объединение: для каждого SKU берём лучшую цену
# Если первый проход дал цену > 3000 (некорректную) — берём из второго
final_prices = {}
for r in pass1:
    n = r['n']
    price = r.get('price_new')
    if price and price <= 3000:  # нормальная цена
        final_prices[n] = {
            'price': price,
            'url': r.get('url', ''),
            'source': 'pass1',
        }

# Дополняем из второго прохода
for n_str, r in pass2.items():
    n = int(n_str)
    if n not in final_prices or not final_prices[n]['price']:
        if r.get('price_new'):
            final_prices[n] = {
                'price': r['price_new'],
                'url': r.get('url', ''),
                'source': 'pass2',
            }

print("=== Финальные цены ===")
for n in sorted(final_prices.keys()):
    r = final_prices[n]
    print(f"  SKU {n}: {r['price']} ₽ ({r['source']})")

# Обновление Excel
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb['Цены было-стало']

# Строки данных: 5-20 (16 SKU)
updated = 0
for ri, n in enumerate(range(1, 17), 5):
    if n in final_prices:
        price = final_prices[n]['price']
        url = final_prices[n]['url']
        ws.cell(ri, 8).value = price  # Колонка H = Цена актуальная
        # URL можно добавить в Notes или заменить search_query
        # Оставим — пользователь может посмотреть в сниппете
        updated += 1
        print(f"  ✓ SKU {n} (row {ri}): {price} ₽")

print(f"\nОбновлено: {updated} из 16")

# Сохранение
wb.save(EXCEL_PATH)
print(f"\nOK: {EXCEL_PATH}")
import os
print(f'Размер: {os.path.getsize(EXCEL_PATH) / 1024:.1f} KB')
