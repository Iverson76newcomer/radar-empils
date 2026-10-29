"""Проверка готового файла БРОЗЭКС."""
from openpyxl import load_workbook
wb = load_workbook('/home/z/my-project/download/TITAN_БРОЗЭКС_анализ_статистики.xlsx', data_only=True)
for sh_name in wb.sheetnames:
    ws = wb[sh_name]
    print(f"\n{'='*80}\nЛИСТ: {sh_name} | dim={ws.dimensions}\n{'='*80}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append('')
            elif isinstance(v, float):
                vals.append(f'{v:.2f}')
            else:
                vals.append(str(v))
        print(' | '.join(vals))
