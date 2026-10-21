import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

# === Загрузить НОВЫЙ прайс-лист от 16.10.2026 ===
PRICE_PATH = "/home/z/my-project/upload/16.10.2026 РОССИЯ с НДС, руб..xlsx"
wb2 = load_workbook(PRICE_PATH, data_only=True)
print(f"=== Новый прайс-лист 16.10.2026 ===")
print(f"Листы ({len(wb2.sheetnames)}):")
for sn in wb2.sheetnames:
    print(f"  {sn}")

# Собрать словарь прайс-листа
price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) \
            and not f_val.startswith('СТО') \
            and not f_val.startswith('ТУ') \
            and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })
print(f"\nЗаписей с ценами: {sum(len(v) for v in price_dict.values())}")

# 8 несопоставленных SKU из ассортимента
unmatched_items = [
    (37,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 вишневая 1.9 кг',        'Расцвет',  1.9),
    (38,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 голубая 0.9 кг',         'Расцвет',  0.9),
    (56,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 тёмно-зелёная 0.9 кг',   'Расцвет',  0.9),
    (61,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 0.9 кг',       'Расцвет',  0.9),
    (62,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 1.9 кг',       'Расцвет',  1.9),
    (63,  'Расцвет',  'Грунт-эмаль "РАСЦВЕТ" 3 в 1 шоколадная 2.7 кг',       'Расцвет',  2.7),
    (150, 'Ореол',    'Грунт-концентрат АКВАСТОП ОРЕОЛ 1 кг',                'Ореол',    1.0),
    (169, 'NEWTONE',  'Краска ВД «NEWTONE ABSOLUTE BLACK» для интерьеров моющаяся глубокоматовая черная RAL 9005 10 л 13.9 кг', 'NEWTONE', 13.9),
]

# === Поиск каждого SKU в новом прайс-листе ===
print("\n" + "=" * 110)
print("ПОИСК 8 РАНЕЕ НЕ СОПОСТАВЛЕННЫХ SKU В НОВОМ ПРАЙС-ЛИСТЕ (16.10.2026)")
print("=" * 110)

# Поиск по ключевым словам
for idx, brand, name, sheet, target_fas in unmatched_items:
    print(f"\n--- #{idx} | {brand} | {name} (фасовка={target_fas} кг) ---")
    
    # Определяем ключевые слова для поиска
    name_lower = name.lower()
    
    # Для РАСЦВЕТ 3 в 1 — ищем с цветом
    if '3 в 1' in name_lower or '3в1' in name_lower:
        # Извлечь цвет
        color_map = {
            'вишневая': ['вишн', 'красная'],
            'голубая': ['голуб'],
            'тёмно-зелёная': ['тёмно-зелён', 'темно-зелен', 'зелён'],
            'шоколадная': ['шоколад'],
        }
        color_keywords = []
        for k, v in color_map.items():
            if k in name_lower:
                color_keywords = v
                break
        
        print(f"  Поиск в листе 'Расцвет' по '3 в 1' + цвет: {color_keywords}")
        for p in price_dict['Расцвет']:
            if p['name'] and '3 в 1' in p['name'].lower():
                for kw in color_keywords:
                    if kw in p['name'].lower():
                        print(f"  ✓ НАЙДЕНО: '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
                        break
    
    # Для АКВАСТОП
    elif 'аквастоп' in name_lower:
        print(f"  Поиск по 'аквастоп' во ВСЕХ листах:")
        for sn in wb2.sheetnames:
            for p in price_dict[sn]:
                if p['name'] and 'аквастоп' in p['name'].lower():
                    print(f"  ✓ [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
        # Also search by 'грунт-концентрат'
        print(f"  Поиск по 'грунт-концентрат' в листе 'Ореол':")
        for p in price_dict['Ореол']:
            if p['name'] and 'грунт-концентрат' in p['name'].lower():
                print(f"  - '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
    
    # Для NEWTONE ABSOLUTE BLACK
    elif 'absolute' in name_lower or 'black' in name_lower:
        print(f"  Поиск по 'absolute' / 'black' во ВСЕХ листах:")
        for sn in wb2.sheetnames:
            for p in price_dict[sn]:
                if p['name'] and ('absolute' in p['name'].lower() or 'black' in p['name'].lower()):
                    print(f"  ✓ [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
        # Also search NEWTONE черная
        print(f"  Поиск по 'чёрная' / 'черная' в NEWTONE:")
        for p in price_dict['NEWTONE']:
            if p['name'] and ('чёрная' in p['name'].lower() or 'черная' in p['name'].lower()):
                print(f"  - '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
