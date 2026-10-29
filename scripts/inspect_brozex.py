"""Изучаем структуру файла БРОЗЭКС и образца (Рыгалов)."""
from openpyxl import load_workbook

print("=" * 80)
print("ФАЙЛ-ИСТОЧНИК: Титан БРОЗЭКС.xlsx")
print("=" * 80)
wb_src = load_workbook('/home/z/my-project/upload/Титан БРОЗЭКС.xlsx', data_only=True)
for sh_name in wb_src.sheetnames:
    ws = wb_src[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions} | max_row={ws.max_row} | max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append('')
            elif isinstance(v, float):
                vals.append(f'{v:.2f}')
            else:
                vals.append(str(v))
        print(' | '.join(vals))

print("\n\n" + "=" * 80)
print("ОБРАЗЕЦ РЕЗУЛЬТАТА: TITAN_Рыгалов_анализ_статистики.xlsx")
print("=" * 80)
wb_res = load_workbook('/home/z/my-project/download/TITAN_Рыгалов_анализ_статистики.xlsx')
for sh_name in wb_res.sheetnames:
    ws = wb_res[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions} | max_row={ws.max_row} | max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append('')
            elif isinstance(v, float):
                vals.append(f'{v:.2f}')
            else:
                vals.append(str(v))
        print(' | '.join(vals))
