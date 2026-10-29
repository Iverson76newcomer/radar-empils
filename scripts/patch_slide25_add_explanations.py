#!/usr/bin/env python3
"""
Дополнение пояснений на слайде 25 в презентации Эмпилс.
Принцип: каждый пункт маркированного списка = заголовок (23pt bold #313131)
         + тире с пояснением (20pt не bold #6B6B6B).
"""

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

SRC = "/home/z/my-project/upload/Презентация_для_новых_партнёров.pptx"
DST = "/home/z/my-project/download/Презентация_для_новых_партнёров_доп.pptx"

# Цвета (как в исходном файле)
ORANGE = RGBColor(0xF4, 0x76, 0x29)
DARK = RGBColor(0x31, 0x31, 0x31)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)

FONT = "Calibri"

# === Карта дополнений ===
# (параграф_индекс, заголовок_текста, пояснение)
# Используем реальные факты Эмпилс из презентации (8% рынка, 2% на трейд-маркетинг и т.д.)
ADDITIONS = {
    0: ("ЛИДЕР НА РЫНКЕ ЛКМ", "ТОП-3 производителей России · 8% рынка в денежном выражении"),
    1: (">100 лет непрерывного развития", "с 1906 года · устойчивость и опыт поколений"),
    4: ("МАРКЕТИНГОВАЯ И трейд-маркетинговая поддержка", "2% от оборота на развитие вашего бизнеса"),
    5: ("ПОДДЕРЖКА ТРЕТИЧНЫХ ПРОДАЖ", "стимулирование спроса на полке и в рознице"),
    6: ("региональная команда", "менеджеры и ЭТП в вашем регионе"),
}

def clear_paragraph_runs(p):
    """Удалить все runs из параграфа."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)

def add_run(p, text, size, bold, color):
    """Добавить run с заданным форматированием."""
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r

def rebuild_paragraph(p, header, explanation):
    """Пересоздать параграф в эталонной структуре:
       ▸  [HEADER]  —  [explanation]
       Маркер: 24pt bold оранжевый
       Заголовок: 23pt bold тёмный
       Пояснение: 20pt не bold серый
    """
    clear_paragraph_runs(p)
    add_run(p, "▸  ", 24, True, ORANGE)
    add_run(p, header, 23, True, DARK)
    add_run(p, "  —  ", 20, False, GRAY)
    add_run(p, explanation, 20, False, GRAY)

def main():
    prs = Presentation(SRC)
    slide = prs.slides[24]  # слайд 25 (0-indexed)

    # Нужная текстовая рамка — shape[2]
    text_shape = slide.shapes[2]
    tf = text_shape.text_frame

    print(f"Слайд 25 — обработка {len(ADDITIONS)} параграфов\n")

    for para_idx, (header, explanation) in ADDITIONS.items():
        para = tf.paragraphs[para_idx]
        old_text = para.text
        rebuild_paragraph(para, header, explanation)
        new_text = para.text
        print(f"  P{para_idx}: '{old_text}'")
        print(f"        → '{new_text}'\n")

    # Сохраняем
    import os
    os.makedirs(os.path.dirname(DST), exist_ok=True)
    prs.save(DST)
    print(f"✅ Сохранено: {DST}")
    print(f"   Размер: {os.path.getsize(DST):,} байт")

if __name__ == "__main__":
    main()
