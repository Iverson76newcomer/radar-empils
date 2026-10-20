import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# Search Расцвет '3 в 1' с цветами
print("=== Расцвет: '3 в 1' вишневая/голубая/светло-серая/тёмно-зелёная/шоколадная ===")
for p in price_dict['Расцвет']:
    if p['name'] and '3 в 1' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# NEWTONE SKY
print("\n=== NEWTONE: 'SKY' ===")
for p in price_dict['NEWTONE']:
    if p['name'] and 'sky' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# NEWTONE ABSOLUTE
print("\n=== NEWTONE: 'ABSOLUTE' ===")
for p in price_dict['NEWTONE']:
    if p['name'] and 'absolute' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Ореол - АКВАСТОП
print("\n=== Ореол: 'АКВАСТОП' / 'аквастоп' ===")
for p in price_dict['Ореол']:
    if p['name'] and 'аквастоп' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Ореол - all 'грунт-концентрат'
print("\n=== Ореол: 'грунт-концентрат' ===")
for p in price_dict['Ореол']:
    if p['name'] and 'грунт-концентрат' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Ореол Дисконт - 'для стен и потолков'
print("\n=== Ореол Дисконт: 'для стен и потолков' ===")
for p in price_dict['Ореол Дисконт']:
    if p['name'] and 'стен' in p['name'].lower() and 'потолк' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']}")

# PROFIWOOD - 'лак'
print("\n=== PROFIWOOD: 'лак' ===")
seen = set()
for p in price_dict['PROFIWOOD']:
    if p['name'] and 'лак' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Ореол - 'шпатлевка'
print("\n=== Ореол: 'шпатлевка' ===")
seen = set()
for p in price_dict['Ореол']:
    if p['name'] and 'шпатлевка' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}' фас={p['fасовка']}")
