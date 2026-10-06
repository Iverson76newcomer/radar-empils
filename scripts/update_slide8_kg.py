"""
Update slide 8 table:
- Replace "Первичка (₽)" with "Первичка (кг)" 
- Add "Вторичка (кг)" column
- Remove ₽ from data
- Keep АКБ
- Update subtitle and totals to use kg instead of ₽

New column structure: #, Новинка, SKU, Фасовки, Цвета, Первичка (кг), Вторичка (кг), АКБ = 8 columns

Also check for table overlapping with the 3-phase section below.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import shutil
import openpyxl
import re

SRC = "/home/z/my-project/upload/Презентация_для_ГД_final.pptx"
OUT = "/home/z/my-project/download/Презентация_для_ГД_final.pptx"
EXCEL_SRC = "/home/z/my-project/upload/Новинки за 2026 год первичка и вторичка.xlsx"

shutil.copy(SRC, OUT)

# ── Palette ──────────────────────────────────────────────────────────
BG          = RGBColor(0xF5, 0xF2, 0xEC)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xD9, 0xD2, 0xC2)
DIVIDER     = RGBColor(0xEF, 0xE9, 0xDD)
INK         = RGBColor(0x2F, 0x3A, 0x48)
MUTE        = RGBColor(0x8A, 0x95, 0xA3)
RED         = RGBColor(0xC0, 0x00, 0x00)
SAND        = RGBColor(0xC9, 0xA7, 0x7C)
BLUE        = RGBColor(0x6B, 0x8F, 0xA8)
GREEN       = RGBColor(0x8B, 0xAA, 0x8A)
HEADER_BG   = RGBColor(0x2F, 0x3A, 0x48)
ROW_ALT     = RGBColor(0xF5, 0xF2, 0xEC)
ROW_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

def cm(v):
    return Emu(int(v * 360000))

def add_text(slide, x, y, w, h, text, *,
             size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
             va=MSO_ANCHOR.TOP, font="Inter", spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = va
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if spacing:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(spacing))
    return tb

def add_rect(slide, x, y, w, h, fill, *,
             line=None, line_w=6350, corner_radius=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner_radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Emu(line_w)
    else:
        s.line.fill.background()
    if corner_radius is not None:
        try: s.adjustments[0] = corner_radius
        except: pass
    s.shadow.inherit = False
    return s


# ── Read Excel data ────────────────────────────────────────────────────
wb_src = openpyxl.load_workbook(EXCEL_SRC, data_only=True)
ws = wb_src['Первичные и вторичные продажи']

novinki = {}
current_n = None
current_data = None

for row in ws.iter_rows(min_row=3, values_only=True):
    if not row or not any(row):
        continue
    n_val = row[0]
    name = row[4] if len(row) > 4 else None
    series = row[2] if len(row) > 2 else None
    p1_kg = row[8] if len(row) > 8 else 0
    p1_rub = row[9] if len(row) > 9 else 0
    p2_kg = row[10] if len(row) > 10 else 0
    p2_rub = row[11] if len(row) > 11 else 0
    akb = row[12] if len(row) > 12 else 0

    if series is None and name is None:
        continue

    if n_val is not None:
        current_n = n_val
        sku_text = str(name) if name else ''
        if 'PROFIWOOD' in sku_text and 'Антисептик' in sku_text:
            short_name = 'PROFIWOOD антисептик кроющий ВД'
        elif 'OSB' in sku_text:
            short_name = 'Ореол краска для плит OSB'
        elif 'ABSOLUTE BLACK' in sku_text:
            short_name = 'NEWTONE ABSOLUTE BLACK'
        elif 'SKY' in sku_text:
            short_name = 'NEWTONE SKY потолочная'
        elif 'резиновая' in sku_text.lower():
            short_name = 'Ореол краска резиновая ВД'
        else:
            short_name = sku_text[:50]
        current_data = {
            'short_name': short_name,
            'skus': [],
            'total_p1_kg': 0, 'total_p1_rub': 0,
            'total_p2_kg': 0, 'total_p2_rub': 0,
            'total_akb': 0, 'sku_count': 0,
        }
        novinki[current_n] = current_data

    if current_data is not None and name:
        name_str = str(name)
        color = ''
        color_keywords = {
            'бежевый': 'бежевый', 'серый туман': 'серый туман',
            'скандинавский серый': 'сканд. серый',
            'горький шоколад': 'горький шоколад',
            'шоколадный мусс': 'шоколад. мусс',
            'белый': 'белый', 'база а': 'база А', 'база с': 'база С',
            'черная': 'чёрный', 'чёрная': 'чёрный',
            'ультрабелая': 'ультрабелый',
        }
        for kw, c in color_keywords.items():
            if kw in name_str.lower():
                color = c
                break
        vol_match = re.search(r'(\d+(?:\.\d+)?)\s*(л|кг)', name_str, re.IGNORECASE)
        volume = vol_match.group(0) if vol_match else ''
        current_data['skus'].append({'color': color, 'volume': volume})
        current_data['total_p1_kg'] += float(p1_kg) if p1_kg else 0
        current_data['total_p2_kg'] += float(p2_kg) if p2_kg else 0
        current_data['total_akb'] += int(akb) if akb else 0
        current_data['sku_count'] += 1

# Build table data
table_data = []
for n in sorted(novinki.keys()):
    d = novinki[n]
    colors = sorted(set(s['color'] for s in d['skus'] if s['color']))
    volumes = sorted(set(s['volume'] for s in d['skus'] if s['volume']))
    table_data.append({
        'n': n,
        'name': d['short_name'],
        'sku': d['sku_count'],
        'volumes': ', '.join(volumes) if volumes else '—',
        'colors': ', '.join(colors) if colors else '—',
        'p1_kg': d['total_p1_kg'],
        'p2_kg': d['total_p2_kg'],
        'akb': d['total_akb'],
    })

total_p1_kg = sum(t['p1_kg'] for t in table_data)
total_p2_kg = sum(t['p2_kg'] for t in table_data)
total_sku = sum(t['sku'] for t in table_data)
total_akb = sum(t['akb'] for t in table_data)

print("Table data:")
for t in table_data:
    print(f"  {t['n']}: P1={t['p1_kg']:.1f} кг, P2={t['p2_kg']:.1f} кг, АКБ={t['akb']}")
print(f"Total: P1={total_p1_kg:.1f} кг, P2={total_p2_kg:.1f} кг, SKU={total_sku}, АКБ={total_akb}")


# ════════════════════════════════════════════════════════════════════════
# Open presentation, find slide 8, clear all shapes
# ════════════════════════════════════════════════════════════════════════
prs = Presentation(OUT)

target_idx = None
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame and 'Запуск новинок' in shape.text_frame.text:
            target_idx = i
            break
    if target_idx is not None:
        break

print(f"\nSlide 'Запуск новинок' at index {target_idx}")
slide8 = list(prs.slides)[target_idx]

# Remove ALL shapes
sp_tree = slide8.shapes._spTree
shapes_to_remove = []
for elem in sp_tree:
    if elem.tag.endswith('}sp') or elem.tag.endswith('}pic') or elem.tag.endswith('}grpSp') or elem.tag.endswith('}cxnSp') or elem.tag.endswith('}graphicFrame'):
        shapes_to_remove.append(elem)
print(f"Removing {len(shapes_to_remove)} shapes")
for elem in shapes_to_remove:
    sp_tree.remove(elem)

bg = slide8.background
bg.fill.solid()
bg.fill.fore_color.rgb = BG


# ════════════════════════════════════════════════════════════════════════
# HEADER
# ════════════════════════════════════════════════════════════════════════
add_text(slide8, cm(0.5), cm(0.5),
         cm(20.0), cm(0.95),
         "Запуск новинок · Стандарт вывода 2026",
         size=20, bold=True, color=INK, font="Inter")
add_text(slide8, cm(0.5), cm(1.55),
         cm(20.0), cm(0.5),
         f"Итоги 5 новинок 2026 · Первичка: {total_p1_kg:,.0f} кг · Вторичка: {total_p2_kg:,.0f} кг · {total_sku} SKU · АКБ {total_akb}".replace(',', ' '),
         size=11, color=MUTE, font="Inter")
# Status pill
add_rect(slide8, cm(20.0), cm(0.45),
         cm(5.0), cm(0.78),
         WHITE, line=SAND, line_w=9525, corner_radius=0.5)
add_text(slide8, cm(20.0), cm(0.58),
         cm(5.0), cm(0.65),
         "СТАНДАРТ · 2026", size=10, bold=True,
         color=SAND, font="Inter", align=PP_ALIGN.CENTER, spacing=150)
# Divider
add_rect(slide8, cm(0.5), cm(2.22),
         cm(24.5), cm(0.08), CARD_BORDER)


# ════════════════════════════════════════════════════════════════════════
# TABLE — 8 columns: #, Новинка, SKU, Фасовки, Цвета, Первичка кг, Вторичка кг, АКБ
# Total width: 24.5 cm
# ════════════════════════════════════════════════════════════════════════

TABLE_X = cm(0.5)
TABLE_Y = cm(2.67)
TABLE_W = cm(24.5)

# 8 columns — total 24.5 cm
COL_WIDTHS_CM = [0.8, 4.8, 1.0, 3.6, 5.6, 2.9, 2.9, 2.9]
COL_WIDTHS = [cm(w) for w in COL_WIDTHS_CM]
print(f"Sum of column widths: {sum(COL_WIDTHS_CM)} cm (should be 24.5)")

HEADER_H = cm(0.7)
ROW_H = cm(0.85)
TOTAL_H = cm(0.7)
rows_count = 7  # header + 5 data + total
table_total_h = HEADER_H + 5 * ROW_H + TOTAL_H

# Create table
table_shape = slide8.shapes.add_table(rows_count, 8, TABLE_X, TABLE_Y, TABLE_W, table_total_h)
table = table_shape.table

# Set column widths
for i, w in enumerate(COL_WIDTHS):
    table.columns[i].width = w

# Set row heights
table.rows[0].height = HEADER_H
for i in range(1, 6):
    table.rows[i].height = ROW_H
table.rows[6].height = TOTAL_H

# Disable default styling
tbl = table._tbl
for elem in tbl.findall(qn('a:tblPr')):
    tbl.remove(elem)
tbl_pr = tbl.makeelement(qn('a:tblPr'), {'firstRow': '0', 'bandRow': '0'})
tbl.insert(0, tbl_pr)

def set_cell(cell, text, *, size=10, bold=False, color=INK, bg_color=WHITE,
             align=PP_ALIGN.LEFT, font="Inter", va=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = va
    cell.margin_left = Emu(50000)
    cell.margin_right = Emu(50000)
    cell.margin_top = Emu(20000)
    cell.margin_bottom = Emu(20000)
    cell.fill.solid()
    cell.fill.fore_color.rgb = bg_color
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for r in p.runs:
        r.text = ''
    r = p.add_run()
    r.text = str(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color

# Headers
HEADERS = ['#', 'Новинка', 'SKU', 'Фасовки', 'Цвета', 'Первичка (кг)', 'Вторичка (кг)', 'АКБ']
for col_idx, h in enumerate(HEADERS):
    align = PP_ALIGN.CENTER if col_idx in (0, 2, 5, 6, 7) else PP_ALIGN.LEFT
    set_cell(table.cell(0, col_idx), h, size=9.5, bold=True, color=WHITE,
             bg_color=HEADER_BG, align=align, font="Inter")

# Data rows
NOVINKI_ACCENTS = [RED, SAND, BLUE, GREEN, RGBColor(0x7B, 0x6F, 0xA8)]

for row_idx, t in enumerate(table_data):
    excel_row = row_idx + 1
    bg_color = ROW_ALT if row_idx % 2 == 0 else ROW_WHITE
    accent = NOVINKI_ACCENTS[row_idx % 5]

    set_cell(table.cell(excel_row, 0), f"{t['n']:02d}", size=10, bold=True,
             color=accent, bg_color=bg_color, align=PP_ALIGN.CENTER, font="Roboto Mono")
    set_cell(table.cell(excel_row, 1), t['name'], size=9.5, bold=True,
             color=INK, bg_color=bg_color, align=PP_ALIGN.LEFT)
    set_cell(table.cell(excel_row, 2), str(t['sku']), size=10, bold=True,
             color=INK, bg_color=bg_color, align=PP_ALIGN.CENTER, font="Roboto Mono")
    set_cell(table.cell(excel_row, 3), t['volumes'], size=8.5,
             color=MUTE, bg_color=bg_color, align=PP_ALIGN.LEFT)
    set_cell(table.cell(excel_row, 4), t['colors'], size=8.5,
             color=MUTE, bg_color=bg_color, align=PP_ALIGN.LEFT)
    # Первичка (кг)
    set_cell(table.cell(excel_row, 5), f"{t['p1_kg']:,.1f}".replace(',', ' '),
             size=10, bold=True, color=accent, bg_color=bg_color,
             align=PP_ALIGN.RIGHT, font="Roboto Mono")
    # Вторичка (кг) — show 0 as "—" for clarity, otherwise actual value
    p2_text = f"{t['p2_kg']:,.1f}".replace(',', ' ') if t['p2_kg'] > 0 else "—"
    p2_color = INK if t['p2_kg'] > 0 else MUTE
    set_cell(table.cell(excel_row, 6), p2_text, size=10, bold=(t['p2_kg'] > 0),
             color=p2_color, bg_color=bg_color, align=PP_ALIGN.RIGHT, font="Roboto Mono")
    # АКБ
    set_cell(table.cell(excel_row, 7), str(t['akb']), size=10, bold=True,
             color=INK, bg_color=bg_color, align=PP_ALIGN.CENTER, font="Roboto Mono")

# Total row
set_cell(table.cell(6, 0), "Σ", size=10, bold=True, color=WHITE,
         bg_color=INK, align=PP_ALIGN.CENTER, font="Roboto Mono")
set_cell(table.cell(6, 1), "ИТОГО", size=10, bold=True, color=WHITE,
         bg_color=INK, align=PP_ALIGN.LEFT)
set_cell(table.cell(6, 2), str(total_sku), size=10, bold=True, color=WHITE,
         bg_color=INK, align=PP_ALIGN.CENTER, font="Roboto Mono")
set_cell(table.cell(6, 3), "", bg_color=INK)
set_cell(table.cell(6, 4), "", bg_color=INK)
set_cell(table.cell(6, 5), f"{total_p1_kg:,.0f}".replace(',', ' '), size=10, bold=True,
         color=WHITE, bg_color=INK, align=PP_ALIGN.RIGHT, font="Roboto Mono")
set_cell(table.cell(6, 6), f"{total_p2_kg:,.0f}".replace(',', ' '), size=10, bold=True,
         color=WHITE, bg_color=INK, align=PP_ALIGN.RIGHT, font="Roboto Mono")
set_cell(table.cell(6, 7), str(total_akb), size=10, bold=True, color=WHITE,
         bg_color=INK, align=PP_ALIGN.CENTER, font="Roboto Mono")


# ════════════════════════════════════════════════════════════════════════
# Check: where does the table end?
# Table starts at y=2.67cm, total height = 0.7 + 5*0.85 + 0.7 = 5.65 cm
# So table ends at y = 2.67 + 5.65 = 8.32 cm
# We need to leave a gap of ~0.4 cm before the phase section
# Phase section title should be at y = 8.7 cm (same as before — OK)
# Phase cards at y = 9.3 cm with height 4.0 cm → end at y = 13.3 cm
# Bottom note at y = 13.5 cm — fits within slide height 14.29 cm
# ════════════════════════════════════════════════════════════════════════

table_end_y = 2.67 + 5.65
print(f"Table ends at y = {table_end_y:.2f} cm")
print(f"Phase section starts at y = 8.70 cm")
print(f"Gap between table and phase section: {8.70 - table_end_y:.2f} cm (OK if positive)")

# Section title for phases
PHASE_TITLE_Y = cm(8.7)
add_text(slide8, cm(0.5), PHASE_TITLE_Y,
         cm(24.5), cm(0.4),
         "СТАНДАРТ ВЫВОДА · 3 ФАЗЫ",
         size=11, bold=True, color=RED, font="Inter", spacing=200)
add_rect(slide8, cm(0.5), PHASE_TITLE_Y + cm(0.45),
         cm(24.5), cm(0.04), DIVIDER)

# 3 phase cards — horizontal row
PHASE_Y = cm(9.3)
PHASE_H = cm(4.0)
PHASE_W = cm(7.8)
PHASE_GAP = cm(0.3)
PHASE_X0 = cm(0.5)

PHASES = [
    {"n": "01", "title": "ПОДГОТОВКА", "duration": "−60 до 0 дней",
     "bullets": ["Анализ рынка и конкурентов", "Ценообразование и MML", "Подготовка POSM и контента"],
     "kpi_value": "60 дней", "kpi_desc": "до запуска", "accent": RED},
    {"n": "02", "title": "ЗАПУСК", "duration": "0 — 30 дней",
     "bullets": ["Старт продаж во всех пилотах", "Обучение ТП: 350 человек", "POSM в 100% пилотных точек"],
     "kpi_value": "100%", "kpi_desc": "наличие POSM", "accent": SAND},
    {"n": "03", "title": "СОПРОВОЖДЕНИЕ", "duration": "30 — 90 дней",
     "bullets": ["Промоакции и мастер-классы", "Куратор на каждую точку", "Аналитика продаж по неделе"],
     "kpi_value": "350 ТП", "kpi_desc": "обучено в 1-й мес.", "accent": BLUE},
]

for i, p in enumerate(PHASES):
    px = PHASE_X0 + i * (PHASE_W + PHASE_GAP)

    add_rect(slide8, px, PHASE_Y, PHASE_W, PHASE_H,
             WHITE, line=CARD_BORDER, line_w=6350, corner_radius=0.03)
    add_rect(slide8, px, PHASE_Y, PHASE_W, cm(0.06), p['accent'])

    add_text(slide8, px + cm(0.4), PHASE_Y + cm(0.3),
             cm(1.5), cm(0.7),
             p['n'], size=24, bold=True, color=p['accent'],
             font="Roboto Mono", va=MSO_ANCHOR.MIDDLE)
    add_text(slide8, px + cm(2.0), PHASE_Y + cm(0.35),
             cm(5.5), cm(0.4),
             p['title'], size=11, bold=True, color=INK,
             font="Inter", spacing=100)
    add_text(slide8, px + cm(2.0), PHASE_Y + cm(0.75),
             cm(5.5), cm(0.3),
             p['duration'], size=9, bold=True, color=p['accent'],
             font="Roboto Mono")

    add_rect(slide8, px + cm(0.4), PHASE_Y + cm(1.3),
             PHASE_W - cm(0.8), cm(0.04), DIVIDER)

    bullet_y = PHASE_Y + cm(1.5)
    for j, b in enumerate(p['bullets']):
        add_text(slide8, px + cm(0.4), bullet_y,
                 cm(0.2), cm(0.3),
                 "·", size=12, bold=True, color=p['accent'],
                 font="Inter")
        add_text(slide8, px + cm(0.65), bullet_y + cm(0.02),
                 PHASE_W - cm(1.1), cm(0.3),
                 b, size=9, color=INK, font="Inter")
        bullet_y += cm(0.38)

    KPI_Y = PHASE_Y + PHASE_H - cm(0.85)
    add_rect(slide8, px + cm(0.4), KPI_Y - cm(0.1),
             PHASE_W - cm(0.8), cm(0.03), DIVIDER)
    add_text(slide8, px + cm(0.4), KPI_Y,
             cm(3.0), cm(0.35),
             "KPI:", size=8, bold=True, color=MUTE,
             font="Inter", va=MSO_ANCHOR.MIDDLE)
    add_text(slide8, px + cm(1.5), KPI_Y,
             cm(3.0), cm(0.35),
             p['kpi_value'], size=12, bold=True, color=p['accent'],
             font="Roboto Mono", va=MSO_ANCHOR.MIDDLE)
    add_text(slide8, px + cm(4.5), KPI_Y,
             PHASE_W - cm(5.0), cm(0.35),
             p['kpi_desc'], size=8.5, color=MUTE,
             font="Inter", va=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)

# Bottom note
add_rect(slide8, cm(0.5), cm(13.5),
         cm(24.5), cm(0.04), DIVIDER)
add_text(slide8, cm(0.5), cm(13.6),
         cm(24.5), cm(0.4),
         "Стандарт обязателен для всех новинок. Без прохождения 3 фаз новинка не вводится в MML дистрибьюторов.",
         size=9, color=MUTE, font="Inter", align=PP_ALIGN.CENTER)

# Save
prs.save(OUT)
print(f"\nSaved: {OUT}")
print(f"Slide 8 updated: 8-column table with kg instead of ₽")
