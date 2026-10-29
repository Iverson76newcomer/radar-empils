"""Анализ структуры загруженного файла."""
from docx import Document

doc = Document('/home/z/my-project/upload/TITAN_Портфолио_объектов_18.09.2026.docx')

print("=== ПАРАГРАФЫ (последние 30) ===")
for i, p in enumerate(doc.paragraphs[-30:]):
    if p.text.strip():
        print(f"[{i}] {p.text[:150]}")

print(f"\n\n=== ТАБЛИЦЫ ({len(doc.tables)}) ===")
for ti, t in enumerate(doc.tables):
    print(f"\n--- ТАБЛИЦА {ti+1}: {len(t.rows)} строк × {len(t.columns)} колонок ---")
    for ri, r in enumerate(t.rows[:5]):
        cells = [c.text.strip().replace('\n', '|')[:80] for c in r.cells]
        print(f"  R{ri}: {' || '.join(cells)}")

# Найдём последний объект
print(f"\n\n=== ПОСЛЕДНИЙ ОБЪЕКТ ===")
last_text = doc.paragraphs[-1].text if doc.paragraphs else ''
print(f"Last para: {last_text}")
