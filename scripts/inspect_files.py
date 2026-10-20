import openpyxl
from openpyxl import load_workbook

# === Ассортимент ===
print("=" * 100)
print("ФАЙЛ 1: Ассортимент_Олимп_300_SKU.xlsx")
print("=" * 100)

wb1 = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
for sn in wb1.sheetnames:
    ws = wb1[sn]
    print(f"\n--- Sheet: '{sn}' (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    # Print first 15 rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
        print(f"  Row {row_idx}: {row}")

# === Прайс-лист ===
print("\n" + "=" * 100)
print("ФАЙЛ 2: 21.08.2026 РОССИЯ с НДС, руб..xlsx (прайс-лист)")
print("=" * 100)

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)
# Print first sheet headers and few rows
for sn in wb2.sheetnames[:3]:
    ws = wb2[sn]
    print(f"\n--- Sheet: '{sn}' (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
        # Print only first 14 columns
        print(f"  Row {row_idx}: {row[:14]}")
