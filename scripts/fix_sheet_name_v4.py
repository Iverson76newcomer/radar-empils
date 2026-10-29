"""Переименовать длинный лист."""
import openpyxl
from openpyxl import load_workbook
import os

OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v4.xlsx"

wb = load_workbook(OUT_PATH)
print(f"Текущие листы: {wb.sheetnames}")

old_name = "Изменения v4 (AvantGarde→новинки)"
new_name = "Изменения v4"
if old_name in wb.sheetnames:
    wb[old_name].title = new_name

wb.save(OUT_PATH)
print(f"\n✓ Файл пересохранён: {OUT_PATH}")
print(f"  Новые листы: {wb.sheetnames}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
