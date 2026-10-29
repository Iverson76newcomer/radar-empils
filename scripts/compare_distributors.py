"""
Create comparison Excel + chart for 3 distributors:
- Центр красок сентябрь 2025: 633.6 кг, 477 943 ₽, 19 РТТ
- Центр красок сентябрь 2026: 565.2 кг, 465 124 ₽, 27 РТТ
- Палитра сентябрь 2026: 694.3 кг, 511 774 ₽, 37 РТТ
"""
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np

fm.fontManager.addfont('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

# Palette
BG     = '#F5F2EC'
INK    = '#2F3A48'
MUTE   = '#8A95A3'
RED    = '#C00000'
SAND   = '#C9A77C'
BLUE   = '#6B8FA8'
GREEN  = '#8BAA8A'
BORDER_C = '#D9D2C2'
DIV    = '#EFE9DD'
WHITE  = '#FFFFFF'

# Data
distributors = [
    {'name': 'Центр красок', 'period': 'сент 2025', 'kg': 633.6, 'rub': 477943, 'akb': 19, 'color': MUTE},
    {'name': 'Центр красок', 'period': 'сент 2026', 'kg': 565.2, 'rub': 465124, 'akb': 27, 'color': SAND},
    {'name': 'Палитра',     'period': 'сент 2026', 'kg': 694.3, 'rub': 511774, 'akb': 37, 'color': RED},
]

# ── Excel ─────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Сравнение дистрибьюторов"

THIN = Side(border_style="thin", color="D9D2C2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill(start_color="2F3A48", end_color="2F3A48", fill_type="solid")
ALT_FILL = PatternFill(start_color="F5F2EC", end_color="F5F2EC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FFF5E6", end_color="FFF5E6", fill_type="solid")

# Title
ws.merge_cells('A1:E1')
ws['A1'] = "Сравнение дистрибьюторов ОРЕОЛ TITAN · сентябрь"
ws['A1'].font = Font(name="Inter", size=14, bold=True, color="2F3A48")
ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

ws.merge_cells('A2:E2')
ws['A2'] = "Центр красок (2025 vs 2026) + Палитра (2026)"
ws['A2'].font = Font(name="Inter", size=10, italic=True, color="8A95A3")

# Headers (row 4)
headers = ['Дистрибьютор', 'Период', 'Товарооборот (кг)', 'Выручка (₽)', 'АКБ (РТТ)']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col)
    cell.value = h
    cell.font = Font(name="Inter", size=10, bold=True, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = BORDER

# Data rows
for i, d in enumerate(distributors):
    row_num = 5 + i
    fill = ALT_FILL if i % 2 == 0 else WHITE_FILL
    # Highlight Палитра row
    if d['name'] == 'Палитра':
        fill = HIGHLIGHT_FILL

    values = [d['name'], d['period'], d['kg'], d['rub'], d['akb']]
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=col)
        cell.value = v
        cell.border = BORDER
        cell.fill = fill
        if col == 1:  # Дистрибьютор
            cell.font = Font(name="Inter", size=10, bold=True, color="2F3A48")
            cell.alignment = Alignment(horizontal='left', vertical='center')
        elif col == 2:  # Период
            cell.font = Font(name="Inter", size=10, color="8A95A3")
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif col == 3:  # кг
            cell.font = Font(name="Roboto Mono", size=11, bold=True, color="C00000")
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '0.0 "кг"'
        elif col == 4:  # ₽
            cell.font = Font(name="Roboto Mono", size=11, bold=True, color="2F3A48")
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "₽"'
        elif col == 5:  # АКБ
            cell.font = Font(name="Roboto Mono", size=11, bold=True, color="6B8FA8")
            cell.alignment = Alignment(horizontal='center', vertical='center')

# Row heights
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 16
ws.row_dimensions[4].height = 32
for r in range(5, 8):
    ws.row_dimensions[r].height = 28

# Column widths
widths = [18, 14, 22, 20, 14]
for col, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = w

ws.freeze_panes = 'A5'

# ── Analysis section ───────────────────────────────────────────────────
ws.merge_cells('A9:E9')
ws['A9'] = "Аналитика"
ws['A9'].font = Font(name="Inter", size=12, bold=True, color="C00000")

# Динамика Центр красок 2025→2026
ws['A10'] = "Центр красок 2025→2026:"
ws['A10'].font = Font(name="Inter", size=10, bold=True, color="2F3A48")
ws['B10'] = f"кг: {distributors[0]['kg']:.1f} → {distributors[1]['kg']:.1f} ({(distributors[1]['kg']/distributors[0]['kg']-1)*100:+.1f}%)"
ws['B10'].font = Font(name="Roboto Mono", size=10, color="2F3A48")
ws.merge_cells('B10:E10')

ws['A11'] = ""
ws['B11'] = f"₽: {distributors[0]['rub']:,} → {distributors[1]['rub']:,} ({(distributors[1]['rub']/distributors[0]['rub']-1)*100:+.1f}%)".replace(',', ' ')
ws['B11'].font = Font(name="Roboto Mono", size=10, color="2F3A48")
ws.merge_cells('B11:E11')

ws['A12'] = ""
ws['B12'] = f"АКБ: {distributors[0]['akb']} → {distributors[1]['akb']} (+{distributors[1]['akb']-distributors[0]['akb']} РТТ, +{(distributors[1]['akb']/distributors[0]['akb']-1)*100:.0f}%)"
ws['B12'].font = Font(name="Roboto Mono", size=10, color="2F3A48")
ws.merge_cells('B12:E12')

# Палитра vs Центр красок 2026
ws['A14'] = "Палитра vs Центр красок 2026:"
ws['A14'].font = Font(name="Inter", size=10, bold=True, color="2F3A48")
ws['B14'] = f"кг: {distributors[2]['kg']:.1f} vs {distributors[1]['kg']:.1f} (+{distributors[2]['kg']-distributors[1]['kg']:.1f} кг, +{(distributors[2]['kg']/distributors[1]['kg']-1)*100:.1f}%)"
ws['B14'].font = Font(name="Roboto Mono", size=10, color="C00000")
ws.merge_cells('B14:E14')

ws['A15'] = ""
ws['B15'] = f"₽: {distributors[2]['rub']:,} vs {distributors[1]['rub']:,} (+{distributors[2]['rub']-distributors[1]['rub']:,} ₽, +{(distributors[2]['rub']/distributors[1]['rub']-1)*100:.1f}%)".replace(',', ' ')
ws['B15'].font = Font(name="Roboto Mono", size=10, color="C00000")
ws.merge_cells('B15:E15')

ws['A16'] = ""
ws['B16'] = f"АКБ: {distributors[2]['akb']} vs {distributors[1]['akb']} (+{distributors[2]['akb']-distributors[1]['akb']} РТТ, +{(distributors[2]['akb']/distributors[1]['akb']-1)*100:.0f}%)"
ws['B16'].font = Font(name="Roboto Mono", size=10, color="C00000")
ws.merge_cells('B16:E16')

# Средний чек на РТТ
ws['A18'] = "Средний чек на РТТ (₽/РТТ):"
ws['A18'].font = Font(name="Inter", size=10, bold=True, color="2F3A48")
for i, d in enumerate(distributors):
    avg = d['rub'] / d['akb']
    ws.cell(row=19+i, column=1).value = f"  {d['name']} {d['period']}"
    ws.cell(row=19+i, column=1).font = Font(name="Inter", size=9, color="8A95A3")
    ws.cell(row=19+i, column=2).value = f"{avg:,.0f} ₽/РТТ".replace(',', ' ')
    ws.cell(row=19+i, column=2).font = Font(name="Roboto Mono", size=10, bold=True, color="2F3A48")

# Средний чек на кг
ws['A23'] = "Средняя цена за кг (₽/кг):"
ws['A23'].font = Font(name="Inter", size=10, bold=True, color="2F3A48")
for i, d in enumerate(distributors):
    avg = d['rub'] / d['kg']
    ws.cell(row=24+i, column=1).value = f"  {d['name']} {d['period']}"
    ws.cell(row=24+i, column=1).font = Font(name="Inter", size=9, color="8A95A3")
    ws.cell(row=24+i, column=2).value = f"{avg:,.0f} ₽/кг".replace(',', ' ')
    ws.cell(row=24+i, column=2).font = Font(name="Roboto Mono", size=10, bold=True, color="2F3A48")

OUT_XLSX = '/home/z/my-project/download/Сравнение_дистрибьюторов_TITAN.xlsx'
wb.save(OUT_XLSX)
print(f"Excel: {OUT_XLSX}")

# ── Chart image ───────────────────────────────────────────────────────
fig, axes = plt.subplots(1, 3, figsize=(16, 6), dpi=200, facecolor=BG, constrained_layout=True)

labels = [f"{d['name']}\n{d['period']}" for d in distributors]
colors = [d['color'] for d in distributors]

# Chart 1: Товарооборот (кг)
ax1 = axes[0]
ax1.set_facecolor(BG)
bars1 = ax1.bar(labels, [d['kg'] for d in distributors], width=0.55, color=colors,
                edgecolor='white', linewidth=1.5, zorder=3)
for bar, val in zip(bars1, [d['kg'] for d in distributors]):
    ax1.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 8,
             f'{val:.1f} кг', ha='center', va='bottom', fontsize=11, fontweight='bold', color=INK)
ax1.set_title('Товарооборот (кг)', fontsize=13, fontweight='bold', color=INK, pad=12)
ax1.set_ylim(0, max(d['kg'] for d in distributors) * 1.25)
ax1.tick_params(axis='x', labelsize=9, colors=INK)
ax1.tick_params(axis='y', labelsize=9, colors=MUTE)
ax1.grid(axis='y', alpha=0.3, color=DIV, linestyle='--', linewidth=0.5)
ax1.set_axisbelow(True)
ax1.spines['top'].set_visible(False)
ax1.spines['right'].set_visible(False)
ax1.spines['left'].set_color(BORDER_C)
ax1.spines['bottom'].set_color(BORDER_C)

# Chart 2: Выручка (₽)
ax2 = axes[1]
ax2.set_facecolor(BG)
bars2 = ax2.bar(labels, [d['rub']/1000 for d in distributors], width=0.55, color=colors,
                edgecolor='white', linewidth=1.5, zorder=3)
for bar, val in zip(bars2, [d['rub'] for d in distributors]):
    ax2.text(bar.get_x() + bar.get_width()/2, bar.get_height()/1000 + 5,
             f'{val/1000:.0f} тыс ₽', ha='center', va='bottom', fontsize=11, fontweight='bold', color=INK)
ax2.set_title('Выручка (тыс ₽)', fontsize=13, fontweight='bold', color=INK, pad=12)
ax2.set_ylim(0, max(d['rub'] for d in distributors)/1000 * 1.25)
ax2.tick_params(axis='x', labelsize=9, colors=INK)
ax2.tick_params(axis='y', labelsize=9, colors=MUTE)
ax2.grid(axis='y', alpha=0.3, color=DIV, linestyle='--', linewidth=0.5)
ax2.set_axisbelow(True)
ax2.spines['top'].set_visible(False)
ax2.spines['right'].set_visible(False)
ax2.spines['left'].set_color(BORDER_C)
ax2.spines['bottom'].set_color(BORDER_C)

# Chart 3: АКБ (РТТ)
ax3 = axes[2]
ax3.set_facecolor(BG)
bars3 = ax3.bar(labels, [d['akb'] for d in distributors], width=0.55, color=colors,
                edgecolor='white', linewidth=1.5, zorder=3)
for bar, val in zip(bars3, [d['akb'] for d in distributors]):
    ax3.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.8,
             f'{val} РТТ', ha='center', va='bottom', fontsize=11, fontweight='bold', color=INK)
ax3.set_title('АКБ (РТТ)', fontsize=13, fontweight='bold', color=INK, pad=12)
ax3.set_ylim(0, max(d['akb'] for d in distributors) * 1.25)
ax3.tick_params(axis='x', labelsize=9, colors=INK)
ax3.tick_params(axis='y', labelsize=9, colors=MUTE)
ax3.grid(axis='y', alpha=0.3, color=DIV, linestyle='--', linewidth=0.5)
ax3.set_axisbelow(True)
ax3.spines['top'].set_visible(False)
ax3.spines['right'].set_visible(False)
ax3.spines['left'].set_color(BORDER_C)
ax3.spines['bottom'].set_color(BORDER_C)

# Suptitle
fig.suptitle('Сравнение дистрибьюторов ОРЕОЛ TITAN · сентябрь',
             fontsize=15, fontweight='bold', color=INK, y=1.02)

OUT_PNG = '/home/z/my-project/download/Сравнение_дистрибьюторов_TITAN.png'
fig.savefig(OUT_PNG, dpi=200, facecolor=BG, bbox_inches='tight', pad_inches=0.3)
print(f"Chart: {OUT_PNG}")
plt.close(fig)
