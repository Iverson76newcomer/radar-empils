"""Чтение обоих файлов: стандарта вывода новинки + акции TITAN ФБ."""
from docx import Document

print("=" * 80)
print("ФАЙЛ 1: Эмпилс_Стандарт_вывода_новинки.docx")
print("=" * 80)
doc1 = Document('/home/z/my-project/upload/Эмпилс_Стандарт_вывода_новинки.docx')
for i, p in enumerate(doc1.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text[:200]}")
print(f"\nТАБЛИЦ: {len(doc1.tables)}")
for ti, t in enumerate(doc1.tables):
    print(f"\n--- Т{ti+1}: {len(t.rows)}×{len(t.columns)} ---")
    for ri, r in enumerate(t.rows[:3]):
        cells = [c.text.strip()[:50] for c in r.cells]
        print(f"  R{ri}: {' | '.join(cells)}")

print("\n\n" + "=" * 80)
print("ФАЙЛ 2: Формат акции TITAN ФБ для ЭТП.docx")
print("=" * 80)
doc2 = Document('/tmp/Формат акции TITAN ФБ для ЭТП.docx')
for i, p in enumerate(doc2.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text[:250]}")
print(f"\nТАБЛИЦ: {len(doc2.tables)}")
for ti, t in enumerate(doc2.tables):
    print(f"\n--- Т{ti+1}: {len(t.rows)}×{len(t.columns)} ---")
    for ri, r in enumerate(t.rows):
        cells = [c.text.strip()[:60] for c in r.cells]
        print(f"  R{ri}: {' | '.join(cells)}")
