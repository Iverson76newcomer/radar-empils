import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

# === Read price list (build dict) ===
wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# === Read assortment ===
wb1 = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
ws1 = wb1['300 SKU — детально']
assortment_items = []
for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, values_only=True):
    if row[0] is None:
        continue
    assortment_items.append({
        'idx': row[0],
        'brand': row[4],
        'name': row[6],
    })

# === Sheet mapping for brands ===
brand_to_sheets = {
    'AvantGarde': ['AvantGarde'],
    'NEWTONE': ['NEWTONE'],
    'PROFIWOOD': ['PROFIWOOD'],
    'ОРЕОЛ TITAN': ['ОРЕОЛ TITAN', 'ОРЕОЛ TITAN(л)'],
    'Ореол': ['Ореол'],
    'Ореол Дисконт': ['Ореол Дисконт'],
    'Простокрашено! БАУ': ['Простокрашено! БАУ'],
    'Расцвет': ['Расцвет'],
}

# === Normalization ===
def norm(s):
    if s is None:
        return ''
    s = str(s).lower().strip()
    s = re.sub(r'\s+', ' ', s)
    return s

# === Try matching ===
matched_count = 0
unmatched = []
multiple = []
for item in assortment_items:
    brand = item['brand']
    name = item['name']
    name_n = norm(name)
    sheets = brand_to_sheets.get(brand, [])
    
    candidates = []
    for sn in sheets:
        for p in price_dict[sn]:
            pn = norm(p['name'])
            if pn == name_n:
                candidates.append((sn, p, 'exact'))
            elif pn and name_n and (pn in name_n or name_n in pn):
                candidates.append((sn, p, 'substring'))
    
    if len(candidates) == 1:
        matched_count += 1
    elif len(candidates) > 1:
        multiple.append((item, candidates))
    else:
        unmatched.append(item)

print(f"Точное/уникальное сопоставление: {matched_count}")
print(f"Множественные совпадения: {len(multiple)}")
print(f"Не сопоставлено: {len(unmatched)}")

print("\n=== Примеры unsopostavlenno ===")
for item in unmatched[:15]:
    print(f"  [{item['idx']}] {item['brand']} | {item['name']}")

print("\n=== Примеры множественных совпадений ===")
for item, cands in multiple[:5]:
    print(f"\n[{item['idx']}] {item['brand']} | {item['name']}")
    for sn, p, mt in cands[:5]:
        print(f"    → {sn}: '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Check ОРЕОЛ TITAN sheet structure
print("\n=== ОРЕОЛ TITAN — sample names ===")
for p in price_dict['ОРЕОЛ TITAN'][:5]:
    print(f"  {p}")
print("\n=== ОРЕОЛ TITAN(л) — sample names ===")
for p in price_dict['ОРЕОЛ TITAN(л)'][:5]:
    print(f"  {p}")
