"""Финальная проверка v4."""
import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v4.xlsx", data_only=True)

print("=" * 80)
print("ПРОВЕРКА v4")
print("=" * 80)

# === Лист 300 SKU с ценами ===
ws = wb["300 SKU с ценами"]
print(f"\n--- Лист '300 SKU с ценами': {ws.max_row - 4} SKU ---")

# Проверим первые 5 строк (теперь это новинки)
print("\n=== Первые 5 SKU (бывшие AvantGarde, теперь новинки) ===")
for r in range(5, 10):
    idx = ws.cell(row=r, column=1).value
    brand = ws.cell(row=r, column=5).value
    name = ws.cell(row=r, column=7).value
    grp_code = ws.cell(row=r, column=8).value
    price = ws.cell(row=r, column=10).value
    base = ws.cell(row=r, column=11).value
    fasovka = ws.cell(row=r, column=12).value
    sheet = ws.cell(row=r, column=13).value
    status = ws.cell(row=r, column=14).value
    print(f"\n  #{idx} | {brand} | гр.код={grp_code}")
    print(f"  Имя: {name}")
    print(f"  Цена: отгруз={price} | базовая={base} | фасовка={fasovka} | лист={sheet}")
    print(f"  Статус: {status}")

# Подсчёт итогов
total = 0
has_price = 0
no_price = 0
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=1).value is not None:
        total += 1
        if ws.cell(row=r, column=10).value is not None:
            has_price += 1
        else:
            no_price += 1

print(f"\n=== ИТОГ ===")
print(f"Всего SKU: {total}")
print(f"С ценой: {has_price}")
print(f"Без цены: {no_price}")

# === Проверка матриц ===
ws_sv = wb["Сводка"]
print("\n--- Лист 'Сводка' ---")
print("\nМАТРИЦА Основа × Сфера:")
for r in range(4, 8):
    vals = [ws_sv.cell(row=r, column=c).value for c in range(1, 9)]
    print(f"  Row {r}: {vals}")

print("\nМАТРИЦА СФЕРА × БРЕНД:")
for r in range(10, 17):
    vals = [ws_sv.cell(row=r, column=c).value for c in range(1, 14)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

print("\nДЕТАЛЬНАЯ КЛАССИФИКАЦИЯ ПО НАЗНАЧЕНИЮ:")
for r in range(19, 36):
    vals = [ws_sv.cell(row=r, column=c).value for c in range(1, 8)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

# === Лист Изменения v4 ===
print("\n\n--- Лист 'Изменения v4' ---")
ws_log = wb["Изменения v4"]
print(f"Строк: {ws_log.max_row}")
for r in range(4, ws_log.max_row + 1):
    vals = [ws_log.cell(row=r, column=c).value for c in range(1, 9)]
    if any(v is not None for v in vals):
        action = vals[1] or ''
        idx_val = vals[0] if vals[0] else ''
        old_name = (vals[2] or '')[:40] if vals[2] else ''
        new_name = (vals[3] or '')[:50] if vals[3] else ''
        price = vals[4] if vals[4] else ''
        fas = vals[5] if vals[5] else ''
        vtor = vals[6] if vals[6] else ''
        print(f"  Row {r}: #{idx_val} | {action} | old='{old_name}' | new='{new_name}' | цена={price} | фас={fas} | втор={vtor}")

# Подсчёт распределения по брендам
ws_det = wb["300 SKU — детально"]
from collections import Counter
brand_counts = Counter()
for r in range(4, ws_det.max_row + 1):
    if ws_det.cell(row=r, column=1).value is not None:
        brand_counts[ws_det.cell(row=r, column=5).value] += 1

print(f"\n=== Распределение по брендам (после замены) ===")
for b, c in sorted(brand_counts.items(), key=lambda x: -x[1]):
    print(f"  {b}: {c}")
