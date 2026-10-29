"""Исправить матрицу СФЕРА × БРЕНД - удалить лишний столбец ИТОГО."""
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter
import os

OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v5.xlsx"
wb = load_workbook(OUT_PATH)

ws_sv = wb["Сводка"]
ws_det = wb["300 SKU — детально"]

# Перечитать ассортимент
items_all = []
for r in range(4, ws_det.max_row + 1):
    if ws_det.cell(row=r, column=1).value is None:
        continue
    items_all.append({
        'sphere': ws_det.cell(row=r, column=2).value,
        'brand': ws_det.cell(row=r, column=5).value,
    })

spheres_order = ['Металл', 'Помещения (интерьер)', 'Дерево', 'Пол', 'Фасад']
brand_counts = Counter(it['brand'] for it in items_all)

# Полный список брендов (10 — без NoName)
original_brands = ['ОРЕОЛ TITAN', 'Ореол', 'Ореол Дисконт', 'NEWTONE', 'AvantGarde',
                    'Расцвет', 'Простокрашено! БАУ', 'PROFIWOOD',
                    'СТРОЙКОМПЛЕКТ', 'EMPILS ЭКОНОМ']
all_brands = original_brands + [b for b in brand_counts if b not in original_brands]
print(f"Бренды в v5: {all_brands}")

# Очистить всю матрицу СФЕРА × БРЕНД (строки 10-16, все колонки до 15)
for r in range(10, 17):
    for c in range(1, 16):  # очистить до 15 колонки включительно
        cell = ws_sv.cell(row=r, column=c)
        cell.value = None
        cell.fill = PatternFill()  # сбросить заливку
        cell.border = Border()  # сбросить границы

# Стили
HEAD_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEAD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_FONT = Font(name="Calibri", size=10)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
TOTAL_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Заголовок строки 10
ws_sv.cell(row=10, column=1, value='СФЕРА / БРЕНД')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=10, column=2 + col_offset, value=brand)
ws_sv.cell(row=10, column=2 + len(all_brands), value='ИТОГО')

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=10, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER

# Матрица
matrix_sphere_brand = {(s, b): 0 for s in spheres_order for b in all_brands}
for it in items_all:
    if it['brand'] in all_brands:
        matrix_sphere_brand[(it['sphere'], it['brand'])] += 1

for row_offset, sphere in enumerate(spheres_order):
    r = 11 + row_offset
    ws_sv.cell(row=r, column=1, value=sphere)
    for col_offset, brand in enumerate(all_brands):
        val = matrix_sphere_brand[(sphere, brand)]
        if val > 0:
            ws_sv.cell(row=r, column=2 + col_offset, value=val)
    total_row = sum(matrix_sphere_brand[(sphere, b)] for b in all_brands)
    ws_sv.cell(row=r, column=2 + len(all_brands), value=total_row)
    
    for c in range(1, 3 + len(all_brands)):
        cell = ws_sv.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c > 1:
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Итого
r = 16
ws_sv.cell(row=r, column=1, value='ИТОГО')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=r, column=2 + col_offset, value=brand_counts[brand])
ws_sv.cell(row=r, column=2 + len(all_brands), value=len(items_all))

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=r, column=c)
    cell.font = Font(name="Calibri", size=10, bold=True)
    cell.fill = TOTAL_FILL
    cell.border = BORDER
    if c > 1:
        cell.alignment = CENTER_ALIGN
    else:
        cell.alignment = Alignment(horizontal="left", vertical="center")

# Ширина колонок (только для нужных колонок)
for col_offset in range(len(all_brands) + 1):
    ws_sv.column_dimensions[get_column_letter(2 + col_offset)].width = 13
ws_sv.column_dimensions[get_column_letter(2 + len(all_brands))].width = 8
ws_sv.column_dimensions['A'].width = 24

# Сбросить ширину лишних колонок (если были в v4)
for c in range(2 + len(all_brands) + 1, 16):
    ws_sv.column_dimensions[get_column_letter(c)].width = 8.43  # default Excel width

wb.save(OUT_PATH)
print(f"✓ Файл пересохранён: {OUT_PATH}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")

# Проверить матрицу
print("\n=== Матрица СФЕРА × БРЕНД после исправления ===")
for r in range(10, 17):
    vals = [ws_sv.cell(row=r, column=c).value for c in range(1, 14)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")
