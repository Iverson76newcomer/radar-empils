#!/usr/bin/env python3
"""Compact partnership program for Центр Красок — half the size, no detailed timings."""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_PATH = '/home/z/my-project/download/Программа_Центр_Красок_Краснодар_2026_сжатая.docx'

P = {
    'primary': RGBColor(0xC6, 0x28, 0x28),
    'dark': RGBColor(0x1A, 0x1F, 0x36),
    'body': RGBColor(0x2D, 0x31, 0x42),
    'secondary': RGBColor(0x5A, 0x60, 0x80),
    'accent': RGBColor(0x16, 0xA3, 0x4A),
}


def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color='CCCCCC', size='4'):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), size)
        border.set(qn('w:color'), color)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_heading(doc, text, level=1, color=None):
    if color is None:
        color = P['primary'] if level == 1 else P['dark']
    h = doc.add_heading('', level=level)
    run = h.add_run(text)
    run.font.color.rgb = color
    run.font.name = 'Calibri'


def add_para(doc, text, size=11, bold=False, italic=False, color=None, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Calibri'


def make_table(doc, headers, rows, col_widths=None, header_color='C62828'):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if col_widths:
        for i, width in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Cm(width)
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = header_cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = 'Calibri'
        set_cell_bg(cell, header_color)
        set_cell_borders(cell)
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_text in enumerate(row_data):
            cell = row_cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(cell_text))
            run.font.size = Pt(10)
            run.font.name = 'Calibri'
            set_cell_borders(cell)
            if r_idx % 2 == 0:
                set_cell_bg(cell, 'FAF5F5')
    return table


# ====== Build Document ======
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.2

section = doc.sections[0]
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)

# ===== COMPACT COVER =====
cover_table = doc.add_table(rows=1, cols=1)
cover_cell = cover_table.rows[0].cells[0]
set_cell_bg(cover_cell, 'C62828')
cover_cell.text = ''
p = cover_cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(10)
run = p.add_run('Empils • TITAN')
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

p2 = cover_cell.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(30)
run = p2.add_run('Программа партнёрского сотрудничества')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run('Дистрибьютор: «Центр Красок», Краснодар (ЮФО)')
run.font.size = Pt(14)
run.font.bold = True
run.font.color.rgb = P['dark']

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Период: 06.07.2026 — 31.12.2026')
run.font.size = Pt(14)
run.font.color.rgb = P['primary']

# Key conditions in one compact block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run('✓ Бонус 8% ежемесячно + 3% полугодовой  ✓ Отсрочка 60 дней  ✓ Выездное обучение  ✓ Экскурсия на завод  ✓ Выкрасы TITAN')
run.font.size = Pt(10)
run.font.bold = True
run.font.color.rgb = P['accent']

doc.add_page_break()

# ===== 1. ОБЩИЕ ПОЛОЖЕНИЯ =====
add_heading(doc, '1. Общие положения', level=1)

add_para(doc,
    'Программа определяет условия сотрудничества между АО «Эмпилс» (бренд TITAN) и '
    'дистрибьютором «Центр Красок» (Краснодар) на территории ЮФО. Программа разработана '
    'на основе протокола встречи от 01.07.2026 и действует с 6 июля по 31 декабря 2026 года.',
    indent=True)

info_data = [
    ['Производитель', 'АО «Эмпилс» (TITAN)'],
    ['Дистрибьютор', '«Центр Красок», Краснодар'],
    ['Территория', 'Южный федеральный округ (ЮФО)'],
    ['Период действия', '06.07.2026 — 31.12.2026'],
    ['Основание', 'Протокол встречи от 01.07.2026'],
]
make_table(doc, ['Параметр', 'Значение'], info_data, col_widths=[5, 12])

# ===== 2. ФИНАНСОВЫЕ УСЛОВИЯ =====
add_heading(doc, '2. Финансовые условия', level=1)

bonus_data = [
    ['Ежемесячный бонус', '8% от объёма закупок', 'Безоговорочный, выплата в течение 10 раб. дней после расчётного месяца'],
    ['Полугодовой бонус', '3% от объёма закупок за полугодие', 'Выплата в течение 20 раб. дней после окончания полугодия'],
    ['Отсрочка платежа', '60 календарных дней', 'При соблюдении графика платежей по предыдущим поставкам'],
]
make_table(doc, ['Условие', 'Размер', 'Порядок'], bonus_data, col_widths=[3.5, 4, 9.5])

add_para(doc, 'Пример: при закупке 1 млн руб./мес. — ежемесячный бонус 80 000 руб., '
              'полугодовой бонус 180 000 руб. Итого за полугодие: 660 000 руб.',
         italic=True, color=P['secondary'])

# ===== 3. ВЫЕЗДНОЕ ОБУЧЕНИЕ =====
add_heading(doc, '3. Выездное обучение', level=1)

add_para(doc,
    '«Эмпилс» проводит выездное обучение торговой и обучающей команды дистрибьютора. '
    'Дата и формат согласовываются сторонами. Обучение включает модули: о компании «Эмпилс», '
    'о продуктах TITAN (с упором на TITAN 3в1), методика ТСО (совокупная стоимость владения) '
    'с примерами для розничных точек и B2B-клиентов, практикум с выкрасами, работа с возражениями, '
    'цифровые инструменты и тестирование с выдачей сертификата.',
    indent=True)

add_para(doc, 'Условия проведения:', bold=True, color=P['primary'])
add_bullet(doc, 'Аудитория: торговые представители, супервайзеры, тренеры (15–25 человек)')
add_bullet(doc, 'Формат: офлайн на территории дистрибьютора, возможно разделение на потоки')
add_bullet(doc, 'Материалы: шпаргалка продавца, сравнительная таблица, калькулятор ТСО, сертификаты — за счёт «Эмпилс»')
add_bullet(doc, 'Командировочные тренера — за счёт «Эмпилс»; аренда помещения — за счёт дистрибьютора')

# ===== 4. ЭКСКУРСИЯ НА ЗАВОД =====
add_heading(doc, '4. Экскурсия на завод «Эмпилс»', level=1)

add_para(doc,
    '«Эмпилс» приглашает делегацию «Центр Красок» (до 10 человек) на экскурсию на завод '
    'в Ростов-на-Дону. Программа включает: экскурсию по производству и лаборатории, '
    'дегустацию продуктов (TITAN 3в1, Avangarde, Newton), презентацию новинок, круглый стол '
    'с руководством и вручение сувениров. Дата согласовывается сторонами (предпочтительно '
    'сентябрь—октябрь 2026). Программа экскурсии, обед и сувениры — за счёт «Эмпилс»; '
    'проезд и проживание — по согласованию (50/50).', indent=True)

# ===== 5. ВЫКРАСЫ TITAN =====
add_heading(doc, '5. Обеспечение выкрасами по TITAN', level=1)

add_para(doc,
    '«Эмпилс» бесплатно обеспечивает дистрибьютора образцами покрытий (выкрасами) по TITAN '
    'для демонстрации покупателям и на обучающих мероприятиях. Базовый набор включает образцы '
    '5 цветов RAL (чёрный, белый, красно-коричневый, серый, зелёный), сравнительные образцы '
    'TITAN vs ПФ-115 и TITAN vs Hammerite, а также образец «до/после через год» для демонстрации '
    'долговечности.', indent=True)

add_para(doc, 'Условия:', bold=True, color=P['primary'])
add_bullet(doc, 'Базовый набор — 1 комплект на каждый филиал; дополнительные — по запросу (до 5 на филиал)')
add_bullet(doc, 'Первая поставка — в течение 14 дней после подписания программы')
add_bullet(doc, 'Стоимость и доставка — за счёт «Эмпилс»')

# ===== 6. КАЛЕНДАРНЫЙ ПЛАН =====
add_heading(doc, '6. Календарный план', level=1)

calendar_data = [
    ['06.07.2026', 'Старт программы'],
    ['До 20.07.2026', 'Первая поставка выкрасов'],
    ['Июль—Авг 2026', 'Выездное обучение торговой команды'],
    ['Ежемесячно', 'Выплата ежемесячного бонуса (8%)'],
    ['Сент—Окт 2026', 'Экскурсия делегации на завод «Эмпилс»'],
    ['31.12.2026', 'Окончание программы, расчёт полугодового бонуса (3%)'],
    ['Январь 2027', 'Анализ результатов, решение о продлении'],
]
make_table(doc, ['Срок', 'Мероприятие'], calendar_data, col_widths=[4, 13])

# ===== 7. ОЖИДАЕМЫЕ РЕЗУЛЬТАТЫ =====
add_heading(doc, '7. Ожидаемые результаты', level=1)

results_data = [
    ['Рост продаж TITAN', '+20–30% к первому полугодию 2026'],
    ['Обучение команды', '≥20 сотрудников прошли обучение и тестирование'],
    ['Применение ТСО', '≥80% торговой команды используют методику'],
    ['Восстановление ассортимента', 'Возврат сокращённых позиций + расширение'],
    ['Экскурсия на завод', 'Делегация 8–10 человек посетила производство'],
]
make_table(doc, ['Показатель', 'Целевое значение'], results_data, col_widths=[7, 9])

# ===== 8. КООРДИНАЦИЯ =====
add_heading(doc, '8. Координация', level=1)

add_para(doc,
    'Стороны назначают ответственных сотрудников для реализации программы. '
    'Ежемесячные созвоны (онлайн) для контроля выполнения, ежеквартальные очные встречи. '
    'Оперативные вопросы решаются в течение 1–2 рабочих дней, эскалация — на уровне '
    'руководителей подразделений.', indent=True)

# ===== 9. ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ =====
add_heading(doc, '9. Заключительные положения', level=1)

add_bullet(doc, 'Программа вступает в силу 06.07.2026 и действует до 31.12.2026')
add_bullet(doc, 'По истечении срока стороны проводят анализ и решают вопрос о продлении')
add_bullet(doc, 'Изменения оформляются письменно и подписываются обеими сторонами')
add_bullet(doc, 'Программа не является публичной офертой')

# Signatures
doc.add_paragraph()
sig_table = doc.add_table(rows=2, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, header in enumerate(['АО «Эмпилс»', '«Центр Красок»']):
    cell = sig_table.rows[0].cells[i]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(header)
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = P['primary']
    set_cell_bg(cell, 'F5F5F5')
    set_cell_borders(cell)
    cell2 = sig_table.rows[1].cells[i]
    cell2.text = ''
    p2 = cell2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(25)
    run = p2.add_run('____________ / ____________\nМ.П.')
    run.font.size = Pt(10)
    run.font.color.rgb = P['secondary']
    set_cell_borders(cell2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(15)
run = p.add_run('Дата: «____» ____________ 2026 г.')
run.font.size = Pt(11)
run.font.color.rgb = P['body']

# Save
doc.save(OUT_PATH)
print(f'Document saved: {OUT_PATH}')
print(f'Size: {os.path.getsize(OUT_PATH)} bytes')
