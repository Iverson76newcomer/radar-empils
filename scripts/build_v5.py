"""
Финальная сборка v5:
1. Заменить NoName SKU (idx 296) на топ-1 новинку PROFIWOOD серый туман
2. Обновить обе вкладки: '300 SKU — детально' и '300 SKU с ценами'
3. Пересчитать матрицы на 'Сводка'
4. Добавить лист с журналом изменений и пояснением про NoName
"""
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict, Counter
import shutil
import os

SRC_PATH = "/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал_v4.xlsx"
OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v5.xlsx"

# Копируем оригинал
shutil.copy(SRC_PATH, OUT_PATH)
wb = load_workbook(OUT_PATH)

# === Стили ===
HEAD_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEAD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
BODY_FONT = Font(name="Calibri", size=10)
BODY_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
NUM_ALIGN = Alignment(horizontal="right", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
PRICE_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E79")
NEW_ROW_FILL = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
REMOVED_FILL = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# === Замена ===
OLD_ITEM = {
    'idx': 296,
    'sphere': 'Дерево',
    'purpose': 'Защита и декор дерева',
    'base': 'Водно-дисперсионная',
    'brand': 'NoName',
    'type': 'Краски',
    'name': 'Краска для садовых деревьев водно-дисперсионная полиакриловая 3 кг',
    'grp_code': 63180,
    'abc': 'А',
    'отгрузочная': 221.27, 'базовая': 276.59, 'фасовка': 3,
    'лист': 'NoName',
    'reason': 'В прайс-листе Эмпилс эта SKU вынесена в отдельный лист "NoName" — продукция завода Эмпилс без указания конкретного бренда (штрихкод 4605365... подтверждает изготовителя Эмпилс). Поскольку требуется, чтобы у всех SKU был чёткий бренд, заменяем на топ-1 новинку PROFIWOOD.',
}

NEW_ITEM = {
    'idx': 296,
    'sphere': 'Дерево',
    'purpose': 'Защита и декор дерева',
    'base': 'Водно-дисперсионная',
    'brand': 'PROFIWOOD',
    'type': 'Декоративные составы',
    'name': 'Антисептик кроющий PROFIWOOD серый туман 0.8 л 0.9 кг',
    'grp_code': 88847,
    'abc': 'А',
    'отгрузочная': 463.60, 'базовая': 579.50, 'фасовка': 0.9,
    'лист': 'PROFIWOOD',
    'reason': 'Топ-1 новинка 2026 года среди непопавших в ассортимент (вторичка 15 157 руб). Та же категория (Дерево / Защита и декор дерева). Чёткий бренд PROFIWOOD. Новая фасовка 0.9 кг, которой не было в ассортименте для "Декоративных составов" PROFIWOOD.',
}

# === 1. Обновить лист "300 SKU — детально" ===
print("=== Обновление листа '300 SKU — детально' ===")
ws_det = wb["300 SKU — детально"]

# Найти строку по idx 296
row_to_update = None
for r in range(4, ws_det.max_row + 1):
    if ws_det.cell(row=r, column=1).value == 296:
        row_to_update = r
        break

if row_to_update:
    # Старые значения для лога
    old_name = ws_det.cell(row=row_to_update, column=7).value
    old_brand = ws_det.cell(row=row_to_update, column=5).value
    print(f"  Старая позиция (строка {row_to_update}): #{296} | {old_brand} | {old_name}")
    
    # Обновить
    ws_det.cell(row=row_to_update, column=1, value=NEW_ITEM['idx'])
    ws_det.cell(row=row_to_update, column=2, value=NEW_ITEM['sphere'])
    ws_det.cell(row=row_to_update, column=3, value=NEW_ITEM['purpose'])
    ws_det.cell(row=row_to_update, column=4, value=NEW_ITEM['base'])
    ws_det.cell(row=row_to_update, column=5, value=NEW_ITEM['brand'])
    ws_det.cell(row=row_to_update, column=6, value=NEW_ITEM['type'])
    ws_det.cell(row=row_to_update, column=7, value=NEW_ITEM['name'])
    ws_det.cell(row=row_to_update, column=8, value=NEW_ITEM['grp_code'])
    ws_det.cell(row=row_to_update, column=9, value=NEW_ITEM['abc'])
    print(f"  Новая позиция: {NEW_ITEM['brand']} | {NEW_ITEM['name']}")

# === 2. Обновить лист "300 SKU с ценами" ===
print("\n=== Обновление листа '300 SKU с ценами' ===")
ws_prices = wb["300 SKU с ценами"]

row_to_update_p = None
for r in range(5, ws_prices.max_row + 1):
    if ws_prices.cell(row=r, column=1).value == 296:
        row_to_update_p = r
        break

if row_to_update_p:
    # Обновить данные
    ws_prices.cell(row=row_to_update_p, column=1, value=NEW_ITEM['idx'])
    ws_prices.cell(row=row_to_update_p, column=2, value=NEW_ITEM['sphere'])
    ws_prices.cell(row=row_to_update_p, column=3, value=NEW_ITEM['purpose'])
    ws_prices.cell(row=row_to_update_p, column=4, value=NEW_ITEM['base'])
    ws_prices.cell(row=row_to_update_p, column=5, value=NEW_ITEM['brand'])
    ws_prices.cell(row=row_to_update_p, column=6, value=NEW_ITEM['type'])
    ws_prices.cell(row=row_to_update_p, column=7, value=NEW_ITEM['name'])
    ws_prices.cell(row=row_to_update_p, column=8, value=NEW_ITEM['grp_code'])
    ws_prices.cell(row=row_to_update_p, column=9, value=NEW_ITEM['abc'])
    ws_prices.cell(row=row_to_update_p, column=10, value=NEW_ITEM['отгрузочная'])
    ws_prices.cell(row=row_to_update_p, column=11, value=NEW_ITEM['базовая'])
    ws_prices.cell(row=row_to_update_p, column=12, value=NEW_ITEM['фасовка'])
    ws_prices.cell(row=row_to_update_p, column=13, value=NEW_ITEM['лист'])
    ws_prices.cell(row=row_to_update_p, column=14, value=f"ЗАМЕНА: NoName → PROFIWOOD (новинка 2026, топ-1)")
    
    # Применить подсветку (новая — зелёным, но не ярко, чтобы отличалась от новинок v4)
    for c in range(1, 15):
        cell = ws_prices.cell(row=row_to_update_p, column=c)
        cell.fill = NEW_ROW_FILL
    
    # Формат цен
    ws_prices.cell(row=row_to_update_p, column=10).number_format = '# ##0.00'
    ws_prices.cell(row=row_to_update_p, column=10).font = PRICE_FONT
    ws_prices.cell(row=row_to_update_p, column=11).number_format = '# ##0.00'
    ws_prices.cell(row=row_to_update_p, column=12).number_format = '0.##'
    
    print(f"  ✓ Заменена строка {row_to_update_p} на листе '300 SKU с ценами'")

# Обновить заголовок (строка 2)
ws_prices["A2"] = ("Источник цен: прайс-лист «16.10.2026 РОССИЯ с НДС, руб.» | "
                    "v5: устранена SKU с брендом NoName → заменена на топ-1 новинку PROFIWOOD серый туман | "
                    "ИТОГО: 300 SKU, у всех чёткий бренд")

# === 3. Пересчёт матриц на листе "Сводка" ===
print("\n=== Пересчёт матриц 'Сводка' ===")
ws_sv = wb["Сводка"]

# Прочитать обновлённый детальный лист
items_all = []
for r in range(4, ws_det.max_row + 1):
    if ws_det.cell(row=r, column=1).value is None:
        continue
    items_all.append({
        'idx': ws_det.cell(row=r, column=1).value,
        'sphere': ws_det.cell(row=r, column=2).value,
        'purpose': ws_det.cell(row=r, column=3).value,
        'base': ws_det.cell(row=r, column=4).value,
        'brand': ws_det.cell(row=r, column=5).value,
        'type': ws_det.cell(row=r, column=6).value,
    })

print(f"  Всего SKU: {len(items_all)}")

# Проверить, что NoName больше нет
brands = Counter(it['brand'] for it in items_all)
print(f"  Бренды: {dict(brands)}")
print(f"  NoName в списке: {'NoName' in brands}")

# === Матрица 1: Основа × Сфера ===
spheres_order = ['Металл', 'Помещения (интерьер)', 'Дерево', 'Пол', 'Фасад']
bases_order = ['Органическая', 'Водно-дисперсионная']

matrix_base_sphere = {(b, s): 0 for b in bases_order for s in spheres_order}
for it in items_all:
    key = (it['base'], it['sphere'])
    if key in matrix_base_sphere:
        matrix_base_sphere[key] += 1

for row_offset, base in enumerate(bases_order):
    r = 5 + row_offset
    ws_sv.cell(row=r, column=1, value=base)
    for col_offset, sphere in enumerate(spheres_order):
        ws_sv.cell(row=r, column=2 + col_offset, value=matrix_base_sphere[(base, sphere)])
    total_row = sum(matrix_base_sphere[(base, s)] for s in spheres_order)
    ws_sv.cell(row=r, column=7, value=total_row)
    pct = total_row / len(items_all)
    ws_sv.cell(row=r, column=8, value=f"{pct*100:.1f}%")

r = 7
ws_sv.cell(row=r, column=1, value='ИТОГО')
for col_offset, sphere in enumerate(spheres_order):
    col_total = sum(matrix_base_sphere[(b, sphere)] for b in bases_order)
    ws_sv.cell(row=r, column=2 + col_offset, value=col_total)
ws_sv.cell(row=r, column=7, value=len(items_all))
ws_sv.cell(row=r, column=8, value='100%')

# === Матрица 2: СФЕРА × БРЕНД ===
brand_counts = Counter(it['brand'] for it in items_all)
# Оригинальные бренды без NoName
original_brands = ['ОРЕОЛ TITAN', 'Ореол', 'Ореол Дисконт', 'NEWTONE', 'AvantGarde',
                    'Расцвет', 'Простокрашено! БАУ', 'PROFIWOOD',
                    'СТРОЙКОМПЛЕКТ', 'EMPILS ЭКОНОМ']  # СТРОЙКОМПЛЕКТ и EMPILS ЭКОНОМ — добавлены в v3
all_brands = original_brands + [b for b in brand_counts if b not in original_brands]

# Очистить и пересоздать
for r in range(10, 17):
    for c in range(1, 15):
        ws_sv.cell(row=r, column=c, value=None)

ws_sv.cell(row=10, column=1, value='СФЕРА / БРЕНД')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=10, column=2 + col_offset, value=brand)
ws_sv.cell(row=10, column=2 + len(all_brands), value='ИТОГО')

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=10, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER

matrix_sphere_brand = {(s, b): 0 for s in spheres_order for b in all_brands}
for it in items_all:
    if it['brand'] in all_brands:
        matrix_sphere_brand[(it['sphere'], it['brand'])] += 1

for row_offset, sphere in enumerate(spheres_order):
    r = 11 + row_offset
    ws_sv.cell(row=r, column=1, value=sphere)
    for col_offset, brand in enumerate(all_brands):
        val = matrix_sphere_brand[(sphere, brand)]
        if val > 0:
            ws_sv.cell(row=r, column=2 + col_offset, value=val)
    total_row = sum(matrix_sphere_brand[(sphere, b)] for b in all_brands)
    ws_sv.cell(row=r, column=2 + len(all_brands), value=total_row)
    
    for c in range(1, 3 + len(all_brands)):
        cell = ws_sv.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c > 1:
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

r = 16
ws_sv.cell(row=r, column=1, value='ИТОГО')
for col_offset, brand in enumerate(all_brands):
    ws_sv.cell(row=r, column=2 + col_offset, value=brand_counts[brand])
ws_sv.cell(row=r, column=2 + len(all_brands), value=len(items_all))

for c in range(1, 3 + len(all_brands)):
    cell = ws_sv.cell(row=r, column=c)
    cell.font = Font(name="Calibri", size=10, bold=True)
    cell.fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    cell.border = BORDER
    if c > 1:
        cell.alignment = CENTER_ALIGN
    else:
        cell.alignment = Alignment(horizontal="left", vertical="center")

# Ширина колонок
for col_offset in range(len(all_brands) + 1):
    ws_sv.column_dimensions[get_column_letter(2 + col_offset)].width = 13
ws_sv.column_dimensions[get_column_letter(2 + len(all_brands))].width = 8
ws_sv.column_dimensions['A'].width = 24

# === Матрица 3: ДЕТАЛЬНАЯ КЛАССИФИКАЦИЯ ПО НАЗНАЧЕНИЮ ===
purposes_order = [
    'Грунт-эмаль по металлу', 'Эмаль по металлу', 'Пол',
    'Стены и потолки (интерьер)', 'Защита и декор дерева',
    'Эмали ВД прочие', 'Лаки прочие', 'Грунтовки', 'Фасад',
    'Прочее', 'Яхтный лак', 'Шпатлёвки', 'Паркетный лак',
    'Молотковая эмаль', 'Потолки', 'Резиновая краска',
]

matrix_purpose_sphere = {(p, s): 0 for p in purposes_order for s in spheres_order}
purpose_counts = Counter()
for it in items_all:
    key = (it['purpose'], it['sphere'])
    if key in matrix_purpose_sphere:
        matrix_purpose_sphere[key] += 1
        purpose_counts[it['purpose']] += 1

purposes_sorted = sorted(purposes_order, key=lambda p: -purpose_counts.get(p, 0))

for r in range(20, 36):
    for c in range(1, 8):
        ws_sv.cell(row=r, column=c, value=None)

ws_sv.cell(row=19, column=1, value='Сфера / Назначение')
sphere_cols = ['Металл', 'Помещения (интерьер)', 'Дерево', 'Пол', 'Фасад']
for col_offset, sp in enumerate(sphere_cols):
    ws_sv.cell(row=19, column=2 + col_offset, value=sp)
ws_sv.cell(row=19, column=7, value='ИТОГО')

for c in range(1, 8):
    cell = ws_sv.cell(row=19, column=c)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER

for row_offset, purpose in enumerate(purposes_sorted):
    r = 20 + row_offset
    ws_sv.cell(row=r, column=1, value=purpose)
    for col_offset, sp in enumerate(sphere_cols):
        val = matrix_purpose_sphere[(purpose, sp)]
        if val > 0:
            ws_sv.cell(row=r, column=2 + col_offset, value=val)
    total_row = purpose_counts.get(purpose, 0)
    if total_row > 0:
        ws_sv.cell(row=r, column=7, value=total_row)
    
    for c in range(1, 8):
        cell = ws_sv.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if c > 1:
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

print(f"  ✓ Матрицы обновлены (NoName удалён из шапки)")

# === 4. Создать лист "Замена NoName (v5)" ===
print("\n=== Создание листа 'Замена NoName (v5)' ===")
sheet_name = "Замена NoName v5"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws_log = wb.create_sheet(sheet_name)

# Заголовок
ws_log.merge_cells("A1:G1")
ws_log["A1"] = "Журнал изменений v5: устранение SKU с брендом NoName"
ws_log["A1"].font = TITLE_FONT
ws_log.row_dimensions[1].height = 28

ws_log.merge_cells("A2:G2")
ws_log["A2"] = ("В файле v4 в колонке «Бренд» была 1 SKU с значением «NoName». "
                "Пояснение: «NoName» — это не ошибка классификации, а реальный лист в прайс-листе Эмпилс "
                "(16.10.2026), куда вынесена продукция завода без указания конкретного бренда. "
                "Штрихкод 4605365... подтверждает изготовителя ЗАО «Эмпилс» (тот же префикс, что у Ореол и AvantGarde). "
                "Поскольку пользователь запросил, чтобы у всех SKU был чёткий бренд, заменили эту SKU "
                "на топ-1 новинку 2026 года — PROFIWOOD серый туман.")
ws_log["A2"].font = SUBTITLE_FONT
ws_log["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws_log.row_dimensions[2].height = 80

# Шапка
log_headers = [
    ("Параметр", 28),
    ("Старая позиция (NoName)", 40),
    ("Новая позиция (PROFIWOOD)", 40),
    ("Комментарий", 50),
]
for col_idx, (title, width) in enumerate(log_headers, 1):
    cell = ws_log.cell(row=4, column=col_idx, value=title)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER
    ws_log.column_dimensions[get_column_letter(col_idx)].width = width
ws_log.row_dimensions[4].height = 38

# Данные
rows_data = [
    ("idx", "296", "296", "Индекс SKU сохранён"),
    ("Сфера", OLD_ITEM['sphere'], NEW_ITEM['sphere'], "Без изменений (Дерево)"),
    ("Назначение", OLD_ITEM['purpose'], NEW_ITEM['purpose'], "Без изменений (Защита и декор дерева)"),
    ("Основа", OLD_ITEM['base'], NEW_ITEM['base'], "Без изменений (Водно-дисперсионная)"),
    ("Бренд", "NoName (лист в прайс-листе)", "PROFIWOOD", "✓ Заменён на чёткий бренд"),
    ("Тип продукции", OLD_ITEM['type'], NEW_ITEM['type'], "Старый — Краски, новый — Декоративные составы"),
    ("Наименование", OLD_ITEM['name'], NEW_ITEM['name'], "Полная замена"),
    ("Гр. код", str(OLD_ITEM['grp_code']), str(NEW_ITEM['grp_code']), "Старый — 63180, новый — 88847"),
    ("ABC", OLD_ITEM['abc'], NEW_ITEM['abc'], "Без изменений (А)"),
    ("Отгрузочная цена, руб", f"{OLD_ITEM['отгрузочная']:.2f}", f"{NEW_ITEM['отгрузочная']:.2f}", "Старая — 221,27, новая — 463,60"),
    ("Базовая цена, руб", f"{OLD_ITEM['базовая']:.2f}", f"{NEW_ITEM['базовая']:.2f}", "Старая — 276,59, новая — 579,50"),
    ("Фасовка, кг", str(OLD_ITEM['фасовка']), str(NEW_ITEM['фасовка']), "Старая — 3 кг, новая — 0.9 кг (новая)"),
    ("Лист прайса", OLD_ITEM['лист'], NEW_ITEM['лист'], "Старый — NoName, новый — PROFIWOOD"),
    ("Вторичка 2026, руб", "— (не новинка)", "15 156,73 руб", "Топ-1 среди непопавших в ассортимент новинок"),
]

for i, (param, old_val, new_val, comment) in enumerate(rows_data):
    r = 5 + i
    ws_log.cell(row=r, column=1, value=param).font = Font(name="Calibri", size=10, bold=True)
    ws_log.cell(row=r, column=1).border = BORDER
    ws_log.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
    
    cell = ws_log.cell(row=r, column=2, value=old_val)
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.fill = REMOVED_FILL
    
    cell = ws_log.cell(row=r, column=3, value=new_val)
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.fill = NEW_ROW_FILL
    
    ws_log.cell(row=r, column=4, value=comment).font = BODY_FONT
    ws_log.cell(row=r, column=4).border = BORDER
    ws_log.cell(row=r, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws_log.row_dimensions[r].height = 32

# Раздел: пояснение про NoName
r = 5 + len(rows_data) + 2
ws_log.merge_cells(f"A{r}:D{r}")
ws_log.cell(row=r, column=1, value="ПОЯСНЕНИЕ ПРО БРЕНД «NoName» В ПРАЙС-ЛИСТЕ ЭМПИЛС").font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
ws_log.cell(row=r, column=1).fill = WARN_FILL

explanation = (
    "В прайс-листе Эмпилс от 16.10.2026 есть отдельный лист с названием «NoName», в котором размещена "
    "продукция завода без указания конкретного бренда. Всего в этом листе 2 SKU:\n\n"
    "1. Краска для садовых деревьев водно-дисперсионная полиакриловая, фасовка 1.5 кг (гр.код 63562)\n"
    "2. Краска для садовых деревьев водно-дисперсионная полиакриловая, фасовка 3 кг (гр.код 63180) ← была в нашем ассортименте\n\n"
    "По штрихкоду 4605365... видно, что изготовитель — ЗАО «Эмпилс» (тот же префикс, что у Ореол, AvantGarde и т.д.). "
    "То есть это собственная продукция Эмпилс, которая либо ещё не получила отдельного бренда, либо намеренно "
    "позиционируется как универсальная «безымянная» краска для садовых деревьев.\n\n"
    "В файле ABC-анализа (лист «Белые пятна», строка 115 — дистрибьютор «Вертикаль») эта SKU тоже отнесена к бренду "
    "«NoName» — то есть это не ошибка нашей классификации, а реальная маркетинговая позиция.\n\n"
    "Поскольку пользователь запросил, чтобы у всех SKU в ассортименте Олимпа был чёткий бренд, мы заменили "
    "«Краску для садовых деревьев» (NoName) на «Антисептик кроющий PROFIWOOD серый туман 0.8 л 0.9 кг» — "
    "топ-1 новинку 2026 года среди непопавших в ассортимент. Это позволяет:\n"
    "  • Сохранить категорию (Дерево / Защита и декор дерева)\n"
    "  • Использовать чёткий бренд PROFIWOOD\n"
    "  • Добавить новую фасовку 0.9 кг (которой не было в ассортименте для «Декоративных составов» PROFIWOOD)\n"
    "  • Опереться на подтверждённый спрос — вторичка 15 157 руб за 9 мес. 2026 г."
)

ws_log.merge_cells(f"A{r+1}:D{r+10}")
ws_log.cell(row=r+1, column=1, value=explanation).font = BODY_FONT
ws_log.cell(row=r+1, column=1).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
ws_log.cell(row=r+1, column=1).fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
for rr in range(r+1, r+11):
    ws_log.row_dimensions[rr].height = 22

ws_log.freeze_panes = "B5"

# === 5. Сохранить ===
wb.properties.creator = "Z.ai"
wb.properties.title = "Ассортимент Олимп 300 SKU — финал v5 (устранён NoName)"
wb.properties.subject = "SKU с брендом NoName заменена на топ-1 новинку PROFIWOOD серый туман"

# Удалить старый лист v4 "Изменения v4" если есть
if "Изменения v4" in wb.sheetnames:
    del wb["Изменения v4"]

print(f"\nЛисты в финальном файле: {wb.sheetnames}")
wb.save(OUT_PATH)
print(f"\n✓ Файл сохранён: {OUT_PATH}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
