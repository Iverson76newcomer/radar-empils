"""Проверка обновлённых Excel и PPTX."""
from openpyxl import load_workbook
from pptx import Presentation

print("=" * 80)
print("EXCEL: проверка вкладки 2 'Рейтинг брендов'")
print("=" * 80)
wb = load_workbook('/home/z/my-project/download/Wordstat_узнаваемость_брендов.xlsx', data_only=True)
ws = wb['2. Рейтинг брендов']
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
    vals = []
    for cell in row:
        v = cell.value if cell.value is not None else ''
        if isinstance(v, (int, float)):
            vals.append(f'{v}')
        else:
            vals.append(str(v))
    line = ' | '.join(vals)
    if line.strip(' |'):
        print(line)

print("\n" + "=" * 80)
print("EXCEL: проверка вкладки 5 'Детализация' (последние строки)")
print("=" * 80)
ws5 = wb['5. Детализация по запросам']
for row in ws5.iter_rows(min_row=ws5.max_row - 10, max_row=ws5.max_row, values_only=False):
    vals = []
    for cell in row:
        v = cell.value if cell.value is not None else ''
        vals.append(str(v))
    line = ' | '.join(vals)
    if line.strip(' |'):
        print(line)

print("\n" + "=" * 80)
print("EXCEL: проверка вкладки 6 'Матрица бренд × запрос'")
print("=" * 80)
ws6 = wb['6. Матрица бренд × запрос']
for row in ws6.iter_rows(min_row=1, max_row=ws6.max_row, values_only=False):
    vals = []
    for cell in row:
        v = cell.value if cell.value is not None else ''
        if isinstance(v, (int, float)):
            vals.append(f'{v}')
        else:
            vals.append(str(v))
    line = ' | '.join(vals)
    if line.strip(' |'):
        print(line)

print("\n" + "=" * 80)
print("PPTX: проверка слайда 3 'Рейтинг узнаваемости брендов грунт-эмалей'")
print("=" * 80)
prs = Presentation('/home/z/my-project/download/Wordstat_полный_анализ_TITAN_ с компанией.pptx')
slide = prs.slides[2]
for shape in slide.shapes:
    if shape.has_table:
        print(f"\nТАБЛИЦА: {len(shape.table.rows)} строк × {len(shape.table.columns)} колонок")
        for r in shape.table.rows:
            cells = [c.text.strip() for c in r.cells]
            print(' | '.join(cells))
