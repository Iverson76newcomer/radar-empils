"""Чтение предыдущего файла с ценами для понимания структуры."""
from docx import Document

doc = Document('/home/z/my-project/upload/TITAN_Актуальные_цены_конкурентов 26.08.2026.docx')

print("=== ПАРАГРАФЫ ===")
for i, p in enumerate(doc.paragraphs[:80]):
    if p.text.strip():
        print(p.text)

print("\n\n=== ТАБЛИЦЫ ===")
for ti, t in enumerate(doc.tables):
    print(f"\n[TABLE {ti+1}] — {len(t.rows)} rows × {len(t.columns)} cols")
    for r in t.rows:
        cells = [c.text.strip().replace('\n', ' | ') for c in r.cells]
        print(' | '.join(cells))
