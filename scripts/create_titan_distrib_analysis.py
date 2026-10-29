#!/usr/bin/env python3
"""
Excel-сводка по анализу дистрибьюторов TITAN:
4 листа:
1. Сводка (3 главные цифры)
2. Группы дистрибьюторов (зелёная/жёлтая/красная)
3. 30 потерянных клиентов (для отработки возврата)
4. 16 новых клиентов 2026 (для тиражирования опыта)
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

SRC = '/home/z/my-project/upload/Титан итого.xlsx'
OUT = '/home/z/my-project/download/TITAN_анализ_дистрибьюторов_для_программы.xlsx'

# Стили
HEADER_FONT = Font(name='Arial', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1E3A5F')
TITLE_FONT = Font(name='Arial', size=16, bold=True, color='1E3A5F')
SUBTITLE_FONT = Font(name='Arial', size=11, bold=True, color='F47920')
NORMAL_FONT = Font(name='Arial', size=10, color='1A1F36')
BOLD_FONT = Font(name='Arial', size=10, bold=True, color='1A1F36')
RED_FONT = Font(name='Arial', size=10, bold=True, color='C62828')
GREEN_FONT = Font(name='Arial', size=10, bold=True, color='16A34A')
ORANGE_FONT = Font(name='Arial', size=10, bold=True, color='F47920')
GRAY_FONT = Font(name='Arial', size=9, color='6B7280')
WHITE_BOLD = Font(name='Arial', size=11, bold=True, color='FFFFFF')

LIGHT_GRAY = PatternFill('solid', fgColor='F3F4F6')
LIGHT_RED = PatternFill('solid', fgColor='FEE2E2')
LIGHT_GREEN = PatternFill('solid', fgColor='DCFCE7')
LIGHT_YELLOW = PatternFill('solid', fgColor='FEF9C3')
LIGHT_ORANGE = PatternFill('solid', fgColor='FFEDD5')
LIGHT_BLUE = PatternFill('solid', fgColor='DBEAFE')

THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)


def style_cell(ws, row, col, value, font=NORMAL_FONT, fill=None, align=CENTER, border=BORDER, num_format=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align
    cell.border = border
    if num_format:
        cell.number_format = num_format
    return cell


# ========================================================
# ЧТЕНИЕ ИСХОДНЫХ ДАННЫХ
# ========================================================
src_wb = openpyxl.load_workbook(SRC, data_only=True)

# --- Лист 1: охват дистрибьюторов по регионам ---
ws_src1 = src_wb['1. Охват дистрибьюторов']
distrib_by_region = defaultdict(list)
current_region = None
for r in range(2, ws_src1.max_row + 1):
    sbu = ws_src1.cell(r, 1).value
    rd = ws_src1.cell(r, 2).value
    rm = ws_src1.cell(r, 3).value
    client = ws_src1.cell(r, 4).value
    zakup = ws_src1.cell(r, 5).value
    if not client:
        continue
    if sbu and sbu.strip():
        current_region = sbu.strip()
    distrib_by_region[current_region].append({
        'name': client, 'zakup': zakup or '', 'rd': rd or '', 'rm': rm or ''
    })

# --- Лист 2: проникновение в розницу ---
ws_src2 = src_wb['2. Проникновение в розницу']
penetration = []
for r in range(3, ws_src2.max_row + 1):
    name = ws_src2.cell(r, 1).value
    if not name:
        continue
    rtt25 = ws_src2.cell(r, 2).value
    rtt26 = ws_src2.cell(r, 3).value
    akb25 = ws_src2.cell(r, 4).value
    akb26 = ws_src2.cell(r, 5).value
    sh25 = ws_src2.cell(r, 6).value
    sh26 = ws_src2.cell(r, 7).value
    dyn = ws_src2.cell(r, 8).value
    penetration.append({
        'name': name, 'rtt25': rtt25, 'rtt26': rtt26,
        'akb25': akb25, 'akb26': akb26,
        'sh25': sh25, 'sh26': sh26, 'dyn': dyn
    })

# --- Лист 3: динамика вторичных продаж ---
ws_src3 = src_wb['3. Динамика вторичных продаж']
dynamics_by_distr = defaultdict(lambda: {'kg25': 0, 'kg26': 0, 'p25': 0, 'p26': 0, 'clients': []})
all_clients = []
for r in range(4, ws_src3.max_row + 1):
    distr = ws_src3.cell(r, 1).value
    client = ws_src3.cell(r, 2).value
    p25 = ws_src3.cell(r, 3).value or 0
    p26 = ws_src3.cell(r, 4).value or 0
    kg25 = ws_src3.cell(r, 5).value or 0
    kg26 = ws_src3.cell(r, 6).value or 0
    if not distr or not isinstance(p25, (int, float)):
        continue
    dynamics_by_distr[distr]['kg25'] += kg25
    dynamics_by_distr[distr]['kg26'] += kg26
    dynamics_by_distr[distr]['p25'] += p25
    dynamics_by_distr[distr]['p26'] += p26
    dynamics_by_distr[distr]['clients'].append({
        'client': client, 'p25': p25, 'p26': p26, 'kg25': kg25, 'kg26': kg26
    })
    all_clients.append({'distr': distr, **{'client': client, 'p25': p25, 'p26': p26, 'kg25': kg25, 'kg26': kg26}})


# ========================================================
# КЛАССИФИКАЦИЯ ДИСТРИБЬЮТОРОВ
# ========================================================
# Объединяем данные листов 1 + 3
distrib_full = []
for region, distrs in distrib_by_region.items():
    for d in distrs:
        name = d['name']
        zakup = d['zakup']
        # ищем в dynamics
        dyn = dynamics_by_distr.get(name, None)
        kg25 = dyn['kg25'] if dyn else 0
        kg26 = dyn['kg26'] if dyn else 0
        p25 = dyn['p25'] if dyn else 0
        p26 = dyn['p26'] if dyn else 0
        # ищем в penetration
        pen = next((p for p in penetration if p['name'] == name), None)
        rtt25 = pen['rtt25'] if pen else None
        rtt26 = pen['rtt26'] if pen else None
        akb25 = pen['akb25'] if pen else None
        akb26 = pen['akb26'] if pen else None
        # группа
        active = zakup == '2025-2026'
        if not active:
            group = '🔴 Не активен'
        elif kg26 > 1000 and kg26 >= kg25 * 0.9:
            group = '🟢 Активный, стабильный/растущий'
        elif kg26 > 100 and kg26 < kg25 * 0.7:
            group = '🟡 Активный, падающий'
        elif kg25 == 0 and kg26 > 50:
            group = '🟢 Активный, новый'
        elif active:
            group = '🟢 Активный, стабильный/растущий' if kg26 >= kg25 else '🟡 Активный, падающий'
        else:
            group = '🔴 Не активен'
        distrib_full.append({
            'region': region, 'name': name, 'rd': d['rd'], 'rm': d['rm'],
            'zakup': zakup, 'kg25': kg25, 'kg26': kg26, 'p25': p25, 'p26': p26,
            'rtt25': rtt25, 'rtt26': rtt26, 'akb25': akb25, 'akb26': akb26,
            'group': group
        })

# Сортировка: сначала зелёные по кг26 убывание, потом жёлтые, потом красные
group_order = {'🟢 Активный, стабильный/растущий': 1, '🟢 Активный, новый': 2, '🟡 Активный, падающий': 3, '🔴 Не активен': 4}
distrib_full.sort(key=lambda x: (group_order.get(x['group'], 5), -(x['kg26'] or 0)))


# ========================================================
# ПОТЕРЯННЫЕ КЛИЕНТЫ (кг25 > 50, кг26 = 0)
# ========================================================
lost_clients = [c for c in all_clients if c['kg25'] > 50 and c['kg26'] == 0]
lost_clients.sort(key=lambda x: -x['kg25'])

# ========================================================
# НОВЫЕ КЛИЕНТЫ 2026 (кг25 = 0, кг26 > 50)
# ========================================================
new_clients = [c for c in all_clients if c['kg25'] == 0 and c['kg26'] > 50]
new_clients.sort(key=lambda x: -x['kg26'])


# ========================================================
# СОЗДАНИЕ EXCEL
# ========================================================
wb = openpyxl.Workbook()

# ========================================================
# ЛИСТ 1: СВОДКА
# ========================================================
ws1 = wb.active
ws1.title = 'Сводка'

ws1.merge_cells('A1:F1')
ws1['A1'] = 'TITAN — Анализ вторичных продаж 2025 vs 2026 (за полугодие)'
ws1['A1'].font = TITLE_FONT
ws1['A1'].alignment = CENTER
ws1.row_dimensions[1].height = 30

ws1.merge_cells('A2:F2')
ws1['A2'] = 'Источник: «Титан итого.xlsx», 3 вкладки — Охват / Проникновение / Динамика'
ws1['A2'].font = GRAY_FONT
ws1['A2'].alignment = CENTER

# Три главные цифры
ws1.merge_cells('A4:F4')
ws1['A4'] = 'ТРИ ГЛАВНЫЕ ЦИФРЫ'
ws1['A4'].font = SUBTITLE_FONT
ws1['A4'].alignment = LEFT

# Заголовки 3 цифр
metrics = [
    ('АКБ розничных клиентов', '6 096', '9 342', '+53%', 'Розница пришла к TITAN', LIGHT_GREEN, GREEN_FONT),
    ('Доля TITAN от ТО в АКБ', '7.6%', '5.4%', '−29%', 'Конкуренты сжимают TITAN с полки', LIGHT_RED, RED_FONT),
    ('Кг TITAN (вторичка)', '30 589', '20 352', '−33%', 'Берут реже и меньше — проблема глубины', LIGHT_RED, RED_FONT),
]
for i, (name, v25, v26, delta, comment, fill, font) in enumerate(metrics):
    r = 5 + i
    style_cell(ws1, r, 1, name, font=BOLD_FONT, fill=LIGHT_GRAY, align=LEFT)
    style_cell(ws1, r, 2, v25, font=NORMAL_FONT, fill=LIGHT_GRAY)
    style_cell(ws1, r, 3, '→', font=NORMAL_FONT, fill=LIGHT_GRAY)
    style_cell(ws1, r, 4, v26, font=BOLD_FONT, fill=LIGHT_GRAY)
    style_cell(ws1, r, 5, delta, font=font, fill=fill)
    style_cell(ws1, r, 6, comment, font=font, fill=fill, align=LEFT)
    ws1.row_dimensions[r].height = 28

# Распределение по группам
ws1.merge_cells('A9:F9')
ws1['A9'] = 'ГРУППЫ ДИСТРИБЬЮТОРОВ'
ws1['A9'].font = SUBTITLE_FONT
ws1['A9'].alignment = LEFT

groups_header = ['Группа', 'Кол-во договоров', 'Кг 2025', 'Кг 2026', 'Δ кг', 'Доля кг 2026']
for ci, h in enumerate(groups_header, 1):
    style_cell(ws1, 10, ci, h, font=HEADER_FONT, fill=HEADER_FILL)
ws1.row_dimensions[10].height = 30

group_stats = defaultdict(lambda: {'count': 0, 'kg25': 0, 'kg26': 0})
for d in distrib_full:
    group_stats[d['group']]['count'] += 1
    group_stats[d['group']]['kg25'] += d['kg25']
    group_stats[d['group']]['kg26'] += d['kg26']

group_order_list = [
    '🟢 Активный, стабильный/растущий',
    '🟢 Активный, новый',
    '🟡 Активный, падающий',
    '🔴 Не активен',
]
total_kg26 = sum(g['kg26'] for g in group_stats.values())
r = 11
for grp in group_order_list:
    g = group_stats[grp]
    if grp.startswith('🟢'):
        fill = LIGHT_GREEN
        font = GREEN_FONT
    elif grp.startswith('🟡'):
        fill = LIGHT_YELLOW
        font = ORANGE_FONT
    else:
        fill = LIGHT_RED
        font = RED_FONT
    delta = f"{(g['kg26']/g['kg25']-1)*100:+.0f}%" if g['kg25'] > 0 else "—"
    share = f"{g['kg26']/total_kg26*100:.0f}%" if total_kg26 > 0 else "—"
    style_cell(ws1, r, 1, grp, font=BOLD_FONT, fill=fill, align=LEFT)
    style_cell(ws1, r, 2, g['count'], font=font, fill=fill)
    style_cell(ws1, r, 3, g['kg25'], font=NORMAL_FONT, fill=fill, num_format='#,##0')
    style_cell(ws1, r, 4, g['kg26'], font=BOLD_FONT, fill=fill, num_format='#,##0')
    style_cell(ws1, r, 5, delta, font=font, fill=fill)
    style_cell(ws1, r, 6, share, font=font, fill=fill)
    r += 1

# Итого
style_cell(ws1, r, 1, 'ИТОГО', font=WHITE_BOLD, fill=HEADER_FILL, align=LEFT)
style_cell(ws1, r, 2, sum(g['count'] for g in group_stats.values()), font=WHITE_BOLD, fill=HEADER_FILL)
style_cell(ws1, r, 3, sum(g['kg25'] for g in group_stats.values()), font=WHITE_BOLD, fill=HEADER_FILL, num_format='#,##0')
style_cell(ws1, r, 4, total_kg26, font=WHITE_BOLD, fill=HEADER_FILL, num_format='#,##0')
_total_kg25 = sum(g['kg25'] for g in group_stats.values())
_delta_total = f"{(total_kg26/_total_kg25-1)*100:+.0f}%" if _total_kg25 > 0 else "—"
style_cell(ws1, r, 5, _delta_total, font=WHITE_BOLD, fill=HEADER_FILL)
style_cell(ws1, r, 6, '100%', font=WHITE_BOLD, fill=HEADER_FILL)

# Тревожные сигналы
r += 2
ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws1.cell(r, 1).value = '🚨 ТРИ ТРЕВОЖНЫХ СИГНАЛА'
ws1.cell(r, 1).font = SUBTITLE_FONT
ws1.cell(r, 1).alignment = LEFT

alarms = [
    '1. Кг упали на 33% при росте АКБ на 53% — клиенты берут «для пробы», средний чек упал',
    '2. Доля TITAN в АКБ снизилась с 7.6% до 5.4% — конкуренты вытесняют TITAN с полки',
    '3. Оберон (−86%), БРОЗЭКС (−81%), Севастьянов (−52%), Центр Красок (−28%) дают 80% падения',
]
for i, a in enumerate(alarms):
    r += 1
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    style_cell(ws1, r, 1, a, font=RED_FONT, fill=LIGHT_RED, align=LEFT)

# Позитивы
r += 2
ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws1.cell(r, 1).value = '✅ ТРИ ПОЗИТИВНЫХ СИГНАЛА'
ws1.cell(r, 1).font = Font(name='Arial', size=11, bold=True, color='16A34A')
ws1.cell(r, 1).alignment = LEFT

positives = [
    '1. АЮСС +867%, Уралстройсервис +55%, СтройТерминал — с нуля 306 кг (3 точки роста)',
    '2. 16 новых розничных клиентов >50 кг — приток есть, даже у падающих (Центр Красок 8 из 16)',
    '3. Держава Трейд (Казахстан) стабильно +11%, 7 065 кг = 35% всей вторички TITAN',
]
for i, p in enumerate(positives):
    r += 1
    ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    style_cell(ws1, r, 1, p, font=GREEN_FONT, fill=LIGHT_GREEN, align=LEFT)

col_widths = [32, 16, 14, 14, 12, 38]
for i, w in enumerate(col_widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w


# ========================================================
# ЛИСТ 2: ГРУППЫ ДИСТРИБЬЮТОРОВ
# ========================================================
ws2 = wb.create_sheet('Дистрибьюторы по группам')

ws2.merge_cells('A1:K1')
ws2['A1'] = 'Все дистрибьюторы TITAN — классификация по 3 группам для программы продвижения'
ws2['A1'].font = TITLE_FONT
ws2['A1'].alignment = CENTER
ws2.row_dimensions[1].height = 30

headers2 = ['Регион', 'Дистрибьютор', 'РД', 'Закуп', 'Группа',
            'Кг 2025', 'Кг 2026', 'Δ кг %', 'РТТ 25', 'РТТ 26', 'АКБ 26']
for ci, h in enumerate(headers2, 1):
    style_cell(ws2, 3, ci, h, font=HEADER_FONT, fill=HEADER_FILL)
ws2.row_dimensions[3].height = 35

r = 4
for d in distrib_full:
    fill = LIGHT_GREEN if d['group'].startswith('🟢') else (LIGHT_YELLOW if d['group'].startswith('🟡') else LIGHT_RED)
    font_color = GREEN_FONT if d['group'].startswith('🟢') else (ORANGE_FONT if d['group'].startswith('🟡') else RED_FONT)
    delta = f"{(d['kg26']/d['kg25']-1)*100:+.0f}%" if d['kg25'] > 0 else ("new" if d['kg26'] > 0 else "—")
    
    style_cell(ws2, r, 1, d['region'], font=NORMAL_FONT, fill=fill, align=LEFT)
    style_cell(ws2, r, 2, d['name'], font=BOLD_FONT, fill=fill, align=LEFT)
    style_cell(ws2, r, 3, d['rd'], font=NORMAL_FONT, fill=fill, align=LEFT)
    style_cell(ws2, r, 4, d['zakup'] or '—', font=NORMAL_FONT, fill=fill)
    style_cell(ws2, r, 5, d['group'], font=font_color, fill=fill, align=LEFT)
    style_cell(ws2, r, 6, d['kg25'] if d['kg25'] else '—', font=NORMAL_FONT, fill=fill, num_format='#,##0', align=RIGHT)
    style_cell(ws2, r, 7, d['kg26'] if d['kg26'] else '—', font=BOLD_FONT, fill=fill, num_format='#,##0', align=RIGHT)
    style_cell(ws2, r, 8, delta, font=font_color, fill=fill)
    style_cell(ws2, r, 9, d['rtt25'] if d['rtt25'] else '—', font=NORMAL_FONT, fill=fill)
    style_cell(ws2, r, 10, d['rtt26'] if d['rtt26'] else '—', font=NORMAL_FONT, fill=fill)
    style_cell(ws2, r, 11, d['akb26'] if d['akb26'] else '—', font=NORMAL_FONT, fill=fill, num_format='#,##0', align=RIGHT)
    r += 1

col_widths2 = [22, 32, 22, 12, 32, 10, 10, 10, 9, 9, 10]
for i, w in enumerate(col_widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w


# ========================================================
# ЛИСТ 3: ПОТЕРЯННЫЕ КЛИЕНТЫ
# ========================================================
ws3 = wb.create_sheet('30 потерянных клиентов')

ws3.merge_cells('A1:F1')
ws3['A1'] = 'Потерянные клиенты 2026 — для таргетированной акции возврата'
ws3['A1'].font = TITLE_FONT
ws3['A1'].alignment = CENTER
ws3.row_dimensions[1].height = 30

ws3.merge_cells('A2:F2')
ws3['A2'] = f'Критерий: закупали >50 кг в 2025, 0 кг в 2026. Всего: {len(lost_clients)} клиентов, потенциал возврата ~{sum(c["kg25"] for c in lost_clients):,.0f} кг'
ws3['A2'].font = GRAY_FONT
ws3['A2'].alignment = CENTER

headers3 = ['№', 'Дистрибьютор', 'Розничный клиент', 'Кг 2025', 'Покупок 2025', 'Приоритет']
for ci, h in enumerate(headers3, 1):
    style_cell(ws3, 4, ci, h, font=HEADER_FONT, fill=HEADER_FILL)
ws3.row_dimensions[4].height = 30

r = 5
for i, c in enumerate(lost_clients, 1):
    priority = '🔴 Высокий' if c['kg25'] >= 800 else ('🟡 Средний' if c['kg25'] >= 200 else '🟢 Низкий')
    pfill = LIGHT_RED if c['kg25'] >= 800 else (LIGHT_YELLOW if c['kg25'] >= 200 else LIGHT_GREEN)
    pfont = RED_FONT if c['kg25'] >= 800 else (ORANGE_FONT if c['kg25'] >= 200 else GREEN_FONT)
    style_cell(ws3, r, 1, i, font=BOLD_FONT, fill=pfill)
    style_cell(ws3, r, 2, c['distr'], font=NORMAL_FONT, fill=pfill, align=LEFT)
    style_cell(ws3, r, 3, c['client'], font=BOLD_FONT, fill=pfill, align=LEFT)
    style_cell(ws3, r, 4, c['kg25'], font=BOLD_FONT, fill=pfill, num_format='#,##0', align=RIGHT)
    style_cell(ws3, r, 5, c['p25'], font=NORMAL_FONT, fill=pfill)
    style_cell(ws3, r, 6, priority, font=pfont, fill=pfill)
    r += 1

# Итого
r += 1
style_cell(ws3, r, 1, '', fill=LIGHT_GRAY)
style_cell(ws3, r, 2, 'ИТОГО', font=WHITE_BOLD, fill=HEADER_FILL, align=LEFT)
style_cell(ws3, r, 3, f'{len(lost_clients)} клиентов', font=WHITE_BOLD, fill=HEADER_FILL, align=LEFT)
style_cell(ws3, r, 4, sum(c['kg25'] for c in lost_clients), font=WHITE_BOLD, fill=HEADER_FILL, num_format='#,##0', align=RIGHT)
style_cell(ws3, r, 5, '', fill=HEADER_FILL)
style_cell(ws3, r, 6, 'потенциал возврата', font=WHITE_BOLD, fill=HEADER_FILL)

col_widths3 = [5, 28, 50, 12, 12, 14]
for i, w in enumerate(col_widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w


# ========================================================
# ЛИСТ 4: НОВЫЕ КЛИЕНТЫ 2026
# ========================================================
ws4 = wb.create_sheet('16 новых клиентов 2026')

ws4.merge_cells('A1:F1')
ws4['A1'] = 'Новые клиенты 2026 — для тиражирования опыта привлечения'
ws4['A1'].font = TITLE_FONT
ws4['A1'].alignment = CENTER
ws4.row_dimensions[1].height = 30

ws4.merge_cells('A2:F2')
ws4['A2'] = f'Критерий: 0 кг в 2025, >50 кг в 2026. Всего: {len(new_clients)} клиентов. Лидеры привлечения: Центр Красок (8), Оберон (5)'
ws4['A2'].font = GRAY_FONT
ws4['A2'].alignment = CENTER

headers4 = ['№', 'Дистрибьютор', 'Розничный клиент', 'Кг 2026', 'Покупок 2026', 'Кейс для тиражирования']
for ci, h in enumerate(headers4, 1):
    style_cell(ws4, 4, ci, h, font=HEADER_FONT, fill=HEADER_FILL)
ws4.row_dimensions[4].height = 30

r = 5
for i, c in enumerate(new_clients, 1):
    style_cell(ws4, r, 1, i, font=BOLD_FONT, fill=LIGHT_GREEN)
    style_cell(ws4, r, 2, c['distr'], font=NORMAL_FONT, fill=LIGHT_GREEN, align=LEFT)
    style_cell(ws4, r, 3, c['client'], font=BOLD_FONT, fill=LIGHT_GREEN, align=LEFT)
    style_cell(ws4, r, 4, c['kg26'], font=BOLD_FONT, fill=LIGHT_GREEN, num_format='#,##0', align=RIGHT)
    style_cell(ws4, r, 5, c['p26'], font=NORMAL_FONT, fill=LIGHT_GREEN)
    # рекомендация
    if c['kg26'] >= 200:
        rec = 'Крупный новый клиент — закрепить, изучить, как пришли'
    elif c['p26'] >= 5:
        rec = 'Частые покупки — лояльный, добавить в MML расширенный'
    else:
        rec = 'Разовые покупки — нужен повторный визит ТП'
    style_cell(ws4, r, 6, rec, font=NORMAL_FONT, fill=LIGHT_GREEN, align=LEFT)
    r += 1

col_widths4 = [5, 28, 50, 12, 12, 40]
for i, w in enumerate(col_widths4, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w


# Сохранение
wb.save(OUT)
print(f'OK: {OUT}')
import os
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
print(f'\nЛисты:')
for sn in wb.sheetnames:
    print(f'  "{sn}"')
print(f'\nПотерянных клиентов: {len(lost_clients)}')
print(f'Новых клиентов: {len(new_clients)}')
print(f'Дистрибьюторов: {len(distrib_full)}')
