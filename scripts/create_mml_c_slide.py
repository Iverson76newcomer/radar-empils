#!/usr/bin/env python3
"""
Слайд: MML TITAN — Было → Стало (для класса C).
Простой, наглядный формат для презентации.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

RED = RGBColor(0xC6, 0x28, 0x28)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK = RGBColor(0x1A, 0x1F, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xE5, 0xE7, 0xEB)
ORANGE = RGBColor(0xF4, 0x79, 0x20)
YELLOW = RGBColor(0xFB, 0xBF, 0x24)
DARK_BLUE = RGBColor(0x1E, 0x3A, 0x5F)
LIGHT_GREEN = RGBColor(0xDC, 0xFC, 0xE7)
LIGHT_RED = RGBColor(0xFE, 0xE2, 0xE2)
LIGHT_ORANGE = RGBColor(0xFF, 0xED, 0xD5)
LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)

OUTPUT = '/home/z/my-project/download/MML_класс_C_было_стало.pptx'

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

RAL_COLORS = {
    'RAL 9005': RGBColor(0x1A, 0x1A, 0x1A),
    'RAL 8017': RGBColor(0x3D, 0x28, 0x17),
    'RAL 7004': RGBColor(0xB0, 0xB0, 0xB0),
    'Чёрный': RGBColor(0x1A, 0x1A, 0x1A),
    'Коричневый': RGBColor(0x4A, 0x2C, 0x17),
    'Серый': RGBColor(0x80, 0x80, 0x80),
}


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK, bold=False, align=PP_ALIGN.LEFT, font_name='Arial',
                anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    return tb


def add_shape(slide, shape_type, left, top, width, height, fill_color=None,
              line_color=None, line_width=None):
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill_color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        if line_width is not None:
            shp.line.width = line_width
    shp.shadow.inherit = False
    return shp


def add_multi_text(slide, left, top, width, height, lines,
                   anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(1); tf.margin_right = Pt(1)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, (text, size, color, bold) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = 'Arial'
    return tb


def draw_can(slide, x, y, w, h, color_rgb, label_top, label_mid, label_bot, effect='gloss'):
    r, g, b = color_rgb[0], color_rgb[1], color_rgb[2]
    brightness = 0.299 * r + 0.587 * g + 0.114 * b
    text_color = DARK if brightness > 140 else WHITE

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
              fill_color=color_rgb, line_color=DARK, line_width=Pt(0.5))
    lid_h = Inches(0.06)
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              x + Inches(0.03), y - lid_h + Inches(0.02),
              w - Inches(0.06), lid_h,
              fill_color=GRAY, line_color=DARK, line_width=Pt(0.3))

    if effect == 'hammer':
        import random
        random.seed(int(x) + int(y))
        for _ in range(3):
            w_in = w / 914400 if isinstance(w, (int, float)) else w.inches
            h_in = h / 914400 if isinstance(h, (int, float)) else h.inches
            dx = random.uniform(0.15, w_in - 0.15)
            dy = random.uniform(0.20, h_in - 0.20)
            dot_r = Inches(0.03)
            add_shape(slide, MSO_SHAPE.OVAL,
                      x + Inches(dx) - dot_r, y + Inches(dy) - dot_r,
                      dot_r * 2, dot_r * 2,
                      fill_color=WHITE, line_color=None)

    tf = slide.shapes.add_textbox(x, y + Inches(0.02), w, h - Inches(0.04)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(0); tf.margin_right = Pt(0)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    for i, (text, size, bold) in enumerate([
        (label_top, 7, True),
        (label_mid, 6, True),
        (label_bot, 6, False),
    ]):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = text_color
        run.font.name = 'Arial'


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H,
              fill_color=WHITE, line_color=None)

    # === ШАПКА ===
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.65),
              fill_color=DARK_BLUE, line_color=None)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(0.65), SLIDE_W, Inches(0.04),
              fill_color=YELLOW, line_color=None)
    add_textbox(slide, Inches(0.4), Inches(0.08), SLIDE_W - Inches(0.8), Inches(0.5),
                'MML КЛАСС C (СУПЕРМАРКЕТ) — БЫЛО → СТАЛО',
                font_size=22, color=WHITE, bold=True,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, SLIDE_W - Inches(5.0), Inches(0.08),
                Inches(4.6), Inches(0.5),
                '2 SKU → 7 SKU · на основе продаж',
                font_size=11, color=YELLOW, bold=False,
                align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # === ЛЕВАЯ ПОЛОВИНА: БЫЛО ===
    half_w = (SLIDE_W - Inches(1.2)) / 2

    add_textbox(slide, Inches(0.4), Inches(0.90),
                half_w, Inches(0.40),
                'БЫЛО — 2 SKU (только молотковая)', font_size=14, color=GRAY, bold=True,
                align=PP_ALIGN.CENTER)

    # Полка «было»
    was_y = Inches(1.40)
    was_h = Inches(2.80)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(0.4), was_y, half_w, was_h,
              fill_color=LIGHT_GRAY, line_color=GRAY, line_width=Pt(1.5))
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              Inches(0.55), was_y + was_h - Inches(0.20) - Inches(0.03),
              half_w - Inches(0.30), Inches(0.20),
              fill_color=RGBColor(0x8B, 0x6F, 0x47), line_color=None)

    # 2 банки
    was_cans = [
        {'ral': 'Чёрный', 'mid': 'МОЛ 0.8кг', 'bot': 'Чёрный', 'effect': 'hammer'},
        {'ral': 'Коричневый', 'mid': 'МОЛ 0.8кг', 'bot': 'Корич', 'effect': 'hammer'},
    ]
    can_w = Inches(1.60)
    can_h = Inches(2.20)
    gap = Inches(0.60)
    total_w = can_w * 2 + gap
    start_x = Inches(0.4) + (half_w - total_w) // 2
    can_y = was_y + was_h - Inches(0.20) - can_h - Inches(0.15)

    for i, can in enumerate(was_cans):
        cx = start_x + (can_w + gap) * i
        color = RAL_COLORS.get(can['ral'], DARK)
        draw_can(slide, cx, can_y, can_w, can_h, color,
                 'TITAN', can['mid'], can['bot'], effect=can.get('effect', 'gloss'))

    # Комментарий
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(0.4), Inches(4.35), half_w, Inches(0.55),
              fill_color=LIGHT_RED, line_color=RED, line_width=Pt(1))
    add_textbox(slide, Inches(0.4), Inches(4.35), half_w, Inches(0.55),
                '⚠  TITAN 3в1 отсутствует. Покупатель уходит с ПФ-115.',
                font_size=11, color=DARK, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # === ПРАВАЯ ПОЛОВИНА: СТАЛО ===
    right_x = Inches(0.4) + half_w + Inches(0.4)

    add_textbox(slide, right_x, Inches(0.90),
                half_w, Inches(0.40),
                'СТАЛО — 7 SKU (3в1 + быстросохнущая + молотковая)', font_size=14, color=GREEN, bold=True,
                align=PP_ALIGN.CENTER)

    # Полка «стало» — 3 ряда
    now_y = Inches(1.40)
    now_h = Inches(1.85)
    now_gap = Inches(0.10)

    # Ряд 1: 3в1 (3 SKU)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              right_x, now_y, half_w, now_h,
              fill_color=LIGHT_GRAY, line_color=GREEN, line_width=Pt(1))
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              right_x + Inches(0.15), now_y + now_h - Inches(0.15) - Inches(0.03),
              half_w - Inches(0.30), Inches(0.15),
              fill_color=RGBColor(0x8B, 0x6F, 0x47), line_color=None)
    add_textbox(slide, right_x + Inches(0.15), now_y + Inches(0.05),
                half_w - Inches(0.30), Inches(0.25),
                '3в1 ПО РЖАВЧИНЕ (3 SKU) — НОВОЕ', font_size=8, color=RED, bold=True,
                align=PP_ALIGN.LEFT)

    cans_3v1 = [
        {'ral': 'RAL 9005', 'mid': '3в1 0.75', 'bot': 'Чёрный'},
        {'ral': 'RAL 8017', 'mid': '3в1 0.75', 'bot': 'Шоколад'},
        {'ral': 'RAL 7004', 'mid': '3в1 0.75', 'bot': 'Серый'},
    ]
    n = 3
    margin = Inches(0.30)
    gap2 = Inches(0.10)
    cw = (half_w - margin * 2 - gap2 * (n - 1)) / n
    ch = Inches(1.20)
    cy = now_y + now_h - Inches(0.15) - ch - Inches(0.10)
    for i, can in enumerate(cans_3v1):
        cx = right_x + margin + (cw + gap2) * i
        color = RAL_COLORS.get(can['ral'], DARK)
        draw_can(slide, cx, cy, cw, ch, color,
                 'TITAN', can['mid'], can['bot'])

    # Ряд 2: быстросохнущая (2 SKU) — НОВОЕ
    fast_y = now_y + now_h + now_gap
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              right_x, fast_y, half_w, now_h,
              fill_color=LIGHT_GRAY, line_color=GREEN, line_width=Pt(1))
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              right_x + Inches(0.15), fast_y + now_h - Inches(0.15) - Inches(0.03),
              half_w - Inches(0.30), Inches(0.15),
              fill_color=RGBColor(0x8B, 0x6F, 0x47), line_color=None)
    add_textbox(slide, right_x + Inches(0.15), fast_y + Inches(0.05),
                half_w - Inches(0.30), Inches(0.25),
                'БЫСТРОСОХНУЩАЯ (2 SKU) — НОВОЕ', font_size=8, color=DARK_BLUE, bold=True,
                align=PP_ALIGN.LEFT)

    cans_fast = [
        {'ral': 'RAL 9005', 'mid': 'БЫСТ 0.75', 'bot': 'Чёрный'},
        {'ral': 'Серый', 'mid': 'БЫСТ 0.75', 'bot': 'Серый'},
    ]
    n = 2
    cw2 = (half_w - margin * 2 - gap2 * (n - 1)) / n
    cy2 = fast_y + now_h - Inches(0.15) - ch - Inches(0.10)
    for i, can in enumerate(cans_fast):
        cx = right_x + margin + (cw2 + gap2) * i
        color = RAL_COLORS.get(can['ral'], DARK)
        draw_can(slide, cx, cy2, cw2, ch, color,
                 'TITAN', can['mid'], can['bot'], effect='matte')

    # Ряд 3: молотковая (2 SKU) — УЖЕ БЫЛО
    ham_y = fast_y + now_h + now_gap
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              right_x, ham_y, half_w, now_h,
              fill_color=LIGHT_GRAY, line_color=GRAY, line_width=Pt(1))
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              right_x + Inches(0.15), ham_y + now_h - Inches(0.15) - Inches(0.03),
              half_w - Inches(0.30), Inches(0.15),
              fill_color=RGBColor(0x8B, 0x6F, 0x47), line_color=None)
    add_textbox(slide, right_x + Inches(0.15), ham_y + Inches(0.05),
                half_w - Inches(0.30), Inches(0.25),
                'МОЛОТКОВАЯ (2 SKU) — уже было', font_size=8, color=GRAY, bold=True,
                align=PP_ALIGN.LEFT)

    cans_ham = [
        {'ral': 'Чёрный', 'mid': 'МОЛ 0.8', 'bot': 'Чёрный', 'effect': 'hammer'},
        {'ral': 'Коричневый', 'mid': 'МОЛ 0.8', 'bot': 'Корич', 'effect': 'hammer'},
    ]
    n = 2
    cy3 = ham_y + now_h - Inches(0.15) - ch - Inches(0.10)
    for i, can in enumerate(cans_ham):
        cx = right_x + margin + (cw2 + gap2) * i
        color = RAL_COLORS.get(can['ral'], DARK)
        draw_can(slide, cx, cy3, cw2, ch, color,
                 'TITAN', can['mid'], can['bot'], effect=can.get('effect', 'hammer'))

    # Комментарий справа
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              right_x, Inches(4.35), half_w, Inches(0.55),
              fill_color=LIGHT_GREEN, line_color=GREEN, line_width=Pt(1))
    add_textbox(slide, right_x, Inches(4.35), half_w, Inches(0.55),
                '✓  TITAN 3в1 + быстросохнущая — теперь в супермаркете',
                font_size=11, color=GREEN, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # === НИЖНЯЯ ЧАСТЬ — таблица обоснования ===
    tbl_y = Inches(5.10)
    add_textbox(slide, Inches(0.4), tbl_y,
                SLIDE_W - Inches(0.8), Inches(0.35),
                'ОБОСНОВАНИЕ ДОБАВЛЕНИЯ — НА ОСНОВЕ РЕАЛЬНЫХ ПРОДАЖ',
                font_size=12, color=DARK_BLUE, bold=True, align=PP_ALIGN.LEFT)

    # Таблица
    headers = ['SKU', 'Категория', 'Цвет', 'RAL', 'Продажи (шт)', 'АКБ', 'Статус']
    col_widths = [Inches(1.5), Inches(2.5), Inches(1.8), Inches(1.5), Inches(1.7), Inches(1.0), Inches(2.3)]
    th_y = Inches(5.50)
    th_h = Inches(0.35)

    cx = Inches(0.4)
    for i, h in enumerate(headers):
        add_shape(slide, MSO_SHAPE.RECTANGLE,
                  cx, th_y, col_widths[i], th_h,
                  fill_color=DARK_BLUE, line_color=WHITE, line_width=Pt(1))
        add_textbox(slide, cx, th_y, col_widths[i], th_h,
                    h, font_size=9, color=WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx += col_widths[i]

    rows = [
        ('87289', '3в1 по ржавчине', 'Чёрный', 'RAL 9005', '7 738', '393', '🆕 Добавить', LIGHT_RED),
        ('87292', '3в1 по ржавчине', 'Шоколад', 'RAL 8017', '4 524', '287', '🆕 Добавить', LIGHT_RED),
        ('87281', '3в1 по ржавчине', 'Серый', 'RAL 7004', '5 071', '247', '🆕 Добавить', LIGHT_RED),
        ('83777', 'Быстросохнущая', 'Чёрный', 'RAL 9005', '5 381', '303', '🆕 Добавить', LIGHT_BLUE),
        ('—', 'Быстросохнущая', 'Серый', 'RAL 7004', '2 695', '199', '🆕 Добавить', LIGHT_BLUE),
        ('77636', 'Молотковая', 'Чёрный', '—', '—', '—', '✅ Уже есть', LIGHT_GREEN),
        ('77660', 'Молотковая', 'Коричневый', '—', '—', '—', '✅ Уже есть', LIGHT_GREEN),
    ]

    row_h = Inches(0.28)
    for ri, (sku, cat, color_name, ral, qty, akb, status, bg) in enumerate(rows):
        ry = th_y + th_h + row_h * ri
        cx = Inches(0.4)
        values = [sku, cat, color_name, ral, qty, akb, status]
        for ci, val in enumerate(values):
            add_shape(slide, MSO_SHAPE.RECTANGLE,
                      cx, ry, col_widths[ci], row_h,
                      fill_color=bg, line_color=GRAY, line_width=Pt(0.3))
            val_color = GREEN if '✅' in str(val) else (RED if '🆕' in str(val) else DARK)
            add_textbox(slide, cx, ry, col_widths[ci], row_h,
                        str(val), font_size=8, color=val_color,
                        bold=('🆕' in str(val) or '✅' in str(val) or ci == 0),
                        align=PP_ALIGN.CENTER,
                        anchor=MSO_ANCHOR.MIDDLE)
            cx += col_widths[ci]

    # Итог
    total_y = th_y + th_h + row_h * len(rows) + Inches(0.05)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(0.4), total_y, SLIDE_W - Inches(0.8), Inches(0.40),
              fill_color=GREEN, line_color=None)
    add_textbox(slide, Inches(0.4), total_y, SLIDE_W - Inches(0.8), Inches(0.40),
                'ИТОГО: 2 SKU → 7 SKU (+5)  ·  Класс C получает TITAN 3в1 и быстросохнущую  ·  Источник: «Титан цвета», 2026',
                font_size=11, color=YELLOW, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Футер
    add_textbox(slide, Inches(0.4), Inches(7.15),
                SLIDE_W - Inches(0.8), Inches(0.25),
                'Empils  •  TITAN  •  MML — Minimal Must List  •  Красишь 1 раз — служит 10 лет',
                font_size=8, color=GRAY, italic=True, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f'OK: {OUTPUT}')
    print(f'Слайдов: {len(prs.slides)}')


if __name__ == '__main__':
    main()
