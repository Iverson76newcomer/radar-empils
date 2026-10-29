"""Изучение предыдущих наработок: Блогеры Титан.xlsx + TITAN_Топ-8_full_идей_вирусных_роликов.docx"""
from openpyxl import load_workbook
from docx import Document

print("=" * 80)
print("ФАЙЛ 1: Блогеры Титан.xlsx")
print("=" * 80)
wb = load_workbook('/home/z/my-project/upload/Блогеры Титан.xlsx', data_only=True)
for sh_name in wb.sheetnames:
    ws = wb[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None: vals.append('')
            elif isinstance(v, float): vals.append(f'{v:.0f}')
            else: vals.append(str(v)[:50])
        line = ' | '.join(vals)
        if line.strip(' |'):
            print(line)

print("\n\n" + "=" * 80)
print("ФАЙЛ 2: TITAN_Топ-8_full_идей_вирусных_роликов.docx")
print("=" * 80)
doc = Document('/home/z/my-project/upload/TITAN_Топ-8_full_идей_вирусных_роликов.docx')
for i, p in enumerate(doc.paragraphs[:200]):
    if p.text.strip():
        print(p.text)
print("\n--- ТАБЛИЦЫ ---")
for ti, t in enumerate(doc.tables):
    print(f"\n[TABLE {ti+1}] — {len(t.rows)} rows × {len(t.columns)} cols")
    for r in t.rows:
        cells = [c.text.strip().replace('\n', ' | ')[:80] for c in r.cells]
        print(' | '.join(cells))
