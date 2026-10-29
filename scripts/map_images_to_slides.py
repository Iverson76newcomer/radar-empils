#!/usr/bin/env python3
"""Map media images to slides via python-pptx."""
from pptx import Presentation
from pptx.util import Emu
import os

PPTX = "/home/z/my-project/download/Программа_развития_Эмпилс_2026-2028_обновлено.pptx"

prs = Presentation(PPTX)

print(f"Total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides, 1):
    print(f"--- Slide {i} ---")
    # Get the layout name
    try:
        print(f"  Layout: {slide.slide_layout.name}")
    except Exception:
        pass

    # Extract text
    text_parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t.strip():
                    text_parts.append(t.strip())
    if text_parts:
        print("  Text:")
        for t in text_parts:
            print(f"    • {t}")

    # Find pictures
    pics = []
    for shape in slide.shapes:
        if shape.shape_type == 13:  # Picture
            try:
                w = Emu(shape.width).inches
                h = Emu(shape.height).inches
            except Exception:
                w = h = 0
            # Find image filename via relationships
            try:
                rid = shape._element.blip_rId
                rel = slide.part.rels[rid]
                target = rel.target_ref
                pics.append((target, w, h))
            except Exception as e:
                pics.append((f"<unknown: {e}>", w, h))
    if pics:
        print("  Pictures:")
        for target, w, h in pics:
            print(f"    📷 {target}  ({w:.1f}\" × {h:.1f}\")")
    print()
