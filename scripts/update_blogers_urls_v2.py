"""Точное обновление TITAN_Региональные_блогеры_коротко.docx
   Заменяем только «через YT поиск» в колонке «Канал» (cell 1),
   оставляя cell 5 (Контакт) нетронутым.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document('/home/z/my-project/upload/TITAN_Региональные_блогеры_коротко.docx')

# Список каналов и их новых URL
updates = {
    'ИДЕИ ДЛЯ ДОМА': 'https://rutube.ru/channel/30125920/',
    'Деревенский кузнец': 'https://www.youtube.com/@ruralsmith',
    'Художественная ковка металла': 'https://rutube.ru/channel/24685954/',
}

def set_cell_text(cell, new_text):
    """Полная замена текста ячейки с сохранением форматирования первого run."""
    # Удаляем все paragraphs кроме первого
    paragraphs = cell.paragraphs
    if len(paragraphs) > 1:
        for p in paragraphs[1:]:
            p._p.getparent().remove(p._p)
    # В первом параграфе удаляем все runs кроме первого
    p = paragraphs[0]
    runs = p.runs
    if len(runs) > 1:
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    if runs:
        # Меняем текст первого run
        runs[0].text = new_text
    else:
        # Если runs нет — добавляем
        r = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = new_text
        r.append(t)
        p._p.append(r)

# Перебираем таблицы и ищем строку с каждым каналом
for table in doc.tables:
    for row in table.rows:
        cells = row.cells
        # Колонка 1 = «Канал» (содержит название канала + URL)
        if len(cells) > 1:
            cell1_text = cells[1].text.strip()
            for channel_name, new_url in updates.items():
                if channel_name in cell1_text:
                    print(f"Найден: {channel_name}")
                    # Заменяем ячейку 1: оставляем имя канала + URL
                    new_cell1_text = f"{channel_name}\n{new_url}"
                    set_cell_text(cells[1], new_cell1_text)
                    print(f"  cell 1 → {new_cell1_text}")
                    # ВАЖНО: cell 5 (Контакт) НЕ ТРОГАЕМ — там должно остаться «через YT»

# Сохраняем
out_path = '/home/z/my-project/download/TITAN_Региональные_блогеры_коротко.docx'
doc.save(out_path)
print(f'\n✅ Сохранено: {out_path}')

# Проверка
doc2 = Document(out_path)
print('\n=== ПРОВЕРКА ===')
for table in doc2.tables:
    for ri, row in enumerate(table.rows):
        cells = row.cells
        if len(cells) > 1:
            text = cells[1].text
            for channel_name in updates.keys():
                if channel_name in text:
                    print(f"\nСтрока {ri}: {channel_name}")
                    print(f"  Канал (cell 1): {cells[1].text.strip()}")
                    print(f"  Подписчики (cell 2): {cells[2].text.strip()}")
                    print(f"  Просмотры (cell 3): {cells[3].text.strip()}")
                    print(f"  Тематика (cell 4): {cells[4].text.strip()[:60]}")
                    print(f"  Контакт (cell 5): {cells[5].text.strip()}")
                    print(f"  Идея (cell 6): {cells[6].text.strip()[:60]}")
