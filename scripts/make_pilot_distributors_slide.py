#!/usr/bin/env python3
"""
Добавить новый слайд с тремя пилотными дистрибьюторами для программы TITAN
в презентацию Titan_Дорожная_карта_АКБ_анализ_v1.pptx (последнюю версию).

Пилотные дистрибьюторы:
1. Палитра — Краснодарский край
2. И.П. Рыгалов — Пензенская область
3. Брозекс — Свердловская область
"""

import os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# Берём последнюю версию с АКБ-анализом
SRC = "/home/z/my-project/download/Titan_Дорожная_карта_АКБ_анализ_v1.pptx"
DST = "/home/z/my-project/download/Titan_Дорожная_карта_пилоты_v1.pptx"
shutil.copy(SRC, DST)

prs = Presentation(DST)
print(f"Слайдов в исходной презентации: {len(prs.slides)}")

# Вставляем новый слайд после слайда 4 (АКБ-анализ) — индекс 4
slide_layout = prs.slide_layouts[6]  # Blank
new_slide = prs.slides.add_slide(slide_layout)

# Перемещаем на позицию 5 (индекс 4) — после АКБ-анализа
sldIdLst = prs.slides._sldIdLst
sldIds = list(sldIdLst)
new_sldId = sldIds[-1]
sldIdLst.remove(new_sldId)
sldIdLst.insert(4, new_sldId)  # позиция 5 (0-indexed = 4)

slide = new_slide

# === Цвета фирменные TITAN/Empils ===
ORANGE = RGBColor(0xF4, 0x76, 0x29)
ORANGE_DARK = RGBColor(0xD5, 0x5E, 0x0A)
ORANGE_LIGHT = RGBColor(0xFB, 0xE6, 0xD5)
DARK = RGBColor(0x31, 0x31, 0x31)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFA, 0xFA, 0xFA)
BLUE = RGBColor(0x21, 0x96, 0xF3)

# === Данные по пилотным дистрибьюторам ===
PILOTS = [
    {
        "num": "01",
        "name": "ПАЛИТРА",
        "region": "Краснодарский край",
        "city": "Краснодар",
        "color": ORANGE,
        "features": [
            "Крупнейший сельскохозяйственный регион Юга России",
            "Высокая концентрация металлоконструкций: фермы, ангары, заборы",
            "Активная стройка в Краснодаре и Сочи",
            "Сильная розничная база DIY",
        ],
    },
    {
        "num": "02",
        "name": "И.П. РЫГАЛОВ",
        "region": "Пензенская область",
        "city": "Пенза",
        "color": ORANGE_DARK,
        "features": [
            "Промышленный регион Поволжья",
            "Производство металлических изделий и конструкций",
            "Дача и частный сектор — высокий спрос на ЛКМ по металлу",
            "Удобная логистика из Ростова-на-Дону",
        ],
    },
    {
        "num": "03",
        "name": "БРОЗЕКС",
        "region": "Свердловская область",
        "city": "Екатеринбург",
        "color": DARK,
        "features": [
            "Крупнейший промышленный хаб Урала",
            "Тяжёлая промышленность: металлургия, машиностроение",
            "30+ городов присутствия Brozex от Москвы до Новосибирска",
            "Собственный автопарк — развитая логистика",
            "Эталонный кейс для других партнёров Эмпилс",
        ],
    },
]

# === Вспомогательные функции ===

def add_rect(slide, left, top, width, height, fill_color, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = line_width
    shape.shadow.inherit = False
    return shape

def add_round_rect(slide, left, top, width, height, fill_color, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = line_width
    shape.shadow.inherit = False
    try:
        shape.adjustments[0] = 0.05
    except:
        pass
    return shape

def add_text(slide, left, top, width, height, text, font_size=10, bold=False,
             color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font_name="Calibri",
             italic=False, line_spacing=1.0):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = Pt(0); tf.margin_right = Pt(0)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.name = font_name
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb

# === Размеры слайда ===
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# === 1. Фон слайда ===
add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)

# === 2. Верхняя оранжевая полоса ===
add_rect(slide, 0, 0, SLIDE_W, Inches(0.12), ORANGE)

# === 3. Заголовок ===
add_text(slide, Inches(0.5), Inches(0.22), Inches(11), Inches(0.45),
         "ПИЛОТНЫЕ ДИСТРИБЬЮТОРЫ ПРОГРАММЫ TITAN",
         font_size=22, bold=True, color=DARK, anchor=MSO_ANCHOR.MIDDLE)

add_text(slide, Inches(0.5), Inches(0.70), Inches(11), Inches(0.3),
         "3 региона · 3 пилота · отработка программы перед масштабированием",
         font_size=10, italic=True, color=GRAY)

# Бренд-марка справа сверху
add_text(slide, Inches(11.0), Inches(0.22), Inches(2.0), Inches(0.45),
         "TITAN · 2026",
         font_size=12, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# === 4. Три большие карточки пилотов ===
CARD_TOP = Inches(1.30)
CARD_HEIGHT = Inches(4.90)
CARD_WIDTH = Inches(4.05)
CARD_GAP = Inches(0.10)
CARD_LEFT_START = Inches(0.5)

for i, pilot in enumerate(PILOTS):
    left = CARD_LEFT_START + i * (CARD_WIDTH + CARD_GAP)
    accent = pilot["color"]
    
    # Фон карточки (белый с тонкой границей)
    add_round_rect(slide, left, CARD_TOP, CARD_WIDTH, CARD_HEIGHT, WHITE,
                   line_color=LIGHT_GRAY, line_width=Pt(0.75))
    
    # Верхняя цветная плашка с номером
    add_rect(slide, left, CARD_TOP, CARD_WIDTH, Inches(0.55), accent)
    add_text(slide, left + Inches(0.2), CARD_TOP, Inches(0.7), Inches(0.55),
             pilot["num"],
             font_size=22, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, font_name="Calibri")
    add_text(slide, left + Inches(1.0), CARD_TOP, CARD_WIDTH - Inches(1.2), Inches(0.55),
             "ПИЛОТ",
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    
    # Название дистрибьютора (большим жирным шрифтом)
    name_top = CARD_TOP + Inches(0.75)
    add_text(slide, left + Inches(0.25), name_top, CARD_WIDTH - Inches(0.5), Inches(0.55),
             pilot["name"],
             font_size=20, bold=True, color=DARK, anchor=MSO_ANCHOR.TOP, line_spacing=1.0)
    
    # Регион
    add_text(slide, left + Inches(0.25), name_top + Inches(0.60), CARD_WIDTH - Inches(0.5), Inches(0.3),
             pilot["region"],
             font_size=13, bold=True, color=accent, anchor=MSO_ANCHOR.TOP)
    
    # Город
    add_text(slide, left + Inches(0.25), name_top + Inches(0.90), CARD_WIDTH - Inches(0.5), Inches(0.25),
             "📍 " + pilot["city"],
             font_size=11, color=GRAY, anchor=MSO_ANCHOR.TOP)
    
    # Разделительная линия
    add_rect(slide, left + Inches(0.25), name_top + Inches(1.30), CARD_WIDTH - Inches(0.5), Pt(1), accent)
    
    # Заголовок "ПОЧЕМУ ВЫБРАН"
    add_text(slide, left + Inches(0.25), name_top + Inches(1.45), CARD_WIDTH - Inches(0.5), Inches(0.3),
             "ПОЧЕМУ ВЫБРАН",
             font_size=9, bold=True, color=accent, anchor=MSO_ANCHOR.TOP)
    
    # Список особенностей
    features_top = name_top + Inches(1.80)
    features_text = "\n".join(["• " + f for f in pilot["features"]])
    add_text(slide, left + Inches(0.25), features_top, CARD_WIDTH - Inches(0.5), Inches(2.0),
             features_text,
             font_size=10, color=DARK, anchor=MSO_ANCHOR.TOP, line_spacing=1.4)

# === 5. Блок «Стратегия пилотирования» (внизу) ===
STRAT_TOP = Inches(6.35)
STRAT_HEIGHT = Inches(0.80)
STRAT_LEFT = Inches(0.5)
STRAT_WIDTH = Inches(12.33)

add_round_rect(slide, STRAT_LEFT, STRAT_TOP, STRAT_WIDTH, STRAT_HEIGHT, ORANGE_LIGHT,
               line_color=ORANGE, line_width=Pt(1))

add_text(slide, STRAT_LEFT + Inches(0.25), STRAT_TOP + Inches(0.10), Inches(2.5), Inches(0.3),
         "🎯 СТРАТЕГИЯ ПИЛОТИРОВАНИЯ",
         font_size=11, bold=True, color=ORANGE_DARK, anchor=MSO_ANCHOR.MIDDLE)

add_text(slide, STRAT_LEFT + Inches(0.25), STRAT_TOP + Inches(0.40), STRAT_WIDTH - Inches(0.5), Inches(0.4),
         "3 региона (Юг · Поволжье · Урал) — покрытие 3 федеральных округов · 3 разных профиля дистрибьютора (сеть DIY · ИП · крупный оптовик) · отработка программы TITAN перед масштабированием на 109 партнёров",
         font_size=10, color=DARK, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)

# === 6. Нижняя оранжевая полоска ===
add_rect(slide, 0, Inches(7.30), SLIDE_W, Inches(0.20), ORANGE)

# Сохраняем
prs.save(DST)
print(f"\n✅ Файл сохранён: {DST}")
print(f"   Размер: {os.path.getsize(DST):,} байт")
print(f"   Слайд вставлен на позицию 5 (после АКБ-анализа)")
print(f"   Всего слайдов: {len(prs.slides)}")
