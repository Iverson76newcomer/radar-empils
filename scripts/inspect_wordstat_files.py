"""Изучаем структуру обоих файлов."""
from openpyxl import load_workbook
from pptx import Presentation

print("=" * 80)
print("EXCEL: Wordstat_узнаваемость_брендов.xlsx")
print("=" * 80)
wb = load_workbook('/home/z/my-project/upload/Wordstat_узнаваемость_брендов.xlsx', data_only=True)
for sh_name in wb.sheetnames:
    ws = wb[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None: vals.append('')
            elif isinstance(v, float): vals.append(f'{v:.2f}')
            else: vals.append(str(v))
        print(' | '.join(vals))

print("\n\n" + "=" * 80)
print("PPTX: Wordstat_полный_анализ_TITAN_ с компанией.pptx")
print("=" * 80)
prs = Presentation('/home/z/my-project/upload/Wordstat_полный_анализ_TITAN_ с компанией.pptx')
for i, slide in enumerate(prs.slides, 1):
    print(f"\n--- СЛАЙД {i} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                print(f"[TEXT shape '{shape.name}' | pos=({shape.left},{shape.top}) | size=({shape.width},{shape.height})]")
                print(txt)
                print('---')
        elif shape.has_table:
            print(f"[TABLE shape '{shape.name}']")
            tbl = shape.table
            for r in tbl.rows:
                cells = [c.text.strip() for c in r.cells]
                print(' | '.join(cells))
            print('---')
