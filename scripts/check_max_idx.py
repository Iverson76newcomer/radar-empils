import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU.xlsx", data_only=True)
ws = wb["300 SKU — детально"]

max_idx = 0
idx_list = []
for r in range(4, ws.max_row + 1):
    idx = ws.cell(row=r, column=1).value
    if idx is not None:
        if isinstance(idx, int):
            idx_list.append(idx)
            if idx > max_idx:
                max_idx = idx

print(f"Max idx в оригинале: {max_idx}")
print(f"Всего SKU: {len(idx_list)}")
print(f"Уникальных idx: {len(set(idx_list))}")
print(f"Дубликаты: {len(idx_list) - len(set(idx_list))}")

# Найдём дубликаты
from collections import Counter
cnt = Counter(idx_list)
dups = [(idx, c) for idx, c in cnt.items() if c > 1]
print(f"\nДубликаты idx (встречаются >1 раза):")
for idx, c in sorted(dups):
    print(f"  idx={idx}: {c} раза")

# Последние 10 idx
print(f"\nПоследние 10 idx в оригинале:")
last_10 = sorted(idx_list)[-10:]
for idx in last_10:
    print(f"  idx={idx}")
