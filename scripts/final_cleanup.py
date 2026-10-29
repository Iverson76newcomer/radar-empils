#!/usr/bin/env python3
"""
Финальная очистка данных:
1. Удалить сайты-справочники
2. Исправить неверные email
3. Заполнить недостающие данные
"""

import openpyxl

dst_file = '/home/z/my-project/download/Задание_на_дополнение_таблицы_заполнено.xlsx'

# Сайты-справочники и нерелевантные сайты → убрать
BAD_SITES = {
    4: None,    # tbank.ru → нет сайта у Дальтех МБ
    6: None,    # checko.ru → нет сайта у ИП Жадаев
    8: None,    # service-online.su → нет сайта у Сибирская Усадьба
    10: None,   # audit-it.ru → нет сайта у Студия красок (ликвидирована)
    17: None,   # checko.ru → нет сайта у Вланж плюс
    20: None,   # audit-it.ru → нет сайта у ИП Ковалев
    21: None,   # checko.ru → нет сайта у ИП Мима
    22: None,   # synapsenet.ru → нет сайта у ИП Минюк
    23: None,   # synapsenet.ru → нет сайта у ИП Миронкин
    24: None,   # checko.ru → нет сайта у ИП Сучков
    25: None,   # saby.ru → нет сайта у ИП Чарыков
    26: None,   # znaemvseh.ru → нет сайта у ИП Шишков
    30: None,   # audit-it.ru → нет сайта у Молоток
    31: None,   # tbank.ru → нет сайта у Сибирь Дистрибуция (закрылись)
    32: None,   # tbank.ru → нет сайта у Сигма
    34: None,   # checko.ru → нет сайта у СТРОЙТЕ С НАМИ
    35: None,   # agro.gov74.ru → нет сайта у Хозстройкомплект
    43: None,   # tbank.ru → нет сайта у Левша
    44: None,   # saby.ru → нет сайта у Стройградторг
    50: None,   # saby.ru → нет сайта у Реал-Маркет (ликвидирован)
    51: None,   # reputation.ru → нет сайта у ИП Сергеев
    54: None,   # tbank.ru → нет сайта у ИП Воронков
    56: None,   # cataloxy.ru → нет сайта у Уютстрой
}

# Неверные email → исправить
WRONG_EMAILS = {
    15: None,   # sales_support@interfax.ru — это email Интерфакса, не Баумаркета
    46: None,   # info@marketpublishers.ru — это email MarketPublishers, не ЭПФ Пилона
}

# Дополнительные корректировки адресов
ADDRESS_FIXES = {
    7: '660077, Красноярский край, г. Красноярск, ул. Молокова, д. 40, помещ. 330',
    9: '140060, Московская обл., г. Люберцы',
    17: '660004, Красноярский край, г. Красноярск, пр-кт Газеты «Красноярский Рабочий», д. 27, офис 43',
    22: '150003, г. Ярославль, Кооперативная ул., д. 11',
    23: '646710, Омская обл., Шербакульский р-н, с. Борисовское',
    27: '620014, Свердловская обл., г. Екатеринбург, ул. 8 Марта, д. 179',
    28: '620014, Свердловская обл., г. Екатеринбург, ул. 8 Марта, д. 179',
    29: '121354, г. Москва, ул. Дорогобужская, д. 14, стр. 7, помещ. 3',
    31: '625002, Тюменская обл., г. Тюмень, Профсоюзная ул., д. 30, кв. 179',
    32: '129327, г. Москва, ул. Лазо, д. 9, помещ. 3/2',
    35: '454085, Челябинская обл., г. Челябинск, Шуменская ул., д. 31, 274',
    36: '623701, Свердловская обл., г. Берёзовский, ул. Гагарина, д. 6',
    37: '660064, Красноярский край, г. Красноярск, ул. Академика Вавилова, зд. 1г/1',
    39: '614025, Пермский край, г. Пермь, ул. Героев Хасана, зд. 109, каб. 218',
    41: '153011, Ивановская обл., г. Иваново, ул. Красных Зорь, д. 34',
    43: '622005, Свердловская обл., г. Нижний Тагил, ул. Индустриальная, д. 28',
    47: '660013, Красноярский край, г. Красноярск, Томская ул., зд. 4, к. 2, стр. 33',
    50: '165300, Архангельская обл., г. Котлас',
    52: '305038, Курская обл., г. Курск, ул. Косухина, д. 51',
    53: '119192, г. Москва, ул. Мосфильмовская, д. 74Б, офис 219',
}

# Дополнительные email из известных источников
EXTRA_EMAILS = {
    16: 'info@vimos.ru',         # Вимос
    27: 'info@gipermarketdom.ru', # ДОМ КорТрейд
    28: 'info@gipermarketdom.ru', # ДОМ Сити строй
    13: 'info@akson.pro',         # Аксон
}

wb = openpyxl.load_workbook(dst_file)
ws = wb['Лист1']

changes = 0

# 1. Убираем плохие сайты
for row, val in BAD_SITES.items():
    current = ws.cell(row=row, column=15).value
    if current:
        ws.cell(row=row, column=15, value=val)
        changes += 1

# 2. Исправляем неверные email
for row, val in WRONG_EMAILS.items():
    current = ws.cell(row=row, column=13).value
    if current:
        ws.cell(row=row, column=13, value=val)
        changes += 1

# 3. Исправляем адреса
for row, val in ADDRESS_FIXES.items():
    ws.cell(row=row, column=16, value=val)
    changes += 1

# 4. Добавляем недостающие email
for row, val in EXTRA_EMAILS.items():
    current = ws.cell(row=row, column=13).value
    if not current and val:
        ws.cell(row=row, column=13, value=val)
        changes += 1

wb.save(dst_file)
print(f"Внесено изменений: {changes}")

# Финальная статистика
filled = {'fo': 0, 'city': 0, 'email': 0, 'website': 0, 'address': 0}
total = 0
for row_num in range(3, ws.max_row + 1):
    if ws.cell(row=row_num, column=4).value:
        total += 1
        if ws.cell(row=row_num, column=2).value: filled['fo'] += 1
        if ws.cell(row=row_num, column=3).value: filled['city'] += 1
        if ws.cell(row=row_num, column=13).value: filled['email'] += 1
        if ws.cell(row=row_num, column=15).value: filled['website'] += 1
        if ws.cell(row=row_num, column=16).value: filled['address'] += 1

print(f"\n📋 ФИНАЛЬНАЯ СВОДКА ({total} компаний):")
labels = {'fo': 'Федеральный округ', 'city': 'Город', 'email': 'Email', 'website': 'Сайт', 'address': 'Почтовый адрес'}
for k, v in filled.items():
    pct = v*100//total
    print(f"  {labels[k]}: {v}/{total} ({pct}%)")
