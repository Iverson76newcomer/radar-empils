"""
Полный анализ файла Knauf_final_v6.docx
Проверка корректности информации и несоответствий
"""
from docx import Document
import re

DOC_PATH = "/home/z/my-project/upload/Knauf_final_v6.docx"
doc = Document(DOC_PATH)

print("=" * 110)
print("ПОЛНЫЙ АНАЛИЗ ФАЙЛА: Knauf_final_v6.docx")
print("=" * 110)

# 1. Document structure
print("\n## СТРУКТУРА ДОКУМЕНТА ##\n")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        style = para.style.name if para.style else "Normal"
        print(f"[P{i:03d}][{style}] {text}")

# 2. Tables
print("\n\n## ТАБЛИЦЫ ##\n")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- ТАБЛИЦА #{t_idx} ({len(table.rows)} строк × {len(table.columns)} колонок) ---")
    for r_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
        print(f"  R{r_idx}: {' || '.join(cells)}")

# 3. Check all known issues
print("\n\n" + "=" * 110)
print("ПРОВЕРКА КЛЮЧЕВЫХ ПАРАМЕТРОВ")
print("=" * 110)

# Issue 1: Employee count for ООО «Кнауф Гипс»
print("\n1. ЧИСЛЕННОСТЬ ООО «КНАУФ ГИПС»:")
emp_patterns = [
    (r"3\s*785|3785", "3 785"),
    (r"~?\s*2\s*100|2\s*099|2\s*296|2\s*027", "Revelio Labs (~2 100)"),
    (r"~?\s*1\s*500|~?\s*1\s*400|~?\s*1\s*450|~?\s*1\s*480", "1 400-1 500"),
    (r"~?\s*4\s*500|4500", "4 500"),
    (r"6\s*000|7\s*000", "6 000-7 000 (общая в РФ)"),
]
for pat, name in emp_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 2: Petrobord financials
print("\n2. ФИНАНСЫ ПЕТРОБОРД:")
petro_patterns = [
    (r"1[,.]?8\s*млрд", "1,8 млрд (старое)"),
    (r"176\s*млн", "176 млн (старое)"),
    (r"15[,.]?23", "15,23 млрд (новое)"),
    (r"1[,.]?33\s*млрд", "1,33 млрд (новое)"),
]
for pat, name in petro_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 3: Flizgrund
print("\n3. ФЛИЗГРУНД (Flizgrund):")
fliz_found = False
for pat in [r"[Фф]лизгрунд", r"Flizgrund", r"flizgrund"]:
    for i, para in enumerate(doc.paragraphs):
        if re.search(pat, para.text, re.IGNORECASE):
            print(f"   P{i}: {para.text}")
            fliz_found = True
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if re.search(pat, cell.text, re.IGNORECASE):
                    print(f"   T{t_idx}R{r_idx}C{c_idx}: {cell.text}")
                    fliz_found = True
if not fliz_found:
    print("   ✅ Полностью удалено")

# Issue 4: Market size 2,13 vs 2-13
print("\n4. РЫНОК ЛКМ:")
ms_patterns = [
    (r"2\s*-\s*13", "2-13 (TYPO)"),
    (r"2[,.]?\s*13\s*млн", "2,13 млн (правильно)"),
]
for pat, name in ms_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 5: Forbes / 100 млрд / 43-48 млрд
print("\n5. ВЫРУЧКА ГРУППЫ / FORBES:")
fb_patterns = [
    (r"100[,.]?\s*млрд", "100 млрд (Forbes)"),
    (r"43[,.]?\s*48\s*млрд|43-48", "43-48 млрд (4 юрлица)"),
    (r"Forbes|Форбс", "Forbes"),
]
for pat, name in fb_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text, re.IGNORECASE):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 6: Profit margin 46,9% / testfirm
print("\n6. НОРМА ПРИБЫЛИ 46,9% / ИСТОЧНИК:")
pm_patterns = [
    (r"46[,.]?9", "46,9%"),
    (r"testfirm", "testfirm.ru"),
    (r"bo\.nalog", "bo.nalog.gov.ru"),
    (r"audit-it", "audit-it.ru"),
]
for pat, name in pm_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text, re.IGNORECASE):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 7: Format consistency "н/д" / "—" / "-"
print("\n7. ФОРМАТЫ ОТСУТСТВИЯ ДАННЫХ:")
all_cells_text = []
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            all_cells_text.append(cell.text.strip())

format_patterns = [
    (r"^н/д$", "н/д (правильно)"),
    (r"^—$", "— (длинное тире)"),
    (r"^-$", "- (дефис)"),
    (r"^$", "(пустая ячейка)"),
]
for pat, name in format_patterns:
    matches = [t for t in all_cells_text if re.search(pat, t)]
    if matches:
        print(f"   '{name}': {len(matches)} ячеек")

# Issue 8: Шпаклёвки count
print("\n8. ШПАКЛЁВКИ:")
shpak_patterns = [
    (r"4\s*продукт|включает\s*4", "4 продукта"),
    (r"1\s*SKU\s*шпаклёв|1 SKU шпаклев", "1 SKU шпаклёв"),
    (r"3\s*сухих|3\s*сыпуч", "3 сухих/сыпучих"),
]
for pat, name in shpak_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 9: Number of plants (20 vs 14)
print("\n9. КОЛИЧЕСТВО ЗАВОДОВ:")
plant_patterns = [
    (r"20\s*производств|20\s*завод", "20 заводов"),
    (r"14\s*завод|14\s*произв", "14 заводов"),
]
for pat, name in plant_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(pat, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-60):m.end()+60]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(pat, cell.text, re.IGNORECASE):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   '{name}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt[:200]}")

# Issue 10: Duplicate tables check
print("\n10. ДУБЛИКАТЫ ТАБЛИЦ:")
table_hashes = []
for t_idx, table in enumerate(doc.tables):
    rows_text = []
    for row in table.rows:
        rows_text.append(" | ".join([cell.text.strip() for cell in row.cells]))
    table_hash = "\n".join(rows_text)
    table_hashes.append((t_idx, table_hash[:300], len(table.rows), len(table.columns)))

seen = {}
duplicates_found = False
for t_idx, hash_str, rows, cols in table_hashes:
    key = hash_str[:150]
    if key in seen:
        print(f"   ⚠️  ДУБЛИКАТ: T{t_idx} совпадает с T{seen[key]}")
        duplicates_found = True
    else:
        seen[key] = t_idx
if not duplicates_found:
    print("   ✅ Дубликатов таблиц не найдено")

# Issue 11: Расчёт нормы прибыли
print("\n11. РАСЧЁТ НОРМЫ ПРИБЫЛИ (проверка):")
calc_tests = [
    ("Кнауф Гипс", 11.4, 24.21, 46.9),
    ("Кунгур", 0.240, 3.54, 6.8),
    ("Петроборд", 1.33, 15.23, 8.7),
]
for name, profit, revenue, expected_pct in calc_tests:
    calc_pct = (profit / revenue) * 100
    delta = abs(calc_pct - expected_pct)
    status = "✅" if delta < 0.5 else "❌"
    print(f"   {status} {name}: {profit} / {revenue} = {calc_pct:.2f}% (ожидается {expected_pct}%)")

print("\n\nАНАЛИЗ ЗАВЕРШЁН")
