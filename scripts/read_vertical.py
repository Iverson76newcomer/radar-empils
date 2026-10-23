import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Вертикаль.xlsx", data_only=True)

# === Лист "Белые пятна" ===
print("=" * 100)
print("ЛИСТ 'Белые пятна' (сводная информация о клиентах)")
print("=" * 100)
ws = wb["Белые пятна"]

# Header (row 1)
print("\n--- Header (row 1) ---")
for col_idx, cell in enumerate(ws[1], 1):
    print(f"  Col {col_idx} ({chr(64+col_idx) if col_idx<=26 else 'A'+chr(64+col_idx-26)}): {cell.value}")

# Find rows containing "Вертикаль" in client column (col D)
print("\n\n--- Все строки с 'Вертикаль' в колонке 'Клиент' (D) ---")
for r in range(2, ws.max_row + 1):
    client = ws.cell(row=r, column=4).value
    if client and 'вертикал' in str(client).lower():
        print(f"\n=== Row {r}: {client} ===")
        for col_idx, cell in enumerate(ws[r], 1):
            header = ws.cell(row=1, column=col_idx).value
            if cell.value is not None and str(cell.value).strip():
                print(f"  {header}: {cell.value}")

# === Лист "Программы" ===
print("\n\n" + "=" * 100)
print("ЛИСТ 'Программы'")
print("=" * 100)
ws2 = wb["Программы"]
for r in range(1, ws2.max_row + 1):
    print(f"\nRow {r}:")
    for col_idx, cell in enumerate(ws2[r], 1):
        if cell.value is not None and str(cell.value).strip():
            print(f"  Col {col_idx}: {cell.value}")
