"""
Финальная правка файла 1: исправляем «5 030 ₽» на «5 954 ₽» в Главном выводе.
Контекст уже был «5 030 ₽ (9 174 − 3 220)» после предыдущей правки цены TITAN,
поэтому исходная замена с «8 250» не сработала.
"""
from docx import Document

OUT = '/home/z/my-project/download/TITAN_ССВ_расчёт_пояснительная_справка.docx'
doc = Document(OUT)

replacements = [
    ('5 030 ₽ (9 174 − 3 220)', '5 954 ₽ (9 174 − 3 220)'),
    # Если есть другие «5 030» без контекста
    ('5 030 ₽', '5 954 ₽'),
]

def replace_text_in_element(elem, old, new):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    t_nodes = elem.findall('.//w:t', ns)
    if not t_nodes:
        return 0
    full_text = ''.join(t.text or '' for t in t_nodes)
    if old not in full_text:
        return 0
    count = full_text.count(old)
    new_text = full_text.replace(old, new)
    if t_nodes:
        t_nodes[0].text = new_text
        for t in t_nodes[1:]:
            t.text = ''
    return count

for old, new in replacements:
    c = 0
    for p in doc.paragraphs:
        c += replace_text_in_element(p._p, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    c += replace_text_in_element(p._p, old, new)
    print(f'  {old!r} → {new!r}: {c}')

doc.save(OUT)
print(f'OK saved: {OUT}')
