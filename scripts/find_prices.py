"""Найти цены для топ-5 новинок в прайс-листе 16.10.2026."""
import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

# === Загрузить прайс-лист 16.10.2026 ===
wb2 = load_workbook("/home/z/my-project/upload/16.10.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) \
            and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name, 'fасовка': h_val,
                'базовая': m_val, 'отгрузочная': n_val,
            })

# === Топ-5 новинок ===
top5 = [
    {'code': 88965, 'brand': 'Ореол', 'name': 'Краска для плит OSB Ореол (база А) 1 кг', 'fasovka': 1.0, 'vtor_rub': 164257},
    {'code': 88967, 'brand': 'Ореол', 'name': 'Краска для плит OSB Ореол (база А) 12 кг', 'fasovka': 12.0, 'vtor_rub': 48656},
    {'code': 88852, 'brand': 'PROFIWOOD', 'name': 'Антисептик кроющий PROFIWOOD горький шоколад 0.8 л 0.9 кг', 'fasovka': 0.9, 'vtor_rub': 19080},
    {'code': 89855, 'brand': 'Ореол', 'name': 'Краска ВД резиновая Ореол (База А) 1.3 кг', 'fasovka': 1.3, 'vtor_rub': 11729},
    {'code': 88857, 'brand': 'PROFIWOOD', 'name': 'Антисептик кроющий PROFIWOOD белый База А 2.1 л 2.6 кг', 'fasovka': 2.6, 'vtor_rub': 8665},
]

# Поиск цен
def norm(s):
    if s is None:
        return ''
    s = re.sub(r'[«»""„"\']', ' ', str(s))
    s = re.sub(r'\s*\d+(?:[.,]\d+)?\s*(?:кг|л|l)\s*', ' ', s, flags=re.IGNORECASE)
    s = re.sub(r'[.,;:!?()\[\]{}-]', ' ', s)
    return s.lower().strip()

print("=" * 110)
print("ПОИСК ЦЕН ДЛЯ ТОП-5 НОВИНОК В ПРАЙС-ЛИСТЕ 16.10.2026")
print("=" * 110)

results = []
for n in top5:
    print(f"\n--- Код {n['code']} | {n['brand']} | {n['name']} (фас={n['fasovka']} кг) ---")
    
    n_norm = norm(n['name'])
    target_fas = n['fasovka']
    
    # Поиск по подстроке
    matches = []
    for sn in wb2.sheetnames:
        if sn not in ['Ореол', 'PROFIWOOD', 'NEWTONE']:
            continue
        for p in price_dict[sn]:
            if p['name'] is None:
                continue
            p_norm = norm(p['name'])
            if not p_norm:
                continue
            
            # Извлечь ключевые слова
            keywords = [w for w in n_norm.split() if len(w) > 3]
            matched_kw = sum(1 for k in keywords if k in p_norm)
            
            if matched_kw >= 2:  # хотя бы 2 ключевых слова совпали
                matches.append((sn, p, matched_kw))
    
    # Sort by number of matched keywords
    matches.sort(key=lambda x: -x[2])
    
    if matches:
        print(f"  Найдено {len(matches)} кандидатов:")
        for sn, p, mk in matches[:5]:
            print(f"    [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']} | совпадений={mk}")
        
        # Выбрать точное совпадение по фасовке
        chosen = None
        for sn, p, mk in matches:
            if isinstance(p['fасовка'], (int, float)) and abs(p['fасовка'] - target_fas) <= 0.3:
                chosen = (sn, p, mk)
                break
        
        # Если нет точного совпадения, взять наименьшую фасовку
        if not chosen:
            sorted_m = sorted([m for m in matches if isinstance(m[1]['fасовка'], (int, float))],
                              key=lambda m: m[1]['fасовка'])
            if sorted_m:
                chosen = sorted_m[0]
        
        if chosen:
            sn, p, mk = chosen
            print(f"\n  ✓ ВЫБРАНО [{sn}]: '{p['name']}'")
            print(f"    фасовка={p['fасовка']} | отгрузочная={p['отгрузочная']} | базовая={p['базовая']}")
            results.append({
                'code': n['code'], 'brand': n['brand'], 'name': n['name'],
                'fasovka_target': target_fas,
                'price_name': p['name'], 'fasovka_price': p['fасовка'],
                'отгрузочная': p['отгрузочная'], 'базовая': p['базовая'],
                'лист': sn, 'vtor_rub': n['vtor_rub'],
            })
        else:
            print(f"  ✗ Не выбрано")
    else:
        print(f"  ✗ Не найдено")
        # Покажем все позиции этого бренда
        for sn in ['Ореол', 'PROFIWOOD', 'NEWTONE']:
            if sn in price_dict:
                relevant = [p for p in price_dict[sn] if p['name'] and any(kw in p['name'].lower() for kw in n_norm.split()[:3] if len(kw) > 2)]
                if relevant:
                    print(f"  Все '{sn}' позиции с похожими именами:")
                    for p in relevant[:5]:
                        print(f"    '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# === ИТОГ ===
print("\n" + "=" * 110)
print("ИТОГ: 5 новинок с ценами")
print("=" * 110)
for r in results:
    print(f"\n  Код {r['code']} | {r['brand']} | втор.руб={r['vtor_rub']:.0f}")
    print(f"  Ассортимент: {r['name']}")
    print(f"  Прайс:       {r['price_name']}")
    print(f"  Фасовка: {r['fasovka_price']} кг | Отгрузочная: {r['отгрузочная']} | Базовая: {r['базовая']} | Лист: {r['лист']}")

# Сохраним в JSON
import json
with open('/home/z/my-project/scripts/top5_novinki.json', 'w', encoding='utf-8') as f:
    json.dump(results, f, ensure_ascii=False, indent=2, default=str)
print(f"\nСохранено в top5_novinki.json")
