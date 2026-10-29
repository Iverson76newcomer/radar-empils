"""Прочитать данные аналитики по Бурятии."""
import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/ДФО + Бурятия ver2.xlsx", data_only=True)

# Лист 'база'
print("=== Лист 'база' ===")
ws = wb["база"]
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

# Лист 'Лист1'
print("\n=== Лист 'Лист1' ===")
ws = wb["Лист1"]
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")
