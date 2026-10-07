"""
Update slide 6 (Белые пятна):
1. Replace the map image (ppt/media/image8.png) with the new version
   that includes 3 new regions (Московская, Воронежская, Липецкая области)
2. Add 3 text rows below the map card with new region info
"""
import zipfile, shutil, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = "/home/z/my-project/upload/Презентация_для_ГД_final.pptx"
OUT = "/home/z/my-project/download/Презентация_для_ГД_final.pptx"
NEW_MAP = "/home/z/my-project/download/slides_titan/assets/white_spots_map_v2.png"

# ── Step 1: Replace map image in PPTX ────────────────────────────────
# The map is stored as ppt/media/image8.png in the PPTX archive
# We need to replace it with the new map

TMP_DIR = "/tmp/pptx_update_map"
if os.path.exists(TMP_DIR):
    shutil.rmtree(TMP_DIR)
os.makedirs(TMP_DIR)

# Extract all files
with zipfile.ZipFile(SRC, 'r') as zin:
    zin.extractall(TMP_DIR)

# Replace the map image
map_path = os.path.join(TMP_DIR, 'ppt', 'media', 'image8.png')
if os.path.exists(map_path):
    orig_size = os.path.getsize(map_path)
    shutil.copy(NEW_MAP, map_path)
    new_size = os.path.getsize(map_path)
    print(f"Map replaced: image8.png  ({orig_size} → {new_size} bytes)")
else:
    print(f"WARNING: image8.png not found!")
    # List all media files
    media_dir = os.path.join(TMP_DIR, 'ppt', 'media')
    if os.path.exists(media_dir):
        for f in sorted(os.listdir(media_dir)):
            print(f"  {f}  ({os.path.getsize(os.path.join(media_dir, f))} bytes)")

# Re-zip into OUT
if os.path.exists(OUT):
    os.remove(OUT)
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for root, _, files in os.walk(TMP_DIR):
        for f in files:
            full = os.path.join(root, f)
            arc = os.path.relpath(full, TMP_DIR).replace(os.sep, '/')
            zout.write(full, arc)
print(f"Saved: {OUT} ({os.path.getsize(OUT)} bytes)")


# ── Step 2: Add 3 text rows below the map using python-pptx ──────────
prs = Presentation(OUT)

# Find slide 6 (Белые пятна)
target_idx = None
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame and 'БЕЛЫЕ ПЯТНА' in shape.text_frame.text:
            target_idx = i
            break
    if target_idx is not None:
        break

print(f"\nSlide 'Белые пятна' at index {target_idx}")
slide6 = list(prs.slides)[target_idx]

# Palette (matches slide)
BG          = RGBColor(0xF5, 0xF2, 0xEC)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xD9, 0xD2, 0xC2)
DIVIDER     = RGBColor(0xEF, 0xE9, 0xDD)
INK         = RGBColor(0x2F, 0x3A, 0x48)
MUTE        = RGBColor(0x8A, 0x95, 0xA3)
RED         = RGBColor(0xC0, 0x00, 0x00)
SAND        = RGBColor(0xC9, 0xA7, 0x7C)
YELLOW      = RGBColor(0xFF, 0xD6, 0x0A)

def cm(v):
    return Emu(int(v * 360000))

def add_text(slide, x, y, w, h, text, *,
             size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
             va=MSO_ANCHOR.TOP, font="Inter", spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = va
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if spacing:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(spacing))
    return tb

def add_rect(slide, x, y, w, h, fill, *,
             line=None, line_w=6350, corner_radius=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner_radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Emu(line_w)
    else:
        s.line.fill.background()
    if corner_radius is not None:
        try: s.adjustments[0] = corner_radius
        except: pass
    s.shadow.inherit = False
    return s

# ── Add 3 new region rows below the map card ─────────────────────────
# Map card ends at y=12.19cm, slide height = 14.29cm
# Space available: y=12.3 to y=14.2 = 1.9cm
# 3 rows × ~0.55cm each = 1.65cm (fits)

# Section label
add_text(slide6, cm(0.5), cm(12.35),
         cm(16.0), cm(0.35),
         "НОВЫЕ РЕГИОНЫ (потенциальные партнёры)",
         size=8, bold=True, color=SAND, font="Inter", spacing=200)

# 3 compact rows
NEW_REGIONS = [
    ("06", "Московская обл.", "ООО Вертикаль, ООО Стройтерминал центр красок", YELLOW),
    ("07", "Воронежская обл.", "ООО Вертикаль", YELLOW),
    ("08", "Липецкая обл.", "ООО Стройбаза-спектр", YELLOW),
]

row_y = cm(12.75)
row_h = cm(0.45)

for num, region, distributor, accent in NEW_REGIONS:
    # Number badge (small)
    add_text(slide6, cm(0.5), row_y,
             cm(0.8), row_h,
             num, size=10, bold=True, color=accent,
             font="Roboto Mono", va=MSO_ANCHOR.MIDDLE)

    # Region name
    add_text(slide6, cm(1.4), row_y,
             cm(4.5), row_h,
             region, size=10, bold=True, color=INK,
             font="Inter", va=MSO_ANCHOR.MIDDLE)

    # Distributor name
    add_text(slide6, cm(6.0), row_y,
             cm(10.5), row_h,
             distributor, size=9, color=MUTE,
             font="Inter", va=MSO_ANCHOR.MIDDLE)

    # Separator line below row
    row_y += row_h
    add_rect(slide6, cm(0.5), row_y - cm(0.02),
             cm(16.0), cm(0.015), DIVIDER)

# Save
prs.save(OUT)
print(f"\nFinal saved: {OUT} ({os.path.getsize(OUT)} bytes)")
print("Slide 6 updated: map replaced + 3 new region rows added")
