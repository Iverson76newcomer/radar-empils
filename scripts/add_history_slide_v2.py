"""
Создаём новый слайд «История компании Эмпилс» с 4 историческими точками.
Стиль: горизонтальная временная линия с крупными серыми годами + оранжевые маркеры +
чередование блоков сверху/снизу. В бренд-цветах Эмпилс.
4 точки: 1923, 1977, 1991, 1998.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SRC = '/home/z/my-project/upload/Презентация_для_новых_партнёров.pptx'
OUT = '/home/z/my-project/download/Презентация_для_новых_партнёров_с_историей_v2.pptx'

# ---------- Цвета ----------
ORANGE      = RGBColor(0xF4, 0x76, 0x29)
DARK        = RGBColor(0x31, 0x31, 0x31)
GREY_LIGHT  = RGBColor(0xC0, 0xC0, 0xC0)
GREY_TEXT   = RGBColor(0x4A, 0x4A, 0x4A)
LINE_GREY   = RGBColor(0xD0, 0xD0, 0xD0)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
GREY_HEADER = RGBColor(0x65, 0x65, 0x65)

FONT_BOLD = 'Inter Extra Bold'
FONT_REG  = 'Inter Regular'

p = Presentation(SRC)
slide_w, slide_h = p.slide_width, p.slide_height  # 20" × 11.25"

blank_layout = p.slide_layouts[4]
slide = p.slides.add_slide(blank_layout)

# ---------- утилиты ----------
def add_rect(slide, left, top, width, height, fill_color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def add_text(slide, left, top, width, height, text, font_size=11, bold=False,
             color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font_name=FONT_REG,
             italic=False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    r = para.add_run()
    r.text = text
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font_name
    r.font.color.rgb = color
    return tb

def add_multiline(slide, left, top, width, height, lines, font_size=11, color=GREY_TEXT,
                  font_name=FONT_REG, line_spacing=1.3, align=PP_ALIGN.CENTER):
    """lines: list of (text, bold, color)"""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, (text, bold, c) in enumerate(lines):
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        r = para.add_run()
        r.text = text
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.name = font_name
        r.font.color.rgb = c
    return tb

def add_line(slide, x1, y1, x2, y2, color, width_pt=1.0, dash=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width_pt)
    if dash:
        ln = line.line._get_or_add_ln()
        prstDash = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(prstDash)
    line.shadow.inherit = False
    return line

# ============================================================
# ЗАГОЛОВОК
# ============================================================
add_text(slide, 0.55, 0.45, 14.0, 0.7,
         'ИСТОРИЯ КОМПАНИИ «ЭМПИЛС»',
         font_size=36, bold=True, color=GREY_HEADER,
         font_name=FONT_BOLD, anchor=MSO_ANCHOR.MIDDLE)

add_text(slide, 0.55, 1.10, 18.0, 0.4,
         'Ключевые вехи развития — от первых цинковых белил до современного производства',
         font_size=14, bold=False, color=GREY_TEXT, italic=True,
         font_name=FONT_REG, anchor=MSO_ANCHOR.TOP)

# Тонкая оранжевая линия под заголовком
add_rect(slide, 0.55, 1.55, 19.0, 0.04, ORANGE)

# ============================================================
# ГОРИЗОНТАЛЬНАЯ ВРЕМЕННАЯ ЛИНИЯ
# ============================================================
LINE_Y = 5.6
LINE_LEFT = 2.0
LINE_RIGHT = 18.0
LINE_WIDTH = LINE_RIGHT - LINE_LEFT

# Главная горизонтальная линия (оранжевая)
add_rect(slide, LINE_LEFT, LINE_Y - 0.025, LINE_WIDTH, 0.05, ORANGE)

# ============================================================
# 4 ИСТОРИЧЕСКИЕ ТОЧКИ
# ============================================================
# Позиции центров X для 4 точек — равномерно распределены по линии
year_positions = [2.0, 7.33, 12.66, 18.0]
year_labels = ['1923', '1977', '1991', '1998']

# Крупные серые годы как декоративные якоря
for x, year in zip(year_positions, year_labels):
    add_text(slide, x - 1.5, LINE_Y - 1.1, 3.0, 1.0,
             year, font_size=72, bold=True, color=GREY_LIGHT,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font_name=FONT_BOLD)

# Точки-маркеры (оранжевые круги на линии)
for x in year_positions:
    marker_size = 0.20
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                 Inches(x - marker_size/2),
                                 Inches(LINE_Y - marker_size/2),
                                 Inches(marker_size),
                                 Inches(marker_size))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ORANGE
    sh.line.color.rgb = WHITE
    sh.line.width = Pt(2.5)
    sh.shadow.inherit = False

# ============================================================
# СОБЫТИЯ — ЧЕРЕДОВАНИЕ СВЕРХУ/СНИЗУ
# ============================================================
events = [
    # 1923 — сверху
    (year_positions[0], True, 'Первые цинковые белила', [
        'Начало производства',
        'цинковых белил —',
        'базового пигмента',
        'для лакокрасочных',
        'материалов',
    ]),
    # 1977 — снизу
    (year_positions[1], False, 'РХПО им. Октябрьской революции', [
        'Создано Ростовское',
        'химическое производственное',
        'объединение им. ',
        'Октябрьской революции —',
        'крупнейшее в Юге РСФСР',
    ]),
    # 1991 — сверху
    (year_positions[2], True, 'Образовано ЗАО «Эмпилс»', [
        'Акционирование завода,',
        'сохранение команды',
        'и технологий,',
        'новый этап развития',
        'как независимой компании',
    ]),
    # 1998 — снизу
    (year_positions[3], False, 'Промсоюз «Новое Содружество»', [
        'Предприятие вошло',
        'в Промышленный Союз',
        '«Новое Содружество» —',
        'укрепление позиций',
        'в промышленности РФ',
    ]),
]

# Координаты блоков
BLOCK_WIDTH = 3.6
TOP_DESC_Y = 2.2
TOP_TITLE_Y = 4.7
BOT_TITLE_Y = 6.0
BOT_DESC_Y = 6.7

# Соединительные линии от точки к заголовку
for x, is_top, title, _ in events:
    if is_top:
        add_line(slide, x, LINE_Y - 0.10, x, TOP_TITLE_Y + 0.2, LINE_GREY, width_pt=0.75, dash=True)
    else:
        add_line(slide, x, LINE_Y + 0.10, x, BOT_TITLE_Y, LINE_GREY, width_pt=0.75, dash=True)

# Размещение заголовков и описаний
for x, is_top, title, desc_lines in events:
    if is_top:
        # Заголовок события (оранжевый, ближе к линии)
        add_text(slide, x - BLOCK_WIDTH/2, TOP_TITLE_Y, BLOCK_WIDTH, 0.6,
                 title, font_size=15, bold=True, color=ORANGE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, font_name=FONT_BOLD)
        # Описание (выше заголовка)
        lines_data = [(line, False, GREY_TEXT) for line in desc_lines]
        add_multiline(slide, x - BLOCK_WIDTH/2, TOP_DESC_Y, BLOCK_WIDTH, 2.5,
                      lines_data, font_size=11, color=GREY_TEXT, line_spacing=1.3)
    else:
        # Заголовок события (оранжевый, сразу под линией)
        add_text(slide, x - BLOCK_WIDTH/2, BOT_TITLE_Y, BLOCK_WIDTH, 0.6,
                 title, font_size=15, bold=True, color=ORANGE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, font_name=FONT_BOLD)
        # Описание (ниже заголовка)
        lines_data = [(line, False, GREY_TEXT) for line in desc_lines]
        add_multiline(slide, x - BLOCK_WIDTH/2, BOT_DESC_Y, BLOCK_WIDTH, 2.5,
                      lines_data, font_size=11, color=GREY_TEXT, line_spacing=1.3)

# ============================================================
# ДЕКОР — мелкие засечки на линии
# ============================================================
minor_ticks = 40
for i in range(1, minor_ticks):
    t = i / minor_ticks
    x = LINE_LEFT + t * LINE_WIDTH
    # Пропускаем близко к маркерам
    too_close = any(abs(x - xp) < 0.4 for xp in year_positions)
    if too_close:
        continue
    add_rect(slide, x - 0.015, LINE_Y - 0.015, 0.03, 0.03, LINE_GREY)

# ============================================================
# ФИНАЛЬНЫЙ БЛОК — диапазон лет + слоган
# ============================================================
# «1923 → 1998» как декоративная надпись
add_text(slide, 0.55, 9.7, 9.0, 0.5,
         '75 лет становления — от первых цинковых белил до вхождения в Промсоюз «Новое Содружество»',
         font_size=11, bold=False, color=GREY_TEXT, italic=True,
         font_name=FONT_REG, anchor=MSO_ANCHOR.MIDDLE)

# Большая оранжевая цифра «75» в правом нижнем углу
add_text(slide, 16.5, 9.3, 3.0, 1.4,
         '75',
         font_size=72, bold=True, color=ORANGE,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, font_name=FONT_BOLD)
add_text(slide, 16.5, 10.5, 3.0, 0.4,
         'лет истории',
         font_size=12, bold=False, color=GREY_TEXT,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP, font_name=FONT_REG)

# ---------- Сохранение ----------
p.save(OUT)
print(f'OK saved: {OUT}')
import os
print(f'Size: {os.path.getsize(OUT)} bytes')
print(f'Total slides: {len(p.slides)}')
