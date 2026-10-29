#!/usr/bin/env python3
"""
Слайд «Антикризисный план Б» — управление рисками дистрибьютора.
Что делать, если через 6 месяцев продаж уход с полки ниже заявленного.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# === COLORS ===
RED = RGBColor(0xC6, 0x28, 0x28)
DARK_RED = RGBColor(0x8E, 0x00, 0x00)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK_GREEN = RGBColor(0x0B, 0x6E, 0x31)
DARK = RGBColor(0x1A, 0x1F, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xE5, 0xE7, 0xEB)
ORANGE = RGBColor(0xF4, 0x79, 0x20)
YELLOW = RGBColor(0xFB, 0xBF, 0x24)
DARK_BLUE = RGBColor(0x1E, 0x3A, 0x5F)
LIGHT_GREEN = RGBColor(0xDC, 0xFC, 0xE7)
LIGHT_ORANGE = RGBColor(0xFF, 0xED, 0xD5)
LIGHT_RED = RGBColor(0xFE, 0xE2, 0xE2)
LIGHT_BLUE = RGBColor(0xDB, 0xEA, 0xFE)

OUTPUT = '/home/z/my-project/download/Слайд_план_Б_управление_рисками.pptx'

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK, bold=False, align=PP_ALIGN.LEFT, font_name='Arial',
                anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    return tb


def add_shape(slide, shape_type, left, top, width, height, fill_color=None,
              line_color=None, line_width=None):
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill_color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        if line_width is not None:
            shp.line.width = line_width
    shp.shadow.inherit = False
    return shp


def add_multi_text(slide, left, top, width, height, lines,
                   anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(1); tf.margin_right = Pt(1)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, (text, size, color, bold) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = 'Arial'
    return tb


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H,
              fill_color=WHITE, line_color=None)

    # === ШАПКА ===
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.65),
              fill_color=DARK_BLUE, line_color=None)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(0.65), SLIDE_W, Inches(0.04),
              fill_color=YELLOW, line_color=None)
    add_textbox(slide, Inches(0.4), Inches(0.08), SLIDE_W - Inches(0.8), Inches(0.5),
                '🛡  ПЛАН «Б» — УПРАВЛЕНИЕ РИСКАМИ',
                font_size=22, color=WHITE, bold=True,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, SLIDE_W - Inches(5.0), Inches(0.08),
                Inches(4.6), Inches(0.5),
                'Если через 6 мес. уход с полки ниже заявленного',
                font_size=11, color=YELLOW, bold=False,
                align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # === ВЕРХНЯЯ ПОЛОСА — обещание Эмпилс ===
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(0.4), Inches(0.85), SLIDE_W - Inches(0.8), Inches(0.65),
              fill_color=GREEN, line_color=None)
    add_textbox(slide, Inches(0.4), Inches(0.85), SLIDE_W - Inches(0.8), Inches(0.65),
                '✓  Эмпилс не бросит вас на полпути. План «Б» включается автоматически при отклонении от KPI > 15%',
                font_size=14, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # === ЛЕВАЯ ЧАСТЬ — 6 инструментов плана «Б» ===
    add_textbox(slide, Inches(0.4), Inches(1.70),
                Inches(7.0), Inches(0.40),
                '6 ИНСТРУМЕНТОВ ПЛАНА «Б» — ЧТО ДЕЛАЕМ',
                font_size=13, color=DARK_BLUE, bold=True, align=PP_ALIGN.LEFT)

    tools = [
        ('1', 'Дополнительное обучение',
         'Повторный выезд тренера Эмпилс к ТП дистрибьютора. Анализ ошибок, отработка возражений «по цене», разбор реальных кейсов из этой же сети.',
         RED),
        ('2', 'Корректировка ассортимента',
         'Анализ продаж по SKU за 6 мес. Замена медленных цветов на ходовые. Снятие неликвида, возврат по программе 90 дней. Подбор оптимальной матрицы под конкретный регион.',
         ORANGE),
        ('3', 'Дополнительные промо-акции',
         'Совместное промо 50/50 с Эмпилс. Скидка 7-10% на TITAN 2-4 недели. Бонус продавцу 100 ₽/банку. Акция «купи 2 — получи кисть в подарок».',
         GREEN),
        ('4', 'Мастер-классы в топовых точках',
         'Выезд РМ + тренера в 5-7 самых крупных РТТ. Демонстрация: TITAN vs ПФ-115 (выкрас, адгезия, скорость). Обучение продавцов прямо на месте, 45 мин.',
         DARK_BLUE),
        ('5', 'Усиление POS-материалов',
         'Дополнительные воблеры, шелфтокеры, выкрасы «до/после». Брендированный ценник «цена за год». Замена изношенных материалов.',
         RED),
        ('6', 'Личный куратор от Эмпилс',
         'Закрепление за проблемным дистрибьютором старшего менеджера. Еженедельные созвоны. Личный визит 1 раз в месяц до выхода на KPI.',
         ORANGE),
    ]

    tool_y = Inches(2.15)
    tool_h = Inches(0.76)
    for i, (num, name, desc, color) in enumerate(tools):
        cy = tool_y + tool_h * i
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        add_shape(slide, MSO_SHAPE.RECTANGLE,
                  Inches(0.4), cy, Inches(7.2), tool_h - Inches(0.04),
                  fill_color=bg, line_color=GRAY, line_width=Pt(0.3))
        # Номер
        add_shape(slide, MSO_SHAPE.OVAL,
                  Inches(0.50), cy + Inches(0.15),
                  Inches(0.40), Inches(0.40),
                  fill_color=color, line_color=None)
        add_textbox(slide, Inches(0.50), cy + Inches(0.15),
                    Inches(0.40), Inches(0.40),
                    num, font_size=14, color=WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # Название + описание
        add_multi_text(slide, Inches(1.05), cy + Inches(0.04),
                       Inches(6.45), tool_h - Inches(0.10),
                       [
                           (name, 11, color, True),
                           (desc, 9, DARK, False),
                       ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

    # === ПРАВАЯ ЧАСТЬ — триггеры + что делают конкуренты ===

    # (1) Триггеры плана «Б» — когда включается
    add_textbox(slide, Inches(7.8), Inches(1.70),
                Inches(5.5), Inches(0.40),
                'КОГДА ВКЛЮЧАЕТСЯ ПЛАН «Б»',
                font_size=13, color=DARK_BLUE, bold=True, align=PP_ALIGN.LEFT)

    triggers = [
        ('🔴 КРАСНЫЙ', 'Уход с полки < 50% от плана', 'Включение всех 6 инструментов + личный куратор', RED, LIGHT_RED),
        ('🟡 ЖЁЛТЫЙ', 'Уход с полки 50-70% от плана', 'Обучение + корректировка ассортимента + промо', ORANGE, LIGHT_ORANGE),
        ('🟢 ЗЕЛЁНЫЙ', 'Уход с полки 70-85% от плана', 'Дополнительные промо-акции + усиление POS', GREEN, LIGHT_GREEN),
    ]
    tr_y = Inches(2.15)
    tr_h = Inches(0.80)
    for i, (level, cond, action, color, bg) in enumerate(triggers):
        cy = tr_y + tr_h * i
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                  Inches(7.8), cy, Inches(5.4), tr_h - Inches(0.06),
                  fill_color=bg, line_color=color, line_width=Pt(1.5))
        # Уровень
        add_shape(slide, MSO_SHAPE.RECTANGLE,
                  Inches(7.8), cy, Inches(1.5), tr_h - Inches(0.06),
                  fill_color=color, line_color=None)
        add_textbox(slide, Inches(7.8), cy, Inches(1.5), tr_h - Inches(0.06),
                    level, font_size=11, color=WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # Условие + действие
        add_multi_text(slide, Inches(9.4), cy + Inches(0.04),
                       Inches(3.7), tr_h - Inches(0.14),
                       [
                           (cond, 10, DARK, True),
                           (action, 9, color, False),
                       ], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

    # (2) Сравнение с конкурентами — что НЕ делают другие
    add_textbox(slide, Inches(7.8), Inches(4.65),
                Inches(5.5), Inches(0.40),
                'ЧЕГО НЕ ДЕЛАЮТ КОНКУРЕНТЫ',
                font_size=13, color=DARK_BLUE, bold=True, align=PP_ALIGN.LEFT)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(7.8), Inches(5.10), Inches(5.4), Inches(2.00),
              fill_color=LIGHT_GRAY, line_color=DARK_BLUE, line_width=Pt(1.5))

    competitors = [
        ('✗  Tikkurila', 'Нет письменного антикризисного плана. Реагируют по обстоятельствам.'),
        ('✗  Dufa', 'Возврат неликвида — только брак. Неликвид по причине «не пошло» — проблема дистрибьютора.'),
        ('✗  Dali', 'Обучение — 1 раз при старте. Повторного нет. Мастер-классов в РТТ не проводят.'),
        ('✗  Husky Klondike', 'Нет личного куратора. Только горячая линия и email-поддержка.'),
        ('✓  TITAN (Эмпилс)', 'Письменный план «Б» с 6 инструментами + 3 уровнями срабатывания.'),
    ]
    comp_y = Inches(5.20)
    comp_h = Inches(0.36)
    for i, (brand, desc) in enumerate(competitors):
        cy = comp_y + comp_h * i
        # Подсветим последнюю строку (TITAN)
        if 'TITAN' in brand:
            add_shape(slide, MSO_SHAPE.RECTANGLE,
                      Inches(7.85), cy, Inches(5.30), comp_h,
                      fill_color=GREEN, line_color=None)
            add_textbox(slide, Inches(7.95), cy, Inches(5.20), comp_h,
                        f'{brand} — {desc}', font_size=10, color=WHITE, bold=True,
                        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_textbox(slide, Inches(7.95), cy, Inches(5.20), comp_h,
                        f'{brand} — {desc}', font_size=9, color=DARK, bold=False,
                        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)

    # === ФУТЕР ===
    add_textbox(slide, Inches(0.4), Inches(7.15),
                SLIDE_W - Inches(0.8), Inches(0.30),
                'Empils  •  TITAN  •  План «Б» — официальная часть программы. Не дополнительная услуга, а гарантия партнёрства.',
                font_size=9, color=GRAY, bold=False,
                align=PP_ALIGN.CENTER, italic=True)

    prs.save(OUTPUT)
    print(f'OK: {OUTPUT}')
    print(f'Слайдов: {len(prs.slides)}')


if __name__ == '__main__':
    main()
