import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v3.xlsx", data_only=True)
ws_det = wb["300 SKU — детально"]
ws_prices = wb["300 SKU с ценами"]

# Найти SKU по имени (не по idx)
print("=== Проверка заменённых SKU по имени ===")
replaced_keywords = [
    ('красная RAL 3020', '3 в 1'),
    ('влагоизолирующий', 'Грунт-концентрат'),
    ('чёрная', 'NEWTONE'),  # есть другие SKU с этим, уточним
]

for r in range(4, ws_det.max_row + 1):
    name = ws_det.cell(row=r, column=7).value or ''
    name_lower = name.lower()
    
    # #37 (ориг): "Грунт-эмаль РАСЦВЕТ 3 в 1 красная RAL 3020"
    if 'красная ral 3020' in name_lower and '3 в 1' in name_lower:
        print(f"\nСтрока {r}: {name[:80]}")
        print(f"  Сфера: {ws_det.cell(row=r, column=2).value}")
        print(f"  Основа: {ws_det.cell(row=r, column=4).value}")
        print(f"  Бренд: {ws_det.cell(row=r, column=5).value}")
        print(f"  Гр.код: {ws_det.cell(row=r, column=8).value}")
    
    # #150 (ориг): АКВАСТОП заменён на влагоизолирующий
    if 'влагоизолирующий' in name_lower and 'грунт-концентрат' in name_lower:
        print(f"\nСтрока {r}: {name[:80]}")
        print(f"  Сфера: {ws_det.cell(row=r, column=2).value}")
        print(f"  Основа: {ws_det.cell(row=r, column=4).value}")
        print(f"  Бренд: {ws_det.cell(row=r, column=5).value}")
        print(f"  Гр.код: {ws_det.cell(row=r, column=8).value}")
    
    # #169 (ориг): ABSOLUTE BLACK заменён на NEWTONE черная RAL 9005
    if 'newtone' in name_lower and 'черная ral 9005' in name_lower and '13.9' in name:
        print(f"\nСтрока {r}: {name[:80]}")
        print(f"  Сфера: {ws_det.cell(row=r, column=2).value}")
        print(f"  Основа: {ws_det.cell(row=r, column=4).value}")
        print(f"  Бренд: {ws_det.cell(row=r, column=5).value}")
        print(f"  Гр.код: {ws_det.cell(row=r, column=8).value}")

# Также найти новые SKU (idx 296-300)
print("\n=== Новые SKU (по grp_code) ===")
new_grp_codes = [63180, 79754, 73122, 7803, 72676]
for r in range(4, ws_det.max_row + 1):
    gc = ws_det.cell(row=r, column=8).value
    if gc in new_grp_codes:
        print(f"\nСтрока {r} (idx={ws_det.cell(row=r, column=1).value}):")
        print(f"  Имя: {ws_det.cell(row=r, column=7).value}")
        print(f"  Бренд: {ws_det.cell(row=r, column=5).value}")
        print(f"  Сфера: {ws_det.cell(row=r, column=2).value}")

# Проверить цены в листе "300 SKU с ценами"
print("\n=== Проверка цен для заменённых и новых SKU ===")
for r in range(5, ws_prices.max_row + 1):
    name = ws_prices.cell(row=r, column=7).value or ''
    name_lower = name.lower()
    
    if ('красная ral 3020' in name_lower and '3 в 1' in name_lower) or \
       ('влагоизолирующий' in name_lower and 'грунт-концентрат' in name_lower) or \
       ('newtone' in name_lower and 'черная ral 9005' in name_lower and '13.9' in name) or \
       (ws_prices.cell(row=r, column=8).value in new_grp_codes):
        idx = ws_prices.cell(row=r, column=1).value
        brand = ws_prices.cell(row=r, column=5).value
        price = ws_prices.cell(row=r, column=10).value
        base = ws_prices.cell(row=r, column=11).value
        fasovka = ws_prices.cell(row=r, column=12).value
        status = ws_prices.cell(row=r, column=14).value
        print(f"\n  Строка {r} | #{idx} | {brand}")
        print(f"  Имя: {name[:75]}")
        print(f"  Отгруз: {price} | Базовая: {base} | Фасовка: {fasovka}")
        print(f"  Статус: {status}")

# Подсчёт итогов
print("\n=== ИТОГИ ===")
total = 0
has_price = 0
for r in range(5, ws_prices.max_row + 1):
    if ws_prices.cell(row=r, column=1).value is not None:
        total += 1
        if ws_prices.cell(row=r, column=10).value is not None:
            has_price += 1
print(f"Всего SKU: {total}")
print(f"С ценой: {has_price} ({has_price*100//total}%)")
print(f"Без цены: {total - has_price}")
