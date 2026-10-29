"""
Analyze Knauf_Россия_аналитическая_справка v4.docx for:
1. Inconsistencies / errors / contradictory information
2. Cross-reference with previously discussed issues:
   - Flizgrund should be removed (not sold in Russia)
   - Employee counts: Кнауф Гипс=3,785, Кунгур=216, Колпино=234
   - Petrobord financials: 15.23 млрд ₽ revenue / 1.33 млрд ₽ profit (2024)
   - Other legal entities: "н/д" for missing data
   - Revenue 100 млрд should be Forbes estimate
   - Profit margin 46.9% source: testfirm.ru
   - Market size: 2,13 млн т/год (not 2-13)
3. SKIP price analysis (user explicitly excluded prices)
"""
from docx import Document
from pathlib import Path
import re
import json

DOC_PATH = "/home/z/my-project/upload/Knauf_Россия_аналитическая_справка v4.docx"

doc = Document(DOC_PATH)

print("=" * 100)
print("АНАЛИЗ ФАЙЛА: Knauf_Россия_аналитическая_справка v4.docx")
print("=" * 100)

# 1. Print full document text structure
print("\n## СТРУКТУРА ДОКУМЕНТА ##\n")

para_idx = 0
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    style = para.style.name if para.style else "Normal"
    print(f"[P{i:03d}][{style}] {text}")

# 2. Print all tables
print("\n\n## ТАБЛИЦЫ ##\n")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- ТАБЛИЦА #{t_idx} ---")
    for r_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
        print(f"  R{r_idx}: {' || '.join(cells)}")

# 3. Search for known issues
print("\n\n## ПРОВЕРКА ИЗВЕСТНЫХ ПРОБЛЕМ ##\n")

issues = []

# Get all text combined
all_text = "\n".join([p.text for p in doc.paragraphs])
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            all_text += "\n" + cell.text

# Issue 1: Flizgrund
print("1. Проверка Flizgrund (Флизгрунд):")
flizgrund_matches = []
for i, para in enumerate(doc.paragraphs):
    if re.search(r"[Фф]лизгрунд|Flizgrund|flizgrund", para.text):
        flizgrund_matches.append((i, para.text))
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if re.search(r"[Фф]лизгрунд|Flizgrund|flizgrund", cell.text):
                flizgrund_matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
if flizgrund_matches:
    print("   НАЙДЕНО:")
    for loc, txt in flizgrund_matches:
        print(f"   - {loc}: {txt}")
else:
    print("   НЕ НАЙДЕНО (удалено)")

# Issue 2: Employee counts
print("\n2. Проверка численности сотрудников:")
print("   Ищем: Кнауф Гипс=3785, Кунгур=216, Колпино=234")
patterns = [
    r"3\s*785|3785",
    r"1\s*500|1500",
    r"216",
    r"234",
    r"700",
    r"203",
]
for p in patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}':")
        for loc, txt in matches[:10]:
            print(f"     {loc}: {txt}")

# Issue 3: Petrobord financials
print("\n3. Проверка Petrobord (Петроборд):")
petrobord_patterns = [r"[Пп]етроборд|Petrobord", r"1[,.]?8\s*млрд", r"15[,.]?23", r"1[,.]?33", r"176"]
for p in petrobord_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}':")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 4: Other legal entities - н/д
print("\n4. Проверка 'н/д' для юрлиц без данных:")
nd_patterns = [r"н/д", r"н\.д", r"нет данных", r"не\s*доступ"]
for p in nd_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text, re.IGNORECASE):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 5: Revenue 100 млрд / Forbes
print("\n5. Проверка выручки 100 млрд / Forbes:")
rev_patterns = [r"100[,.]?\s*млрд|~?100", r"Forbes|Форбс", r"RuWiki|РуВики"]
for p in rev_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}':")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 6: Profit margin 46.9% / testfirm
print("\n6. Проверка нормы прибыли 46,9% / testfirm:")
pm_patterns = [r"46[,.]?9|46\.9", r"testfirm|тестфирм", r"bo\.nalog|налог"]
for p in pm_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text, re.IGNORECASE):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text, re.IGNORECASE):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}':")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 7: Market size 2,13 vs 2-13
print("\n7. Проверка объёма рынка 2,13 млн т/год:")
ms_patterns = [r"2[,\-]?\s*13", r"2\s*-\s*13", r"млн\s*т/год|млн\s*тонн"]
for p in ms_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}':")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 8: Other Кнауф entities
print("\n8. Проверка юрлиц Кнауф Инсулейшн/Кубань/Дзержинск:")
inst_patterns = [r"Инсулейшн", r"Кубань", r"Дзержинск"]
for p in inst_patterns:
    matches = []
    for i, para in enumerate(doc.paragraphs):
        for m in re.finditer(p, para.text):
            matches.append((f"P{i}", para.text[max(0, m.start()-50):m.end()+50]))
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for m in re.finditer(p, cell.text):
                    matches.append((f"T{t_idx}R{r_idx}C{c_idx}", cell.text))
    if matches:
        print(f"   Паттерн '{p}': {len(matches)} совпадений")
        for loc, txt in matches[:5]:
            print(f"     {loc}: {txt}")

# Issue 9: Duplicate tables check
print("\n9. Проверка дубликатов таблиц:")
table_hashes = []
for t_idx, table in enumerate(doc.tables):
    rows_text = []
    for row in table.rows:
        rows_text.append(" | ".join([cell.text.strip() for cell in row.cells]))
    table_hash = "\n".join(rows_text)
    table_hashes.append((t_idx, table_hash, len(table.rows), len(table.columns)))

# Find duplicates
seen = {}
for t_idx, hash_str, rows, cols in table_hashes:
    key = hash_str[:200]
    if key in seen:
        print(f"   ВОЗМОЖНЫЙ ДУБЛИКАТ: T{t_idx} совпадает с T{seen[key]} (первые 200 символов)")
    else:
        seen[key] = t_idx
        print(f"   T{t_idx}: {rows} строк × {cols} колонок, первые символы: {key[:80]}")

print("\n\nАНАЛИЗ ЗАВЕРШЁН")
