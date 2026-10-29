#!/usr/bin/env python3
"""
TITAN 3в1 — Инструкция за 5 шагов в формате PowerPoint.
Версия с Unicode-эмодзи для пиктограмм (более крупные и узнаваемые).
Слоган: "10 лет без перекраски — 10 лет экономии!"
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

# === COLORS ===
RED = RGBColor(0xC6, 0x28, 0x28)
DARK_RED = RGBColor(0x8E, 0x00, 0x00)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK_GREEN = RGBColor(0x0B, 0x6E, 0x31)
DARK = RGBColor(0x1A, 0x1F, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xF3, 0xF4, 0xF6)
YELLOW = RGBColor(0xFB, 0xBF, 0x24)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
BLUE = RGBColor(0x25, 0x63, 0xEB)
BLACK = RGBColor(0x00, 0x00, 0x00)

OUTPUT = '/home/z/my-project/download/Инструкция_TITAN_3в1_5_шагов.pptx'

# Slide size 16:9 (13.333 x 7.5 inches)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Пиктограммы (Unicode-эмодзи) — большие, яркие, узнаваемые
PICTOGRAMS = [
    ('🧹', 'ОЧИСТИТЬ'),       # метла
    ('🧴', 'ОБЕЗЖИРИТЬ'),     # флакон с распылителем
    ('🥄', 'ПЕРЕМЕШАТЬ'),     # ложка
    ('🖌️', 'ПОКРАСИТЬ'),      # кисть
    ('✅', 'ГОТОВО!'),        # галочка
]


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK, bold=False, align=PP_ALIGN.LEFT, font_name='Arial',
                anchor=MSO_ANCHOR.TOP):
    """Add a textbox with given text."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font_name
    return tb


def add_emoji_box(slide, left, top, width, height, emoji, font_size=120,
                  align=PP_ALIGN.CENTER):
    """Add a textbox with a single large emoji (uses Segoe UI Emoji on Windows)."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = emoji
    run.font.size = Pt(font_size)
    run.font.name = 'Segoe UI Emoji'
    return tb


def add_shape(slide, shape_type, left, top, width, height, fill_color=None,
              line_color=None, line_width=None):
    """Add a shape with fill and line."""
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill_color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        if line_width is not None:
            shp.line.width = line_width
    shp.shadow.inherit = False
    return shp


# ==============================================================
# СЛАЙД 1: ТИТУЛЬНЫЙ
# ==============================================================

def slide_title(prs):
    """Слайд 1: Титульный."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Белый фон
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H,
              fill_color=WHITE, line_color=None)

    # Красная верхняя плашка-заголовок
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(2.2),
              fill_color=RED, line_color=None)

    # Тонкая жёлтая полоса под красной плашкой
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(2.2), SLIDE_W, Inches(0.08),
              fill_color=YELLOW, line_color=None)

    # TITAN 3в1
    add_textbox(slide, Inches(0.5), Inches(0.4), SLIDE_W - Inches(1.0), Inches(1.0),
                'TITAN 3в1', font_size=54, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # Подзаголовок
    add_textbox(slide, Inches(0.5), Inches(1.45), SLIDE_W - Inches(1.0), Inches(0.6),
                'Инструкция за 5 шагов — 10 лет защиты',
                font_size=24, color=WHITE, bold=False,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # Большой слоган в центре
    add_textbox(slide, Inches(0.8), Inches(2.7), SLIDE_W - Inches(1.6), Inches(1.2),
                '10 лет без перекраски — 10 лет экономии!',
                font_size=40, color=GREEN, bold=True,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # Подпись под слоганом
    add_textbox(slide, Inches(1.0), Inches(3.9), SLIDE_W - Inches(2.0), Inches(0.6),
                'Соблюдайте 5 простых шагов — и покрытие прослужит 10 лет',
                font_size=18, color=DARK, bold=False,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # 5 кругов с цифрами и эмодзи в ряд
    picto_y = Inches(4.8)
    spacing = Inches(2.3)
    start_x = (SLIDE_W - spacing * 4 - Inches(1.4)) // 2  # центрирование

    for i, (emoji, name) in enumerate(PICTOGRAMS):
        cx = start_x + spacing * i + Inches(0.7)

        # Красный круг
        add_shape(slide, MSO_SHAPE.OVAL,
                  cx - Inches(0.7), picto_y,
                  Inches(1.4), Inches(1.4),
                  fill_color=RED, line_color=DARK_RED, line_width=Pt(2))

        # Цифра вверху круга
        add_textbox(slide, cx - Inches(0.7), picto_y + Inches(0.02),
                    Inches(1.4), Inches(0.30),
                    str(i + 1), font_size=14, color=WHITE, bold=True,
                    align=PP_ALIGN.CENTER)

        # Эмодзи в центре круга (большое)
        add_emoji_box(slide, cx - Inches(0.6), picto_y + Inches(0.30),
                      Inches(1.2), Inches(1.0), emoji, font_size=50)

        # Название шага под кругом
        add_textbox(slide, cx - Inches(1.0), picto_y + Inches(1.5),
                    Inches(2.0), Inches(0.4),
                    name, font_size=12, color=DARK, bold=True,
                    align=PP_ALIGN.CENTER)

    # Футер
    add_textbox(slide, Inches(0.5), Inches(7.0), SLIDE_W - Inches(1.0), Inches(0.4),
                'Empils  •  TITAN  •  Грунт-эмаль по ржавчине 3-в-1',
                font_size=11, color=GRAY, bold=False,
                align=PP_ALIGN.CENTER, font_name='Arial')

    return slide


# ==============================================================
# СЛАЙДЫ ШАГОВ 2-6
# ==============================================================

def slide_step(prs, step_num, emoji, title, lines, tip_text, table_rows=None):
    """Слайд для одного шага.
    step_num: 1..5
    emoji: '🧹'
    title: 'ОЧИСТИТЬ'
    lines: список строк описания (с префиксом 'GREEN:' для зелёного текста)
    tip_text: текст жёлтой подсказки
    table_rows: опционально — список (label, value) для таблицы
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Белый фон
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H,
              fill_color=WHITE, line_color=None)

    # Тонкая красная плашка вверху
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.6),
              fill_color=RED, line_color=None)

    # Название TITAN 3в1 слева вверху
    add_textbox(slide, Inches(0.4), Inches(0.10), Inches(4.0), Inches(0.4),
                'TITAN 3в1  •  Инструкция',
                font_size=14, color=WHITE, bold=True,
                align=PP_ALIGN.LEFT, font_name='Arial')

    # Номер шага справа вверху
    add_textbox(slide, SLIDE_W - Inches(2.0), Inches(0.10),
                Inches(1.6), Inches(0.4),
                f'Шаг {step_num} из 5',
                font_size=14, color=WHITE, bold=True,
                align=PP_ALIGN.RIGHT, font_name='Arial')

    # === ЛЕВАЯ ЧАСТЬ: Большой круг с эмодзи ===
    circle_cx = Inches(3.0)
    circle_cy = Inches(4.1)
    circle_size = Inches(3.8)

    # Красный круг
    add_shape(slide, MSO_SHAPE.OVAL,
              circle_cx - circle_size // 2,
              circle_cy - circle_size // 2,
              circle_size, circle_size,
              fill_color=RED, line_color=DARK_RED, line_width=Pt(3))

    # Цифра шага (большая, вверху круга)
    add_textbox(slide,
                circle_cx - circle_size // 2,
                circle_cy - circle_size // 2 + Inches(0.15),
                circle_size, Inches(0.7),
                str(step_num), font_size=36, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER)

    # Большой эмодзи в центре круга
    add_emoji_box(slide,
                  circle_cx - circle_size // 2 + Inches(0.3),
                  circle_cy - Inches(0.50),
                  circle_size - Inches(0.6), Inches(2.2),
                  emoji, font_size=180)

    # === ПРАВАЯ ЧАСТЬ: Заголовок + описание + подсказка ===
    right_x = Inches(5.5)
    right_w = SLIDE_W - right_x - Inches(0.5)

    # Заголовок шага (большой красный текст)
    title_color = GREEN if step_num == 5 else RED
    add_textbox(slide, right_x, Inches(1.0), right_w, Inches(0.9),
                title, font_size=40, color=title_color, bold=True,
                align=PP_ALIGN.LEFT, font_name='Arial')

    # Жёлтая разделительная полоса
    add_shape(slide, MSO_SHAPE.RECTANGLE,
              right_x, Inches(1.85),
              Inches(2.0), Inches(0.06),
              fill_color=YELLOW, line_color=None)

    # Описание (строки)
    desc_y = Inches(2.1)
    line_height = Inches(0.45)
    for i, line in enumerate(lines):
        is_green = line.startswith('GREEN:')
        clean_line = line[6:] if is_green else line
        add_textbox(slide, right_x, desc_y + line_height * i,
                    right_w, line_height,
                    clean_line, font_size=18,
                    color=GREEN if is_green else DARK,
                    bold=is_green,
                    align=PP_ALIGN.LEFT, font_name='Arial')

    # Жёлтая подсказка
    if tip_text:
        tip_y = desc_y + line_height * len(lines) + Inches(0.15)
        tip_h = Inches(0.75)
        # Жёлтый фон
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                  right_x, tip_y, right_w, tip_h,
                  fill_color=YELLOW, line_color=None)
        # Текст подсказки
        tb = slide.shapes.add_textbox(right_x + Inches(0.15),
                                      tip_y + Inches(0.05),
                                      right_w - Inches(0.30), tip_h - Inches(0.10))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(4); tf.margin_right = Pt(4)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = '!  ' + tip_text
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = DARK
        r.font.name = 'Arial'

    # Таблица (опционально)
    if table_rows:
        tbl_y = desc_y + line_height * len(lines) + Inches(1.10)
        tbl_h = Inches(0.65)
        col_w = right_w // len(table_rows)
        for i, (label, value) in enumerate(table_rows):
            x = right_x + col_w * i
            # Светло-серая плашка
            add_shape(slide, MSO_SHAPE.RECTANGLE,
                      x, tbl_y, col_w, tbl_h,
                      fill_color=LIGHT_GRAY, line_color=GRAY, line_width=Pt(0.5))
            # Текст: метка + значение
            tb = slide.shapes.add_textbox(x + Inches(0.05),
                                          tbl_y + Inches(0.05),
                                          col_w - Inches(0.10), tbl_h - Inches(0.10))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r1 = p.add_run()
            r1.text = label + '\n'
            r1.font.size = Pt(11)
            r1.font.bold = True
            r1.font.color.rgb = DARK
            r1.font.name = 'Arial'
            r2 = p.add_run()
            r2.text = value
            r2.font.size = Pt(13)
            r2.font.bold = True
            r2.font.color.rgb = RED
            r2.font.name = 'Arial'

    # Стрелка вниз (между шагами) — внизу справа
    if step_num < 5:
        add_shape(slide, MSO_SHAPE.DOWN_ARROW,
                  SLIDE_W - Inches(1.2), Inches(6.7),
                  Inches(0.5), Inches(0.4),
                  fill_color=GRAY, line_color=None)
        add_textbox(slide, SLIDE_W - Inches(2.0), Inches(6.75),
                    Inches(0.8), Inches(0.35),
                    'далее', font_size=10, color=GRAY, bold=False,
                    align=PP_ALIGN.RIGHT)

    # Футер
    add_textbox(slide, Inches(0.5), Inches(7.1), SLIDE_W - Inches(1.0), Inches(0.35),
                'Empils  •  TITAN  •  Грунт-эмаль 3-в-1  •  10 лет без перекраски — 10 лет экономии!',
                font_size=10, color=GRAY, bold=False,
                align=PP_ALIGN.CENTER, font_name='Arial')

    return slide


# ==============================================================
# СЛАЙД 7: ТСО — ИТОГ
# ==============================================================

def slide_tso(prs):
    """Слайд 7: ТСО — итоговый слоган."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Светло-зелёный фон
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H,
              fill_color=RGBColor(0xEC, 0xFD, 0xF5), line_color=None)

    # Зелёная плашка-заголовок
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(2.0),
              fill_color=GREEN, line_color=None)

    # Жёлтая полоса
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(2.0), SLIDE_W, Inches(0.08),
              fill_color=YELLOW, line_color=None)

    # TITAN 3в1 вверху
    add_textbox(slide, Inches(0.5), Inches(0.3), SLIDE_W - Inches(1.0), Inches(0.7),
                'TITAN 3в1  •  Итог', font_size=22, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # Большой слоган
    add_textbox(slide, Inches(0.5), Inches(0.9), SLIDE_W - Inches(1.0), Inches(1.0),
                '10 лет без перекраски — 10 лет экономии!',
                font_size=40, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, font_name='Arial')

    # === 3 КРУПНЫЕ ЦИФРЫ ===
    card1_x = Inches(0.8)
    card_y = Inches(2.6)
    card_w = Inches(3.7)
    card_h = Inches(3.5)

    # 1) 1 банка вместо 3
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              card1_x, card_y, card_w, card_h,
              fill_color=WHITE, line_color=GREEN, line_width=Pt(2))
    add_textbox(slide, card1_x, card_y + Inches(0.3),
                card_w, Inches(1.4),
                '1', font_size=120, color=GREEN, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card1_x, card_y + Inches(1.7),
                card_w, Inches(0.5),
                'банка вместо 3', font_size=22, color=DARK, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card1_x, card_y + Inches(2.3),
                card_w, Inches(1.0),
                'грунт + эмаль +\nпреобразователь ржавчины\nв одной банке',
                font_size=14, color=GRAY, bold=False,
                align=PP_ALIGN.CENTER)

    # 2) 61% экономии
    card2_x = Inches(4.8)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              card2_x, card_y, card_w, card_h,
              fill_color=WHITE, line_color=GREEN, line_width=Pt(2))
    add_textbox(slide, card2_x, card_y + Inches(0.3),
                card_w, Inches(1.4),
                '61%', font_size=100, color=GREEN, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card2_x, card_y + Inches(1.7),
                card_w, Inches(0.5),
                'экономия за 10 лет', font_size=22, color=DARK, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card2_x, card_y + Inches(2.3),
                card_w, Inches(1.0),
                'покрасил 1 раз —\nзабыл на 10 лет',
                font_size=14, color=GRAY, bold=False,
                align=PP_ALIGN.CENTER)

    # 3) 10 лет защиты
    card3_x = Inches(8.8)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              card3_x, card_y, card_w, card_h,
              fill_color=GREEN, line_color=DARK_GREEN, line_width=Pt(2))
    add_textbox(slide, card3_x, card_y + Inches(0.3),
                card_w, Inches(1.4),
                '10', font_size=120, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card3_x, card_y + Inches(1.7),
                card_w, Inches(0.5),
                'лет защиты', font_size=22, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER)
    add_textbox(slide, card3_x, card_y + Inches(2.3),
                card_w, Inches(1.0),
                'срок службы\nпокрытия без\nперекраски',
                font_size=14, color=WHITE, bold=False,
                align=PP_ALIGN.CENTER)

    # Итоговый слоган внизу
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
              Inches(0.8), Inches(6.4),
              SLIDE_W - Inches(1.6), Inches(0.8),
              fill_color=GREEN, line_color=None)
    add_textbox(slide, Inches(0.8), Inches(6.45),
                SLIDE_W - Inches(1.6), Inches(0.7),
                'TITAN 3в1 — 10 лет без перекраски, 10 лет экономии!',
                font_size=22, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, font_name='Arial',
                anchor=MSO_ANCHOR.MIDDLE)

    return slide


# ==============================================================
# СБОРКА ПРЕЗЕНТАЦИИ
# ==============================================================

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Слайд 1: Титульный
    slide_title(prs)

    # Слайд 2: Шаг 1 — Очистить
    slide_step(prs, 1, PICTOGRAMS[0][0], 'ОЧИСТИТЬ',
               ['Удалите рыхлую ржавчину и старую',
                'краску металлической щёткой или',
                'наждачной бумагой (Р40-Р80).'],
               'Плотная ржавчина — можно оставить. TITAN красит по ржавчине.')

    # Слайд 3: Шаг 2 — Обезжирить
    slide_step(prs, 2, PICTOGRAMS[1][0], 'ОБЕЗЖИРИТЬ',
               ['Протрите поверхность тряпкой с',
                'растворителем (ацетон, ксилол).',
                'Поверхность должна быть сухой.'],
               'Не используйте уайт-спирит. После обезжиривания не трогайте руками.')

    # Слайд 4: Шаг 3 — Перемешать
    slide_step(prs, 3, PICTOGRAMS[2][0], 'ПЕРЕМЕШАТЬ',
               ['Открыть банку и тщательно перемешать',
                'до однородного цвета. Со дна вверх.',
                'GREEN:Готово к нанесению — грунтовать не нужно!'],
               'TITAN 3в1 = грунт + эмаль + преобразователь ржавчины в одной банке.')

    # Слайд 5: Шаг 4 — Покрасить
    slide_step(prs, 4, PICTOGRAMS[3][0], 'ПОКРАСИТЬ',
               ['Нанесите 2-3 слоя кистью, валиком',
                'или распылителем. Расход: 80-100 г/м².',
                'Между слоями: 5 часов при +20 °C.'],
               'Не красьте под прямым солнцем и при влажности выше 80%.',
               table_rows=[
                   ('Кисть / валик', '5 ч'),
                   ('Распылитель', '8 ч'),
                   ('Температура', '+5...+30 °C'),
               ])

    # Слайд 6: Шаг 5 — Готово
    slide_step(prs, 5, PICTOGRAMS[4][0], 'ГОТОВО! 10 ЛЕТ',
               ['Покрытие сохнет 5 часов. Не кантовать',
                'до полного высыхания. Эксплуатация —',
                'GREEN:через 7 дней. Срок службы: 10 лет.'],
               'Первые 7 дней покрытие набирает прочность — избегайте царапин.')

    # Слайд 7: ТСО
    slide_tso(prs)

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    prs.save(OUTPUT)
    print(f'OK: {OUTPUT}')
    print(f'Слайдов: {len(prs.slides)}')


if __name__ == '__main__':
    main()
