"""Переименуем длинный лист и пересохраним файл."""
import openpyxl
from openpyxl import load_workbook
import os

OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_с_отгрузочными_ценами_v2.xlsx"

wb = load_workbook(OUT_PATH)
print(f"Текущие листы: {wb.sheetnames}")

# Переименовать длинный лист
old_name = "8 несопоставленных — альтернативы"
new_name = "8 SKU — альтернативы"
if old_name in wb.sheetnames:
    wb[old_name].title = new_name

wb.save(OUT_PATH)
print(f"\n✓ Файл пересохранён: {OUT_PATH}")
print(f"  Новые листы: {wb.sheetnames}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
