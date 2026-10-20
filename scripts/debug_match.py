import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
                'sheet': sn,
            })

# Test normalization
QUOTE_RE = re.compile(r'[«»""„"\']')
def strip_quotes(s):
    return QUOTE_RE.sub(' ', s)

FASOVKA_RE = re.compile(r'\s+\d+(?:\.\d+)?\s*(?:кг|л|l)\s*', re.IGNORECASE)

def norm_for_match(name):
    if name is None:
        return ''
    s = strip_quotes(str(name))
    s = FASOVKA_RE.sub(' ', s)
    s = s.lower()
    s = re.sub(r'\s+', ' ', s)
    return s.strip()

# Test cases
test_names = [
    'Эмаль с молотковым эффектом ОРЕОЛ TITAN коричневая 0.8 кг',
    'Грунт-эмаль "РАСЦВЕТ" 3 в 1 белая 0.9 кг',
    'Грунт-эмаль "РАСЦВЕТ" 3 в 1 белая 1.9 кг',
]
for n in test_names:
    print(f"\nОриг: '{n}'")
    print(f"Норм: '{norm_for_match(n)}'")
    # Tokenize
    tokens = set(re.findall(r'[a-zа-яё]+(?:-[a-zа-яё]+)?|\d+', norm_for_match(n)))
    print(f"Токены: {tokens}")

# Check if matching fails for these
print("\n=== Поиск 'Эмаль с молотковым эффектом ОРЕОЛ TITAN коричневая 0.8 кг' ===")
assort_norm = norm_for_match('Эмаль с молотковым эффектом ОРЕОЛ TITAN коричневая 0.8 кг')
print(f"Ассортимент норм: '{assort_norm}'")
a_tokens = set(re.findall(r'[a-zа-яё]+(?:-[a-zа-яё]+)?|\d+', assort_norm))
print(f"Ассортимент токены: {a_tokens}")

for sn in ['ОРЕОЛ TITAN', 'ОРЕОЛ TITAN(л)']:
    for p in price_dict[sn]:
        if p['name'] and 'молотк' in p['name'].lower():
            pn = norm_for_match(p['name'])
            b_tokens = set(re.findall(r'[a-zа-яё]+(?:-[a-zа-яё]+)?|\d+', pn))
            overlap = a_tokens & b_tokens
            smaller = min(len(a_tokens), len(b_tokens))
            ratio = len(overlap) / smaller if smaller else 0
            print(f"  Цена: '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")
            print(f"  Цена норм: '{pn}'")
            print(f"  Цена токены: {b_tokens}")
            print(f"  Overlap={overlap} ({len(overlap)}/{smaller}={ratio:.2f})")
            print(f"  substr: assort_in_price={assort_norm in pn}, price_in_assort={pn in assort_norm}")

# Check 'РАСЦВЕТ 3 в 1 белая'
print("\n=== Поиск 'Грунт-эмаль РАСЦВЕТ 3 в 1 белая' ===")
assort_norm = norm_for_match('Грунт-эмаль "РАСЦВЕТ" 3 в 1 белая 0.9 кг')
print(f"Ассортимент норм: '{assort_norm}'")
for p in price_dict['Расцвет']:
    if p['name'] and '3 в 1' in p['name'].lower():
        pn = norm_for_match(p['name'])
        print(f"  Цена: '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']} -> норм='{pn}'")
