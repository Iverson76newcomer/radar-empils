import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# Search for ПФ-115 in Расцвет
print("=== Расцвет: 'ПФ-115' ===")
for p in price_dict['Расцвет']:
    if p['name'] and 'пф-115' in p['name'].lower():
        print(f"  '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Search for ПФ-115 across all sheets
print("\n=== All sheets: 'ПФ-115' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'пф-115' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Search for "ПФ 115" (without hyphen)
print("\n=== All sheets: 'ПФ 115' (space) ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'пф 115' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Search for "эмаль пф"
print("\n=== All sheets: 'эмаль пф' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'эмаль пф' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Show all names in Расцвет that include "эмаль"
print("\n=== Расцвет: all names containing 'эмаль' ===")
seen = set()
for p in price_dict['Расцвет']:
    if p['name'] and 'эмаль' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}'")

# Search '3 в 1' across ALL sheets
print("\n=== All sheets: '3 в 1' OR '3в1' ===")
seen = set()
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and ('3 в 1' in p['name'].lower() or '3в1' in p['name'].lower()) and p['name'] not in seen:
            seen.add(p['name'])
            print(f"  [{sn}] '{p['name']}'")

# Check Простокрашено - look for ПФ-115
print("\n=== Простокрашено! БАУ: 'ПФ' ===")
seen = set()
for p in price_dict['Простокрашено! БАУ']:
    if p['name'] and 'пф' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}'")

# Check 'грунт-эмаль' in Расцвет
print("\n=== Расцвет: 'грунт-эмаль' ===")
seen = set()
for p in price_dict['Расцвет']:
    if p['name'] and 'грунт-эмаль' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}'")
