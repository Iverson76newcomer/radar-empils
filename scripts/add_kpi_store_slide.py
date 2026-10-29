"""
Статистика Лемана Про — добавляем 3-й слайд: KPI для одного магазина федеральной сети.
Открываем исходный pptx, добавляем слайд в том же визуальном языке (прямоугольники + TextBox),
сохраняем в новый файл в /home/z/my-project/download/.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = '/home/z/my-project/upload/Статистика Лемана Про.pptx'
OUT = '/home/z/my-project/download/Статистика_Лемана_Про_KPI_магазина.pptx'

# ---------- Палитра (из исходного слайда) ----------
C_BG       = RGBColor(0xFF, 0xFF, 0xFF)
C_HEADER   = RGBColor(0x1E, 0x3A, 0x5F)   # тёмно-синий хедер
C_LINE     = RGBColor(0xFB, 0xBF, 0x24)   # жёлтая линия под хедером
C_SECTION  = RGBColor(0xDB, 0xEA, 0xFE)   # светло-синий блок заголовка секции
C_SECTION_TXT = RGBColor(0x1E, 0x3A, 0x5F)
C_RED      = RGBColor(0xC6, 0x28, 0x28)   # этап 1 — критично
C_ORANGE   = RGBColor(0xF4, 0x79, 0x20)   # этап 2 — важно
C_GREEN    = RGBColor(0x16, 0xA3, 0x4A)   # этап 3 — желательно
C_GREY_LT  = RGBColor(0xE5, 0xE7, 0xEB)   # светло-серый (zebra)
C_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
C_TEXT     = RGBColor(0x1A, 0x1F, 0x36)   # тёмно-серый текст
C_PERIOD   = RGBColor(0x9C, 0xA3, 0xAF)   # серый текст периодичности
C_FOOTNOTE = RGBColor(0x6B, 0x72, 0x80)   # серый текст подписи

FONT_NAME = 'Arial'

p = Presentation(SRC)
slide_w, slide_h = p.slide_width, p.slide_height  # 10.0" x 6.25"

# Создаём пустой слайд (используем blank layout)
blank_layout = p.slide_layouts[6]
slide = p.slides.add_slide(blank_layout)

# ---------- Утилиты ----------
def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill_color
    if line_color is None:
        sh.line.fill.background()  # нет границы
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(0.5)
    sh.shadow.inherit = False
    return sh

def add_text(slide, left, top, width, height, text, font_size=11, bold=False,
             color=C_TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, italic=False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT_NAME
    r.font.color.rgb = color
    return tb

# ---------- 1. Фон слайда ----------
add_rect(slide, 0, 0, 10.0, 6.25, C_BG)

# ---------- 2. Верхняя тёмная полоса ----------
add_rect(slide, 0, 0, 10.0, 0.54, C_HEADER)
add_rect(slide, 0, 0.54, 10.0, 0.03, C_LINE)

# ---------- 3. Заголовок слайда ----------
add_text(slide, 0.30, 0.06, 6.0, 0.43,
         'KPI МАГАЗИНА СЕТИ', font_size=22, bold=True, color=C_WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
add_text(slide, 5.05, 0.04, 4.65, 0.46,
         '9 метрик по 3 этапам  •  для одной ТТ', font_size=12,
         color=C_WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 4. Подзаголовок-блок секции ----------
add_rect(slide, 0.30, 1.10, 9.40, 0.34, C_SECTION)
add_text(slide, 0.30, 1.08, 9.40, 0.33,
         'Параметры оценки одного магазина Лемана Про', font_size=16, bold=True,
         color=C_SECTION_TXT, anchor=MSO_ANCHOR.MIDDLE)

# Краткое пояснение под подзаголовком
add_text(slide, 0.30, 1.50, 9.40, 0.32,
         'Метрики применяются к одной торговой точке Лемана Про (не к сети в целом). '
         'Не включают показатели дистрибьютора: объём закупок, кол-во РТТ, охват ТП — они на отдельном слайде.',
         font_size=9, italic=True, color=C_FOOTNOTE, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 5. Таблица KPI ----------
# Сетка (как в исходном слайде 1):
# Колонка 1 (Этап):     left=0.30  width=0.60
# Колонка 2 (Метрика):  left=0.90  width=3.75
# Колонка 3 (Цель):     left=4.65  width=3.00
# Колонка 4 (Период.):  left=7.65  width=2.02
# Шапка на Y=2.00, высота 0.38
# 9 строк данных на Y=2.42, 2.84, 3.26, 3.68, 4.10, 4.52, 4.94, 5.36, 5.78
# Высота строк 0.42

HEAD_Y = 2.00
HEAD_H = 0.38
ROW_H = 0.42
ROWS_START = HEAD_Y + HEAD_H + 0.04  # = 2.42

# Шапка
header_cells = [
    (0.30, 0.60, 'Этап'),
    (0.90, 3.75, 'Метрика'),
    (4.65, 3.00, 'Целевое значение'),
    (7.65, 2.02, 'Периодичность'),
]
for left, w, title in header_cells:
    add_rect(slide, left, HEAD_Y, w, HEAD_H, C_HEADER)
    add_text(slide, left, HEAD_Y, w, HEAD_H, title, font_size=11, bold=True,
             color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Данные — 9 метрик
# Формат: (этап_число, цвет_этапа, метрика, цель, периодичность, цвет_цели)
metrics = [
    # Этап 1 — Выкладка и наличие (КРИТИЧНО)
    ('1', C_RED,    'Соблюдение планограммы TITAN',      '95% SKU по эталону',         'Еженедельно',  C_RED),
    ('1', C_RED,    'Наличие топ-5 цветов RAL',           '5 из 5 SKU в наличии',       'Еженедельно',  C_RED),
    ('1', C_RED,    'Out-of-Stock по топ-SKU',            '0 SKU отсутствует',          'Ежедневно',    C_RED),
    # Этап 2 — POSm, цена и промо (ВАЖНО)
    ('2', C_ORANGE, 'Комплект POSm на полке',             '5 из 5 видов POSm',          'Ежемесячно',   C_ORANGE),
    ('2', C_ORANGE, 'Соблюдение РРЦ (±10%)',              '100% SKU в коридоре',        'Ежедневно',    C_ORANGE),
    ('2', C_ORANGE, 'Внедрение промо сети',               '100% акций в срок',          'По акциям',    C_ORANGE),
    # Этап 3 — Активность и продажи (ЖЕЛАТЕЛЬНО)
    ('3', C_GREEN,  'Сканирования QR с шелфтокера',       '___ сканов / мес',           'Ежемесячно',   C_GREEN),
    ('3', C_GREEN,  'Оборот TITAN в магазине',            '+__ % к базовому периоду',   'Ежемесячно',   C_GREEN),
    ('3', C_GREEN,  'Обученность продавца TITAN',         '100 % тест сдан',            'К 2-му месяцу', C_GREEN),
]

for i, (stage, stage_color, metric, target, period, target_color) in enumerate(metrics):
    y = ROWS_START + i * ROW_H
    # Zebra: чётные строки светло-серые, нечётные белые
    row_bg = C_GREY_LT if i % 2 == 0 else C_WHITE
    # Колонка этапа — цветная с номером
    add_rect(slide, 0.30, y, 0.60, ROW_H, stage_color)
    add_text(slide, 0.30, y, 0.60, ROW_H, stage, font_size=14, bold=True,
             color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Колонка метрики
    add_rect(slide, 0.90, y, 3.75, ROW_H, row_bg)
    add_text(slide, 0.97, y + 0.10, 3.60, 0.22, metric, font_size=10, bold=True,
             color=C_TEXT, anchor=MSO_ANCHOR.MIDDLE)
    # Колонка цели
    add_rect(slide, 4.65, y, 3.00, ROW_H, row_bg)
    add_text(slide, 4.72, y + 0.10, 2.85, 0.22, target, font_size=10, bold=True,
             color=target_color, anchor=MSO_ANCHOR.MIDDLE)
    # Колонка периодичности
    add_rect(slide, 7.65, y, 2.02, ROW_H, row_bg)
    add_text(slide, 7.72, y + 0.10, 1.88, 0.22, period, font_size=9,
             color=C_PERIOD, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 6. Подвал: источник данных ----------
add_text(slide, 0.30, 5.78 + 0.42 + 0.02, 9.40, 0.30,
         'Источник данных: фотоотчёт ТП по графику  ·  CIS-Link по SKU  ·  отчёт Лемана Про по кэшбеку и промо  ·  тест продавца в LMS',
         font_size=8, italic=True, color=C_FOOTNOTE, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE)

# ---------- Сохранение ----------
p.save(OUT)
print(f'OK saved: {OUT}')
import os
print(f'Size: {os.path.getsize(OUT)} bytes')
print(f'Total slides: {len(p.slides)}')
