"""Сравнить партнёров из файла Москвы с общей таблицей."""
import openpyxl
from openpyxl import load_workbook

# === 1. Загрузить общую таблицу ===
wb_general = load_workbook("/home/z/my-project/upload/Belye_pyatna_obschaya_tablitsa_21_09_2026.xlsx", data_only=True)
ws_general = wb_general["Белые пятна"]

# Получить заголовки
headers_general = [ws_general.cell(row=1, column=c).value for c in range(1, ws_general.max_column + 1)]
print("=== Заголовки общей таблицы ===")
for i, h in enumerate(headers_general, 1):
    print(f"  Col {i}: {h}")

# Прочитать все строки
general_partners = []
for r in range(2, ws_general.max_row + 1):
    row_data = {}
    for c, h in enumerate(headers_general, 1):
        row_data[h] = ws_general.cell(row=r, column=c).value
    general_partners.append(row_data)

print(f"\nВсего партнёров в общей таблице: {len(general_partners)}")

# === 2. Загрузить файл Москвы ===
wb_msk = load_workbook("/home/z/my-project/upload/Белые пятна - Москва.xlsx", data_only=True)
ws_msk = wb_msk["Лист1"]

# Получить заголовки
headers_msk = []
for c in range(1, ws_msk.max_column + 1):
    val = ws_msk.cell(row=1, column=c).value
    headers_msk.append(val)
print("\n=== Заголовки файла Москвы ===")
for i, h in enumerate(headers_msk, 1):
    print(f"  Col {i}: {h}")

# Прочитать все строки
msk_partners = []
for r in range(2, ws_msk.max_row + 1):
    row_data = {}
    has_data = False
    for c, h in enumerate(headers_msk, 1):
        val = ws_msk.cell(row=r, column=c).value
        row_data[c] = val
        if val is not None:
            has_data = True
    if has_data:
        msk_partners.append(row_data)

print(f"\nВсего партнёров в файле Москвы: {len(msk_partners)}")

# === 3. Вывести содержимое файла Москвы ===
print("\n=== Содержимое файла Москвы ===")
print(f"{'#':<4} {'Col1':<40} {'Col2':<15} {'Col3':<15} {'Col4':<15} {'Col5':<20} {'Col6':<15} {'Col7':<15}")
for i, p in enumerate(msk_partners, 1):
    print(f"{i:<4} {str(p.get(1, ''))[:40]:<40} {str(p.get(2, ''))[:15]:<15} {str(p.get(3, ''))[:15]:<15} {str(p.get(4, ''))[:15]:<15} {str(p.get(5, ''))[:20]:<20} {str(p.get(6, ''))[:15]:<15} {str(p.get(7, ''))[:15]:<15}")

