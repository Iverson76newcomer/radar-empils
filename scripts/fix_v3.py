"""
Исправление: привязка цен по grp_code, а не по idx.
"""
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
import json
import os

OUT_PATH = "/home/z/my-project/download/Ассортимент_Олимп_300_SKU_финал_v3.xlsx"

# Загрузить предыдущие результаты
with open('/home/z/my-project/scripts/results.json', 'r', encoding='utf-8') as f:
    prev_results = json.load(f)

# Построить маппинг grp_code → price
prev_results_by_grp_code = {}
for r in prev_results:
    gc = r.get('grp_code')
    if gc and gc != 0 and r.get('отгрузочная') is not None:
        prev_results_by_grp_code[int(gc)] = r

# Гр.коды для замен
REPLACEMENTS_BY_GRP = {
    65199: {  # #37 - вишневая → красная RAL 3020
        'отгрузочная': 610.78, 'базовая': 763.48, 'фасовка': 1.9,
        'лист': 'Расцвет', 'match_type': 'replacement',
        'ценовое_имя': 'Грунт-эмаль "РАСЦВЕТ" 3 в 1 красная RAL 3020',
        'status': 'ЗАМЕНА: вишневая → красная RAL 3020 (та же фасовка 1.9 кг, ближайший цвет)',
    },
    66403: {  # #150 - АКВАСТОП → влагоизолирующий
        'отгрузочная': 315.58, 'базовая': 394.47, 'фасовка': 1,
        'лист': 'Ореол', 'match_type': 'replacement',
        'ценовое_имя': 'Грунт-концентрат влагоизолирующий водно-дисперсионный для наружных и внутренних работ "ОРЕОЛ"',
        'status': 'ЗАМЕНА: АКВАСТОП → влагоизолирующий грунт-концентрат ОРЕОЛ (та же фасовка 1 кг)',
    },
    89508: {  # #169 - ABSOLUTE BLACK → NEWTONE черная
        'отгрузочная': 2436.34, 'базовая': 3045.43, 'фасовка': 13.9,
        'лист': 'NEWTONE', 'match_type': 'replacement',
        'ценовое_имя': 'Краска для интерьеров моющаяся глубокоматовая «NEWTONE» черная RAL 9005',
        'status': 'ЗАМЕНА: ABSOLUTE BLACK → NEWTONE черная RAL 9005 (та же фасовка 13.9 кг)',
    },
}

# Гр.коды для новых SKU
NEW_ITEMS_BY_GRP = {
    63180: {  # #296
        'отгрузочная': 221.27, 'базовая': 276.59, 'фасовка': 3,
        'лист': 'NoName', 'match_type': 'new_from_abc',
        'ценовое_имя': 'Краска для садовых деревьев водно-дисперсионная полиакриловая',
        'status': 'НОВАЯ: топ-1 ABC-анализа, новая позиция (гр.код 63180)',
    },
    79754: {  # #297
        'отгрузочная': 179.02, 'базовая': 223.78, 'фасовка': 0.8,
        'лист': 'Простокрашено! БАУ', 'match_type': 'new_from_abc',
        'ценовое_имя': 'Эмаль ПФ-115 "ПРОСТОКРАШЕНО!" синяя БАУ',
        'status': 'НОВАЯ: топ-2 ABC-анализа (№Name Эмаль ПФ-115 синяя) — использована альтернатива Простокрашено! БАУ 0.8 кг',
    },
    73122: {  # #298
        'отгрузочная': 1503.92, 'базовая': 1879.9, 'фасовка': 13,
        'лист': 'СТРОЙКОМПЛЕКТ', 'match_type': 'new_from_abc',
        'ценовое_имя': 'Краска фасадная атмосферостойкая водно-дисперсионная полиакриловая СТРОЙКОМПЛЕКТ',
        'status': 'НОВАЯ: топ-3 ABC-анализа, новая позиция (новый бренд СТРОЙКОМПЛЕКТ)',
    },
    7803: {  # #299
        'отгрузочная': 2578.87, 'базовая': 3223.59, 'фасовка': 20,
        'лист': 'EMPILS ЭКОНОМ', 'match_type': 'new_from_abc',
        'ценовое_имя': 'Эмаль EMPILS ПФ-115 «ЭКОНОМ» чёрная',
        'status': 'НОВАЯ: топ-4 ABC-анализа, новая позиция (новый бренд EMPILS ЭКОНОМ)',
    },
    72676: {  # #300
        'отгрузочная': 1640.57, 'базовая': 2050.71, 'фасовка': 10,
        'лист': 'СТРОЙКОМПЛЕКТ', 'match_type': 'new_from_abc',
        'ценовое_имя': 'Эмаль ПФ-115 СТРОЙКОМПЛЕКТ чёрная',
        'status': 'НОВАЯ: топ-5 ABC-анализа, новая позиция (бренд СТРОЙКОМПЛЕКТ)',
    },
}

# Открыть файл и обновить лист "300 SKU с ценами"
wb = load_workbook(OUT_PATH)
ws_det = wb["300 SKU — детально"]

# Сначала удалить старый лист "300 SKU с ценами" и пересоздать
if "300 SKU с ценами" in wb.sheetnames:
    del wb["300 SKU с ценами"]
ws_prices = wb.create_sheet("300 SKU с ценами", 2)  # вставить после "300 SKU — детально"

# Стили
HEAD_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEAD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
BODY_FONT = Font(name="Calibri", size=10)
BODY_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
NUM_ALIGN = Alignment(horizontal="right", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
PRICE_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E79")
NEW_ROW_FILL = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
REPLACED_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Заголовок
ws_prices.merge_cells("A1:N1")
ws_prices["A1"] = "Стартовый ассортимент Олимп — 300 SKU с отгрузочными ценами (ФИНАЛ v3)"
ws_prices["A1"].font = TITLE_FONT
ws_prices["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws_prices.row_dimensions[1].height = 28

ws_prices.merge_cells("A2:N2")
ws_prices["A2"] = ("Источник цен: прайс-лист «16.10.2026 РОССИЯ с НДС, руб.» | "
                   "Удалено: 5 SKU РАСЦВЕТ без альтернатив | Заменено на альтернативы: 3 SKU | "
                   "Добавлено из ABC-анализа: 5 SKU | ИТОГО: 300 SKU, все с ценами")
ws_prices["A2"].font = SUBTITLE_FONT
ws_prices["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws_prices.row_dimensions[2].height = 18

# Шапка (строка 4)
HEADERS = [
    ("#", 5), ("Сфера", 14), ("Назначение", 28), ("Основа", 16),
    ("Бренд", 16), ("Тип продукции", 22), ("Наименование", 60),
    ("Гр. код", 9), ("ABC", 6),
    ("Отгрузочная цена, руб", 14), ("Базовая цена, руб", 13),
    ("Фасовка в прайсе, кг", 11), ("Лист прайса", 16), ("Статус сопоставления", 30),
]
for col_idx, (title, width) in enumerate(HEADERS, 1):
    cell = ws_prices.cell(row=4, column=col_idx, value=title)
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = HEAD_ALIGN
    cell.border = BORDER
    ws_prices.column_dimensions[get_column_letter(col_idx)].width = width
ws_prices.row_dimensions[4].height = 38

# Заполнить строки с ценами, ища по grp_code
row_i = 5
new_grps = set(NEW_ITEMS_BY_GRP.keys())
replaced_grps = set(REPLACEMENTS_BY_GRP.keys())

# Подсчёт для статистики
stats = {'matched_original': 0, 'replaced': 0, 'new': 0, 'no_price': 0}

for r in range(4, ws_det.max_row + 1):
    idx = ws_det.cell(row=r, column=1).value
    if idx is None:
        continue
    
    grp_code = ws_det.cell(row=r, column=8).value
    if grp_code is not None:
        grp_code = int(grp_code)
    
    item = {
        'idx': idx,
        'sphere': ws_det.cell(row=r, column=2).value,
        'purpose': ws_det.cell(row=r, column=3).value,
        'base': ws_det.cell(row=r, column=4).value,
        'brand': ws_det.cell(row=r, column=5).value,
        'type': ws_det.cell(row=r, column=6).value,
        'name': ws_det.cell(row=r, column=7).value,
        'grp_code': grp_code,
        'abc': ws_det.cell(row=r, column=9).value,
    }
    
    # Ищем цену по grp_code
    price_info = None
    
    # 1. Сначала проверяем замены
    if grp_code in REPLACEMENTS_BY_GRP:
        price_info = REPLACEMENTS_BY_GRP[grp_code]
        stats['replaced'] += 1
        is_replaced = True
    # 2. Потом новые SKU
    elif grp_code in NEW_ITEMS_BY_GRP:
        price_info = NEW_ITEMS_BY_GRP[grp_code]
        stats['new'] += 1
        is_replaced = False
    # 3. Потом оригинальные цены
    elif grp_code in prev_results_by_grp_code:
        prev = prev_results_by_grp_code[grp_code]
        price_info = {
            'отгрузочная': prev['отгрузочная'],
            'базовая': prev['базовая'],
            'фасовка': prev.get('фасовка_прайс'),
            'лист': prev.get('лист', ''),
            'match_type': prev.get('match_type', 'matched'),
            'ценовое_имя': prev.get('ценовое_имя', ''),
            'status': f"Цена: {prev.get('match_type', 'matched')}",
        }
        stats['matched_original'] += 1
        is_replaced = False
    else:
        # Нет цены
        price_info = None
        stats['no_price'] += 1
        is_replaced = False
    
    is_new = grp_code in new_grps
    
    vals = [
        item['idx'], item['sphere'], item['purpose'], item['base'],
        item['brand'], item['type'], item['name'],
        item['grp_code'], item['abc'],
        price_info['отгрузочная'] if price_info else None,
        price_info['базовая'] if price_info else None,
        price_info['фасовка'] if price_info else None,
        price_info['лист'] if price_info else '',
        price_info['status'] if price_info else 'не найдено в прайс-листе',
    ]
    
    fill = NEW_ROW_FILL if is_new else (REPLACED_FILL if is_replaced else None)
    
    for col_idx, val in enumerate(vals, 1):
        cell = ws_prices.cell(row=row_i, column=col_idx, value=val)
        cell.border = BORDER
        cell.font = BODY_FONT
        if col_idx == 1:
            cell.alignment = CENTER_ALIGN
        elif col_idx in (10, 11, 12):
            cell.alignment = NUM_ALIGN
            if col_idx == 10 and val is not None:
                cell.font = PRICE_FONT
                cell.number_format = '# ##0.00'
            elif col_idx == 11 and val is not None:
                cell.number_format = '# ##0.00'
            elif col_idx == 12 and val is not None:
                cell.number_format = '0.##'
        elif col_idx == 7:
            cell.alignment = BODY_ALIGN
        elif col_idx == 14:
            cell.alignment = BODY_ALIGN
            if is_new:
                cell.font = Font(name="Calibri", size=9, italic=True, color="375623")
            elif is_replaced:
                cell.font = Font(name="Calibri", size=9, italic=True, color="7F6000")
        else:
            cell.alignment = BODY_ALIGN
        
        if fill:
            cell.fill = fill
    
    if is_new:
        ws_prices.row_dimensions[row_i].height = 40
    elif is_replaced:
        ws_prices.row_dimensions[row_i].height = 36
    
    row_i += 1

# Закрепить заголовок и автофильтр
ws_prices.freeze_panes = "B5"
last_col = get_column_letter(14)
ws_prices.auto_filter.ref = f"A4:{last_col}{row_i - 1}"

print(f"Заполнено {row_i - 5} строк с ценами")
print(f"Статистика: {stats}")

# Сохранить
wb.save(OUT_PATH)
print(f"\n✓ Файл сохранён: {OUT_PATH}")
print(f"  Размер: {os.path.getsize(OUT_PATH):,} байт")
