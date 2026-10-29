"""Изучение файлов Wordstat."""
from openpyxl import load_workbook
from pptx import Presentation

print("=" * 100)
print("ФАЙЛ 1: Wordstat_узнаваемость_брендов.xlsx")
print("=" * 100)
wb1 = load_workbook('/home/z/my-project/upload/Wordstat_узнаваемость_брендов.xlsx', data_only=True)
for sh_name in wb1.sheetnames:
    ws = wb1[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions} | max_row={ws.max_row} | max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append('')
            elif isinstance(v, float):
                vals.append(f'{v:.2f}')
            else:
                vals.append(str(v))
        print(' | '.join(vals))

print("\n\n" + "=" * 100)
print("ФАЙЛ 2: Wordstat_TITAN.xlsx")
print("=" * 100)
wb2 = load_workbook('/home/z/my-project/upload/Wordstat_TITAN.xlsx', data_only=True)
for sh_name in wb2.sheetnames:
    ws = wb2[sh_name]
    print(f"\n--- Лист: {sh_name} | dim={ws.dimensions} | max_row={ws.max_row} | max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=False):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append('')
            elif isinstance(v, float):
                vals.append(f'{v:.2f}')
            else:
                vals.append(str(v))
        print(' | '.join(vals))

print("\n\n" + "=" * 100)
print("ФАЙЛ 3: Wordstat_полный_анализ_TITAN_ без компании.pptx")
print("=" * 100)
prs = Presentation('/home/z/my-project/upload/Wordstat_полный_анализ_TITAN_ без компании.pptx')
for i, slide in enumerate(prs.slides, 1):
    print(f"\n--- СЛАЙД {i} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                txt = para.text.strip()
                if txt:
                    print(txt)
        elif shape.has_table:
            tbl = shape.table
            for r in tbl.rows:
                cells = [c.text.strip() for c in r.cells]
                print(' | '.join(cells))
