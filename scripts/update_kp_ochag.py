"""
Обновить КП Очаг v2 с новыми ТО: Год 1=27, Год 2=50, Год 3=80.
"""
from docx import Document
from docx.shared import Pt, RGBColor
import os, shutil

SRC = "/home/z/my-project/upload/КП_Очаг_коммерческое_предложение_v2.docx"
OUT = "/home/z/my-project/download/КП_Очаг_коммерческое_предложение_v3.docx"

shutil.copy(SRC, OUT)
doc = Document(OUT)

def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        fn = p.runs[0].font.name
        fs = p.runs[0].font.size
        fb = p.runs[0].font.bold
        fc = p.runs[0].font.color.rgb if p.runs[0].font.color and p.runs[0].font.color.rgb else None
        for run in p.runs:
            run.text = ''
        lines = text.split('\n')
        p.runs[0].text = lines[0]
        if fn: p.runs[0].font.name = fn
        if fs: p.runs[0].font.size = fs
        if fb is not None: p.runs[0].font.bold = fb
        if fc: p.runs[0].font.color.rgb = fc
        for line in lines[1:]:
            run = p.add_run()
            run.add_break()
            run.text = line
            if fn: run.font.name = fn
            if fs: run.font.size = fs
            if fb is not None: run.font.bold = fb
            if fc: run.font.color.rgb = fc
    else:
        cell.text = text

# === Таблица 1 (index 0): PROFIT CARDS — 2×4 ===
# Row 0: numbers, Row 1: labels
t = doc.tables[0]
set_cell(t.rows[0].cells[0], '11,4 \nмлн ₽')   # Год 1 доход
set_cell(t.rows[0].cells[1], '23,0 \nмлн ₽')   # Год 2 доход
set_cell(t.rows[0].cells[2], '36,8 \nмлн ₽')   # Год 3 доход
set_cell(t.rows[0].cells[3], '71,2 \nмлн ₽')   # Итого доход

# === Таблица 2 (index 1): ТОВАРООБОРОТ ПО ГОДАМ — 5×4 ===
t = doc.tables[1]
set_cell(t.rows[1].cells[1], '27 млн ₽')
set_cell(t.rows[1].cells[2], '9 млн ₽')
set_cell(t.rows[2].cells[1], '50 млн ₽')
set_cell(t.rows[2].cells[2], '18 млн ₽')
set_cell(t.rows[3].cells[1], '80 млн ₽')
set_cell(t.rows[3].cells[2], '28 млн ₽')
set_cell(t.rows[3].cells[3], 'Полное покрытие 4% потенциала. \n+2 ЭТП')
set_cell(t.rows[4].cells[1], '157 млн ₽')
set_cell(t.rows[4].cells[2], '55 млн ₽')

# === Таблица 3 (index 2): БОНУСИРОВАНИЕ — 5×7 ===
t = doc.tables[2]
# Год 1
set_cell(t.rows[1].cells[1], '27 млн ₽')
set_cell(t.rows[1].cells[2], '1,08 млн ₽')
set_cell(t.rows[1].cells[3], '1,08 млн ₽')
set_cell(t.rows[1].cells[4], '0,27 млн ₽ (1%)')
set_cell(t.rows[1].cells[5], '2,43 млн ₽ (9%)')
set_cell(t.rows[1].cells[6], '11,4 млн ₽')
# Год 2
set_cell(t.rows[2].cells[1], '50 млн ₽')
set_cell(t.rows[2].cells[2], '2,0 млн ₽')
set_cell(t.rows[2].cells[3], '2,0 млн ₽')
set_cell(t.rows[2].cells[4], '1,0 млн ₽ (2%)')
set_cell(t.rows[2].cells[5], '5,0 млн ₽ (10%)')
set_cell(t.rows[2].cells[6], '23,0 млн ₽')
# Год 3
set_cell(t.rows[3].cells[1], '80 млн ₽')
set_cell(t.rows[3].cells[2], '3,2 млн ₽')
set_cell(t.rows[3].cells[3], '3,2 млн ₽')
set_cell(t.rows[3].cells[4], '2,4 млн ₽ (3%)')
set_cell(t.rows[3].cells[5], '8,8 млн ₽ (11%)')
set_cell(t.rows[3].cells[6], '36,8 млн ₽')
# ИТОГО
set_cell(t.rows[4].cells[1], '157 млн ₽')
set_cell(t.rows[4].cells[2], '6,28 млн ₽')
set_cell(t.rows[4].cells[3], '6,28 млн ₽')
set_cell(t.rows[4].cells[4], '3,67 млн ₽')
set_cell(t.rows[4].cells[5], '16,23 млн ₽')
set_cell(t.rows[4].cells[6], '71,2 млн ₽')

# === Таблица 6 (index 5): СВОДКА — 2×4 ===
t = doc.tables[5]
set_cell(t.rows[0].cells[3], 'ДОЛЯ ЭМПИЛС 4%')
set_cell(t.rows[1].cells[3], '80 млн ₽\nцель на 3 года')

# === Таблица 9 (index 8): ДОЛЯ ЭМПИЛС — 3×3 ===
t = doc.tables[8]
set_cell(t.rows[2].cells[0], 'Доля Эмпилс (4%)')
set_cell(t.rows[2].cells[1], '80 млн ₽')
set_cell(t.rows[2].cells[2], '1 030 т')

# === Таблица 11 (index 10): ЭТП — 6×2 ===
t = doc.tables[10]
set_cell(t.rows[3].cells[1], '16 млн ₽/год')
set_cell(t.rows[4].cells[1], '11 млн ₽/год')
set_cell(t.rows[5].cells[1], '27 млн ₽')

# === Обновить текстовые абзацы ===
for p in doc.paragraphs:
    if '170 млн' in p.text and '8%' in p.text and 'целевой' in p.text:
        for run in p.runs:
            run.text = ''
        p.runs[0].text = '80 млн ₽ — целевой товарооборот Эмпилс в Год 3 = 4% потенциала региона. Заработок партнёра: 28 млн ₽/год при наценке 35%.'
        break

for p in doc.paragraphs:
    if '700 РТТ' in p.text:
        for run in p.runs:
            run.text = ''
        p.runs[0].text = 'Покрытие АКБ: 330 РТТ. Дистрибьютор: 2 ЭТП + собственная команда торговых представителей.'
        break

for p in doc.paragraphs:
    if 'ДОЛЯ ЭМПИЛС' in p.text and '8%' in p.text:
        for run in p.runs:
            run.text = ''
        p.runs[0].text = '4. ДОЛЯ ЭМПИЛС — 4% РЫНКА'
        break

for p in doc.paragraphs:
    if '170 млн' in p.text and '700 РТТ' in p.text:
        for run in p.runs:
            run.text = ''
        p.runs[0].text = '80 млн ₽ — целевой товарооборот Эмпилс, достигаемый за 3 года. Покрытие АКБ: 330 РТТ.'
        break

# Also update forecast table (Таблица 12, index 11) — if it exists
# Check if there's a forecast table with growth rates
for ti, t in enumerate(doc.tables):
    for row in t.rows:
        for cell in row.cells:
            if '×2' in cell.text and '×4,9' in cell.text:
                # This is the forecast table
                set_cell(t.rows[2].cells[1], '50 млн ₽')
                set_cell(t.rows[2].cells[2], '×1,9')
                set_cell(t.rows[3].cells[1], '80 млн ₽')
                set_cell(t.rows[3].cells[2], '×1,6')
                print(f"  ✓ Обновлена таблица прогноза (index {ti})")
                break

# === Сохранить ===
doc.save(OUT)
print(f"\n✓ Файл сохранён: {OUT}")
print(f"  Размер: {os.path.getsize(OUT):,} байт")

# === Проверка ===
doc2 = Document(OUT)
print("\n=== ПРОВЕРКА ===\n")

print("Таблица 1 (Profit cards):")
for row in doc2.tables[0].rows:
    print(f"  {[c.text[:30] for c in row.cells]}")

print("\nТаблица 2 (ТО по годам):")
for row in doc2.tables[1].rows:
    print(f"  {[c.text[:30] for c in row.cells]}")

print("\nТаблица 3 (Бонусы):")
for row in doc2.tables[2].rows:
    print(f"  {[c.text[:25] for c in row.cells]}")

print("\nТаблица 6 (Сводка):")
for row in doc2.tables[5].rows:
    print(f"  {[c.text[:30] for c in row.cells]}")

print("\nТаблица 9 (Доля Эмпилс):")
for row in doc2.tables[8].rows:
    print(f"  {[c.text[:30] for c in row.cells]}")

print("\nТаблица 11 (ЭТП):")
for row in doc2.tables[10].rows:
    print(f"  {[c.text[:30] for c in row.cells]}")

# Check for forecast table with growth
for ti, t in enumerate(doc2.tables):
    for row in t.rows:
        for cell in row.cells:
            if '×1,9' in cell.text or '×1,6' in cell.text:
                print(f"\nТаблица {ti+1} (Прогноз):")
                for row in t.rows:
                    print(f"  {[c.text[:30] for c in row.cells]}")
                break

print("\nКлючевые абзацы:")
for p in doc2.paragraphs:
    if '80 млн' in p.text or '4%' in p.text or '330 РТТ' in p.text:
        print(f"  ✓ {p.text[:120]}")
