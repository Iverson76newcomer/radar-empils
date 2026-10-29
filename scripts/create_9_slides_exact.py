#!/usr/bin/env python3
"""
9 слайдов по дизайну из PNG-макетов.
Точный перенос дизайна в редактируемый PowerPoint.
"""

from pptx import Presentation
from pptx.util import Pt, Cm, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

OUT = '/home/z/my-project/download/Презентация_9_слайдов_точная.pptx'

# Цвета из макетов
ORANGE = RGBColor(0xF4, 0x76, 0x29)
DARK = RGBColor(0x31, 0x31, 0x31)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)

# Цвета карточек брендов
C_AVANT = RGBColor(0xE0, 0xE7, 0xFF)
C_NEWT = RGBColor(0xDB, 0xEA, 0xFE)
C_OREOL = RGBColor(0xFE, 0xF9, 0xC3)
C_TITAN = RGBColor(0xFE, 0xE2, 0xE2)
C_PROFI = RGBColor(0xDC, 0xFC, 0xE7)
C_RASCV = RGBColor(0xFF, 0xED, 0xD5)
C_OPTIMA = RGBColor(0xCC, 0xFB, 0xF1)
C_DISC = RGBColor(0xFE, 0xF3, 0xC7)
C_PROST = RGBColor(0xF3, 0xF4, 0xF6)
C_EMP = RGBColor(0xD1, 0xFA, 0xE5)
C_STROY = RGBColor(0xFE, 0xE5, 0xD3)
C_LIGHT_BG = RGBColor(0xF8, 0xF8, 0xF8)
C_DARK_BG = RGBColor(0x1A, 0x1A, 0x1A)

prs = Presentation()
prs.slide_width = Cm(33.87)
prs.slide_height = Cm(19.05)

def add_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])

def white_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background(); bg.shadow.inherit = False

def dark_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = C_DARK_BG
    bg.line.fill.background(); bg.shadow.inherit = False

def empils_logo(slide, x=Cm(26), y=Cm(0.8)):
    """Логотип EMPILS в правом верхнем углу с серой линией."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Cm(0.3), Cm(2.5), Cm(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = LIGHT_GRAY
    line.line.fill.background(); line.shadow.inherit = False
    tb = slide.shapes.add_textbox(x + Cm(2.7), y, Cm(5), Cm(0.8))
    tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = "EMPILS"
    r.font.name = 'Arial'; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = ORANGE

def add_text(slide, x, y, w, h, text, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT, font='Arial'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def add_card(slide, x, y, w, h, fill_color):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid(); card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = ORANGE; card.line.width = Pt(1)
    card.shadow.inherit = False
    return card

# ============================================================
# СЛАЙД 1: 2 900 т — Автоматизированный резервуарный парк
# ============================================================
slide = add_slide()
white_bg(slide)
empils_logo(slide)

# Левая часть — фото (используем заглушку с серым фоном)
photo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Cm(18), Cm(19.05))
photo.fill.solid(); photo.fill.fore_color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
photo.line.fill.background(); photo.shadow.inherit = False
add_text(slide, Cm(2), Cm(8), Cm(14), Cm(2), "[Фото резервуарного парка]", 14, False, GRAY, PP_ALIGN.CENTER)

# Правая часть — текст
add_text(slide, Cm(19.5), Cm(6), Cm(13), Cm(2), "2 900 т", 60, True, ORANGE, PP_ALIGN.LEFT)
add_text(slide, Cm(19.5), Cm(9), Cm(13), Cm(1.2), "Автоматизированный резервуарный парк", 22, True, BLACK, PP_ALIGN.LEFT)
add_text(slide, Cm(19.5), Cm(11), Cm(13), Cm(3),
    "Собственный резервуарный парк сырья гарантирует дистрибьюторам бесперебойные поставки — мы выполняем обязательства независимо от внешних факторов.",
    12, False, RGBColor(0x33, 0x33, 0x33), PP_ALIGN.LEFT)

# ============================================================
# СЛАЙД 2: Разделитель — ГЕОГРАФИЯ БИЗНЕСА
# ============================================================
slide = add_slide()
dark_bg(slide)
empils_logo(slide)
add_text(slide, Cm(3), Cm(7), Cm(28), Cm(2), "ГЕОГРАФИЯ БИЗНЕСА", 40, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(3), Cm(10), Cm(28), Cm(1.5), "Дистрибуция в России и экспорт в страны СНГ", 16, False, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 3: Разделитель — АССОРТИМЕНТ И ЦЕНОВОЕ ПОЗИЦИОНИРОВАНИЕ
# ============================================================
slide = add_slide()
dark_bg(slide)
empils_logo(slide)
add_text(slide, Cm(3), Cm(6.5), Cm(28), Cm(3), "АССОРТИМЕНТ\nИ ЦЕНОВОЕ\nПОЗИЦИОНИРОВАНИЕ", 36, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(3), Cm(12), Cm(28), Cm(1.5), "11 брендов · от эконом до премиум", 16, False, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 4: АССОРТИМЕНТНАЯ МАТРИЦА — 11 брендов
# ============================================================
slide = add_slide()
white_bg(slide)
empils_logo(slide)

add_text(slide, Cm(1), Cm(0.5), Cm(25), Cm(1), "АССОРТИМЕНТНАЯ МАТРИЦА", 28, True, DARK, PP_ALIGN.LEFT)
add_text(slide, Cm(1), Cm(1.8), Cm(28), Cm(0.6), "11 брендов · от эконом до премиум · полный спектр ЛКМ для любых поверхностей и задач", 11, False, GRAY, PP_ALIGN.LEFT)

brands = [
    ('AVANTGARDE', 'ПРЕМИУМ', 'Гибридные технологии: экологичность водных + износостойкость органических', ['Краска интерьерная моющаяся (база A/C)', 'Краска для фасадов (база A/C)', 'Лак яхтный / паркетный'], C_AVANT),
    ('NEWTONE', 'СРЕДНИЙ+', 'Передовые технологии и современные материалы', ['Краска интерьерная / фасадная (база A/C)', 'Грунтовка глубокого проникновения', 'Шпатлёвка суперфинишная'], C_NEWT),
    ('ОРЕОЛ', 'СРЕДНИЙ+', 'Краски, лаки, грунтовки — надёжный помощник в ремонте', ['Краска для стен и потолков моющаяся', 'Краска для плит OSB / фасадная', 'Грунтовка бетоноконтакт / шпатлёвка', 'Клей, олифа, растворитель'], C_OREOL),
    ('TITAN', 'СРЕДНИЙ+', 'Грунт-эмаль 3в1 — грунт + эмаль + преобразователь ржавчины', ['3в1 по ржавчине (10 цветов, RAL)', 'Быстросохнущая матовая (7 цветов)', 'Молотковая алкидно-стирольная (6 цветов)', '«Красишь 1 раз — служит 10 лет»'], C_TITAN),
    ('PROFIWOOD', 'СРЕДНИЙ', 'Профессиональная защита и декор древесины', ['Антисептик кроющий / лаковый', 'Защитно-декоративные покрытия', 'Лак паркетный / по дереву и камню', 'Огнебиозащитные составы'], C_PROFI),
    ('РАСЦВЕТ', 'СРЕДНИЙ / ЭКОНОМ', 'Классические ЛКМ по ГОСТ — эмали, грунт-эмали, грунтовки', ['Эмаль ПФ-115 (ГОСТ)', 'Грунт-эмаль по ржавчине', 'Эмаль с молотковым эффектом', 'Грунтовка ГФ-021 / растворители'], C_RASCV),
    ('ОПТИМА', 'ЭКСПОРТ', 'Покрытия на органической основе для стран ближнего зарубежья', ['Эмали цветные (яркие оттенки)', 'Грунтовка ГФ-021', 'Грунт-эмали', 'Оптимальное соотношение цены и качества'], C_OPTIMA),
    ('ОРЕОЛ Дисконт', 'ЭКОНОМ', 'Водно-дисперсионные покрытия по доступной цене', ['Краска для стен и потолков', 'Шпатлёвка масляно-клеевая', 'Грунтовка', '«Покрасил и забыл!»'], C_DISC),
    ('ПРОСТОКРАШЕНО!', 'ЭКОНОМ', 'Базовые ЛКМ для массового потребителя', ['Эмаль ПФ-115 (БАУЦЕНТР)', 'Краска водоэмульсионная', 'Грунтовка / олифа / лак', 'Растворитель'], C_PROST),
    ('EMPILS', 'ПРОМЫШЛЕННЫЕ', 'Покрытия промышленного назначения — для крупных предприятий', ['Огнебиозащитные составы', 'Защитно-декоративные покрытия', 'Краски-базы под колеровку', 'Партнёр: «Ростсельмаш»'], C_EMP),
    ('СТРОЙКОМПЛЕКТ', 'СТРОИТ. ФАСОВКА', 'Крупная фасовка для строительных объектов большой площади', ['Эмаль ПФ-115 (крупная фасовка)', 'Краска интерьерная / фасадная', 'Универсальность · оптимальная цена', 'Яркая цветовая палитра'], C_STROY),
]

cw = Cm(7.6); ch = Cm(6.5); gx = Cm(0.3); gy = Cm(0.3)
sx = Cm(0.8); sy = Cm(2.8)

for i, (name, seg, desc, prods, clr) in enumerate(brands):
    row = i // 4; col = i % 4
    if row == 2: col_offset = Cm(1.0)
    else: col_offset = Cm(0)
    x = sx + col_offset + col * (cw + gx)
    y = sy + row * (ch + gy)
    add_card(slide, x, y, cw, ch, clr)
    # Name
    tb = slide.shapes.add_textbox(x + Cm(0.2), y + Cm(0.15), cw - Cm(0.4), Cm(0.8))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = name
    r.font.name = 'Arial'; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = DARK
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = seg
    r2.font.name = 'Arial'; r2.font.size = Pt(8); r2.font.bold = True; r2.font.color.rgb = ORANGE
    # Underline
    ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Cm(0.2), y + Cm(1.2), Cm(2.5), Cm(0.03))
    ul.fill.solid(); ul.fill.fore_color.rgb = ORANGE; ul.line.fill.background(); ul.shadow.inherit = False
    # Desc
    add_text(slide, x + Cm(0.2), y + Cm(1.4), cw - Cm(0.4), Cm(1.0), desc, 8, False, BLACK)
    # Products
    tb = slide.shapes.add_textbox(x + Cm(0.2), y + Cm(2.5), cw - Cm(0.4), Cm(4.0))
    tf = tb.text_frame; tf.word_wrap = True
    for pi, prod in enumerate(prods):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.space_before = Pt(1); p.space_after = Pt(1)
        r = p.add_run(); r.text = f"•  {prod}"
        r.font.name = 'Arial'; r.font.size = Pt(8); r.font.color.rgb = BLACK

# Footer
ft = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Cm(18.0), prs.slide_width, Cm(1.0))
ft.fill.solid(); ft.fill.fore_color.rgb = DARK; ft.line.fill.background(); ft.shadow.inherit = False
add_text(slide, Cm(1), Cm(17.9), Cm(32), Cm(1), "11 брендов покрывают весь спектр: от эконом до премиум · металл · дерево · интерьер · фасад · промышленные · экспорт   |   230 ₽/кг — зарабатываем вместе.", 10, True, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 5: Разделитель — ПОРТФОЛИО: ОБЪЕКТЫ И B2B
# ============================================================
slide = add_slide()
dark_bg(slide)
empils_logo(slide)
add_text(slide, Cm(3), Cm(6.5), Cm(28), Cm(2), "ПОРТФОЛИО:\nОБЪЕКТЫ И B2B", 36, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(3), Cm(11), Cm(28), Cm(1.5), "Реальные кейсы и промышленные партнёры", 16, False, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 6: Разделитель — ПОДДЕРЖКА
# ============================================================
slide = add_slide()
dark_bg(slide)
empils_logo(slide)
add_text(slide, Cm(3), Cm(7.5), Cm(28), Cm(2), "ПОДДЕРЖКА", 44, True, ORANGE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 7: КОМАНДА ЭМПИЛС НА ВАШЕЙ ТЕРРИТОРИИ
# ============================================================
slide = add_slide()
white_bg(slide)
empils_logo(slide)

add_text(slide, Cm(1), Cm(0.5), Cm(25), Cm(1), "КОМАНДА ЭМПИЛС НА ВАШЕЙ ТЕРРИТОРИИ", 24, True, DARK, PP_ALIGN.LEFT)
add_text(slide, Cm(1), Cm(1.7), Cm(30), Cm(0.6), "Начните работать с Эмпилс — и наша команда поможет вам масштабировать продажи", 11, False, GRAY, PP_ALIGN.LEFT)

# Left: Structure
add_text(slide, Cm(1), Cm(2.8), Cm(15), Cm(0.5), "СТРУКТУРА КОМАНДЫ", 12, True, ORANGE)
team = [
    ('Региональный директор', 'Ответственный за территорию. Стратегическое планирование, согласование условий.', 'Курирует всех дистрибьюторов на территории', C_NEWT),
    ('Региональный менеджер', 'Оперативное управление. Контроль отгрузок, MML, маркетинговые программы, аналитика.', 'Точка контакта для дистрибьютора на ежедневной основе', C_PROFI),
    ('ЭТП — Эксклюзивный торговый представитель', 'Полевая работа с АКБ. Визиты в РТТ, обучение ТП, выкладка, POS-материалы.', '120 РТТ на 1 ЭТП', C_RASCV),
]
for i, (role, desc, detail, clr) in enumerate(team):
    y = Cm(3.5 + i * 4.2)
    add_card(slide, Cm(1), y, Cm(15), Cm(3.7), clr)
    add_text(slide, Cm(1.3), y + Cm(0.2), Cm(14), Cm(0.7), role, 12, True, DARK)
    ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.3), y + Cm(1.0), Cm(2.5), Cm(0.03))
    ul.fill.solid(); ul.fill.fore_color.rgb = ORANGE; ul.line.fill.background(); ul.shadow.inherit = False
    add_text(slide, Cm(1.3), y + Cm(1.2), Cm(14), Cm(1.5), desc, 9, False, BLACK)
    add_text(slide, Cm(1.3), y + Cm(2.8), Cm(14), Cm(0.6), detail, 9, True, ORANGE)
    if i < 2:
        add_text(slide, Cm(8), y + Cm(3.6), Cm(1), Cm(0.8), "↓", 18, True, ORANGE, PP_ALIGN.CENTER)

# Right: Scaling
add_text(slide, Cm(17.5), Cm(2.8), Cm(15), Cm(0.5), "ПРИНЦИП МАСШТАБИРОВАНИЯ", 12, True, ORANGE)
# Dark card
dc = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(17.5), Cm(3.5), Cm(15), Cm(4.0), )
dc.fill.solid(); dc.fill.fore_color.rgb = DARK; dc.line.fill.background(); dc.shadow.inherit = False
add_text(slide, Cm(18), Cm(3.8), Cm(14), Cm(1), "Больше АКБ с Эмпилс", 16, True, WHITE, PP_ALIGN.CENTER)
add_text(slide, Cm(18), Cm(5.0), Cm(14), Cm(0.8), "↓", 22, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(18), Cm(5.8), Cm(14), Cm(1.2), "Больше наша команда\nна вашей территории", 16, True, ORANGE, PP_ALIGN.CENTER)
# 120 РТТ
add_card(slide, Cm(17.5), Cm(8.0), Cm(15), Cm(3.5), C_RASCV)
add_text(slide, Cm(18), Cm(8.2), Cm(14), Cm(1.5), "120 РТТ", 36, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(18), Cm(9.8), Cm(14), Cm(0.6), "на 1 ЭТП", 13, False, DARK, PP_ALIGN.CENTER)
add_text(slide, Cm(18), Cm(10.5), Cm(14), Cm(0.8), "Каждые 120 активных торговых точек = +1 ЭТП на вашей территории", 9, False, BLACK, PP_ALIGN.CENTER)
# 3 cards
add_text(slide, Cm(17.5), Cm(12.2), Cm(15), Cm(0.5), "ЧТО ДЕЛАЕТ КОМАНДА:", 12, True, ORANGE)
supp = [('📋', 'Контроль MML', 'Слежение за присутствием\nвсех обязательных SKU'), ('🎓', 'Обучение ТП', 'Тренинги по продукту,\nпродажам, TCO'), ('📍', 'Выкладка и POS', 'Контроль полки,\nустановка POS-материалов')]
for i, (ic, t, d) in enumerate(supp):
    x = Cm(17.5 + i * 5.2)
    add_card(slide, x, Cm(13.0), Cm(4.7), Cm(3.5), C_LIGHT_BG)
    add_text(slide, x, Cm(13.2), Cm(4.7), Cm(0.8), ic, 18, False, BLACK, PP_ALIGN.CENTER)
    add_text(slide, x + Cm(0.2), Cm(14.1), Cm(4.3), Cm(0.6), t, 11, True, DARK, PP_ALIGN.CENTER)
    add_text(slide, x + Cm(0.2), Cm(14.8), Cm(4.3), Cm(1.5), d, 9, False, BLACK, PP_ALIGN.CENTER)

# Footer
ft = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Cm(18.0), prs.slide_width, Cm(1.0))
ft.fill.solid(); ft.fill.fore_color.rgb = DARK; ft.line.fill.background(); ft.shadow.inherit = False
add_text(slide, Cm(1), Cm(17.6), Cm(32), Cm(1), "Начните работать с Эмпилс — и наша команда поможет вам масштабировать продажи. Чем больше АКБ, тем больше команда на вашей стороне.", 10, True, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 8: ТРЕЙД-МАРКЕТИНГ
# ============================================================
slide = add_slide()
white_bg(slide)
empils_logo(slide)

add_text(slide, Cm(1), Cm(0.5), Cm(25), Cm(1), "ТРЕЙД-МАРКЕТИНГ", 24, True, DARK, PP_ALIGN.LEFT)
add_text(slide, Cm(1), Cm(1.7), Cm(30), Cm(0.6), "2% от оборота — на совместное продвижение. Мы инвестируем в ваши продажи.", 11, False, GRAY, PP_ALIGN.LEFT)

# Left: 2%
dc = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1), Cm(3.0), Cm(8), Cm(7.0), )
dc.fill.solid(); dc.fill.fore_color.rgb = DARK; dc.line.fill.background(); dc.shadow.inherit = False
add_text(slide, Cm(1.5), Cm(3.5), Cm(7), Cm(2.5), "2%", 60, True, ORANGE, PP_ALIGN.CENTER)
add_text(slide, Cm(1.5), Cm(6.2), Cm(7), Cm(0.8), "от оборота", 14, False, WHITE, PP_ALIGN.CENTER)
add_text(slide, Cm(1.5), Cm(7.5), Cm(7), Cm(1), "Эмпилс выделяет на трейд-маркетинг дистрибьютора", 10, False, WHITE, PP_ALIGN.CENTER)
add_text(slide, Cm(1.5), Cm(8.8), Cm(7), Cm(0.8), "230 ₽/кг — зарабатываем вместе", 12, True, ORANGE, PP_ALIGN.CENTER)

# Right: 9 directions
add_text(slide, Cm(10), Cm(2.5), Cm(23), Cm(0.5), "НАПРАВЛЕНИЯ ПОДДЕРЖКИ", 12, True, ORANGE)
acts = [
    ('🏆', 'Мотивация ТП', 'Конкурсы и премии для ТП дистрибьютора за активные продажи. Призы, бонусы за KPI.', C_EMP),
    ('🏪', 'Мотивация розницы', 'Программы лояльности для РТТ: накопительные бонусы, призы за выкладку, рейтинги.', C_OPTIMA),
    ('🟡', 'Жёлтый ценник и спеццены', 'Снижение цены на полке на ключевые SKU. Сезонное ценовое продвижение.', C_DISC),
    ('🎁', 'Акции «Подарок за покупку»', 'При закупке объёма — бонусные банки бесплатно. Стимулирует увеличение заказа.', C_TITAN),
    ('🎓', 'Обучение и мотивация ТП', 'Тренинги по продукту. Мастер-классы. Сравнительные выкрасы.', C_OREOL),
    ('📍', 'POS-материалы', 'Шелфтокеры, воблеры, ценникодержатели, планограммы, баннеры — бесплатно.', C_RASCV),
    ('📺', 'Контент и digital', 'Материалы для соцсетей, видеоинструкции, фото выкрасов. Поддержка интернет-магазина.', C_LIGHT_BG),
    ('🚐', 'Брендирование и реклама', 'Брендирование машин, входных групп в РТТ, баннерная реклама. Узнаваемость.', C_STROY),
    ('💰', 'Ретробонусы', 'Возврат части средств за выполнение плана закупок по итогам периода.', C_PROFI),
]
cw2 = Cm(7.5); ch2 = Cm(4.5); gx2 = Cm(0.2); gy2 = Cm(0.3)
sx2 = Cm(10); sy2 = Cm(3.2)
for i, (ic, t, d, clr) in enumerate(acts):
    row = i // 3; col = i % 3
    x = sx2 + col * (cw2 + gx2); y = sy2 + row * (ch2 + gy2)
    add_card(slide, x, y, cw2, ch2, clr)
    add_text(slide, x + Cm(0.2), y + Cm(0.15), cw2 - Cm(0.4), Cm(0.7), f"{ic}  {t}", 10, True, DARK)
    ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Cm(0.2), y + Cm(0.9), Cm(2.0), Cm(0.03))
    ul.fill.solid(); ul.fill.fore_color.rgb = ORANGE; ul.line.fill.background(); ul.shadow.inherit = False
    add_text(slide, x + Cm(0.2), y + Cm(1.1), cw2 - Cm(0.4), Cm(3.0), d, 8.5, False, BLACK)

# Footer
ft = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Cm(18.0), prs.slide_width, Cm(1.0))
ft.fill.solid(); ft.fill.fore_color.rgb = DARK; ft.line.fill.background(); ft.shadow.inherit = False
add_text(slide, Cm(1), Cm(17.9), Cm(32), Cm(1), "2% от оборота — ваш бюджет на рост. Мы не просто поставляем продукт — мы инвестируем в ваши продажи.", 10, True, WHITE, PP_ALIGN.CENTER)

# ============================================================
# СЛАЙД 9: НОВИНКИ 2027
# ============================================================
slide = add_slide()
white_bg(slide)
empils_logo(slide)

add_text(slide, Cm(1), Cm(0.5), Cm(25), Cm(1), "НОВИНКИ 2027", 28, True, DARK, PP_ALIGN.LEFT)
add_text(slide, Cm(1), Cm(1.8), Cm(30), Cm(0.6), "Расширяем ассортимент — 28 новых продуктов по 6 брендам", 11, False, GRAY, PP_ALIGN.LEFT)

nov = [
    ('AVANTGARDE', '+3', ['Гидроизолирующая грунтовка', 'Пропитка-лак для финишного покрытия', 'Наливные полы'], C_AVANT),
    ('TITAN', '+4', ['Термостойкая эмаль', 'Преобразователь ржавчины', 'Очиститель поверхности', 'Растворитель «TITAN»'], C_TITAN),
    ('NEWTONE', '+2', ['Грунт-концентрат', 'Краска интерьерная в готовых цветах'], C_NEWT),
    ('ОРЕОЛ', '+8', ['Грунтовка глуб. проникновения с антисептиком', 'Краска для ванной и кухни', 'Краска для спальни и гостиной', 'Краска для обоев', 'Краска резиновая', 'Краска для детских комнат', 'Краска для мебели и декора', 'Декоративная штукатурка «Короед»'], C_OREOL),
    ('PROFIWOOD', '+10', ['Масло для террас', 'Концентрат «Антижук»', 'Концентрат «Антиплесень»', 'Концентрат трудновымываемый антисептик', 'Отбеливатель для древесины', 'Средство для удаления высолов', 'Средство для удаления старого покрытия', 'Гидрофобизатор', 'Изолирующий грунт', 'Антисмольный изолятор'], C_PROFI),
    ('РАСЦВЕТ', '+1', ['Смывка-гель для ЛКМ'], C_RASCV),
]
cw3 = Cm(10.5); ch3 = Cm(6.5); gx3 = Cm(0.4); gy3 = Cm(0.4)
sx3 = Cm(1); sy3 = Cm(3.0)
for i, (name, cnt, prods, clr) in enumerate(nov):
    row = i // 3; col = i % 3
    x = sx3 + col * (cw3 + gx3); y = sy3 + row * (ch3 + gy3)
    add_card(slide, x, y, cw3, ch3, clr)
    tb = slide.shapes.add_textbox(x + Cm(0.3), y + Cm(0.2), cw3 - Cm(0.6), Cm(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = name
    r.font.name = 'Arial'; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = DARK
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = f"{cnt} новинок"
    r2.font.name = 'Arial'; r2.font.size = Pt(9); r2.font.bold = True; r2.font.color.rgb = ORANGE
    ul = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Cm(0.3), y + Cm(1.3), Cm(2.5), Cm(0.03))
    ul.fill.solid(); ul.fill.fore_color.rgb = ORANGE; ul.line.fill.background(); ul.shadow.inherit = False
    # Products
    tb = slide.shapes.add_textbox(x + Cm(0.3), y + Cm(1.6), cw3 - Cm(0.6), Cm(4.5))
    tf = tb.text_frame; tf.word_wrap = True
    fs = 7.5 if len(prods) > 5 else 9
    for pi, prod in enumerate(prods):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.space_before = Pt(1); p.space_after = Pt(1)
        r = p.add_run(); r.text = f"•  {prod}"
        r.font.name = 'Arial'; r.font.size = Pt(fs); r.font.color.rgb = BLACK

# Footer
ft = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Cm(18.0), prs.slide_width, Cm(1.0))
ft.fill.solid(); ft.fill.fore_color.rgb = DARK; ft.line.fill.background(); ft.shadow.inherit = False
add_text(slide, Cm(1), Cm(17.9), Cm(32), Cm(1), "28 новых продуктов в 2027 — расширяем категории, закрываем потребности, растём вместе с партнёрами.", 10, True, WHITE, PP_ALIGN.CENTER)

prs.save(OUT)
print(f'OK: {OUT}')
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
print(f'Слайдов: {len(prs.slides)}')
