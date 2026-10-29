"""Извлечь ключевые данные из аналитики."""
import openpyxl
from openpyxl import load_workbook

# === Красноярский край ===
wb1 = load_workbook("/home/z/my-project/upload/ДФО + В Сиб Красноярский край ver2.xlsx", data_only=True)
ws1 = wb1["Красноярский край"]

print("=== КРАСНОЯРСКИЙ КРАЙ ===\n")

# Производители (rows 3-25)
print("ТОП-10 производителей:")
total_market_kr = 0
for r in range(3, 26):
    name = ws1.cell(row=r, column=3).value
    volume = ws1.cell(row=r, column=4).value
    share = ws1.cell(row=r, column=5).value
    if name and volume:
        if name == 'ИТОГО':
            total_market_kr = volume
            print(f"  ИТОГО: {volume:,.0f} руб ({volume/1e6:.1f} млн)")
        else:
            print(f"  {name}: {volume:,.0f} руб ({volume/1e6:.1f} млн, {share*100:.1f}%)")

# Дистрибьюторы (rows 3-67, columns 7-9)
print(f"\nТОП-12 дистрибьюторов:")
for r in range(3, 15):
    name = ws1.cell(row=r, column=7).value
    volume = ws1.cell(row=r, column=8).value
    share = ws1.cell(row=r, column=9).value
    if name and volume:
        print(f"  {name}: {volume:,.0f} руб ({volume/1e6:.1f} млн, {share*100:.1f}%)")

# Эмпилс в Красноярском крае
print(f"\nДоля Эмпилс в Красноярском крае:")
for r in range(3, 26):
    name = ws1.cell(row=r, column=3).value
    if name and 'Эмпилс' in str(name):
        volume = ws1.cell(row=r, column=4).value
        share = ws1.cell(row=r, column=5).value
        print(f"  {name}: {volume:,.0f} руб ({volume/1e6:.1f} млн, {share*100:.1f}%)")

# === Хакасия ===
wb2 = load_workbook("/home/z/my-project/upload/ДФО + В Сиб Хакасия республика ver2.xlsx", data_only=True)
ws2 = wb2["Хакасия республика"]

print(f"\n\n=== ХАКАСИЯ ===\n")

print("Производители:")
total_market_hk = 0
for r in range(3, 14):
    name = ws2.cell(row=r, column=2).value
    volume = ws2.cell(row=r, column=3).value
    share = ws2.cell(row=r, column=4).value
    if name and volume:
        if name == 'ИТОГО':
            total_market_hk = volume
            print(f"  ИТОГО: {volume:,.0f} руб ({volume/1e6:.1f} млн)")
        else:
            print(f"  {name}: {volume:,.0f} руб ({volume/1e6:.1f} млн, {share*100:.1f}%)")

print(f"\nДистрибьюторы:")
for r in range(3, 13):
    name = ws2.cell(row=r, column=6).value
    volume = ws2.cell(row=r, column=7).value
    share = ws2.cell(row=r, column=8).value
    if name and volume:
        if name == 'ИТОГО':
            print(f"  ИТОГО: {volume:,.0f} руб ({volume/1e6:.1f} млн)")
        else:
            print(f"  {name}: {volume:,.0f} руб ({volume/1e6:.1f} млн, {share*100:.1f}%)")

# Сводка
print(f"\n\n=== СВОДКА ===")
print(f"Красноярский край: {total_market_kr/1e6:.1f} млн руб")
print(f"Хакасия: {total_market_hk/1e6:.1f} млн руб")
print(f"ИТОГО: {(total_market_kr + total_market_hk)/1e6:.1f} млн руб")
