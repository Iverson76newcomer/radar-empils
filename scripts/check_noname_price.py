"""Проверить прайс-лист - что в листе 'NoName' и какой реальный бренд у 'Краски для садовых деревьев'."""
import openpyxl
from openpyxl import load_workbook

# Загрузить прайс-лист
wb = load_workbook("/home/z/my-project/upload/16.10.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

# Проверить лист 'NoName'
print("=== Лист 'NoName' в прайс-листе 16.10.2026 ===")
ws = wb["NoName"]
print(f"Строк: {ws.max_row}, колонок: {ws.max_column}")

# Прочитать все строки
current_name = None
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

# Также проверить все листы на наличие 'садовых деревьев'
print("\n\n=== Поиск 'садовых деревьев' во всех листах прайс-листа ===")
for sn in wb.sheetnames:
    ws = wb[sn]
    current_name = None
    for r in range(14, ws.max_row + 1):
        f_val = ws.cell(row=r, column=6).value
        h_val = ws.cell(row=r, column=8).value
        m_val = ws.cell(row=r, column=13).value
        n_val = ws.cell(row=r, column=14).value
        
        if f_val and isinstance(f_val, str) and 'садовых деревьев' in f_val.lower():
            print(f"\n  Лист '{sn}', row {r}:")
            print(f"    Наименование: {f_val}")
            print(f"    Фасовка (H): {h_val}")
            print(f"    Базовая (M): {m_val}")
            print(f"    Отгрузочная (N): {n_val}")
        
        # Также проверить штрихкод и др.
        b_val = ws.cell(row=r, column=1).value
        if b_val and 'садовых деревьев' in str(b_val).lower():
            print(f"  [{sn}] Row {r} - в колонке 1: {b_val}")

# Проверить, какой бренд/производитель указан для садовых деревьев в прайс-листе
print("\n\n=== Все строки листа 'NoName' с ценами ===")
ws = wb["NoName"]
current_name = None
for r in range(14, ws.max_row + 1):
    f_val = ws.cell(row=r, column=6).value
    h_val = ws.cell(row=r, column=8).value
    m_val = ws.cell(row=r, column=13).value
    n_val = ws.cell(row=r, column=14).value
    
    if f_val and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
        current_name = f_val.strip()
    
    if (m_val and isinstance(m_val, (int, float))) or (n_val and isinstance(n_val, (int, float))):
        print(f"  Row {r}: name='{current_name}', фас={h_val}, базовая={m_val}, отгруз={n_val}")
