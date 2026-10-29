"""
Анализ новинок:
1. Какие новинки уже есть в ассортименте (по гр.коду)?
2. Какие новинки популярны по вторичке, но ещё не в ассортименте?
3. Какие фасовки новых SKU?
"""
import openpyxl
from openpyxl import load_workbook
import re

# === Загрузить ассортимент 300 SKU ===
wb1 = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал.xlsx", data_only=True)
ws1 = wb1["300 SKU — детально"]

assortment = []
for row in ws1.iter_rows(min_row=4, max_row=ws1.max_row, values_only=True):
    if row[0] is None:
        continue
    assortment.append({
        'idx': row[0], 'sphere': row[1], 'purpose': row[2], 'base': row[3],
        'brand': row[4], 'type': row[5], 'name': row[6],
        'grp_code': row[7], 'abc': row[8]
    })

print(f"Ассортимент: {len(assortment)} SKU")

# Список групповых кодов и имён
assortment_codes = set()
assortment_names_lower = set()
for it in assortment:
    if it['grp_code']:
        assortment_codes.add(int(it['grp_code']))
    if it['name']:
        assortment_names_lower.add(it['name'].lower())

# Также загрузить 5 жёлтых позиций AvantGarde
wb_p = load_workbook("/home/z/my-project/upload/Ассортимент_Олимп_300_SKU_финал.xlsx", data_only=True)
ws_p = wb_p["300 SKU с ценами"]
yellow_items = []
for r in range(5, 10):  # rows 5-9 (idx 1-5)
    yellow_items.append({
        'idx': ws_p.cell(row=r, column=1).value,
        'name': ws_p.cell(row=r, column=7).value,
        'brand': ws_p.cell(row=r, column=5).value,
        'fasovka': ws_p.cell(row=r, column=12).value,
        'sphere': ws_p.cell(row=r, column=2).value,
        'purpose': ws_p.cell(row=r, column=3).value,
        'base': ws_p.cell(row=r, column=4).value,
        'type': ws_p.cell(row=r, column=6).value,
    })

print(f"\n=== 5 ЖЁЛТЫХ ПОЗИЦИЙ AvantGarde (для замены) ===")
for yi in yellow_items:
    print(f"  #{yi['idx']} | {yi['name']} | фас={yi['fasovka']} кг")

# === Загрузить новинки ===
wb2 = load_workbook("/home/z/my-project/upload/Новинки за 2026 год первичка и вторичка.xlsx", data_only=True)
ws2 = wb2["Первичные и вторичные продажи"]

novinki = []
for row in ws2.iter_rows(min_row=3, max_row=ws2.max_row, values_only=True):
    if row[3] is None or not isinstance(row[3], int):
        continue
    novinki.append({
        'num': row[0], 'period': row[1], 'brand': row[2], 'code': row[3],
        'name': row[4], 'status': row[5], 'vid': row[6], 'type': row[7],
        'pervichka_kg': row[8] or 0,
        'pervichka_rub': row[9] or 0,
        'vtorichka_kg': row[10] or 0,
        'vtorichka_rub': row[11] or 0,
        'akb': row[12] or 0,
    })

print(f"\n=== Новинки: {len(novinki)} SKU ===")
print("\n=== Анализ: какие новинки есть в ассортименте, каких нет ===")

for n in novinki:
    in_assortment = n['code'] in assortment_codes
    # Проверка по имени
    n_name_lower = n['name'].lower() if n['name'] else ''
    name_match = False
    for an in assortment_names_lower:
        if an and (n_name_lower in an or an in n_name_lower):
            name_match = True
            break
    n['in_assortment'] = in_assortment or name_match
    n['name_match'] = name_match

# Сортировка по вторичке (руб) - убывание
novinki_sorted = sorted(novinki, key=lambda x: -x['vtorichka_rub'])

print("\n=== Топ-15 новинок по вторичке (руб) ===")
print(f"{'Код':<7} {'Бренд':<12} {'Перв.кг':<10} {'Перв.руб':<12} {'Втор.кг':<10} {'Втор.руб':<12} {'АКБ':<5} {'В ассорт?':<10} {'Номенклатура'}")
for n in novinki_sorted[:15]:
    print(f"{n['code']:<7} {n['brand']:<12} {n['pervichka_kg']:<10} {n['pervichka_rub']:<12.0f} {n['vtorichka_kg']:<10} {n['vtorichka_rub']:<12.0f} {n['akb']:<5} {'да' if n['in_assortment'] else 'НЕТ':<10} {n['name'][:55]}")

# Новинки, которых НЕТ в ассортименте
new_not_in_assortment = [n for n in novinki_sorted if not n['in_assortment']]
print(f"\n=== Новинок нет в ассортименте: {len(new_not_in_assortment)} ===")
for n in new_not_in_assortment:
    print(f"  Код {n['code']} | {n['brand']} | {n['name'][:65]} | втор.руб={n['vtorichka_rub']:.0f}")

# Извлечь фасовку из имени новинок
def extract_fasovka(name):
    if not name:
        return None
    m = re.search(r'(\d+(?:\.\d+)?)\s*кг', name, re.IGNORECASE)
    return float(m.group(1)) if m else None

def extract_l(name):
    if not name:
        return None
    m = re.search(r'(\d+(?:\.\d+)?)\s*л', name, re.IGNORECASE)
    return float(m.group(1)) if m else None

# Проверим какие фасовки уже есть в ассортименте по брендам и типам
print("\n=== Существующие фасовки в ассортименте по бренду/типу ===")
from collections import defaultdict
brand_type_fasovki = defaultdict(set)
for it in assortment:
    if not it['name']:
        continue
    kg = extract_fasovka(it['name'])
    l = extract_l(it['name'])
    key = (it['brand'], it['type'] or '')
    if kg:
        brand_type_fasovki[key].add(f"кг={kg}")
    if l:
        brand_type_fasovki[key].add(f"л={l}")

# Выведем для интересных брендов
for key in sorted(brand_type_fasovki.keys()):
    brand, type_p = key
    if brand in ['PROFIWOOD', 'Ореол', 'NEWTONE', 'AvantGarde']:
        print(f"  {brand} / {type_p}: {sorted(brand_type_fasovki[key])}")
