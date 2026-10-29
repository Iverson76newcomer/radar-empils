#!/usr/bin/env python3
"""
Обновление презентации: заполняем пустые разделительные слайды.
Слайд 8 → ГЕОГРАФИЯ БИЗНЕСА
Слайд 11 → АССОРТИМЕНТ И ЦЕНОВЫЕ СЕГМЕНТЫ
Слайд 14 → ПОРТФОЛИО: ОБЪЕКТЫ И B2B
Удаляем лишний пустой слайд 15.
"""

from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

SRC = '/home/z/my-project/upload/Презентация_партнёрам_с_ассортиментом_v2.pptx'
OUT = '/home/z/my-project/download/Презентация_партнёрам_разделители.pptx'

COL_ORANGE = RGBColor(0xF4, 0x76, 0x29)
COL_DARK = RGBColor(0x31, 0x31, 0x31)
COL_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(SRC)
sw = prs.slide_width
sh = prs.slide_height

def fill_divider(slide, title, subtitle):
    """Заполняет пустой слайд как разделитель."""
    # Тёмный фон
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
    bg.fill.solid(); bg.fill.fore_color.rgb = COL_DARK
    bg.line.fill.background(); bg.shadow.inherit = False
    
    # Оранжевые полосы
    for y in [0, sh - Cm(0.4)]:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, sw, Cm(0.4))
        bar.fill.solid(); bar.fill.fore_color.rgb = COL_ORANGE
        bar.line.fill.background(); bar.shadow.inherit = False
    
    # Заголовок
    tb = slide.shapes.add_textbox(Cm(2), Cm(10.5), sw - Cm(4), Cm(3))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.name = 'Proxima Nova Regular'; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = COL_ORANGE
    
    # Подзаголовок
    sb = slide.shapes.add_textbox(Cm(2), Cm(14.5), sw - Cm(4), Cm(1.5))
    tf = sb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = subtitle
    r.font.name = 'Inter Regular'; r.font.size = Pt(16); r.font.color.rgb = COL_WHITE

# Заполняем 3 пустых разделительных слайда
slides = list(prs.slides)

# Слайд 8 (индекс 7) — ГЕОГРАФИЯ БИЗНЕСА
fill_divider(slides[7], 'ГЕОГРАФИЯ БИЗНЕСА', 'Дистрибуция в России и экспорт в страны СНГ')

# Слайд 11 (индекс 10) — АССОРТИМЕНТ И ЦЕНОВЫЕ СЕГМЕНТЫ
fill_divider(slides[10], 'АССОРТИМЕНТ И ЦЕНОВЫЕ СЕГМЕНТЫ', '11 брендов · от эконом до премиум')

# Слайд 14 (индекс 13) — ПОРТФОЛИО: ОБЪЕКТЫ И B2B
fill_divider(slides[13], 'ПОРТФОЛИО: ОБЪЕКТЫ И B2B', 'Реальные кейсы и промышленные партнёры')

# Удаляем слайд 15 (индекс 14) — лишний пустой
sldIdLst = prs.slides._sldIdLst
all_elems = list(sldIdLst)
sldIdLst.remove(all_elems[14])  # удаляем 15-й слайд (индекс 14)

prs.save(OUT)
print(f'OK: {OUT}')
print(f'Размер: {os.path.getsize(OUT) / 1024:.1f} KB')
print(f'Слайдов: {len(prs.slides)}\n')

# Проверка
prs2 = Presentation(OUT)
for i, s in enumerate(prs2.slides, 1):
    texts = []
    for shape in s.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t and len(t) > 2:
                texts.append(t[:80])
    label = ' | '.join(texts[:2])[:120] if texts else '(пустой)'
    print(f"  Слайд {i:2d}: {label}")
