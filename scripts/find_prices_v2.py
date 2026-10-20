"""
Расширенный поиск цен — учитывает частичные совпадения.
"""
import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

PRICE_PATH = "/home/z/my-project/upload/16.10.2026 РОССИЯ с НДС, руб..xlsx"
wb2 = load_workbook(PRICE_PATH, data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) \
            and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name, 'fасовка': h_val, 'базовая': m_val, 'отгрузочная': n_val,
            })

# === Топ-5 кандидатов ===
top5 = [
    {'grp_code': 63180, 'brand': 'NoName', 'type': 'Краски', 'name': 'Краска ВД для садовых деревьев 3 кг', 'cumulative': 0.641, 'sphere': 'Дерево', 'purpose': 'Антисептик для древесины', 'base': 'Водно-дисперсионная'},
    {'grp_code': 79754, 'brand': 'NoName', 'type': 'Эмали Алкидные', 'name': 'Эмаль ПФ-115 синяя 0.7 кг', 'cumulative': 0.647, 'sphere': 'Металл', 'purpose': 'Эмаль ПФ-115', 'base': 'Органическая'},
    {'grp_code': 73122, 'brand': 'СТРОЙКОМПЛЕКТ', 'type': 'Краски', 'name': 'Краска фасадная СТРОЙКОМПЛЕКТ 13 кг', 'cumulative': 0.653, 'sphere': 'Фасад', 'purpose': 'Краска фасадная', 'base': 'Водно-дисперсионная'},
    {'grp_code': 7803, 'brand': 'EMPILS ЭКОНОМ', 'type': 'Эмали Алкидные', 'name': 'Эмаль EMPILS ПФ-115 ЭКОНОМ чёрная 20 кг', 'cumulative': 0.655, 'sphere': 'Металл', 'purpose': 'Эмаль ПФ-115', 'base': 'Органическая'},
    {'grp_code': 72676, 'brand': 'СТРОЙКОМПЛЕКТ', 'type': 'Эмали Алкидные', 'name': 'Эмаль ПФ-115 СТРОЙКОМПЛЕКТ чёрная 10 кг', 'cumulative': 0.673, 'sphere': 'Металл', 'purpose': 'Эмаль ПФ-115', 'base': 'Органическая'},
]

# Поиск по ключевым словам
def search_in_all_sheets(keywords, target_fas=None):
    """Search all sheets for entries containing all keywords."""
    results = []
    for sn in wb2.sheetnames:
        for p in price_dict[sn]:
            if p['name'] is None:
                continue
            name_lower = p['name'].lower()
            if all(kw in name_lower for kw in keywords):
                if target_fas is not None and isinstance(p['fасовка'], (int, float)):
                    if abs(p['fасовка'] - target_fas) > 0.5:
                        continue
                results.append((sn, p))
    return results

# === Поиск для каждого ===
print("=" * 100)
print("ПОИСК ЦЕН ДЛЯ 5 НОВЫХ SKU (расширенный поиск)")
print("=" * 100)

new_sku_with_prices = []

# 1. Краска ВД для садовых деревьев 3 кг
print("\n--- 1. Краска ВД для садовых деревьев 3 кг ---")
matches = search_in_all_sheets(['садовых деревьев'], target_fas=3)
for sn, p in matches:
    print(f"  ✓ [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
if matches:
    chosen = matches[0]
    new_sku_with_prices.append({
        'info': top5[0],
        'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
    })

# 2. Эмаль ПФ-115 синяя 0.7 кг (NoName)
print("\n--- 2. Эмаль ПФ-115 синяя 0.7 кг (NoName) ---")
matches = search_in_all_sheets(['пф-115', 'синяя'], target_fas=0.7)
if not matches:
    matches = search_in_all_sheets(['пф-115', 'синяя'])
for sn, p in matches:
    print(f"  - [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
# If 0.7 фасовка not found, take smallest available
if matches and not any(isinstance(m[1]['fасовка'], (int, float)) and abs(m[1]['fасовка'] - 0.7) <= 0.5 for m in matches):
    print("  → Точной фасовки 0.7 кг нет. Берём наименьшую доступную:")
    sorted_matches = sorted([m for m in matches if isinstance(m[1]['fасовка'], (int, float))], key=lambda m: m[1]['fасовка'])
    if sorted_matches:
        chosen = sorted_matches[0]
        print(f"  ✓ [{chosen[0]}] '{chosen[1]['name']}' | фас={chosen[1]['fасовка']} | отгруз={chosen[1]['отгрузочная']}")
        new_sku_with_prices.append({
            'info': top5[1],
            'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
        })
elif matches:
    chosen = matches[0]
    new_sku_with_prices.append({
        'info': top5[1],
        'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
    })

# 3. Краска фасадная СТРОЙКОМПЛЕКТ 13 кг
print("\n--- 3. Краска фасадная СТРОЙКОМПЛЕКТ 13 кг ---")
matches = search_in_all_sheets(['фасадная', 'стройкомплект'], target_fas=13)
for sn, p in matches:
    print(f"  ✓ [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
if matches:
    chosen = matches[0]
    new_sku_with_prices.append({
        'info': top5[2],
        'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
    })

# 4. Эмаль EMPILS ПФ-115 ЭКОНОМ чёрная 20 кг
print("\n--- 4. Эмаль EMPILS ПФ-115 ЭКОНОМ чёрная 20 кг ---")
matches = search_in_all_sheets(['empils', 'пф-115', 'чёрная'], target_fas=20)
if not matches:
    matches = search_in_all_sheets(['empils', 'пф-115', 'чёрная'])
for sn, p in matches:
    print(f"  - [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
if matches:
    chosen = matches[0]
    new_sku_with_prices.append({
        'info': top5[3],
        'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
    })

# 5. Эмаль ПФ-115 СТРОЙКОМПЛЕКТ чёрная 10 кг
print("\n--- 5. Эмаль ПФ-115 СТРОЙКОМПЛЕКТ чёрная 10 кг ---")
matches = search_in_all_sheets(['пф-115', 'стройкомплект', 'чёрная'], target_fas=10)
for sn, p in matches:
    print(f"  ✓ [{sn}] '{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']} | базовая={p['базовая']}")
if matches:
    chosen = matches[0]
    new_sku_with_prices.append({
        'info': top5[4],
        'match': {'name': chosen[1]['name'], 'fасовка': chosen[1]['fасовка'], 'базовая': chosen[1]['базовая'], 'отгрузочная': chosen[1]['отгрузочная'], 'лист': chosen[0]}
    })

# === Итог ===
print("\n" + "=" * 100)
print(f"ИТОГ: найдено цен для {len(new_sku_with_prices)} из 5 SKU")
print("=" * 100)
for r in new_sku_with_prices:
    c = r['info']
    m = r['match']
    print(f"\n  Код {c['grp_code']} | {c['brand']} | {c['name']}")
    print(f"    Сфера: {c['sphere']} | Назначение: {c['purpose']} | Основа: {c['base']}")
    print(f"    → [{m['лист']}] {m['name']}")
    print(f"    фасовка={m['fасовка']} | отгрузочная={m['отгрузочная']} | базовая={m['базовая']}")

# Сохраним результат в JSON
import json
with open('/home/z/my-project/scripts/new_sku.json', 'w', encoding='utf-8') as f:
    json.dump(new_sku_with_prices, f, ensure_ascii=False, indent=2, default=str)
print(f"\nСохранено: /home/z/my-project/scripts/new_sku.json")
