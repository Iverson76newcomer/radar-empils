import openpyxl
from openpyxl import load_workbook
from collections import defaultdict

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# Ореол Дисконт: all "краска" entries
print("=== Ореол Дисконт: 'краска' ===")
seen = set()
for p in price_dict['Ореол Дисконт']:
    if p['name'] and 'краска' in p['name'].lower() and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Ореол Дисконт: search 'потолк' or 'стен'
print("\n=== Ореол Дисконт: 'потолк' or 'стен' ===")
for p in price_dict['Ореол Дисконт']:
    if p['name'] and ('потолк' in p['name'].lower() or 'стен' in p['name'].lower()):
        print(f"  '{p['name']}' фас={p['fасовка']}")

# Search all sheets for 'стен и потолков'
print("\n=== All sheets: 'стен и потолков' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'стен' in p['name'].lower() and 'потолк' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")
