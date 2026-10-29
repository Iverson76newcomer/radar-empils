"""Анализ документа — поиск всех упоминаний ПФ-115, ГФ, 320, 1280, 2780, 9174."""
from docx import Document
import re

doc = Document('/home/z/my-project/upload/TITAN_ССВ_расчёт_пояснительная_справка.docx')

print("=== ПАРАГРАФЫ (с упоминанием цены/расчёта) ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if any(kw in t for kw in ['320', '1 280', '1280', '2 780', '2780', '9 174', '9174', 'ГФ', 'грунт', 'ПФ-115', '440', '820', '4 банки', '4,8']):
        print(f"[{i}] {t[:250]}")

print("\n\n=== ТАБЛИЦЫ ===")
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' | ')[:100] for c in row.cells]
        line = ' || '.join(cells)
        if any(kw in line for kw in ['320', '1 280', '1280', '2 780', '2780', '9 174', '9174', 'ГФ', 'грунт', 'ПФ-115', '440', '820', '4 банки', '4,8']):
            print(f"  T{ti+1} R{ri}: {line}")
