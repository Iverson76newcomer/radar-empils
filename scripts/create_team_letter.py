#!/usr/bin/env python3
"""
Финальная версия документа с двумя письмами для каскадной рассылки.
Письмо №1 — текст пользователя (короткий, по делу, с цифрами).
Письмо №2 — готовый шаблон для пересылки дистрибьюторам.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = '/home/z/my-project/download/Письмо_торговой_команде_соцсети.docx'

EMPILS_ORANGE = RGBColor(0xF4, 0x79, 0x20)
DARK = RGBColor(0x1A, 0x1F, 0x36)
GRAY = RGBColor(0x60, 0x60, 0x60)


def set_para_format(p, line_spacing=1.3, space_after=6, space_before=0,
                    first_line_indent=0, alignment=None):
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if first_line_indent:
        pf.first_line_indent = Cm(first_line_indent)
    if alignment is not None:
        p.alignment = alignment


def add_para(doc, text='', size=11, bold=False, color=DARK,
             alignment=WD_ALIGN_PARAGRAPH.LEFT, italic=False,
             line_spacing=1.3, space_after=6, space_before=0,
             first_line_indent=0, font_name='Arial'):
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=line_spacing, space_after=space_after,
                    space_before=space_before,
                    first_line_indent=first_line_indent, alignment=alignment)
    if text:
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font_name
    return p


def add_divider(doc, color='F47920'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_section_label(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.2, space_after=4, space_before=8)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'
    rPr = run._element.get_or_add_rPr()
    caps = OxmlElement('w:caps')
    caps.set(qn('w:val'), 'true')
    rPr.append(caps)
    return p


def add_letter_header(doc, from_text, to_text, subject_text):
    add_section_label(doc, 'От кого')
    add_para(doc, from_text, size=11, color=DARK, space_after=8)

    add_section_label(doc, 'Кому')
    add_para(doc, to_text, size=11, color=DARK, space_after=8)

    add_section_label(doc, 'Тема')
    add_para(doc, subject_text, size=11, bold=True, color=DARK, space_after=10)


# ==============================================================
# СОЗДАНИЕ ДОКУМЕНТА
# ==============================================================

def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # === ЗАГОЛОВОК ДОКУМЕНТА ===
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.2, space_after=4, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run('Empils')
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    add_para(doc, 'Письмо для торговой команды — вступление в социальные сети',
             size=14, color=DARK, bold=True, space_after=4)
    add_para(doc, 'Комплект: 2 письма для каскадной рассылки',
             size=10, color=GRAY, italic=True, space_after=2)

    add_divider(doc)

    # Пояснение
    add_para(doc,
             'Документ содержит два письма для каскадной рассылки по цепочке:',
             size=10, color=GRAY, italic=True, space_after=4)
    add_para(doc,
             'Эмпилс → региональные представители → дистрибьюторы → торговые представители дистрибьюторов → розничные точки.',
             size=10, color=GRAY, italic=True, space_after=10)

    # Инструкция
    add_section_label(doc, 'Как использовать')
    instructions = [
        'Письмо №1 — отправляется региональным представителям Эмпилс. В нём ставится задача: переслать письмо №2 своим дистрибьюторам.',
        'Письмо №2 — готовый шаблон. Дистрибьюторы пересылают его своим торговым представителям. К письму прикрепляется картинка с QR-кодами (файл «Empils_SMM.jpg»).',
        'QR-код можно считывать прямо с картинки со смартфона — никакого дополнительного ПО не требуется.',
    ]
    for i, txt in enumerate(instructions, 1):
        p = doc.add_paragraph()
        set_para_format(p, line_spacing=1.3, space_after=4)
        p.paragraph_format.left_indent = Cm(0.5)
        run_num = p.add_run(f'{i}. ')
        run_num.font.size = Pt(11)
        run_num.font.bold = True
        run_num.font.color.rgb = EMPILS_ORANGE
        run_num.font.name = 'Arial'
        run = p.add_run(txt)
        run.font.size = Pt(11)
        run.font.color.rgb = DARK
        run.font.name = 'Arial'

    # ============== ПИСЬМО №1 ==============
    doc.add_page_break()

    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.2, space_after=4)
    run = p.add_run('Письмо №1')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    add_para(doc, 'Для региональных представителей Эмпилс',
             size=11, color=GRAY, italic=True, space_after=10)

    add_divider(doc)

    add_letter_header(
        doc,
        from_text='Отдел маркетинга и развития продаж Empils',
        to_text='Региональным представителям Empils',
        subject_text='Каскадная рассылка: приглашение торговых представителей дистрибьюторов в наши социальные сети'
    )

    add_divider(doc)

    # === ТЕКСТ ПИСЬМА №1 (текст пользователя) ===
    add_para(doc, 'Коллеги, добрый день!', size=11, bold=True, space_after=10)

    add_para(doc,
             'Запускаем каскадную рассылку, чтобы привлечь в наши социальные '
             'сети торговых представителей дистрибьюторов — тех самых людей, '
             'которые ежедневно посещают розничные точки и напрямую общаются '
             'с продавцами и владельцами магазинов. Это наша самая «полевая» '
             'аудитория: они первыми должны узнавать о новинках, акциях и т.д.',
             size=11, first_line_indent=1.0, space_after=10)

    add_para(doc,
             'Цель — чтобы торговые представители дистрибьюторов подписались '
             'на наши группы в ВКонтакте, Telegram и МАХ.',
             size=11, first_line_indent=1.0, space_after=10)

    # Ключевая фраза про цифры — выделенная
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.3, space_after=10, first_line_indent=1.0)
    # Оранжевая плашка через заливку абзаца
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'FFF4E6')  # светло-оранжевый фон
    pPr.append(shd)
    # Левые/правые отступы для плашки
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)

    run = p.add_run(
        'У нас порядка 100 дистрибьюторов и, наверное, порядка 500 торговых '
        'представителей общего прайса (включая экспорт) — а всего 92 подписчика '
        'в МАХ.'
    )
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    add_para(doc, '', space_after=6)

    add_para(doc,
             'Файл для пересылки во вложении. Он содержит QR-код и письменное '
             'написание. QR можно считывать прямо с картинки со смартфона.',
             size=11, first_line_indent=1.0, space_after=10)

    add_para(doc, 'Спасибо.', size=11, bold=True, first_line_indent=1.0,
             space_after=12)

    add_divider(doc)

    add_para(doc, 'С уважением,', size=11, space_after=2)
    add_para(doc, 'Отдел маркетинга и развития продаж Empils', size=11, bold=True,
             color=EMPILS_ORANGE, space_after=2)
    add_para(doc, 'marketing@empils.ru  •  www.titanpaints.ru', size=10,
             color=GRAY, italic=True)

    # ============== ПИСЬМО №2 ==============
    doc.add_page_break()

    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.2, space_after=4)
    run = p.add_run('Письмо №2')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    add_para(doc, 'Готовый шаблон для пересылки дистрибьюторам → торговым представителям',
             size=11, color=GRAY, italic=True, space_after=10)

    add_divider(doc)

    add_letter_header(
        doc,
        from_text='Региональный представитель Empils',
        to_text='Дистрибьютор → для пересылки торговым представителям',
        subject_text='Empils приглашает ваших торговых представителей в наши соцсети — первыми узнавать об акциях и новинках'
    )

    add_divider(doc)

    # === ТЕКСТ ПИСЬМА №2 ===
    add_para(doc, 'Уважаемые партнёры!', size=12, bold=True, color=EMPILS_ORANGE,
             space_after=8)

    add_para(doc,
             'Мы ценим наше сотрудничество и приглашаем ваших торговых '
             'представителей стать ближе к жизни компании Empils. Те, кто '
             'ежедневно посещает розничные точки и работает с продавцами '
             'магазинов — наша самая важная «полевая» команда. Нам очень '
             'хочется быть с вами на связи напрямую.',
             size=11, first_line_indent=1.0, space_after=8)

    add_para(doc,
             'Просим вас переслать это письмо вашим торговым представителям '
             'и пригласить их подписаться на наши социальные сети. Это '
             'бесплатный канал, который помогает продавать больше:',
             size=11, first_line_indent=1.0, space_after=8)

    # Что получают торговые представители
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.3, space_after=4)
    run = p.add_run('Что получают торговые представители:')
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    benefits2 = [
        ('Первые новости об акциях и новинках', ' — узнавать раньше, чем официальные анонсы дойдут до розницы.'),
        ('Видеоинструкции по продукту', ' — короткие ролики по TITAN 3в1, удобно пересылать продавцам в магазинах.'),
        ('Аргументы для отработки цены', ' — готовые ответы на «почему TITAN дороже ПФ-115» и «зачем платить больше».'),
        ('Календарь сезонных активностей', ' — напоминания о выкладке, акциях, POS-материалах в нужный момент.'),
        ('Прямой канал с производителем', ' — можно задать вопрос по продукту, не дожидаясь встречи с региональным представителем.'),
    ]
    for title, desc in benefits2:
        p = doc.add_paragraph()
        set_para_format(p, line_spacing=1.3, space_after=4)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run('•  ')
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = EMPILS_ORANGE
        run.font.name = 'Arial'
        run_t = p.add_run(title)
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = DARK
        run_t.font.name = 'Arial'
        run_d = p.add_run(desc)
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = DARK
        run_d.font.name = 'Arial'

    add_para(doc, '', space_after=6)

    add_para(doc,
             'QR-коды и ссылки на наши социальные сети — во вложении к этому '
             'письму (картинка с тремя QR-кодами: ВКонтакте, Telegram, МАХ). '
             'QR можно считывать прямо с картинки со смартфона — никакого '
             'дополнительного приложения не нужно.',
             size=11, first_line_indent=1.0, space_after=8)

    # Ссылки на соцсети
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.3, space_after=4)
    run = p.add_run('Наши социальные сети:')
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    socials = [
        ('ВКонтакте', 'https://vk.com/empils_brands'),
        ('Telegram', 'https://t.me/empils_lkm'),
        ('МАХ', 'https://max.ru/id6167008343_biz'),
    ]
    for name, url in socials:
        p = doc.add_paragraph()
        set_para_format(p, line_spacing=1.3, space_after=3)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run('•  ')
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = EMPILS_ORANGE
        run.font.name = 'Arial'
        run_n = p.add_run(f'{name}:  ')
        run_n.font.size = Pt(11)
        run_n.font.bold = True
        run_n.font.color.rgb = DARK
        run_n.font.name = 'Arial'
        run_u = p.add_run(url)
        run_u.font.size = Pt(11)
        run_u.font.color.rgb = GRAY
        run_u.font.name = 'Consolas'

    add_para(doc, '', space_after=8)

    # Просьба
    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.3, space_after=8, first_line_indent=1.0)
    run = p.add_run('Просим вас ')
    run.font.size = Pt(11)
    run.font.color.rgb = DARK
    run.font.name = 'Arial'
    run_b = p.add_run('переслать это письмо вашим торговым представителям')
    run_b.font.size = Pt(11)
    run_b.font.bold = True
    run_b.font.color.rgb = EMPILS_ORANGE
    run_b.font.name = 'Arial'
    run2 = p.add_run(
        ' и попросить их подписаться. Это займёт у них 30 секунд, '
        'а польза — на каждый рабочий день: первыми узнавать об акциях, '
        'иметь под рукой видеоинструкции и аргументы для продаж.'
    )
    run2.font.size = Pt(11)
    run2.font.color.rgb = DARK
    run2.font.name = 'Arial'

    add_para(doc,
             'Подписывайтесь! Будем рады видеть вас и ваших коллег в нашем кругу.',
             size=11, bold=True, color=EMPILS_ORANGE, first_line_indent=1.0,
             space_after=12)

    add_divider(doc)

    add_para(doc, 'С уважением,', size=11, space_after=2)
    add_para(doc, 'Команда Empils', size=11, bold=True, color=EMPILS_ORANGE,
             space_after=2)
    add_para(doc,
             'ВКонтакте: vk.com/empils_brands  •  Telegram: t.me/empils_lkm  •  МАХ: max.ru/id6167008343_biz',
             size=10, color=GRAY, italic=True)
    add_para(doc, 'www.titanpaints.ru', size=10, color=GRAY, italic=True)

    # ============== ПРИЛОЖЕНИЕ ==============
    doc.add_page_break()

    p = doc.add_paragraph()
    set_para_format(p, line_spacing=1.2, space_after=4)
    run = p.add_run('Приложение')
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = EMPILS_ORANGE
    run.font.name = 'Arial'

    add_para(doc, 'Картинка с QR-кодами для рассылки',
             size=11, color=GRAY, italic=True, space_after=10)

    add_divider(doc)

    # Встраиваем картинку (JPEG — финальная версия)
    import os
    img_path = '/home/z/my-project/download/Empils_SMM.jpg'
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Cm(14))

        add_para(doc,
                 'Файл: Empils_SMM.jpg',
                 size=10, color=GRAY, italic=True,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        add_para(doc,
                 'Прикрепите этот файл к письму №2 при отправке дистрибьюторам.',
                 size=10, color=GRAY, italic=True,
                 alignment=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        add_para(doc, '[Картинка не найдена — прикрепите файл Empils_SMM.jpg вручную]',
                 size=10, color=GRAY, italic=True)

    doc.save(OUT)
    print(f'OK: {OUT}')


if __name__ == '__main__':
    main()
