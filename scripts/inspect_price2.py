import openpyxl
from openpyxl import load_workbook

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

# Find a sheet with actual prices — look at AvantGarde rows 18-30 fully
ws = wb2['AvantGarde']
print("=== AvantGarde — find rows with prices ===")
print(f"Merged cells: {len(ws.merged_cells.ranges)} ranges")
for mr in list(ws.merged_cells.ranges)[:20]:
    print(f"  {mr}")

# Print rows 18-40 with all 14 cols
print("\n=== Rows 18-40 ===")
for row_idx, row in enumerate(ws.iter_rows(min_row=18, max_row=40, values_only=True), 18):
    if any(c is not None for c in row):
        print(f"Row {row_idx}: {row}")

# Find rows where M (index 12) and N (index 13) have values
print("\n=== Rows with prices (col M or N not None) ===")
count = 0
for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=80, values_only=True), 14):
    if row[12] is not None or row[13] is not None:
        print(f"Row {row_idx}: name={row[5]!r}, фасовка={row[10]}, ед={row[11]}, базовая={row[12]}, отгруз={row[13]}")
        count += 1
        if count >= 15:
            break
