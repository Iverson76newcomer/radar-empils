#!/usr/bin/env python3
"""Инвентаризация слайда 25 — найти все текстовые рамки и их структуру."""
from pptx import Presentation
from pptx.util import Emu

PPTX = "/home/z/my-project/upload/Презентация_для_новых_партнёров.pptx"
prs = Presentation(PPTX)

# Слайд 25 (индекс 24)
slide = prs.slides[24]

print(f"=== Slide 25 ===")
print(f"Total shapes: {len(slide.shapes)}\n")

def walk(shapes, indent=0):
    prefix = "  " * indent
    for i, sh in enumerate(shapes):
        stype = sh.shape_type
        try:
            l = Emu(sh.left).inches if sh.left else 0
            t = Emu(sh.top).inches if sh.top else 0
            w = Emu(sh.width).inches if sh.width else 0
            h = Emu(sh.height).inches if sh.height else 0
        except Exception:
            l=t=w=h=0
        print(f"{prefix}[{i}] type={stype} pos=({l:.2f},{t:.2f}) size=({w:.2f}×{h:.2f})")
        if sh.has_text_frame:
            tf = sh.text_frame
            for j, p in enumerate(tf.paragraphs):
                text = "".join(r.text for r in p.runs)
                if text.strip():
                    # Print each run with formatting info
                    runs_info = []
                    for r in p.runs:
                        if r.text.strip():
                            sz = r.font.size.pt if r.font.size else "?"
                            b = "B" if r.font.bold else "-"
                            col = "?"
                            try:
                                if r.font.color and r.font.color.rgb:
                                    col = str(r.font.color.rgb)
                            except Exception:
                                pass
                            runs_info.append(f"'{r.text}'[sz={sz},b={b},col={col}]")
                    print(f"{prefix}    P{j}: {text!r}")
                    if runs_info:
                        print(f"{prefix}       runs: {' '.join(runs_info)}")
        if sh.shape_type == 6:  # GROUP
            walk(sh.shapes, indent+1)

walk(slide.shapes)
