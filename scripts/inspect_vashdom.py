"""Изучаем структуру файла Ваш Дом (.xls — старый формат)."""
import subprocess
# .xls файл — нужно конвертировать в xlsx или читать xlrd
try:
    import xlrd
    wb = xlrd.open_workbook('/home/z/my-project/upload/Титан Ваш Дом 09.09.2026.xls')
    for sh_name in wb.sheet_names():
        ws = wb.sheet_by_name(sh_name)
        print(f"\n--- Лист: {sh_name} | rows={ws.nrows} | cols={ws.ncols}")
        for ri in range(min(ws.nrows, 50)):
            row = [str(ws.cell_value(ri, ci)) if ws.cell_value(ri, ci) != '' else '' for ci in range(ws.ncols)]
            print(' | '.join(row))
except ImportError:
    print("xlrd не установлен, пробуем libreoffice конвертацию")
    subprocess.run(['libreoffice', '--headless', '--convert-to', 'xlsx',
                    '/home/z/my-project/upload/Титан Ваш Дом 09.09.2026.xls',
                    '--outdir', '/home/z/my-project/scripts/'], check=True)
    from openpyxl import load_workbook
    wb = load_workbook('/home/z/my-project/scripts/Титан Ваш Дом 09.09.2026.xlsx', data_only=True)
    for sh_name in wb.sheetnames:
        ws = wb[sh_name]
        print(f"\n--- Лист: {sh_name} | dim={ws.dimensions}")
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50), values_only=False):
            vals = [str(c.value) if c.value is not None else '' for c in row]
            print(' | '.join(vals))
