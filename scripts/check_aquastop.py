import openpyxl
from openpyxl import load_workbook
from collections import defaultdict

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО') and not f_val.startswith('ТУ') and not f_val.startswith('ГОСТ'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# Search all sheets for "аквастоп" or "аква"
print("=== All sheets: 'аквастоп' or 'аква' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and ('аквастоп' in p['name'].lower() or 'аква' in p['name'].lower()):
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Search all sheets for 'absolute'
print("\n=== All sheets: 'absolute' or 'black' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and ('absolute' in p['name'].lower() or 'black' in p['name'].lower()):
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Search all sheets for 'modern'
print("\n=== All sheets: 'modern' ===")
for sn in wb2.sheetnames:
    for p in price_dict[sn]:
        if p['name'] and 'modern' in p['name'].lower():
            print(f"  [{sn}] '{p['name']}' фас={p['fасовка']} отгруз={p['отгрузочная']}")

# Check all NEWTONE names
print("\n=== NEWTONE: ALL unique names ===")
seen = set()
for p in price_dict['NEWTONE']:
    if p['name'] and p['name'] not in seen:
        seen.add(p['name'])
        print(f"  '{p['name']}'")
