#!/usr/bin/env python3
"""Analyze Презентация_для_новых_партнёров.pptx — extract structure & images."""
import os
import shutil
import zipfile
from pptx import Presentation
from pptx.util import Emu

PPTX = "/home/z/my-project/upload/Презентация_для_новых_партнёров.pptx"
EXTRACT_DIR = "/home/z/my-project/extracted_images/partners_pres"
MEDIA_DIR = os.path.join(EXTRACT_DIR, "ppt/media")

# Clean & extract
if os.path.exists(EXTRACT_DIR):
    shutil.rmtree(EXTRACT_DIR)
os.makedirs(EXTRACT_DIR, exist_ok=True)

with zipfile.ZipFile(PPTX) as zf:
    zf.extractall(EXTRACT_DIR)

# List media
print("=== MEDIA FILES ===")
if os.path.isdir(MEDIA_DIR):
    media_files = sorted(os.listdir(MEDIA_DIR), key=lambda x: (len(x), x))
    for f in media_files:
        size = os.path.getsize(os.path.join(MEDIA_DIR, f))
        print(f"  {f}  ({size:,} bytes)")
print()

# Map images to slides
prs = Presentation(PPTX)
print(f"=== SLIDES: {len(prs.slides)} ===\n")

# Build mapping of slide_num -> list of image refs
slides_info = []
for i, slide in enumerate(prs.slides, 1):
    info = {"num": i, "text_parts": [], "pictures": []}
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t.strip():
                    info["text_parts"].append(t.strip())
        if shape.shape_type == 13:  # Picture
            try:
                w = Emu(shape.width).inches
                h = Emu(shape.height).inches
            except Exception:
                w = h = 0
            try:
                rid = shape._element.blip_rId
                rel = slide.part.rels[rid]
                target = rel.target_ref
                # Get filename
                fname = os.path.basename(target)
                info["pictures"].append((fname, w, h))
            except Exception as e:
                info["pictures"].append((f"<unknown: {e}>", w, h))
    slides_info.append(info)

# Print summary
for info in slides_info:
    print(f"--- Slide {info['num']} ---")
    if info["text_parts"]:
        # Print first 15 lines of text
        print("  Text:")
        for t in info["text_parts"][:15]:
            print(f"    • {t[:200]}")
        if len(info["text_parts"]) > 15:
            print(f"    ... (+{len(info['text_parts'])-15} more lines)")
    if info["pictures"]:
        print("  Pictures:")
        for fname, w, h in info["pictures"]:
            print(f"    📷 {fname}  ({w:.1f}\" × {h:.1f}\")")
    # Heuristic: slide is "image-only" if very little text + a big picture
    total_text_chars = sum(len(t) for t in info["text_parts"])
    big_pics = [p for p in info["pictures"] if p[1] >= 5 and p[2] >= 3]
    if big_pics and total_text_chars < 50:
        print(f"  ⚠️ LIKELY IMAGE-BASED SLIDE (text={total_text_chars} chars, big_pics={len(big_pics)})")
    print()
