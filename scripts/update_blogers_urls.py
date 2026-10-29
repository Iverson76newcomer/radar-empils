"""Обновление TITAN_Региональные_блогеры_коротко.docx
   Заполняем URL для 3 каналов, у которых было «через YT поиск»:
   1. ИДЕИ ДЛЯ ДОМА → https://rutube.ru/channel/30125920/ (Rutube)
   2. Деревенский кузнец → https://www.youtube.com/@ruralsmith (YouTube)
   3. Художественная ковка металла → https://rutube.ru/channel/24685954/ (Rutube, ArtMetalLab)
"""
from docx import Document
from copy import deepcopy

doc = Document('/home/z/my-project/upload/TITAN_Региональные_блогеры_коротко.docx')

# Список каналов и их новых URL
updates = {
    'ИДЕИ ДЛЯ ДОМА': 'https://rutube.ru/channel/30125920/',
    'Деревенский кузнец': 'https://www.youtube.com/@ruralsmith',
    'Художественная ковка металла': 'https://rutube.ru/channel/24685954/',
}

# Перебираем все таблицы
for table in doc.tables:
    for row in table.rows:
        # Ищем ячейку с названием канала (обычно колонка 1 — «Канал»)
        for ci, cell in enumerate(row.cells):
            text = cell.text.strip()
            for channel_name, new_url in updates.items():
                if channel_name in text:
                    # В этой же ячейке нужно заменить «через YT поиск» на URL
                    # Если URL в отдельной ячейке — найдём её
                    print(f"Найден канал: {channel_name} в строке, ячейка {ci}")
                    # Ищем ячейку с «через YT поиск» в этой строке
                    for ci2, cell2 in enumerate(row.cells):
                        if 'через YT поиск' in cell2.text or 'через YT' in cell2.text:
                            print(f"  → Заменяем в ячейке {ci2}: '{cell2.text}' → '{new_url}'")
                            # Очищаем ячейку и записываем новый URL
                            for paragraph in cell2.paragraphs:
                                for run in paragraph.runs:
                                    run.text = ''
                            # Записываем новый URL в первый параграф
                            if cell2.paragraphs:
                                if cell2.paragraphs[0].runs:
                                    cell2.paragraphs[0].runs[0].text = new_url
                                else:
                                    # Если нет runs, добавляем новый
                                    from docx.oxml.ns import qn
                                    from docx.oxml import OxmlElement
                                    r = OxmlElement('w:r')
                                    t = OxmlElement('w:t')
                                    t.text = new_url
                                    r.append(t)
                                    cell2.paragraphs[0]._p.append(r)

# Сохраняем
out_path = '/home/z/my-project/download/TITAN_Региональные_блогеры_коротко.docx'
doc.save(out_path)
print(f'\n✅ Сохранено: {out_path}')

# Проверка
doc2 = Document(out_path)
print('\n=== ПРОВЕРКА ===')
for table in doc2.tables:
    for row in table.rows:
        for cell in row.cells:
            text = cell.text.strip()
            for channel_name in updates.keys():
                if channel_name in text:
                    # Найдём URL в этой строке
                    for cell2 in row.cells:
                        if 'http' in cell2.text:
                            print(f"  {channel_name}: {cell2.text.strip()}")
