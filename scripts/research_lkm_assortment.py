#!/usr/bin/env python3
"""
Script to research LCM (лакокрасочные материалы) assortment for each client
in the Empils distributor table. Uses web search and page reader via z-ai CLI.
"""
import openpyxl
import subprocess
import json
import re
import time
import os
import sys

# ===== CONFIG =====
INPUT_FILE = '/home/z/my-project/upload/Белые пятна + ушедшие клиенты_общая_таблица_24.06.26.xlsx'
OUTPUT_FILE = '/home/z/my-project/download/Белые пятна + ушедшие клиенты_обновлённая.xlsx'
CACHE_FILE = '/home/z/my-project/scripts/lkm_research_cache.json'

# Known LCM brands to search for in text
LCM_BRANDS = [
    # Russian manufacturers
    'Лакра', 'Рогнеда', 'Декарт', 'Престиж', 'Элкон', 'Ижсинтез',
    'Ярославские краски', 'Радуга Маляр', 'Радуга Малер', 'Царицыно', 'Царицинские краски',
    'НБХ', 'Проремонт', 'Профилюкс', 'Новоколор', 'Простокрашено', 'Профивуд',
    'Ореол', 'Крафор', 'Славянка', 'Оптимист', 'Светлица', 'Грида', 'Хват',
    'ВГТ', 'VGT', 'Палиж', 'Palizh', 'Неомид', 'Neomid', 'Расцвет',
    'Алия', 'Маршал', 'Marshall', 'Эмпилс', 'Empils', 'Титан', 'Titan',
    'Текс', 'Tex', 'АКЦ', 'Альтакс', 'Altax', 'Колорит',
    # STMs
    'Баупроф', 'Строймарка',
    # International
    'АВС Farben', 'АВС Фарбен', 'ABC Farben', 'Тиккурила', 'Tikkurila',
    'Dulux', 'Дюлакс', 'AkzoNobel', 'АкзоНобель', 'Акзо', 'Akzo',
    'Alpina', 'Альпина', 'Caparol', 'Капарол', 'Dufa', 'Дюфа',
    'FLUGGER', 'Flugger', 'Флюггер', 'Pinotex', 'Пинотекс',
    'Sadolin', 'Садолин', 'Beckers', 'Беккерс', 'Hammerite', 'Хаммерайт',
    'Jotun', 'Йотун', 'Hempel', 'Хемпел', 'Nippon', 'Kansai',
    'Marshall', 'Eskaro', 'Эскаро', 'Sniezka', 'Снежка',
    'Magnat', 'Магнат', 'Oswal', 'Berger', 'Sirca', 'Indigo',
    # Niche / special
    'Altax', 'Powertex', 'Ardex', 'Deccan',
]

# Simplified mapping for standardization of Дистр/сеть column
DISTANDARD_MAP = {
    'дистр': 'Дистрибьютор',
    'дитс/сеть': 'Дистрибьютор+Сеть',
    'дистр/сеть': 'Дистрибьютор+Сеть',
    'дистр/сеть': 'Дистрибьютор+Сеть',
    'сеть': 'Сеть',
    'сеть, 38 тт': 'Сеть',
    'субдистр': 'Субдистрибьютор',
    'сеть с торговой командой': 'Сеть',
}

def search_web(query, num=5):
    """Search web using z-ai CLI"""
    try:
        cmd = ['z-ai', 'function', '-n', 'web_search', '-a', json.dumps({"query": query, "num": num})]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return []
        data = json.loads(result.stdout)
        return data if isinstance(data, list) else []
    except Exception as e:
        print(f"  Search error: {e}", file=sys.stderr)
        return []

def read_page(url):
    """Read web page using z-ai CLI"""
    try:
        out_file = f'/tmp/page_read_{int(time.time()*1000)}.json'
        cmd = ['z-ai', 'function', '-n', 'page_reader', '-a', json.dumps({"url": url}), '-o', out_file]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        if result.returncode != 0:
            return ""
        with open(out_file, 'r') as f:
            data = json.load(f)
        if isinstance(data, dict) and 'data' in data:
            html = data['data'].get('html', '')
        elif isinstance(data, dict):
            html = data.get('html', '')
        else:
            html = ''
        # Strip HTML tags
        text = re.sub(r'<[^>]+>', ' ', html)
        text = re.sub(r'\s+', ' ', text).strip()
        # Clean up temp file
        try:
            os.unlink(out_file)
        except:
            pass
        return text[:5000]  # Limit text size
    except Exception as e:
        print(f"  Page read error for {url}: {e}", file=sys.stderr)
        return ""

def extract_brands_from_text(text):
    """Extract known LCM brand names from text"""
    found = []
    text_lower = text.lower()
    for brand in LCM_BRANDS:
        if brand.lower() in text_lower:
            # Avoid duplicates (case-insensitive)
            if not any(b.lower() == brand.lower() for b in found):
                found.append(brand)
    return found

def categorize_assortment(brands, company_name=''):
    """Split brands into main and additional assortment"""
    if not brands:
        return '', ''
    
    # Categorize by brand prominence
    main_brands = []
    additional_brands = []
    
    # Major domestic brands that are typically main assortment
    major_brands = ['Лакра', 'Рогнеда', 'Декарт', 'Текс', 'Новоколор', 'Профилюкс',
                    'АВС Farben', 'АВС Фарбен', 'ABC Farben', 'Престиж', 'Элкон',
                    'Тиккурила', 'Tikkurila', 'Dulux', 'Caparol', 'Alpina',
                    'Царицыно', 'Царицинские краски', 'Ярославские краски',
                    'ВГТ', 'VGT', 'Flugger', 'FLUGGER', 'Sadolin',
                    'Marshall', 'Sniezka', 'Magnat', 'Эмпилс', 'Титан',
                    'Радуга Маляр', 'Радуга Малер', 'Неомид', 'Neomid',
                    'Палиж', 'Palizh', 'Крафор', 'Ореол', 'Ижсинтез',
                    'Оптимист', 'Славянка']
    
    # STM brands
    stm_brands = ['Баупроф', 'Строймарка']
    
    for brand in brands:
        if brand in stm_brands:
            additional_brands.append(brand + ' (СТМ)')
        elif brand in major_brands:
            main_brands.append(brand)
        else:
            additional_brands.append(brand)
    
    # If no main brands found, put first 2-3 brands as main
    if not main_brands and brands:
        main_brands = brands[:min(3, len(brands))]
        additional_brands = brands[min(3, len(brands)):]
    
    return ', '.join(main_brands), ', '.join(additional_brands)

def load_cache():
    """Load research cache"""
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def save_cache(cache):
    """Save research cache"""
    with open(CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

def main():
    print("Loading Excel file...")
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active
    
    # Load cache
    cache = load_cache()
    print(f"Loaded cache with {len(cache)} entries")
    
    # Extract unique clients
    clients = {}
    for row in range(2, ws.max_row + 1):
        klient = ws.cell(row=row, column=4).value
        if not klient:
            continue
        klient = klient.strip()
        inn = ws.cell(row=row, column=5).value
        inn_key = str(inn).strip() if inn else klient
        
        if inn_key not in clients:
            clients[inn_key] = {
                'name': klient,
                'inn': inn_key,
                'rows': [row],
                'site': ws.cell(row=row, column=14).value,
                'address': ws.cell(row=row, column=15).value,
                'comments': ws.cell(row=row, column=8).value or '',
                'comment2': ws.cell(row=row, column=9).value or '',
                'distr': ws.cell(row=row, column=17).value,
            }
        else:
            clients[inn_key]['rows'].append(row)
            # Fill in missing data from duplicate rows
            if not clients[inn_key]['site'] and ws.cell(row=row, column=14).value:
                clients[inn_key]['site'] = ws.cell(row=row, column=14).value
            if not clients[inn_key]['address'] and ws.cell(row=row, column=15).value:
                clients[inn_key]['address'] = ws.cell(row=row, column=15).value
    
    print(f"Found {len(clients)} unique clients")
    
    # ===== TASK 2: Standardize Дистр/сеть column =====
    print("\n=== TASK 2: Standardizing 'Дистр/сеть' column ===")
    standardization_count = 0
    for inn_key, data in clients.items():
        current_val = data.get('distr', '')
        if current_val:
            val_stripped = str(current_val).strip()
            val_lower = val_stripped.lower()
            if val_lower in DISTANDARD_MAP:
                new_val = DISTANDARD_MAP[val_lower]
                for row in data['rows']:
                    ws.cell(row=row, column=17).value = new_val
                standardization_count += 1
                print(f"  Standardized: '{val_stripped}' -> '{new_val}' (client: {data['name']})")
            elif val_stripped == 'Дистр':
                for row in data['rows']:
                    ws.cell(row=row, column=17).value = 'Дистрибьютор'
                standardization_count += 1
                print(f"  Standardized: 'Дистр' -> 'Дистрибьютор' (client: {data['name']})")
            elif val_stripped == 'Дистр/сеть':
                for row in data['rows']:
                    ws.cell(row=row, column=17).value = 'Дистрибьютор+Сеть'
                standardization_count += 1
                print(f"  Standardized: 'Дистр/сеть' -> 'Дистрибьютор+Сеть' (client: {data['name']})")
    
    print(f"Standardized {standardization_count} client entries")
    
    # ===== TASK 1: Research LCM assortment =====
    print("\n=== TASK 1: Researching LCM assortment ===")
    
    # Add new column headers
    # Find the next available column after existing data (col 26 is the last real header)
    new_col_main = 27  # Основной ассортимент ЛКМ
    new_col_add = 28   # Дополнительный ассортимент ЛКМ
    ws.cell(row=1, column=new_col_main).value = 'Основной ассортимент ЛКМ'
    ws.cell(row=1, column=new_col_add).value = 'Дополнительный ассортимент ЛКМ'
    
    # Style the headers
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font_white = Font(bold=True, size=10, color='FFFFFF')
    ws.cell(row=1, column=new_col_main).font = header_font_white
    ws.cell(row=1, column=new_col_main).fill = header_fill
    ws.cell(row=1, column=new_col_main).alignment = Alignment(horizontal='center', wrap_text=True)
    ws.cell(row=1, column=new_col_add).font = header_font_white
    ws.cell(row=1, column=new_col_add).fill = header_fill
    ws.cell(row=1, column=new_col_add).alignment = Alignment(horizontal='center', wrap_text=True)
    
    # Set column widths
    ws.column_dimensions[openpyxl.utils.get_column_letter(new_col_main)].width = 35
    ws.column_dimensions[openpyxl.utils.get_column_letter(new_col_add)].width = 35
    
    processed = 0
    skipped = 0
    
    for inn_key, data in clients.items():
        client_name = data['name']
        
        # Skip if already cached
        if inn_key in cache and cache[inn_key].get('main') is not None:
            main_assortment = cache[inn_key].get('main', '')
            add_assortment = cache[inn_key].get('additional', '')
            for row in data['rows']:
                ws.cell(row=row, column=new_col_main).value = main_assortment
                ws.cell(row=row, column=new_col_add).value = add_assortment
            processed += 1
            continue
        
        # Step 1: Extract brands from comments
        all_comments = f"{data['comments']} {data['comment2']}"
        brands_from_comments = extract_brands_from_text(all_comments)
        
        # Step 2: Search web for client's LCM assortment
        brands_from_web = []
        site_url = data.get('site', '')
        
        # Clean up website URL
        if site_url and site_url not in ['-', 'NO SITE', '', None]:
            # Remove any non-URL text from site field
            site_clean = site_url.strip()
            # If site contains multiple URLs, take the first one
            if '\n' in site_clean:
                site_clean = site_clean.split('\n')[0].strip()
            if ' ' in site_clean and not site_clean.startswith('http'):
                site_clean = site_clean.split(' ')[0].strip()
        else:
            site_clean = ''
        
        # Search query: company name + LCM
        search_query = f"{client_name} лакокрасочные материалы бренды ассортимент"
        print(f"\n[{processed+1}/{len(clients)}] Researching: {client_name}")
        print(f"  Query: {search_query}")
        
        search_results = search_web(search_query, num=5)
        
        if search_results:
            # Extract brands from search snippets
            all_snippets = ' '.join([r.get('snippet', '') for r in search_results])
            all_titles = ' '.join([r.get('name', '') for r in search_results])
            brands_from_web = extract_brands_from_text(f"{all_titles} {all_snippets}")
            print(f"  Found from search snippets: {brands_from_web}")
            
            # Try to read the most relevant result (client's own website)
            for result in search_results[:2]:
                url = result.get('url', '')
                if url and any(domain in url.lower() for domain in 
                              [site_clean.replace('www.', '').replace('https://', '').replace('http://', '').lower()]):
                    print(f"  Reading page: {url}")
                    page_text = read_page(url)
                    if page_text:
                        brands_from_page = extract_brands_from_text(page_text)
                        if brands_from_page:
                            brands_from_web = list(set(brands_from_web + brands_from_page))
                            print(f"  Found from website: {brands_from_page}")
                    break
        
        # Also try to read client's website directly if we have it
        if site_clean and len(brands_from_web) < 2:
            # Normalize URL
            if not site_clean.startswith('http'):
                site_clean_url = f"https://{site_clean}"
            else:
                site_clean_url = site_clean
            
            print(f"  Trying direct website: {site_clean_url}")
            page_text = read_page(site_clean_url)
            if page_text:
                brands_from_page = extract_brands_from_text(page_text)
                if brands_from_page:
                    brands_from_web = list(set(brands_from_web + brands_from_page))
                    print(f"  Found from direct website: {brands_from_page}")
            
            # Also try searching for paint catalog on their site
            if len(brands_from_web) < 2:
                catalog_query = f"site:{site_clean} краски ЛКМ"
                print(f"  Searching site catalog: {catalog_query}")
                site_results = search_web(catalog_query, num=3)
                if site_results:
                    site_snippets = ' '.join([r.get('snippet', '') for r in site_results])
                    brands_from_site_search = extract_brands_from_text(site_snippets)
                    if brands_from_site_search:
                        brands_from_web = list(set(brands_from_web + brands_from_site_search))
                        print(f"  Found from site search: {brands_from_site_search}")
                    
                    # Read the most relevant page
                    for sr in site_results[:1]:
                        sr_url = sr.get('url', '')
                        if sr_url:
                            sr_text = read_page(sr_url)
                            if sr_text:
                                sr_brands = extract_brands_from_text(sr_text)
                                if sr_brands:
                                    brands_from_web = list(set(brands_from_web + sr_brands))
                                    print(f"  Found from site page: {sr_brands}")
        
        # Combine all found brands
        all_brands = list(set(brands_from_comments + brands_from_web))
        
        # Categorize into main and additional
        main_assortment, add_assortment = categorize_assortment(all_brands, client_name)
        
        # Special handling for companies that are clearly not LCM-focused
        if not all_brands:
            # Check if comments indicate this is not an LCM company
            comments_lower = all_comments.lower()
            if any(phrase in comments_lower for phrase in 
                   ['перестал заниматься', 'прекратил деятельность', 'компания ликвидирована',
                    'закрылись', 'не занимается больше продажей']):
                main_assortment = 'Неактивный клиент'
                add_assortment = ''
            elif 'ушедший клиент' in comments_lower:
                main_assortment = 'Нет данных (ушедший клиент)'
                add_assortment = ''
            else:
                main_assortment = 'Нет данных'
                add_assortment = ''
        
        print(f"  RESULT: Main: {main_assortment} | Additional: {add_assortment}")
        
        # Cache the result
        cache[inn_key] = {
            'name': client_name,
            'main': main_assortment,
            'additional': add_assortment,
            'brands_found': all_brands,
            'source': 'comments+web' if all_brands else 'none'
        }
        
        # Save cache periodically (every 10 clients)
        processed += 1
        if processed % 10 == 0:
            save_cache(cache)
            print(f"\n  [Cache saved at {processed} clients]")
        
        # Write to Excel
        for row in data['rows']:
            ws.cell(row=row, column=new_col_main).value = main_assortment
            ws.cell(row=row, column=new_col_add).value = add_assortment
        
        # Rate limiting: small delay between searches
        time.sleep(0.5)
    
    # Final cache save
    save_cache(cache)
    
    # Save the updated Excel
    print(f"\nSaving updated Excel to {OUTPUT_FILE}...")
    wb.save(OUTPUT_FILE)
    print("Done!")
    
    # Print summary
    with_data = sum(1 for v in cache.values() if v.get('main') and v['main'] not in ['Нет данных', 'Нет данных (ушедший клиент)', 'Неактивный клиент'])
    without_data = sum(1 for v in cache.values() if v.get('main') in ['Нет данных', 'Нет данных (ушедший клиент)', 'Неактивный клиент', ''])
    print(f"\nSummary:")
    print(f"  Clients with LCM data: {with_data}")
    print(f"  Clients without LCM data: {without_data}")
    print(f"  Standardized 'Дистр/сеть' entries: {standardization_count}")

if __name__ == '__main__':
    main()
