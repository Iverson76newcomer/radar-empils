#!/usr/bin/env python3
"""
Заполнить колонку "Сайт" в Distributors_2026.xlsx на основе 4 файлов коллег.
Для каждой компании извлекаем БАЗОВЫЙ URL (главную страницу), а не ссылку на товар.
"""

import os
import re
import shutil
from urllib.parse import urlparse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SRC = "/home/z/my-project/upload/Distributors_2026.xlsx"
DST = "/home/z/my-project/download/Distributors_2026_с_сайтами.xlsx"

# Копируем исходный файл
shutil.copy(SRC, DST)

# === Утилиты ===

def normalize(name):
    """Нормализуем название для сравнения: нижний регистр, без лишних пробелов, без кавычек."""
    if not name:
        return ""
    s = str(name).lower().strip()
    # Удаляем лишние пробелы и неразрывные пробелы
    s = re.sub(r'\s+', ' ', s.replace('\xa0', ' '))
    # Удаляем кавычки
    s = re.sub(r'["«»\'\"]+', '', s)
    # Удаляем лишние знаки в конце
    s = s.rstrip(' ,.')
    return s

def extract_base_url(url):
    """Извлечь базовый URL (scheme + host), без path/query.
    Пример: http://www.uss74.ru/catalog?product_id=23546 → http://www.uss74.ru
    """
    if not url:
        return None
    s = str(url).strip()
    if not s:
        return None
    # Если URL без схемы — добавим
    if not s.startswith(('http://', 'https://')):
        s = 'http://' + s
    try:
        parsed = urlparse(s)
        if not parsed.netloc:
            return None
        # Возвращаем scheme + netloc
        base = f"{parsed.scheme}://{parsed.netloc}"
        return base
    except Exception:
        return None

# === 1. Считываем все 4 файла коллег и строим маппинг имя → сайт ===

mapping = {}  # normalized_name → base_url

def add_mapping(name, url):
    """Добавляет в маппинг, нормализуя имя и URL."""
    if not name or not url:
        return
    base = extract_base_url(url)
    if not base:
        return
    norm = normalize(name)
    if not norm:
        return
    # Не перезаписываем, если уже есть (первый файл — приоритет)
    if norm not in mapping:
        mapping[norm] = base
        # Также добавим вариант с удалением организационно-правовой формы
        # (ООО, ИП, АО, ЗАО и т.д.)
        without_opf = re.sub(r'\b(ооо|оао|зао|иооо|ип|ао|пао)\b\s*', '', norm).strip()
        if without_opf and without_opf != norm:
            if without_opf not in mapping:
                mapping[without_opf] = base

# Файл 1: Сайты дистрибьюторов ДВ и ВС_2024.xlsx
# Структура: A=Регион, B=№, C=Юр.наименование, D=Фирменное (коммерческое) наименование, E=Сайт
wb1 = load_workbook("/home/z/my-project/upload/Сайты дистрибьюторов ДВ и ВС_2024.xlsx", data_only=True)
ws1 = wb1["Лист1"]
for row in ws1.iter_rows(min_row=3, values_only=True):
    region, num, legal_name, comm_name, site, ftm, city = row[:7]
    if comm_name and site:
        add_mapping(comm_name, site)
    if legal_name and site:
        add_mapping(legal_name, site)

# Файлы 2-4: Структура: B=SBU, C=Наименование, D=Город, E=ТМ, F=Ссылка
# Имя появляется только при смене дистрибьютора, ниже идут строки с ТМ и ссылками.
# Нужно "fill down" имя и сопоставить любую непустую ссылку к этому имени.

files_234 = [
    "/home/z/my-project/upload/Ссылки на сайты дистрибьюторов Москва, Центр, Северо-Запад.xlsx",
    "/home/z/my-project/upload/Ссылки на сайты дистрибьюторов Поволжье Юг.xlsx",
    "/home/z/my-project/upload/Ссылки на сайты дистрибьюторов Урал-Западная Сибирь.xlsx",
]

for fpath in files_234:
    wb = load_workbook(fpath, data_only=True)
    ws = wb["Лист1"]
    current_name = None
    current_city = None

    for row in ws.iter_rows(min_row=3, values_only=True):
        # Структура: A=empty, B=SBU(только на новой), C=Наименование(только на новой), D=Город(только на новой), E=ТМ, F=Ссылка
        sbu = row[1] if len(row) > 1 else None
        name = row[2] if len(row) > 2 else None
        city = row[3] if len(row) > 3 else None
        tm = row[4] if len(row) > 4 else None
        link = row[5] if len(row) > 5 else None

        if name and str(name).strip():
            current_name = str(name).strip()
            current_city = city
        if link and str(link).strip() and current_name:
            add_mapping(current_name, link)

print(f"=== Построен маппинг: {len(mapping)} записей ===\n")
print("=== Содержимое маппинга ===")
for k, v in sorted(mapping.items()):
    print(f"  {k!r:50} → {v}")

# === 2. Открываем основной файл и заполняем сайт ===

wb_main = load_workbook(DST)
ws_main = wb_main["Лист1"]

# Подсветка для заполненных ячеек
fill_matched = PatternFill("solid", fgColor="C8E6C9")  # светло-зелёный
fill_searched = PatternFill("solid", fgColor="FFF9C4")  # светло-жёлтый (web search)
fill_unmatched = PatternFill("solid", fgColor="FFCDD2")  # светло-красный

# Стиль для URL
font_url = Font(name="Calibri", size=10, color="0563C1", underline="single")

matched_count = 0
searched_count = 0
unmatched_count = 0
unmatched_names = []

for row_idx in range(2, ws_main.max_row + 1):
    client = ws_main.cell(row=row_idx, column=4).value
    if not client or not str(client).strip():
        continue

    norm = normalize(client)
    if norm in mapping:
        # Точное совпадение
        ws_main.cell(row=row_idx, column=5, value=mapping[norm])
        ws_main.cell(row=row_idx, column=5).font = font_url
        ws_main.cell(row=row_idx, column=5).fill = fill_matched
        ws_main.cell(row=row_idx, column=5).alignment = Alignment(horizontal="left", vertical="center")
        matched_count += 1
    else:
        # Пробуем fuzzy: ищем в mapping ключи, которые содержат важные токены из client
        # Например, "Уралстройсервис-опт ООО" → ищем "уралстройсервис"
        # Или "Спектр ООО" → ищем "спектр"
        # Удаляем OPF и общие слова
        cleaned = re.sub(r'\b(ооо|оао|зао|ип|ао|пао|сеть|компания|торговый дом|тд)\b\s*', '', norm).strip()
        cleaned = re.sub(r'[^a-zа-яё0-9\s-]', '', cleaned)
        tokens = [t for t in re.split(r'[\s\-]+', cleaned) if len(t) >= 3]

        best_match = None
        best_score = 0
        for key, url in mapping.items():
            key_cleaned = re.sub(r'\b(ооо|оао|зао|ип|ао|пао|сеть|компания|торговый дом|тд)\b\s*', '', key).strip()
            key_cleaned = re.sub(r'[^a-zа-яё0-9\s-]', '', key_cleaned)
            key_tokens = set(t for t in re.split(r'[\s\-]+', key_cleaned) if len(t) >= 3)
            # Считаем совпадение токенов
            common = tokens and key_tokens and set(tokens) & key_tokens
            if common:
                score = len(common)
                if score > best_score:
                    best_score = score
                    best_match = url

        if best_match:
            ws_main.cell(row=row_idx, column=5, value=best_match)
            ws_main.cell(row=row_idx, column=5).font = font_url
            ws_main.cell(row=row_idx, column=5).fill = fill_matched
            ws_main.cell(row=row_idx, column=5).alignment = Alignment(horizontal="left", vertical="center")
            matched_count += 1
        else:
            # Не нашли в маппинге
            unmatched_names.append((row_idx, client))
            unmatched_count += 1

# Сохраняем
wb_main.save(DST)

print(f"\n{'='*60}")
print(f"=== РЕЗУЛЬТАТ ===")
print(f"{'='*60}")
print(f"✅ Найдено в файлах коллег:        {matched_count}")
print(f"❌ Не найдено (требуют web-поиска): {unmatched_count}")
print(f"   Всего клиентов в файле:         {matched_count + unmatched_count}")
print(f"\n=== НЕ НАЙДЕННЫЕ КЛИЕНТЫ (для web-поиска) ===")
for row_idx, name in unmatched_names:
    print(f"  R{row_idx}: {name}")

print(f"\n✅ Файл сохранён: {DST}")
print(f"   Размер: {os.path.getsize(DST):,} байт")
