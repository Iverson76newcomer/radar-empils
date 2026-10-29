"""Прочитать текущее состояние листа 'Кассовый поток (детально)' и понять, какие данные есть."""
import openpyxl
from openpyxl import load_workbook

wb = load_workbook("/home/z/my-project/upload/Кассовый_разрыв_Олимп_модель_v2.xlsx", data_only=True)
ws = wb["Кассовый поток (детально)"]

print(f"=== Лист 'Кассовый поток (детально)': {ws.max_row} строк × {ws.max_column} колонок ===\n")

# Шапка (строка 4)
print("Шапка (строка 4):")
for c in range(1, ws.max_column + 1):
    cell = ws.cell(row=4, column=c)
    print(f"  Col {c} ({cell.coordinate}): {cell.value}")

# Все строки данных
print("\n\n=== Все строки данных ===")
for r in range(5, ws.max_row + 1):
    print(f"\n--- Строка {r} ---")
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=4, column=c).value
        val = ws.cell(row=r, column=c).value
        if val is not None:
            print(f"  {header}: {val}")

# Также посмотреть ширины колонок
print("\n\n=== Ширины колонок ===")
for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O']:
    dim = ws.column_dimensions[col_letter]
    print(f"  {col_letter}: width={dim.width}")
