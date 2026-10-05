"""
Fix cover slide — recreate it with correct EMU values.
1 cm = 360000 EMU, so 13.25 cm = 4770000 EMU (not 13250000).
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import shutil

SRC = "/home/z/my-project/download/Презентация_для_ГД_final.pptx"
OUT = "/home/z/my-project/download/Презентация_для_ГД_final.pptx"

# Palette
BG          = RGBColor(0xF5, 0xF2, 0xEC)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
INK         = RGBColor(0x2F, 0x3A, 0x48)
MUTE        = RGBColor(0x8A, 0x95, 0xA3)
RED         = RGBColor(0xC0, 0x00, 0x00)
SAND        = RGBColor(0xC9, 0xA7, 0x7C)
DARK_BG     = RGBColor(0x2F, 0x3A, 0x48)
DARKER_BG   = RGBColor(0x44, 0x4A, 0x5A)

def add_text(slide, x, y, w, h, text, *,
             size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
             va=MSO_ANCHOR.TOP, font="Inter", spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = va
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if spacing:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(spacing))
    return tb

def add_rect(slide, x, y, w, h, fill, *,
             line=None, line_w=6350, corner_radius=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner_radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Emu(line_w)
    else:
        s.line.fill.background()
    if corner_radius is not None:
        try: s.adjustments[0] = corner_radius
        except: pass
    s.shadow.inherit = False
    return s


# Open and find the cover slide (slide9.xml is at position 1 in sldIdLst)
prs = Presentation(SRC)

# Identify cover slide by content (it has "Дорога к лидерству" text)
cover_idx = None
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame and 'Дорога к лидерству' in shape.text_frame.text:
            cover_idx = i
            break
    if cover_idx is not None:
        break

print(f"Cover slide found at index {cover_idx}")

# Remove all shapes from cover slide
cover_slide = list(prs.slides)[cover_idx]
# Get a list of all shape elements and remove them
sp_tree = cover_slide.shapes._spTree
shapes_to_remove = []
for elem in sp_tree:
    if elem.tag.endswith('}sp') or elem.tag.endswith('}pic') or elem.tag.endswith('}grpSp') or elem.tag.endswith('}cxnSp'):
        shapes_to_remove.append(elem)
for elem in shapes_to_remove:
    sp_tree.remove(elem)
print(f"Removed {len(shapes_to_remove)} shapes from cover slide")

# Reset background to dark
bg = cover_slide.background
bg.fill.solid()
bg.fill.fore_color.rgb = DARK_BG

# Helper to convert cm → EMU correctly
def cm(v):
    return Emu(int(v * 360000))

# ── Rebuild cover slide with CORRECT EMU values ────────────────────────
# Slide is 25.40 x 14.29 cm

# Top accent line
add_rect(cover_slide, Emu(0), Emu(0), prs.slide_width, Emu(40000), RED)

# TITAN brand mark (top-left)
add_text(cover_slide, cm(0.56), cm(0.56),
         cm(4.17), cm(1.00),
         "TITAN", size=14, bold=True, color=RED,
         font="Inter", spacing=400)
# Vertical divider
add_rect(cover_slide, cm(5.97), cm(0.61),
         cm(0.06), cm(0.89), MUTE)
# Eyebrow text
add_text(cover_slide, cm(6.39), cm(0.67),
         cm(19.44), cm(0.78),
         "ПРОГРАММА РАЗВИТИЯ · 2026", size=10, bold=True,
         color=SAND, font="Inter", spacing=300)

# Decorative red vertical accent bar (left of title)
add_rect(cover_slide, cm(0.78), cm(4.95),
         cm(0.22), cm(6.39), RED)

# Main title
add_text(cover_slide, cm(1.56), cm(4.95),
         cm(20.83), cm(2.50),
         "Дорога к лидерству",
         size=44, bold=True, color=WHITE,
         font="Inter", align=PP_ALIGN.LEFT)

# Subtitle
add_text(cover_slide, cm(1.56), cm(7.50),
         cm(20.83), cm(0.83),
         "Программа развития бренда ЭМПИЛС в сегменте ЛКМ",
         size=18, color=RGBColor(0xCC, 0xCC, 0xCC),
         font="Inter", align=PP_ALIGN.LEFT)

# Sub-subtitle: brands line
add_text(cover_slide, cm(1.56), cm(8.50),
         cm(20.83), cm(0.56),
         "TITAN · ОРЕОЛ · NEWTONE · AVANTGARDE",
         size=12, color=MUTE,
         font="Inter", align=PP_ALIGN.LEFT, spacing=200)

# Bottom info bar (just a divider line)
add_rect(cover_slide, Emu(0), cm(12.50),
         prs.slide_width, cm(0.04), DARKER_BG)

# Left side of bottom bar
add_text(cover_slide, cm(0.56), cm(12.75),
         cm(12.50), cm(0.50),
         "Доклад для ГД · Октябрь 2026", size=11, bold=True,
         color=WHITE, font="Inter")
add_text(cover_slide, cm(0.56), cm(13.30),
         cm(12.50), cm(0.40),
         "ЗАО «Эмпилс» · Ростов-на-Дону", size=9, color=MUTE,
         font="Inter")

# Right side of bottom bar
add_text(cover_slide, cm(14.00), cm(12.75),
         cm(11.00), cm(0.50),
         "Программа TITAN · Стандартизация · Рост",
         size=11, bold=True, color=SAND,
         font="Inter", align=PP_ALIGN.RIGHT, spacing=200)
add_text(cover_slide, cm(14.00), cm(13.30),
         cm(11.00), cm(0.40),
         "8 слайдов · 5 ключевых проектов",
         size=9, color=MUTE, font="Inter", align=PP_ALIGN.RIGHT)

# Save
prs.save(OUT)
print(f"Saved: {OUT}")

# Verify shapes positions
prs2 = Presentation(OUT)
cover_slide2 = list(prs2.slides)[0]
import re
print(f"\nCover slide shapes after fix:")
for shape in cover_slide2.shapes:
    try:
        x = shape.left / 360000
        y = shape.top / 360000
        w = shape.width / 360000
        h = shape.height / 360000
        text = shape.text_frame.text[:40] if shape.has_text_frame else ''
        bottom = y + h
        oob = " ⚠️" if bottom > 14.29 else ""
        print(f"  x={x:5.2f} y={y:5.2f} w={w:5.2f} h={h:5.2f} bot={bottom:5.2f}cm  text={text!r}{oob}")
    except:
        pass
