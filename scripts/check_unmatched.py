import openpyxl
from openpyxl import load_workbook
from collections import defaultdict
import re

wb2 = load_workbook("/home/z/my-project/upload/21.08.2026 РОССИЯ с НДС, руб..xlsx", data_only=True)

price_dict = defaultdict(list)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    current_name = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True), 14):
        f_val = row[5] if len(row) > 5 else None
        m_val = row[12] if len(row) > 12 else None
        n_val = row[13] if len(row) > 13 else None
        h_val = row[7] if len(row) > 7 else None
        if f_val is not None and isinstance(f_val, str) and not f_val.startswith('СТО'):
            current_name = f_val.strip()
        if (m_val is not None and isinstance(m_val, (int, float))) or (n_val is not None and isinstance(n_val, (int, float))):
            price_dict[sn].append({
                'name': current_name,
                'fасовка': h_val,
                'базовая': m_val,
                'отгрузочная': n_val,
            })

# Search Расцвет sheet for "3 в 1 белая"
print("=== Расцвет: '3 в 1 белая' ===")
for p in price_dict['Расцвет']:
    if p['name'] and '3 в 1' in p['name'].lower() and 'белая' in p['name'].lower():
        print(f"  фас={p['fасовка']}, отгруз={p['отгрузочная']}, name='{p['name']}'")

print("\n=== Расцвет: contains 'белая 3 в 1' ===")
for p in price_dict['Расцвет']:
    if p['name'] and 'белая' in p['name'].lower() and '3 в 1' in p['name'].lower():
        print(f"  фас={p['fасовка']}, отгруз={p['отгрузочная']}, name='{p['name']}'")

# Search ОРЕОЛ TITAN sheets for 'молотков'
print("\n=== ОРЕОЛ TITAN: 'молотк' ===")
for p in price_dict['ОРЕОЛ TITAN']:
    if p['name'] and 'молотк' in p['name'].lower():
        print(f"  фас={p['fасовка']}, отгруз={p['отгрузочная']}, name='{p['name']}'")

print("\n=== ОРЕОЛ TITAN(л): 'молотк' ===")
for p in price_dict['ОРЕОЛ TITAN(л)']:
    if p['name'] and 'молотк' in p['name'].lower():
        print(f"  фас={p['fасовка']}, отгруз={p['отгрузочная']}, name='{p['name']}'")

# Check Простокрашено
print("\n=== Простокрашено! БАУ: sample ===")
for p in price_dict['Простокрашено! БАУ'][:10]:
    print(f"  name='{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Check Ореол
print("\n=== Ореол: sample ===")
for p in price_dict['Ореол'][:10]:
    print(f"  name='{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")

# Check Ореол Дисконт
print("\n=== Ореол Дисконт: sample ===")
for p in price_dict['Ореол Дисконт'][:10]:
    print(f"  name='{p['name']}' | фас={p['fасовка']} | отгруз={p['отгрузочная']}")
